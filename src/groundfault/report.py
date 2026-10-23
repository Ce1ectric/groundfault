# -*- coding: utf-8 -*-
"""Figures and tables for the A4 test report (Word/PDF).

Design rules (uniform, scientific typesetting)
----------------------------------------------
* **No figure titles** – the caption is written in the report.
* Figures are created in their **final size** (type area of A4 with 2.54 cm margins → 15.9 cm width, half width
  7.7 cm) and inserted without scaling → font size in the figure = font size in the PDF.
* One font for everything (default: Aptos → Arial → Liberation Sans → DejaVu Sans), also for formula symbols
  (mathtext); axis labels in the form "quantity in unit" (DIN 461).
* One y-axis per diagram (no twin axes) – different quantities → panels stacked on a common time axis and
  labelled with panel letters (a), (b), …
* Fixed colour per quantity (L1, L2, L3, …), thin lines, subtle grid, a legend from two data series on.
* Export as PDF (vector, embedded font), SVG (text stays text) and PNG (300 dpi).

Tables: CSV (semicolon; decimal comma in German, decimal point in English), XLSX, Markdown, LaTeX and – if
``python-docx`` is installed – DOCX with ready-formatted tables to copy into the report.

Language
--------
Texts in figures and tables (axis labels, legends, annotations, table headers and row labels) are rendered in the output
language of [`groundfault.i18n`][groundfault.i18n] that is active when the figure or table is created. Numbers are
written with a decimal comma in German and a decimal point in English; ``ReportStyle.decimal_comma`` overrides this.
Exported tables get their column names and code values from [`groundfault.i18n.localize`][groundfault.i18n.localize]
(German names in German, the English names in English). Formula symbols are the same in both languages, except for
subscripts that are words (coil current ``I_Sp``/``I_coil``, LV network ``I_ON``/``I_LV``, sequence systems
``U_Mit``/``U_pos`` …).

Examples
--------
>>> from groundfault import i18n, report
>>> with i18n.language("de"):
...     report.number(0.012345, 3)
'0,0123'
>>> report.number(0.012345, 3, lang="en")
'0.0123'
"""
from __future__ import annotations

import contextlib
import importlib
import math
import os
import re
import warnings
from dataclasses import dataclass, replace
from decimal import ROUND_HALF_UP, Decimal
from typing import Dict, Iterable, List, Mapping, Optional, Sequence, Tuple, Union

import numpy as np
import pandas as pd

import matplotlib
import matplotlib.pyplot as plt
from matplotlib import ticker

from . import i18n
from .i18n import tr
from .earthing import QUANTITY_EPR, QUANTITY_TOUCH, SOURCE_FAULT_TEST, SOURCE_GROUNDMEAS_FAULT, SOURCE_INJECTION
from .station import MAX_PROBE

__all__ = ["ReportStyle", "STYLE", "set_style", "used_font", "activate_style", "figure", "save",
           "CM", "COLORS", "GREY", "GRID", "ROLE_COLORS", "INTERVAL_COLORS", "COLOR_DISTURBED", "COLOR_MAX",
           "COLOR_TYPICAL",
           "instantaneous", "rms_trend", "harmonics_trend", "spectrum", "spectrum_comparison",
           "voltage_spectrum", "phasor_diagram", "source_fit", "source_plane", "ires_over_ze",
           "pq_series", "pq_duration", "pq_profile", "touch_voltage_cases",
           "v_curve", "station_voltages", "station_frequency", "station_background", "station_colors", "probe_label",
           "earthing_comparison", "reduction_factor", "typical_band", "lv_network", "transformer_power",
           "transformer_impedance", "number",
           "format_table", "export_table", "export_tables", "docx_tables", "report_tables", "table_title",
           "all_figures", "slug"]

CM = 1 / 2.54
"""Inch per cm (matplotlib sizes are given in inch)."""

COLORS = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300", "#4a3aa7", "#e34948"]
"""Categorical colour sequence (fixed order, checked for colour-vision deficiencies)."""
GREY = "#6b6b6b"
"""Neutral grey (reference lines, secondary data)."""
GRID = "#d9d9d9"
"""Grid colour."""
ROLE_COLORS = {"u_l1e": COLORS[0], "u_l2e": COLORS[1], "u_l3e": COLORS[2], "i_l1": COLORS[0],
               "i_l2": COLORS[1], "i_l3": COLORS[2], "i_res": COLORS[6], "i_coil": COLORS[3],
               "u_e": COLORS[6], "u_t": COLORS[1], "i_shield": COLORS[2], "i_pen": COLORS[4], "u_ne": COLORS[7]}
"""Fixed colour per quantity (the colour follows the quantity, not the order in the figure)."""
INTERVAL_COLORS = {"pre_fault": "#9ec5ee", "steady_state": "#a8e0c8", "post_fault": "#f7c3a6"}
"""Background colours of the analysis intervals."""
COLOR_DISTURBED = "#c0392b"
"""Disturbed states in the steady-state interval (hatched; active residual current increase, network change)."""
COLOR_MAX = "#111111"
"""Probe maximum max(UE) – the same colour in all figures."""
COLOR_TYPICAL = "#c9c9c9"
"""Typical range of the reduction factor."""

i18n.register({
    # ---- panels: one per physical quantity
    "report.group_line_to_earth": {"en": "Line-to-earth voltage", "de": "Leiter-Erde-Spannung"},
    "report.group_epr_touch": {"en": "Earth potential rise/touch voltage", "de": "Erdungs-/Berührungsspannung"},
    "report.group_fault_current": {"en": "Fault current", "de": "Fehlerstrom"},
    "report.group_coil_current": {"en": "Coil current", "de": "Spulenstrom"},
    "report.group_line_current": {"en": "Line current", "de": "Leiterstrom"},
    "report.group_shield_pen": {"en": "Shield/PEN current", "de": "Schirm-/PEN-Strom"},
    "report.group_voltage": {"en": "Voltage", "de": "Spannung"},
    "report.group_current": {"en": "Current", "de": "Strom"},
    # ---- analysis intervals and the phases of the fault
    "report.interval_pre_fault": {"en": "pre-fault", "de": "vor"},
    "report.interval_steady_state": {"en": "steady state", "de": "stationär"},
    "report.interval_post_fault": {"en": "post-fault", "de": "nach"},
    "report.interval_disturbed": {"en": "disturbed", "de": "gestört"},
    "report.before_fault": {"en": "before the fault", "de": "vor dem Fehler"},
    "report.during_fault": {"en": "during the fault", "de": "Erdschluss"},
    "report.after_fault": {"en": "after the fault", "de": "nach dem Fehler"},
    # ---- subscripts of formula symbols that are words
    "report.sub_coil": {"en": "coil", "de": "Sp"},
    "report.sub_pos": {"en": "pos", "de": "Mit"},
    "report.sub_neg": {"en": "neg", "de": "Gegen"},
    "report.sub_zero": {"en": "zero", "de": "Null"},
    "report.sub_lv_current": {"en": "LV", "de": "ON"},
    "report.sub_upstream": {"en": "up", "de": "vor"},
    "report.sub_hv": {"en": "HV", "de": "HS"},
    "report.not_significant": {"en": "not significant", "de": "nicht signifikant"},
    "report.transformer_share": {"en": "transformer", "de": "Trafo"},
    "report.hv_share": {"en": "HV network", "de": "HS-Netz"},
    "report.no_upstream": {"en": "No upstream impedance in the result (pre-fault or post-fault interval missing).",
                           "de": "Keine Vorimpedanz im Ergebnis (Vor- oder Nachfehlerbereich fehlt)."},
    "report.sub_lv_impedance": {"en": "LV", "de": "Ortsnetz"},
    # ---- axis labels, legend entries and annotations
    "report.rms": {"en": "RMS value", "de": "Effektivwert"},
    "report.rms_short": {"en": "RMS", "de": "Eff."},
    "report.touch_voltage_axis": {"en": "$U_\\mathrm{{T}}$ in V", "de": "$U_\\mathrm{{T}}$ in V"},
    "report.permissible_touch_voltage": {"en": "permissible touch voltage $U_\\mathrm{{Tp}}$ = {value} V",
                                         "de": "zulässige Berührungsspannung $U_\\mathrm{{Tp}}$ = {value} V"},
    "report.harmonic": {"en": "H{order} ({f} Hz)", "de": "{order}. Harm. ({f} Hz)"},
    "report.order_axis": {"en": "Order {symbol}", "de": "Ordnung {symbol}"},
    "report.measured_range": {"en": "measured range", "de": "Messbereich"},
    "report.scaled": {"en": "{symbol} scaled", "de": "{symbol} skaliert"},
    "report.assessed_curve": {"en": "$I_\\mathrm{{RES,ass}}$ (with margins)",
                              "de": "$I_\\mathrm{{RES,Bew}}$ (mit Zuschlägen)"},
    "report.source_measured": {"en": "measured source", "de": "gemessene Quelle"},
    "report.source_fitted": {"en": "fitted source", "de": "geschätzte Quelle"},
    "report.with_source": {"en": "{symbol} ({source})", "de": "{symbol} ({source})"},
    "report.single_tests": {"en": "single tests", "de": "einzelne Versuche"},
    "report.remainder_not_fitted": {"en": "remainder (not fitted)", "de": "Rest (nicht gefittet)"},
    # ---- power-quality measurement
    "report.pq_largest_phase": {"en": "largest phase per 10 min", "de": "größter Leiter je 10 min"},
    "report.pq_quantile": {"en": "{q} % quantile", "de": "{q}-%-Quantil"},
    "report.pq_with_factor": {"en": "{q} % quantile × {k}", "de": "{q}-%-Quantil × {k}"},
    "report.pq_time_exceeded": {"en": "Time exceeded in %", "de": "Überschreitungsdauer in %"},
    "report.pq_hour": {"en": "Hour of the day", "de": "Stunde des Tages"},
    "report.pq_workdays": {"en": "Mon–Fri", "de": "Mo–Fr"},
    "report.pq_saturday": {"en": "Saturday", "de": "Samstag"},
    "report.pq_sunday": {"en": "Sunday", "de": "Sonntag"},
    "report.per_100_a": {"en": "{symbol} in V per 100 A", "de": "{symbol} in V je 100 A"},
    "report.epr": {"en": "earth potential rise", "de": "Erdungsspannung"},
    "report.touch_voltage": {"en": "touch voltage", "de": "Berührungsspannung"},
    "report.fault_component_span": {"en": "fault-related component (span of the methods)",
                                    "de": "fehlerbedingter Anteil (Spanne der Verfahren)"},
    "report.uncertain_background_methods": {"en": "uncertain (background/methods)",
                                            "de": "unsicher (Grundlast/Verfahren)"},
    "report.uncertain": {"en": "uncertain", "de": "unsicher"},
    "report.measured_directly": {"en": "measured directly (incl. background)",
                                 "de": "direkt gemessen (inkl. Grundlast)"},
    "report.measured_during_fault": {"en": "measured during the fault", "de": "im Fehler gemessen"},
    "report.background_before_fault": {"en": "background before the fault", "de": "Grundlast vor dem Fehler"},
    "report.fault_component": {"en": "fault-related component", "de": "fehlerbedingter Anteil"},
    "report.injection_at_distance": {"en": "Earthing measurement {distance} (ID {id}{date})",
                                     "de": "Erdungsmessung {distance} (ID {id}{date})"},
    "report.injection_id": {"en": "Earthing measurement ID {id}{distance}", "de": "Erdungsmessung ID {id}{distance}"},
    "report.injection_interpolated": {"en": "Earthing measurement interpolated", "de": "Erdungsmessung interpoliert"},
    "report.fault_probe": {"en": "Fault test {channel}{distance}, fault-related",
                           "de": "Erdschluss {channel}{distance}, fehlerbedingt"},
    "report.fault_probe_symbol": {"en": "Fault test {channel}{distance}: {symbol}, fault-related",
                                  "de": "Erdschluss {channel}{distance}: {symbol}, fehlerbedingt"},
    "report.fault_groundmeas": {"en": "Fault test from groundmeas (ID {id}{distance})",
                                "de": "Erdschluss aus groundmeas (ID {id}{distance})"},
    "report.fault_quotient": {"en": "Fault test {name}: {quotient}", "de": "Erdschluss {name}: {quotient}"},
    "report.probe_maximum": {"en": "Probe maximum{probes}", "de": "Sondenmaximum{probes}"},
    "report.fault_to_injection": {"en": "Fault test / earthing measurement", "de": "Erdschluss / Erdungsmessung"},
    "report.reduction_factor_axis": {"en": "Reduction factor {symbol}", "de": "Reduktionsfaktor {symbol}"},
    "report.from_shield_current": {"en": "from the shield current {quotient}", "de": "aus dem Schirmstrom {quotient}"},
    "report.rms_over_harmonics": {"en": "RMS value over the harmonics", "de": "Effektivwert über die Harmonischen"},
    "report.typical_at_50hz": {"en": "typical at 50 Hz ({low}–{high})", "de": "üblich bei 50 Hz ({low}–{high})"},
    "report.fault_rze": {"en": "Fault test {symbol} = {quotient}", "de": "Erdschluss {symbol} = {quotient}"},
    "report.trend_r": {"en": "trend {formula}, b = {b} (orders {orders})",
                       "de": "Verlauf {formula}, b = {b} (Ordnungen {orders})"},
    "report.r_direct": {"en": "r direct (as measured, with background)", "de": "r direkt (wie gemessen, mit Grundlast)"},
    "report.span_methods": {"en": "span of the superposition methods", "de": "Spanne der Überlagerungsverfahren"},
    "report.injection_ze": {"en": "{label}: {symbol}", "de": "{label}: {symbol}"},
    "report.lv_current": {"en": "{symbol} into the LV network (fault-related)",
                          "de": "{symbol} ins Ortsnetz (fehlerbedingt)"},
    "report.before_star_point": {"en": "before the transformer star point (LV feeders)",
                                 "de": "vor dem Trafoknoten (NS-Abgänge)"},
    "report.after_star_point": {"en": "after the transformer star point (to the main earthing busbar)",
                                "de": "nach dem Trafoknoten (zur HES)"},
    "report.fault_current": {"en": "Fault current {symbol}", "de": "Fehlerstrom {symbol}"},
    "report.earth_current": {"en": "Earth current {symbol}", "de": "Erdungsstrom {symbol}"},
    # ---- errors
    "report.no_recording_path": {"en": "Pass the recording (rec) – its path is not stored in the result.",
                                 "de": "Aufzeichnung (rec) übergeben – Pfad nicht im Ergebnis gespeichert."},
    "report.no_voltage_analysis": {
        "en": "No voltage analysis in the result (three line-to-earth voltages needed).",
        "de": "Keine Spannungsanalyse im Ergebnis (drei Leiter-Erde-Spannungen nötig)."},
    "report.no_angle_method": {"en": "Angle method not available.", "de": "Winkelverfahren nicht verfügbar."},
    "report.no_station": {"en": "No station evaluation in the result.", "de": "Keine Stationsauswertung im Ergebnis."},
    "report.no_transformer": {"en": "No transformer evaluation in the result.",
                              "de": "Keine Trafo-Auswertung im Ergebnis."},
    "report.no_station_frequency": {
        "en": "No frequency-selective station evaluation (table {table}) in the result.",
        "de": "Keine frequenzselektive Stationsauswertung (Tabelle {table}) im Ergebnis."},
    "report.no_values_to_compare": {"en": "No values to compare.", "de": "Keine Werte zum Vergleichen."},
    "report.no_reduction_factor": {
        "en": "No reduction factor (no earthing measurement at the frequencies of the fault test).",
        "de": "Kein Reduktionsfaktor (keine Erdungsmessung bei den Frequenzen des Erdschlusses)."},
    "report.no_lv_network": {"en": "No LV network evaluation (PEN currents) in the result.",
                             "de": "Keine Ortsnetz-Auswertung (PEN-Ströme) im Ergebnis."},
    "report.unknown_table_format": {"en": "Unknown table format {format!r}",
                                    "de": "Unbekanntes Tabellenformat {format!r}"},
    "report.export_library_missing": {
        "en": "The {format} export needs {library}, which is not installed: pip install \"groundfault[export]\"",
        "de": "Der {format}-Export braucht {library}, das nicht installiert ist: pip install \"groundfault[export]\""},
    "report.export_skipped": {
        "en": "{file} not written: {library} is not installed (pip install \"groundfault[export]\").",
        "de": "{file} nicht geschrieben: {library} ist nicht installiert (pip install \"groundfault[export]\")."},
    # ---- report tables: headers and row labels
    "report.quantity_in": {"en": "Quantity in {unit}", "de": "Größe in {unit}"},
    "report.pre_fault_of": {"en": "pre-fault {name}", "de": "vor {name}"},
    "report.post_fault_of": {"en": "post-fault {name}", "de": "nach {name}"},
    "report.positive_sequence": {"en": "positive sequence", "de": "Mitsystem"},
    "report.negative_sequence": {"en": "negative sequence", "de": "Gegensystem"},
    "report.thd_max_loop": {"en": "THD max. loop in %", "de": "THD max. Schleife in %"},
    "report.harmonic_content": {"en": "Harmonic content", "de": "Oberschwingungsanteil"},
    "report.active_current": {"en": "Active current Iw", "de": "Wirkstrom Iw"},
    "report.detuning_current": {"en": "Detuning current Iv", "de": "Verstimmungsstrom Iv"},
    "report.coil_current": {"en": "Coil current {channel}", "de": "Spulenstrom {channel}"},
    "report.rze_per_100_a": {"en": "r·Z_E in V per 100 A", "de": "r·Z_E in V je 100 A"},
    "report.label_pre_fault": {"en": "pre-fault", "de": "Vorfehler"},
    "report.label_fault": {"en": "fault", "de": "Fehler"},
    "report.label_post_fault": {"en": "post-fault", "de": "Nachfehler"},
    "report.i_negative_sequence": {"en": "I neg. seq. in A", "de": "I Gegen in A"},
    "report.i_positive_sequence": {"en": "I pos. seq. in A", "de": "I Mit in A"},
    "report.measured_value_of": {"en": "{name} measured value", "de": "{name} Messwert"},
    "report.assessment_value_of": {"en": "{name} assessment value", "de": "{name} Bewertung"},
    "report.residual_in_pct": {"en": "Residual in %", "de": "Residuum in %"},
    "report.uq_to_faulted_phase": {"en": "U_q/U_faulted phase,pre", "de": "U_q/U_Fehlerleiter,vor"},
    "report.uq_fitted_to_faulted_phase": {"en": "U_q,fitted/U_faulted phase,pre",
                                          "de": "U_q,geschätzt/U_Fehlerleiter,vor"},
    "report.angle_in_deg": {"en": "Angle in °", "de": "Winkel in °"},
    "report.current_at_ze": {"en": "I at Z_E = {impedance} Ω in A", "de": "I bei Z_E = {impedance} Ω in A"},
    "report.current_at_ze_source": {"en": "I at Z_E = {impedance} Ω ({source}) in A",
                                    "de": "I bei Z_E = {impedance} Ω ({source}) in A"},
    "report.source_column": {"en": "Source", "de": "Quelle"},
    "report.source_value_measured": {"en": "measured", "de": "gemessen"},
    "report.source_value_fitted": {"en": "fitted", "de": "geschätzt"},
    "report.remainder": {"en": "Remainder", "de": "Rest"},
    # ---- table titles (DOCX headings) and sheet names of the export
    "report.title_voltages": {"en": "Voltages", "de": "Spannungen"},
    "report.title_residual_current": {"en": "Residual earth-fault current", "de": "Erdschlussreststrom"},
    "report.title_split_50hz": {"en": "Split of the 50 Hz current", "de": "Aufteilung 50 Hz"},
    "report.title_epr": {"en": "Earth potential rise", "de": "Erdungsspannungen"},
    "report.title_shield_current": {"en": "Shield current", "de": "Schirmstrom"},
    "report.title_transformer": {"en": "Transformer", "de": "Trafo"},
    "report.title_assessment_quantities": {"en": "Assessment quantities", "de": "Bewertungsgrößen"},
    "report.title_equivalent_sources": {"en": "Equivalent sources", "de": "Ersatzquellen"},
    "report.title_overview": {"en": "Overview of the measurements", "de": "Übersicht der Messungen"},
    "report.title_station_frequency": {"en": "Earth potential rise and touch voltage per frequency",
                                       "de": "Erdungs- und Berührungsspannung je Frequenz"},
    "report.title_earthing_comparison": {"en": "Fault test and earthing measurement per frequency",
                                         "de": "Erdschluss und Erdungsmessung je Frequenz"},
    "report.title_earthing_ratio": {"en": "Ratio fault test / earthing measurement per frequency",
                                    "de": "Verhältnis Erdschluss / Erdungsmessung je Frequenz"},
    "report.title_station_lv_network": {"en": "LV network: current and earthing impedance per frequency",
                                        "de": "Ortsnetz: Strom und Erdungsimpedanz je Frequenz"},
    "report.title_station_lv_network_eff": {"en": "LV network: current and earthing impedance (RMS values)",
                                            "de": "Ortsnetz: Strom und Erdungsimpedanz (Effektivwerte)"},
    "report.title_reduction_factor": {"en": "Reduction factor per harmonic", "de": "Reduktionsfaktor je Harmonische"},
    "report.title_reduction_factor_eff": {"en": "Reduction factor (RMS value)", "de": "Reduktionsfaktor (effektiv)"},
    "report.title_lv_network_check": {"en": "LV network and earthing measurement per frequency",
                                      "de": "Ortsnetz und Erdungsmessung je Frequenz"},
    "report.title_lv_network_check_eff": {"en": "LV network and earthing measurement (RMS values)",
                                          "de": "Ortsnetz und Erdungsmessung (Effektivwerte)"},
    "report.title_plausibility": {"en": "Plausibility checks", "de": "Plausibilitätsprüfungen"},
    "report.title_statistics": {"en": "Statistics over the steady-state interval",
                                "de": "Statistik über den stationären Bereich"},
    "report.title_steady_state_segments": {"en": "States in the steady-state interval (undisturbed and disturbed)",
                                           "de": "Zustände im stationären Bereich (ungestört und gestört)"},
    "report.title_harmonics": {"en": "Harmonics in the max window", "de": "Harmonische im Max-Fenster"},
    "report.title_coil_currents": {"en": "Coil currents", "de": "Spulenströme"},
    "report.title_voltage_metrics": {"en": "Voltages: characteristic values", "de": "Spannungen: Kennwerte"},
    "report.title_voltage_pre_fault": {"en": "Voltages before the fault", "de": "Spannungen vor dem Fehler"},
    "report.title_voltage_fault": {"en": "Voltages during the fault", "de": "Spannungen im Fehler"},
    "report.title_voltage_post_fault": {"en": "Voltages after the fault", "de": "Spannungen nach dem Fehler"},
    "report.title_station_metrics": {"en": "Station: characteristic values", "de": "Station: Kennwerte"},
    "report.title_station_voltages": {"en": "Station: earth potential rise and touch voltage",
                                      "de": "Station: Erdungs- und Berührungsspannung"},
    "report.title_station_currents": {"en": "Station: PEN and shield currents",
                                      "de": "Station: PEN- und Schirmströme"},
    "report.title_station_currents_frequency": {"en": "PEN and shield currents per frequency",
                                                "de": "PEN- und Schirmströme je Frequenz"},
    "report.title_station_extrapolation": {"en": "Extrapolation to the limits of the residual current",
                                           "de": "Hochrechnung auf die Grenzwerte des Reststroms"},
    "report.title_transformer_power": {"en": "Transformer: harmonic power and sequence currents",
                                       "de": "Trafo: Oberschwingungsleistung und Sequenzströme"},
    "report.title_transformer_summary": {"en": "Transformer node: upstream impedance and power change",
                                         "de": "Trafoknoten: Vorimpedanz und Leistungsänderung"},
    "report.tr_quantity": {"en": "Quantity", "de": "Größe"},
    "report.tr_polarity": {"en": "Transformer currents", "de": "Trafoströme"},
    "report.tr_consumer": {"en": "consumer reference arrow system", "de": "Verbraucherzählpfeilsystem"},
    "report.tr_generator": {"en": "generator reference arrow system (reversed)",
                            "de": "Erzeugerzählpfeilsystem (umgekehrt)"},
    "report.tr_x1_up": {"en": "X_up/m in Ω (transformer + HV network)", "de": "X_vor/m in Ω (Trafo + HS-Netz)"},
    "report.tr_x1_count": {"en": "significant orders/sequences", "de": "signifikante Ordnungen/Systeme"},
    "report.tr_x1_transformer": {"en": "X_T/m in Ω (transformer)", "de": "X_T/m in Ω (Trafo)"},
    "report.tr_x1_hv": {"en": "X_HV/m in Ω (HV network)", "de": "X_HS/m in Ω (HS-Netz)"},
    "report.tr_dp1": {"en": "ΔP at 50 Hz in kW (fault − pre-fault)", "de": "ΔP bei 50 Hz in kW (Erdschluss − Vorfehler)"},
    "report.tr_dq1": {"en": "ΔQ at 50 Hz in kvar (fault − pre-fault)",
                      "de": "ΔQ bei 50 Hz in kvar (Erdschluss − Vorfehler)"},
    "report.title_transformer_node": {"en": "Transformer node per state, order and sequence",
                                      "de": "Trafoknoten je Zustand, Ordnung und System"},
    "report.title_transformer_power_change": {"en": "Transformer node: power change during and after the fault",
                                              "de": "Trafoknoten: Leistungsänderung im und nach dem Erdschluss"},
    "report.title_transformer_impedance": {
        "en": "Transformer node: upstream impedance from the fault-related change",
        "de": "Trafoknoten: Vorimpedanz aus der fehlerbedingten Änderung"},
    "report.title_transformer_divider": {
        "en": "HV share of the upstream impedance from HV and MV voltages (voltage divider)",
        "de": "HS-Anteil der Vorimpedanz aus HS- und MS-Spannungen (Spannungsteiler)"},
    "report.title_of": {"en": "{title} – {name}", "de": "{title} – {name}"},
    "report.sheet_table": {"en": "Table", "de": "Tabelle"},
})


@dataclass
class ReportStyle:
    """Report style: fonts, sizes, line width, decimal separator and export formats.

    Change it with [`set_style`][groundfault.report.set_style]; all figure functions use the global
    [`STYLE`][groundfault.report.STYLE].

    Attributes
    ----------
    fonts : tuple of str
        Font families in order of preference; the first installed one is used
        ([`used_font`][groundfault.report.used_font]).
    font_size : float
        Font size of axis labels and panel letters in pt.
    font_size_small : float
        Font size of tick labels, legends and annotations in pt.
    full_width_cm, half_width_cm : float
        Width of full-width and half-width figures in cm (type area of A4 with 2.54 cm margins).
    height_cm : float
        Default figure height in cm.
    line_width : float
        Default line width in pt.
    decimal_comma : bool or None
        Decimal separator of figures and tables: ``None`` (default) follows the output language (comma in German,
        point in English), ``True``/``False`` always writes a comma/point.
    formats : tuple of str
        Default export formats of [`save`][groundfault.report.save].
    dpi : int
        Resolution of raster formats (PNG).
    panel_letters : bool
        Label the panels of figures with several panels with (a), (b), …
    font_file : str, optional
        Path of a font file (e.g. ``Aptos.ttf``) that [`set_style`][groundfault.report.set_style] registers with
        matplotlib.
    """

    fonts: Tuple[str, ...] = ("Aptos", "Arial", "Helvetica", "Liberation Sans", "DejaVu Sans")
    font_size: float = 9.0          # axis labels [pt]
    font_size_small: float = 8.0    # ticks, legend
    full_width_cm: float = 15.9
    half_width_cm: float = 7.7
    height_cm: float = 7.0
    line_width: float = 0.9
    decimal_comma: Optional[bool] = None
    formats: Tuple[str, ...] = ("pdf", "svg", "png")
    dpi: int = 300
    panel_letters: bool = True
    font_file: Optional[str] = None

    def rc(self) -> dict:
        """matplotlib rcParams of the style (font, sizes, colours, grid, export settings)."""
        fam = list(self.fonts)
        return {
            "font.family": "sans-serif", "font.sans-serif": fam, "font.size": self.font_size,
            "axes.labelsize": self.font_size, "axes.titlesize": self.font_size,
            "xtick.labelsize": self.font_size_small, "ytick.labelsize": self.font_size_small,
            "legend.fontsize": self.font_size_small, "legend.frameon": False,
            "legend.handlelength": 1.6, "legend.borderaxespad": 0.3, "legend.columnspacing": 1.0,
            "mathtext.fontset": "custom", "mathtext.rm": fam[0], "mathtext.it": fam[0] + ":italic",
            "mathtext.bf": fam[0] + ":bold", "mathtext.cal": fam[0], "mathtext.sf": fam[0],
            "mathtext.default": "it",
            "axes.linewidth": 0.6, "axes.edgecolor": "#333333", "axes.labelcolor": "#111111",
            "axes.spines.top": False, "axes.spines.right": False, "axes.grid": True,
            "grid.color": GRID, "grid.linewidth": 0.4, "grid.linestyle": "-",
            "axes.axisbelow": True, "axes.prop_cycle": matplotlib.cycler(color=COLORS),
            "lines.linewidth": self.line_width, "lines.markersize": 4.0,
            "xtick.color": "#333333", "ytick.color": "#333333", "xtick.direction": "out",
            "ytick.direction": "out", "xtick.major.width": 0.6, "ytick.major.width": 0.6,
            "xtick.major.size": 3, "ytick.major.size": 3,
            "figure.dpi": 110, "savefig.dpi": self.dpi, "savefig.bbox": None,
            "pdf.fonttype": 42, "ps.fonttype": 42, "svg.fonttype": "none",
            "axes.formatter.useoffset": False, "axes.formatter.limits": (-4, 5),
            "figure.constrained_layout.use": True,
            "figure.constrained_layout.h_pad": 2 / 72, "figure.constrained_layout.w_pad": 2 / 72,
        }


STYLE = ReportStyle()
"""The active report style of all figure and table functions (change it with
[`set_style`][groundfault.report.set_style])."""


def set_style(**kw) -> ReportStyle:
    """Change the global report style.

    Parameters
    ----------
    **kw
        Fields of [`ReportStyle`][groundfault.report.ReportStyle], e.g. ``set_style(font_size=10, formats=("pdf",))``. A
        ``font_file`` is registered with matplotlib and its font put first in ``fonts``.

    Returns
    -------
    ReportStyle
        The new global style [`STYLE`][groundfault.report.STYLE].
    """
    global STYLE
    STYLE = replace(STYLE, **kw)
    if STYLE.font_file and os.path.exists(STYLE.font_file):
        from matplotlib import font_manager
        font_manager.fontManager.addfont(STYLE.font_file)
        name = font_manager.FontProperties(fname=STYLE.font_file).get_name()
        STYLE = replace(STYLE, fonts=(name,) + tuple(s for s in STYLE.fonts if s != name))
    return STYLE


def used_font() -> str:
    """Name of the font actually used (the first installed one of ``STYLE.fonts``, otherwise DejaVu Sans)."""
    from matplotlib import font_manager
    existing = {f.name for f in font_manager.fontManager.ttflist}
    for s in STYLE.fonts:
        if s in existing:
            return s
    return "DejaVu Sans"


def _comma(lang: Optional[str] = None) -> bool:
    """True if numbers are written with a decimal comma: an explicit ``lang`` decides, otherwise
    ``STYLE.decimal_comma`` if set, otherwise the active output language (German → comma)."""
    if lang is None and STYLE.decimal_comma is not None:
        return bool(STYLE.decimal_comma)
    return i18n.normalize_language(lang) == "de"


class _LocalizedFormatter(ticker.ScalarFormatter):
    """Scalar tick formatter with the decimal separator chosen when the figure was created."""

    def __init__(self, comma: bool, **kw):
        super().__init__(**kw)
        self.comma = comma

    def __call__(self, x, pos=None):
        s = super().__call__(x, pos)
        return s.replace(".", ",") if self.comma else s


def _format_axes(ax, comma: bool) -> None:
    """Tick labels of the linear axes of ``ax`` without offset and with the given decimal separator."""
    for a in (ax.xaxis, ax.yaxis):
        if isinstance(a.get_major_formatter(), ticker.ScalarFormatter) and a.get_scale() == "linear":
            f = _LocalizedFormatter(comma, useOffset=False)
            f.set_powerlimits((-4, 5))
            a.set_major_formatter(f)


def _rc_dict() -> dict:
    """rcParams of the style with the font actually installed first."""
    rc = STYLE.rc()
    rc["font.sans-serif"] = [used_font()] + [s for s in STYLE.fonts if s != used_font()]
    rc["mathtext.rm"] = rc["font.sans-serif"][0]
    rc["mathtext.it"] = rc["font.sans-serif"][0] + ":italic"
    rc["mathtext.bf"] = rc["font.sans-serif"][0] + ":bold"
    rc["mathtext.cal"] = rc["mathtext.sf"] = rc["font.sans-serif"][0]
    return rc


@contextlib.contextmanager
def _rc():
    """Context with the rcParams of the report style."""
    with matplotlib.rc_context(_rc_dict()):
        yield


def activate_style() -> None:
    """Set the report style globally (``matplotlib.rcParams``).

    The figure functions work in their own style context anyway; set globally, the inline preview in the notebook
    also draws with the same font as the exported files.
    """
    matplotlib.rcParams.update(_rc_dict())


def _width(b) -> float:
    """Figure width in cm: ``"full"``, ``"half"`` or a number (cm)."""
    if b == "full":
        return STYLE.full_width_cm
    if b == "half":
        return STYLE.half_width_cm
    return float(b)


def figure(width: Union[str, float] = "full", height_cm: Optional[float] = None, nrows: int = 1, ncols: int = 1,
           sharex: bool = False, sharey: bool = False, **kw):
    """New figure in final size with the report style.

    Parameters
    ----------
    width : {"full", "half"} or float
        Full or half width of the type area, or a width in cm.
    height_cm : float, optional
        Height in cm (default: ``STYLE.height_cm``).
    nrows, ncols : int
        Number of panels.
    sharex, sharey : bool
        Shared axes (see ``matplotlib.pyplot.subplots``).
    **kw
        Further arguments of ``matplotlib.pyplot.subplots``.

    Returns
    -------
    fig : matplotlib.figure.Figure
    axes : numpy.ndarray of matplotlib.axes.Axes
        Always two-dimensional (``nrows`` × ``ncols``).
    """
    with _rc():
        fig, axes = plt.subplots(nrows, ncols, figsize=(_width(width) * CM, (height_cm or STYLE.height_cm) * CM),
                                 sharex=sharex, sharey=sharey, squeeze=False, **kw)
    fig._report_style = True
    return fig, axes


def _finish(fig, axes, panel_letters: Optional[bool] = None):
    """Tick formats and panel letters (a), (b), … of a finished figure."""
    kb = STYLE.panel_letters if panel_letters is None else panel_letters
    flat = [a for a in np.ravel(axes) if a.get_visible()]
    comma = _comma()
    with _rc():
        for a in flat:
            _format_axes(a, comma)
        if kb and len(flat) > 1:
            for i, a in enumerate(flat):
                a.text(-0.01, 1.02, "(%s)" % "abcdefghijklmnopqrstuvwxyz"[i], transform=a.transAxes,
                       ha="right", va="bottom", fontsize=STYLE.font_size, fontweight="normal")
    return fig


def save(fig, path_stem: str, formats: Optional[Sequence[str]] = None, close: bool = False) -> List[str]:
    """Save a figure in all formats (final size, no cropping).

    Parameters
    ----------
    fig : matplotlib.figure.Figure
        The figure.
    path_stem : str
        Path without extension; the folder is created if necessary.
    formats : sequence of str, optional
        File formats (default: ``STYLE.formats``, i.e. PDF, SVG and PNG).
    close : bool
        Close the figure afterwards (no second display in the notebook).

    Returns
    -------
    list of str
        Paths of the written files.
    """
    formats = formats or STYLE.formats
    os.makedirs(os.path.dirname(os.path.abspath(path_stem)) or ".", exist_ok=True)
    paths = []
    with _rc():
        for fmt in formats:
            p = "%s.%s" % (path_stem, fmt)
            fig.savefig(p, format=fmt, dpi=STYLE.dpi, metadata=_metadata(fmt))
            paths.append(p)
    if close:
        plt.close(fig)
    return paths


def _metadata(fmt):
    """File metadata of the vector formats."""
    if fmt == "pdf":
        return {"Creator": "groundfault (Python)", "Title": None}
    if fmt == "svg":
        return {"Creator": "groundfault (Python)"}
    return None


def _source_label(series) -> str:
    """Legend text of the source of an equivalent-source series: measured or fitted source voltage."""
    measured = (getattr(series, "settings", None) or {}).get("source") == "measured"
    return tr("report.source_measured" if measured else "report.source_fitted")


def _legend(ax, n: int, **kw):
    """Legend from two data series on – default: above the axes, right-aligned (covers nothing); with long labels
    fewer columns, so that it stays within the figure."""
    if n >= 2:
        above = "loc" not in kw and "bbox_to_anchor" not in kw
        if above:
            kw.update(loc="lower right", bbox_to_anchor=(1.0, 1.0), borderaxespad=0.15)
        kw.setdefault("ncol", min(n, 5))
        leg = ax.legend(**kw)
        if above:
            try:
                renderer = ax.figure.canvas.get_renderer()
                width = 0.98 * ax.figure.bbox.width
                while kw["ncol"] > 1 and leg.get_window_extent(renderer).width > width:
                    kw["ncol"] -= 1
                    leg = ax.legend(**kw)
            except (AttributeError, TypeError, ValueError):     # a canvas without a renderer: as it is
                pass


def _group_colors(ks) -> List[str]:
    """Colours of the channels of a group: the fixed colour of the role; with several channels of the same role
    (e.g. IF1/IF2, ISp21/ISp22) the colour sequence in channel order."""
    roles = [k.role for k in ks]
    if len(set(roles)) == len(roles):
        return [ROLE_COLORS.get(r, COLORS[i % 8]) for i, r in enumerate(roles)]
    return [COLORS[i % 8] for i in range(len(ks))]


def _scale_prefix(values: np.ndarray, unit: str) -> Tuple[float, str]:
    """Factor and unit with the prefix k for large values (V → kV, A → kA)."""
    m = float(np.nanmax(np.abs(values))) if np.size(values) else 0.0
    if unit in ("V", "A", "W", "var", "VA") and m >= 2000:
        return 1e-3, "k" + unit
    return 1.0, unit


def _coil_symbol(lower_case: bool = False) -> str:
    """Formula symbol of the coil current (``I_Sp`` in German, ``I_coil`` in English; ``i`` for instantaneous
    values)."""
    return r"$%s_\mathrm{%s}$" % ("i" if lower_case else "I", tr("report.sub_coil"))


# =========================================================================== figures
def _groups(roles, channels: Optional[Sequence[int]]):
    """Group the channels by physical quantity (one panel per group) → {panel label: [channels]}."""
    ks = [k for k in roles.channels if (channels is None or k.index in channels)]
    groups: Dict[str, list] = {}
    for k in ks:
        r = k.role or ""
        key = ("report.group_line_to_earth" if r.startswith(("u_l", "u_hv")) else
               "report.group_epr_touch" if r in ("u_e", "u_t") else
               "report.group_fault_current" if r == "i_res" else
               "report.group_coil_current" if r == "i_coil" else
               "report.group_line_current" if r in ("i_l1", "i_l2", "i_l3") else
               "report.group_shield_pen" if r in ("i_shield", "i_pen") else
               ("report.group_voltage" if (k.unit or "").upper().startswith("V") else "report.group_current"))
        groups.setdefault(tr(key), []).append(k)
    return groups


def instantaneous(rec, roles, t0: float, duration: float = 0.08, channels: Optional[Sequence[int]] = None,
                  width="full", height_per_panel_cm: float = 4.2, time_ms: bool = True, remove_dc: bool = False):
    """Instantaneous values (waveforms) in the excerpt [t0, t0 + duration] – one panel per quantity.

    Parameters
    ----------
    rec : Recording
        The recording.
    roles : RoleAssignment
        Channel roles (``result.roles`` or [`groundfault.detect_roles`][groundfault.channels.detect_roles]).
    t0 : float
        Start of the excerpt in s.
    duration : float
        Length of the excerpt in s.
    channels : sequence of int, optional
        Channel indices to show (default: all channels of ``roles``).
    width : {"full", "half"} or float
        Figure width (see [`figure`][groundfault.report.figure]).
    height_per_panel_cm : float
        Height per panel in cm.
    time_ms : bool
        Time axis in ms instead of s (relative to ``t0``).
    remove_dc : bool
        Show every channel without its DC component (the median of the excerpt is subtracted, as the evaluation
        does per window).

    Returns
    -------
    matplotlib.figure.Figure
    """
    t, x = rec.read_time(t0, t0 + duration)
    if remove_dc and len(x):
        x = np.asarray(x, dtype=float)
        x = x - np.nanmedian(x, axis=0, keepdims=True)
    groups = _groups(roles, channels)
    n = len(groups)
    fig, axes = figure(width, height_per_panel_cm * n + 0.6, nrows=n, sharex=True)
    tt = (t - t0) * (1e3 if time_ms else 1.0)
    with _rc():
        for ax, (g, ks) in zip(axes[:, 0], groups.items()):
            k_scal, eh = _scale_prefix(x[:, [k.index for k in ks]], ks[0].unit)
            for k, c in zip(ks, _group_colors(ks)):
                ax.plot(tt, x[:, k.index] * k_scal, color=c, label=k.short, lw=STYLE.line_width)
            sym = "u" if (ks[0].unit or "").upper().startswith("V") else "i"
            ax.set_ylabel("$%s$ in %s" % (sym, eh))
            _legend(ax, len(ks))
        axes[-1, 0].set_xlabel("$t$ in %s" % ("ms" if time_ms else "s"))
        axes[-1, 0].set_xlim(tt[0], tt[-1])
    return _finish(fig, axes)


#: labels of the analysis intervals in the RMS trend (message keys)
_INTERVAL_SHORT = {"pre_fault": "report.interval_pre_fault", "steady_state": "report.interval_steady_state",
                   "post_fault": "report.interval_post_fault"}
#: legend labels of the intervals before, during and after the fault (message keys)
_INTERVAL_LEGEND = {"pre_fault": "report.before_fault", "fault": "report.during_fault",
                    "post_fault": "report.after_fault"}


def rms_trend(res, rec=None, channels: Optional[Sequence[str]] = ("i_res", "i_coil"),
              window_s: float = 0.02, t_range: Optional[Tuple[float, float]] = None,
              width="full", height_per_panel_cm: float = 4.5, show_intervals: bool = True):
    """True RMS value over time (window ``window_s``) with the analysis intervals and the max window.

    Parameters
    ----------
    res : Result
        Evaluation result (roles, analysis intervals, max window).
    rec : Recording, optional
        The recording (default: opened from the path stored in the result).
    channels : sequence of str
        Roles, one panel per role (e.g. fault current at the top, coil currents below).
    window_s : float
        RMS window in s.
    t_range : (float, float), optional
        Time axis range in s (default: the whole recording).
    width : {"full", "half"} or float
        Figure width (see [`figure`][groundfault.report.figure]).
    height_per_panel_cm : float
        Height per panel in cm.
    show_intervals : bool
        Shade the analysis intervals (pre-fault, steady state, post-fault) and the max window; disturbed states of
        the steady-state interval ([`groundfault.disturbances`][groundfault.disturbances]) are hatched.

    Returns
    -------
    matplotlib.figure.Figure

    Raises
    ------
    ValueError
        If ``rec`` is not given and the result does not store the path of the recording.
    """
    rec = rec or _recording_from(res)
    roles = res.roles
    role_list = [r for r in channels if roles.has(r)]
    n = max(1, len(role_list))
    fig, axes = figure(width, height_per_panel_cm * n + 0.6, nrows=n, sharex=True)
    idx = sorted({k.index for r in role_list for k in roles.with_role(r)})
    t_r, R = rec.rms_trend(channels=idx, window_s=window_s)
    pos = {j: s for s, j in enumerate(idx)}
    disturbed_handle = None
    with _rc():
        for ax, r in zip(axes[:, 0], role_list):
            ks = roles.with_role(r)
            for k, color in zip(ks, _group_colors(ks)):
                ax.plot(t_r, R[:, pos[k.index]], color=color, label=k.short, lw=STYLE.line_width)
            if show_intervals:
                for name in ("pre_fault", "steady_state", "post_fault"):
                    b = getattr(res.intervals, name)
                    if b:
                        ax.axvspan(b[0], b[1], color=INTERVAL_COLORS[name], alpha=0.35, lw=0, zorder=0)
                        if ax is axes[0, 0]:
                            ax.text(0.5 * (b[0] + b[1]), 0.97, tr(_INTERVAL_SHORT[name]),
                                    transform=ax.get_xaxis_transform(), ha="center", va="top",
                                    fontsize=STYLE.font_size_small, color="#333333")
                handle = _shade_disturbed(ax, _disturbed_ranges(res))
                if handle is not None and ax is axes[0, 0]:
                    disturbed_handle = handle
                if r == "i_res":
                    tm = res.t_max_window; T = res.settings.T_win
                    ax.axvspan(tm - T / 2, tm + T / 2, color="#111111", alpha=0.18, lw=0, zorder=0.5)
            ax.set_ylabel(r"%s in A" % {"i_res": r"$I_\mathrm{RES}$", "i_coil": _coil_symbol(),
                                        "i_shield": r"$I_\mathrm{S}$"}.get(r, "$I$")
                          if not r.startswith("u") else r"$U$ in V")
            if ax is axes[0, 0] and disturbed_handle is not None:
                lines, labels = ax.get_legend_handles_labels()
                _legend(ax, 2, handles=lines + [disturbed_handle], labels=labels + [disturbed_handle.get_label()],
                        ncol=min(len(lines) + 1, 5))
            else:
                _legend(ax, len(ks))
        axes[-1, 0].set_xlim(*(t_range or (0.0, rec.duration_s)))
        axes[-1, 0].set_xlabel("$t$ in s")
    return _finish(fig, axes)


def _recording_from(res):
    """Open the recording of a result from the path stored in it."""
    from .comtrade import open_recording
    p = res.result_set.get("meta", {}).get("cfg_path") or res.recording_info.get("path")
    if not p:
        raise ValueError(tr("report.no_recording_path"))
    return open_recording(p, emit_warnings=False)


def _disturbed_ranges(res) -> List[Tuple[float, float]]:
    """Disturbed time ranges (t0, t1) of the steady-state interval of a result, adjacent ranges merged (empty
    without)."""
    d = getattr(res, "disturbances", None)
    if d is None or not d.found:
        return []
    out: List[List[float]] = []
    for t0, t1, _ in sorted(d.excluded_intervals()):
        if out and t0 <= out[-1][1] + 1e-9:
            out[-1][1] = max(out[-1][1], t1)
        else:
            out.append([t0, t1])
    return [(a, b) for a, b in out]


def _shade_disturbed(ax, ranges) -> Optional[object]:
    """Hatch the disturbed time ranges; returns a legend handle (None without ranges)."""
    import matplotlib.patches as mpatches
    for t0, t1 in ranges:
        ax.axvspan(t0, t1, facecolor="none", edgecolor=COLOR_DISTURBED, hatch="///", lw=0, alpha=0.3, zorder=0.4)
    if not ranges:
        return None
    return mpatches.Patch(facecolor="none", edgecolor=COLOR_DISTURBED, hatch="///", alpha=0.6, lw=0,
                          label=tr("report.interval_disturbed"))


def harmonics_trend(res, role: str = "i_res", orders: Sequence[int] = (1, 3, 5, 7),
                    width="full", height_cm: float = 6.5):
    """RMS value and harmonics per 200 ms window in the steady-state interval.

    Parameters
    ----------
    res : Result
        Evaluation result.
    role : str
        Role of the channel (its main channel is shown).
    orders : sequence of int
        Harmonic orders.
    width : {"full", "half"} or float
        Figure width (see [`figure`][groundfault.report.figure]).
    height_cm : float
        Figure height in cm.

    Returns
    -------
    matplotlib.figure.Figure
        The dashed line marks the max window; disturbed states (not part of the reference) are hatched.
    """
    st = res.steady_state; k = res.roles.main(role)
    fig, axes = figure(width, height_cm)
    ax = axes[0, 0]
    ok = ~st.nan_window
    with _rc():
        disturbed_handle = _shade_disturbed(ax, _disturbed_ranges(res))
        ax.plot(st.t_mid[ok], st.rms[ok, k.index], color="#111111", lw=STYLE.line_width,
                label=tr("report.rms"), marker="o", ms=2.5)
        for i, m in enumerate(o for o in orders if o in st.orders):
            ax.plot(st.t_mid[ok], np.abs(st.X[ok, k.index, st.ix(m)]), color=COLORS[i % 8],
                    lw=STYLE.line_width, marker="o", ms=2.5,
                    label=tr("report.harmonic", order=m, f=round(m * st.f1["f1_used"])))
        ax.axvline(res.t_max_window, color=GREY, lw=0.6, ls="--")
        ax.set_xlabel("$t$ in s"); ax.set_ylabel(r"$I$ in A")
        ax.set_ylim(bottom=0)
        if disturbed_handle is not None:
            lines, labels = ax.get_legend_handles_labels()
            _legend(ax, 2, ncol=3, handles=lines + [disturbed_handle], labels=labels + [disturbed_handle.get_label()])
        else:
            _legend(ax, 2, ncol=3)
    return _finish(fig, axes)


def spectrum(res, role: str = "i_res", orders: Optional[Sequence[int]] = None, width="full",
             height_cm: float = 6.0, frequency_axis: bool = True):
    """RMS values of the harmonics in the max window (bars).

    Parameters
    ----------
    res : Result
        Evaluation result.
    role : str
        Role of the channel (its main channel is shown).
    orders : sequence of int, optional
        Harmonic orders (default: 1…13 as far as evaluated).
    width : {"full", "half"} or float
        Figure width (see [`figure`][groundfault.report.figure]).
    height_cm : float
        Figure height in cm.
    frequency_axis : bool
        Label the bars with the frequency instead of the order.

    Returns
    -------
    matplotlib.figure.Figure
    """
    st = res.steady_state; k = res.roles.main(role); wi = res.max_index
    orders = [o for o in (orders or range(1, 14)) if o in st.orders]
    values = [abs(st.X[wi, k.index, st.ix(m)]) for m in orders]
    fig, axes = figure(width, height_cm); ax = axes[0, 0]
    with _rc():
        x = np.arange(len(orders))
        ax.bar(x, values, width=0.6, color=ROLE_COLORS.get(role, COLORS[0]), edgecolor="none")
        ax.set_xticks(x)
        ax.set_xticklabels(["%g" % round(m * st.f1["f1_used"]) if frequency_axis else str(m) for m in orders])
        ax.set_xlabel("$f$ in Hz" if frequency_axis else tr("report.order_axis", symbol="$m$"))
        ax.set_ylabel(r"$I_m$ in A" if role.startswith("i") else r"$U_m$ in V")
        ax.grid(axis="x", visible=False)
    return _finish(fig, axes)


def spectrum_comparison(results: Sequence, names: Sequence[str], role: str = "i_res",
                        orders: Sequence[int] = (1, 3, 5, 7, 9, 11, 13), width="full", height_cm: float = 6.5,
                        include_rms: bool = True):
    """Grouped bars: harmonics (and RMS value) of several measurements (e.g. 0/2.5/5 Ω).

    Parameters
    ----------
    results : sequence of Result or ResultFile
        The measurements.
    names : sequence of str
        Legend entry per measurement.
    role : str
        Role of the channel (its main channel is shown).
    orders : sequence of int
        Harmonic orders.
    width : {"full", "half"} or float
        Figure width (see [`figure`][groundfault.report.figure]).
    height_cm : float
        Figure height in cm.
    include_rms : bool
        Add the RMS value as the first group.

    Returns
    -------
    matplotlib.figure.Figure
    """
    fig, axes = figure(width, height_cm); ax = axes[0, 0]
    cat = (["RMS"] if include_rms else []) + ["%d" % m for m in orders]
    x = np.arange(len(cat)); n = len(results); b = 0.8 / max(n, 1)
    with _rc():
        for i, (e, nm) in enumerate(zip(results, names)):
            hz = _main_channel_harmonics(e, role)
            v = ([e.values.get("IRES", np.nan)] if include_rms else []) + [hz.get(m, np.nan) for m in orders]
            ax.bar(x - 0.4 + b * (i + 0.5), v, width=b * 0.92, color=COLORS[i % 8], label=nm, edgecolor="none")
        fn = results[0].settings.f_nominal
        ax.set_xticks(x)
        ax.set_xticklabels(([tr("report.rms_short")] if include_rms else []) + ["%g" % round(m * fn) for m in orders])
        ax.set_xlabel("$f$ in Hz")
        ax.set_ylabel(r"$I$ in A" if role.startswith("i") else r"$U$ in V")
        ax.grid(axis="x", visible=False)
        _legend(ax, n)
    return _finish(fig, axes)


def voltage_spectrum(res, quantity: str = "U_neg", orders: Sequence[int] = (5, 7, 11, 13),
                     width="full", height_cm: float = 6.0):
    """Harmonics of a voltage quantity before, during and after the fault.

    Parameters
    ----------
    res : Result
        Evaluation result with the voltage analysis (three line-to-earth voltages).
    quantity : str
        ``"U_pos"``, ``"U_neg"``, ``"U_zero"`` (sequence systems), ``"U_L12"`` … (line-to-line) or ``"U_L1E"`` …
    orders : sequence of int
        Harmonic orders.
    width : {"full", "half"} or float
        Figure width (see [`figure`][groundfault.report.figure]).
    height_cm : float
        Figure height in cm.

    Returns
    -------
    matplotlib.figure.Figure

    Raises
    ------
    ValueError
        If the result has no voltage analysis.
    """
    sp = res.values.get("voltage")
    if not sp:
        raise ValueError(tr("report.no_voltage_analysis"))
    ords = sp["orders"]
    fig, axes = figure(width, height_cm); ax = axes[0, 0]
    bl = [b for b in ("pre_fault", "fault", "post_fault") if b in sp]
    x = np.arange(len(orders)); bw = 0.8 / len(bl)
    colors = {"pre_fault": COLORS[0], "fault": COLORS[7], "post_fault": COLORS[1]}
    with _rc():
        for i, b in enumerate(bl):
            v = [sp[b][quantity][ords.index(m)] if m in ords else np.nan for m in orders]
            ax.bar(x - 0.4 + bw * (i + 0.5), v, width=bw * 0.92, color=colors[b], label=tr(_INTERVAL_LEGEND[b]),
                   edgecolor="none")
        f1 = res.steady_state.f1["f1_used"]
        ax.set_xticks(x); ax.set_xticklabels(["%g" % round(m * f1) for m in orders])
        ax.set_xlabel("$f$ in Hz")
        sym = {"U_pos": r"$U_\mathrm{%s}$" % tr("report.sub_pos"), "U_neg": r"$U_\mathrm{%s}$" % tr("report.sub_neg"),
               "U_zero": r"$U_\mathrm{%s}$" % tr("report.sub_zero")}.get(quantity,
                                                                          "$U_{%s}$" % quantity.replace("U_", ""))
        ax.set_ylabel("%s in V" % sym)
        ax.grid(axis="x", visible=False)
        _legend(ax, len(bl))
    return _finish(fig, axes)


def phasor_diagram(res, width="full", height_cm: float = 7.0, coils: bool = False):
    """Phasor diagram of the fundamental in the reference frame of U_L1-L2 (angle method).

    (a) Line-to-earth voltages and neutral displacement voltage U0 in the max window, (b) fault current I_H1 split
    into the active residual current Iw (in the direction of U0) and the detuning Iv (perpendicular to it). Two
    panels because voltage and current have different units.

    Parameters
    ----------
    res : Result
        Evaluation result with the angle method (``U0`` and ``I_phasors`` in ``res.values``).
    width : {"full", "half"} or float
        Figure width (see [`figure`][groundfault.report.figure]).
    height_cm : float
        Figure height in cm.
    coils : bool
        Also draw the directions of the coil currents (normalised length, magnitude in the label).

    Returns
    -------
    matplotlib.figure.Figure

    Raises
    ------
    ValueError
        If the angle method is not available.
    """
    w = res.values
    if "U0" not in w or "I_phasors" not in w:
        raise ValueError(tr("report.no_angle_method"))
    phi = w.get("phi_LL1") or 0.0
    rot = np.exp(-1j * phi)
    U0 = w["U0"] * rot; I1 = w["I_phasors"][1] * rot
    e0 = U0 / abs(U0)
    st = res.steady_state; wi = res.max_index
    le = res.roles.line_to_earth
    fig, axes = figure(width, height_cm, ncols=2)
    with _rc():
        def arrow(ax, z0, z1, color, ls="-", lw=1.2):
            ax.annotate("", xy=(z1.real, z1.imag), xytext=(z0.real, z0.imag),
                        arrowprops=dict(arrowstyle="-|>", color=color, lw=lw, ls=ls, mutation_scale=9,
                                        shrinkA=0, shrinkB=0))

        def label(ax, z, text, color, r, large=True):
            # text slightly outside the arrow head, shifted radially
            d = z / abs(z) if abs(z) > 0 else 1
            p = z + 0.09 * r * d
            ax.text(p.real, p.imag, text, color=color, ha="left" if p.real >= 0 else "right",
                    va="bottom" if p.imag >= 0 else "top",
                    fontsize=STYLE.font_size if large else STYLE.font_size_small)

        def frame(ax, r, unit, sym):
            ax.axhline(0, color="#999999", lw=0.5, zorder=0); ax.axvline(0, color="#999999", lw=0.5, zorder=0)
            ax.set_aspect("equal"); ax.set_xlim(-1.3 * r, 1.3 * r); ax.set_ylim(-1.3 * r, 1.3 * r)
            ax.set_xlabel("Re(%s) in %s" % (sym, unit)); ax.set_ylabel("Im(%s) in %s" % (sym, unit))

        # ---- (a) voltages
        ax = axes[0, 0]
        voltages = [(U0, ROLE_COLORS["u_ne"], r"$U_\mathrm{0}$")]
        if le:
            for k, r_ in zip(le, ("u_l1e", "u_l2e", "u_l3e")):
                z = st.X[wi, k, st.ix(1)] * rot
                voltages.append((z, ROLE_COLORS[r_], r"$U_\mathrm{%s}$" % r_[2:].upper()))
        k_u, eh_u = _scale_prefix(np.array([abs(z) for z, _, _ in voltages]), "V")
        r_u = max(abs(z) for z, _, _ in voltages) * k_u
        for z, c, t in voltages:
            if abs(z) * k_u < 0.03 * r_u:            # faulted phase ≈ 0: only a dot
                ax.plot([z.real * k_u], [z.imag * k_u], marker="o", ms=3, color=c)
                ax.text(z.real * k_u, z.imag * k_u, "  " + t + " ≈ 0", color=c, fontsize=STYLE.font_size_small,
                        va="bottom")
                continue
            arrow(ax, 0, z * k_u, c)
            label(ax, z * k_u, t, c, r_u)
        frame(ax, r_u, eh_u, "$U$")
        # ---- (b) current
        ax = axes[0, 1]
        iw = w["Iw"] * e0
        r_i = max(abs(I1), abs(iw))
        arrow(ax, 0, I1, ROLE_COLORS["i_res"])
        arrow(ax, 0, iw, COLORS[2], ls="--", lw=1.0)
        arrow(ax, iw, I1, COLORS[3], ls="--", lw=1.0)
        # I_H1 at the middle of the arrow, on the side away from I_w (the head stays free for the I_v label)
        p_ = 1j * I1 / abs(I1) if abs(I1) > 0 else 1j
        if (p_.conjugate() * iw).real > 0:
            p_ = -p_
        q_ = 0.55 * I1 + 0.08 * r_i * p_
        ax.text(q_.real, q_.imag, r"$I_\mathrm{H1}$", color=ROLE_COLORS["i_res"], fontsize=STYLE.font_size,
                ha="left" if p_.real >= 0 else "right", va="center")
        ax.text(*(0.5 * iw.real, 0.5 * iw.imag), r"  $I_\mathrm{w}$ = %s A" % number(w["Iw"], 3), color=COLORS[2],
                fontsize=STYLE.font_size_small, ha="left", va="center")
        m = 0.5 * (iw + I1); n_ = (I1 - iw) / abs(I1 - iw) if abs(I1 - iw) > 0 else 1
        cross = 1j * n_ * (0.09 * r_i)
        if abs(m - cross) > abs(m + cross):           # shift away from the origin
            cross = -cross
        ax.text((m + cross).real, (m + cross).imag, r"$I_\mathrm{v}$ = %s A" % number(w["Iv"], 3), color=COLORS[3],
                fontsize=STYLE.font_size_small, ha="center", va="center")
        if coils:
            for j, k in enumerate(res.roles.with_role("i_coil")):
                z = st.X[wi, k.index, st.ix(1)] * rot
                zs = z / abs(z) * r_i * 0.9
                arrow(ax, 0, zs, COLORS[(4 + j) % 8], ls=":", lw=1.0)
                ax.text(zs.real, zs.imag, " %s (%s A)" % (k.short, number(abs(z), 3)), fontsize=STYLE.font_size_small,
                        color=COLORS[(4 + j) % 8])
        frame(ax, r_i, "A", "$I$")
    return _finish(fig, axes)


def source_fit(series, orders: Optional[Sequence[int]] = None, Z_max: Optional[float] = None,
               width="full", height_cm: float = 6.5, Z_E: Optional[float] = None, comparison=None):
    """|I_m| over the additional impedance: measuring points (normalised) and the two-terminal (Thevenin) model
    |U_q/(Z_q + Z)| per order.

    Parameters
    ----------
    series : SourceSeries
        Equivalent sources
        ([`groundfault.equivalent_source.fit_equivalent_source`][groundfault.equivalent_source.fit_equivalent_source] or
        [`groundfault.equivalent_source.measured_source`][groundfault.equivalent_source.measured_source]).
    orders : sequence of int, optional
        Harmonic orders (default: all fitted ones).
    Z_max : float, optional
        End of the impedance axis in Ω (default: 1.25 × the largest of the measured impedances and ``Z_E``).
    width : {"full", "half"} or float
        Figure width (see [`figure`][groundfault.report.figure]).
    height_cm : float
        Figure height in cm.
    Z_E : float, optional
        Earthing impedance in Ω (dashed vertical line).
    comparison : SourceSeries, optional
        A second series of the same measurements, e.g. the fitted source when ``series`` is the measured one: its
        model curves dashed in the same colours. Within the tested range both pass through the measuring points,
        beyond it they show how far the two source voltages carry.

    Returns
    -------
    matplotlib.figure.Figure
    """
    ords = list(orders or sorted(series.fits))
    fig, axes = figure(width, height_cm); ax = axes[0, 0]
    zz = np.array([abs(m.Z) for m in series.measurements])
    Z_max = Z_max or max(zz.max() * 1.25, (Z_E or 0) * 1.25, 1.0)
    zg = np.linspace(0, Z_max, 200)
    with _rc():
        for i, m in enumerate(ords):
            F = series.fits[m]; c = COLORS[i % 8]
            ax.plot(zg, np.abs(F.Uq / (F.Zq + zg)), color=c, lw=STYLE.line_width,
                    label="%g Hz" % round(F.f_Hz))
            G = comparison.fits.get(m) if comparison is not None else None
            if G is not None:
                ax.plot(zg, np.abs(G.Uq / (G.Zq + zg)), color=c, lw=0.8 * STYLE.line_width, ls="--")
            ax.plot(np.abs(F.Z), np.abs(F.I_fit), ls="none", marker="o", ms=4, mfc="white", mec=c, mew=1.0)
        n = len(ords)
        if comparison is not None:
            from matplotlib.lines import Line2D
            for label, ls in ((_source_label(series), "-"), (_source_label(comparison), "--")):
                ax.add_line(Line2D([], [], color="#333333", lw=STYLE.line_width, ls=ls, label=label)); n += 1
        ax.set_xlim(0, Z_max); ax.set_ylim(bottom=0)
        if Z_E is not None:
            ax.axvline(Z_E, color=GREY, lw=0.7, ls="--")
            ax.text(Z_E, 0.02, r" $Z_\mathrm{E}$", transform=ax.get_xaxis_transform(), va="bottom", ha="left",
                    color=GREY, fontsize=STYLE.font_size_small)
        ax.set_xlabel("$Z$ in Ω"); ax.set_ylabel(r"$I_m$ in A")
        _legend(ax, n)
    return _finish(fig, axes)


def source_plane(series, mc: Optional[dict] = None, width="half", height_cm: float = 7.0):
    """Source impedances Z_q per order in the complex plane (●), including the pairwise solutions of two
    measurements each (×) and, with ``mc``, the Monte Carlo spread of Z_q (shaded box). For the measured source
    ([`groundfault.equivalent_source.measured_source`][groundfault.equivalent_source.measured_source]) also the source
    impedance of every single test, U_q/I_k − Z_k (+).

    Parameters
    ----------
    series : SourceSeries
        Equivalent sources (fitted or measured).
    mc : dict, optional
        Result of [`groundfault.equivalent_source.monte_carlo`][groundfault.equivalent_source.monte_carlo] for the same
        series. Its entry ``Zq`` (per order ``Rq_P5``, ``Rq_P95``, ``Xq_P5``, ``Xq_P95``) is drawn as a shaded box per
        order: the 5–95 % range of R_q and X_q over the Monte Carlo runs. Orders without an entry are drawn without a
        box.
    width : {"full", "half"} or float
        Figure width (see [`figure`][groundfault.report.figure]).
    height_cm : float
        Figure height in cm.

    Returns
    -------
    matplotlib.figure.Figure
    """
    from matplotlib.patches import Rectangle
    fig, axes = figure(width, height_cm); ax = axes[0, 0]
    pw = series.pairwise_table()
    single = series.single_test_table() if (series.settings or {}).get("source") == "measured" else None
    spread = (mc or {}).get("Zq") or {}
    with _rc():
        for i, (m, F) in enumerate(sorted(series.fits.items())):
            c = COLORS[i % 8]
            box = spread.get(m)
            if box:
                r5, r95, x5, x95 = (float(box[k]) for k in ("Rq_P5", "Rq_P95", "Xq_P5", "Xq_P95"))
                if np.all(np.isfinite([r5, r95, x5, x95])):
                    ax.add_patch(Rectangle((r5, x5), r95 - r5, x95 - x5, facecolor=c, edgecolor="none", alpha=0.18,
                                           zorder=1))
            sub = pw[pw["order"] == m] if not pw.empty else pw
            if len(sub):
                ax.plot(sub["Rq_Ohm"], sub["Xq_Ohm"], ls="none", marker="x", ms=4, color=c, mew=0.9)
            if single is not None and F.method == "measured":
                one = single[single["order"] == m]
                ax.plot(one["Rq_Ohm"], one["Xq_Ohm"], ls="none", marker="+", ms=5, color=c, mew=0.9)
            ax.plot([F.Zq.real], [F.Zq.imag], ls="none", marker="o", ms=5, color=c, label="%g Hz" % round(F.f_Hz))
        ax.axhline(0, color="#999999", lw=0.5); ax.axvline(0, color="#999999", lw=0.5)
        ax.set_xlabel(r"$R_\mathrm{q}$ in Ω"); ax.set_ylabel(r"$X_\mathrm{q}$ in Ω")
        _legend(ax, len(series.fits), ncol=min(3, len(series.fits)))
    return _finish(fig, axes)


def ires_over_ze(series, Z_E_max: float = 10.0, limits: Sequence[float] = (),
                 IG: Union[str, float] = "max", scaling: Optional[dict] = None, width="full",
                 height_cm: float = 7.0, parts: bool = True, comparison=None,
                 limit_ranges: Optional[Mapping[float, float]] = None, Z_E: Optional[float] = None,
                 margins: Optional[Mapping[str, float]] = None):
    """Residual current I_RES over the earthing impedance Z_E from the equivalent sources (equivalent-source
    method), with horizontal lines at the limits given by the user; optionally scaled to the 95 % quantile of the
    harmonic voltage.

    I_RES = √(I_G² + Σ_m |U_q,m/(Z_q,m + Z_E)|² + I_rest²). Grey background: range of the measured additional
    impedances (interpolation there, extrapolation outside).

    Parameters
    ----------
    series : SourceSeries
        Equivalent sources (fitted, or measured with
        [`groundfault.equivalent_source.measured_source`][groundfault.equivalent_source.measured_source]).
    Z_E_max : float
        End of the impedance axis in Ω.
    limits : sequence of float
        Limits of the residual current in A, defined by the user (dotted horizontal lines); empty (default): no
        lines.
    IG : {"max", "median"} or float
        Fundamental component I_G: largest or median measured value, or a value in A.
    scaling : dict, optional
        Keyword arguments of
        [`groundfault.equivalent_source.ires_at_impedance`][groundfault.equivalent_source.ires_at_impedance] for the
        scaling of the harmonic currents (``u95``, ``thd95_percent``, ``thd_measured_percent`` …) → additional curve
        "scaled".
    width : {"full", "half"} or float
        Figure width (see [`figure`][groundfault.report.figure]).
    height_cm : float
        Figure height in cm.
    parts : bool
        Also draw I_G, the current per harmonic order and the remainder that is not fitted.
    comparison : SourceSeries, optional
        A second series of the same measurements (e.g. the fitted source when ``series`` is the measured one): its
        I_RES (and I_RES scaled) dashed, for comparison beyond the tested range.
    limit_ranges : mapping of float to float, optional
        Largest earthing impedance in Ω up to which a limit holds (``{limit: Z_E}``, e.g. a limit that needs all
        earthing impedances below 1.7 Ω): a mark on its line there.
    Z_E : float, optional
        Earthing impedance of the assessment (vertical line, in the legend).
    margins : mapping, optional
        ``uncertainty_I`` and ``uncertainty_U`` → additional curve of the assessment value I_RES,ass (scaled, with the
        margins).

    Returns
    -------
    matplotlib.figure.Figure
    """
    from .equivalent_source import ires_at_impedance
    zg = np.linspace(0, Z_E_max, 201)
    res = [ires_at_impedance(series, z, IG=IG, **(scaling or {})) for z in zg]
    key_a = "IRES_scaled_assessed" if scaling else "IRES_assessed"
    res_a = [ires_at_impedance(series, z, IG=IG, **(scaling or {}), **dict(margins))[key_a] for z in zg] \
        if margins else []
    other = [ires_at_impedance(comparison, z, IG=IG, **(scaling or {})) for z in zg] if comparison is not None else []
    fig, axes = figure(width, height_cm); ax = axes[0, 0]
    zmin, zmax = series.Z_range
    n = 0
    sym = r"$I_\mathrm{RES}$"
    lab = (lambda s, x: tr("report.with_source", symbol=s, source=_source_label(x))) if other else (lambda s, x: s)
    with _rc():
        ax.axvspan(zmin, zmax, color="#eeeeee", lw=0, zorder=0)
        ax.text(zmin + 0.01 * Z_E_max, 0.98, tr("report.measured_range"), transform=ax.get_xaxis_transform(),
                ha="left", va="top", fontsize=STYLE.font_size_small, color=GREY)
        ax.plot(zg, [r["IRES"] for r in res], color="#111111", lw=1.3, label=lab(sym, series)); n += 1
        if other:
            ax.plot(zg, [r["IRES"] for r in other], color="#111111", lw=1.0, ls="--", label=lab(sym, comparison)); n += 1
        if scaling:
            ax.plot(zg, [r["IRES_scaled"] for r in res], color=COLORS[7], lw=1.3,
                    label=lab(tr("report.scaled", symbol=sym), series)); n += 1
            if other:
                ax.plot(zg, [r["IRES_scaled"] for r in other], color=COLORS[7], lw=1.0, ls="--",
                        label=lab(tr("report.scaled", symbol=sym), comparison)); n += 1
        if res_a:
            ax.plot(zg, res_a, color=COLORS[7], lw=1.1, ls="-.",
                    label=tr("report.assessed_curve")); n += 1
        if parts:
            ax.plot(zg, [r["IG"] for r in res], color=GREY, lw=0.8, ls="--", label=r"$I_\mathrm{G}$"); n += 1
            for i, m in enumerate(sorted(series.fits)):
                ax.plot(zg, [r["parts"][m] for r in res], color=COLORS[i % 8], lw=0.8,
                        label="%g Hz" % round(series.fits[m].f_Hz)); n += 1
            if res and res[0].get("IH_rest", 0) > 0.05:
                ax.plot(zg, [r["IH_rest"] for r in res], color=GREY, lw=0.8, ls=":",
                        label=tr("report.remainder_not_fitted")); n += 1
        ymax = max(max(r["IRES_scaled"] if scaling else r["IRES"] for r in res), max(res_a or [0]),
                   max(limits or [0])) * 1.08
        ranges = {float(k): float(v) for k, v in (limit_ranges or {}).items()}
        for g in limits:
            ax.axhline(g, color="#333333", lw=0.6, ls=":")
            ax.text(Z_E_max, g, "%s A " % _g(g), va="bottom", ha="right", fontsize=STYLE.font_size_small,
                    color="#333333")
            zr = next((v for k, v in ranges.items() if abs(k - float(g)) < 1e-9), None)
            if zr is not None and 0 < zr < Z_E_max:         # the largest earthing impedance of its condition
                ax.plot([zr], [g], marker="|", ms=9, mew=1.3, color="#333333", ls="none")
        if Z_E is not None and 0 < abs(complex(Z_E)) < Z_E_max:
            ax.axvline(abs(complex(Z_E)), color=GREY, lw=0.8, ls="--",
                       label=r"$Z_\mathrm{E}$ = %s Ω" % _g(abs(complex(Z_E)))); n += 1
        ax.set_xlim(0, Z_E_max); ax.set_ylim(0, ymax)
        ax.set_xlabel(r"$Z_\mathrm{E}$ in Ω"); ax.set_ylabel(r"$I$ in A")
        _legend(ax, n, ncol=4)
    return _finish(fig, axes)


def _pq_orders(assessment, orders, n: int = 3) -> List[int]:
    """Orders of a PQ figure: the given ones or the ``n`` largest quantiles above the fundamental."""
    if orders is not None:
        return [int(m) for m in orders]
    t = assessment.table()
    t = t[t["order"] >= 2].sort_values("U_quantile_V", ascending=False)
    return [int(m) for m in t["order"].head(n)]


def _pq_q(assessment) -> str:
    return _g(100 * assessment.quantile)


def pq_series(assessment, orders: Optional[Sequence[int]] = None, width="full", height_cm: float = 6.5):
    """10-min values of a power-quality measurement over the measuring period – per order the largest of the three
    phases – with the quantile (dashed) and, with a factor, the quantile times the factor (dotted).

    Parameters
    ----------
    assessment : PQAssessment
        Result of [`groundfault.pq.assess`][groundfault.pq.assess].
    orders : sequence of int, optional
        Orders (default: the three largest quantiles).
    width : {"full", "half"} or float
        Figure width (see [`figure`][groundfault.report.figure]).
    height_cm : float
        Figure height in cm.

    Returns
    -------
    matplotlib.figure.Figure
    """
    import matplotlib.dates as mdates
    ords = sorted(_pq_orders(assessment, orders))
    fig, axes = figure(width, height_cm); ax = axes[0, 0]
    n = 0
    with _rc():
        step = pd.Timedelta(seconds=float(assessment.data.interval_s))
        for i, m in enumerate(ords):
            s = assessment.largest(m)
            if len(s) > 1:                                       # no line across a gap (e.g. between two periods)
                gap = (s.index[1:] - s.index[:-1]) > 1.5 * step
                if gap.any():
                    s = pd.concat([s, pd.Series(np.nan, index=s.index[:-1][gap] + step)]).sort_index()
            t = s.index.tz_localize(None) if s.index.tz is not None else s.index      # wall-clock time of the clock
            color = COLORS[i % 8]
            ax.plot(t, s.to_numpy(), color=color, lw=0.6, label="%g Hz" % round(m * assessment.f1)); n += 1
            q = assessment.u95()[m]
            ax.axhline(q, color=color, lw=0.9, ls="--")
            if np.isfinite(assessment.factor) and assessment.factor != 1.0:
                ax.axhline(q * assessment.factor, color=color, lw=0.9, ls=":")
        ax.plot([], [], color="#333333", lw=0.9, ls="--", label=tr("report.pq_quantile", q=_pq_q(assessment))); n += 1
        if np.isfinite(assessment.factor) and assessment.factor != 1.0:
            ax.plot([], [], color="#333333", lw=0.9, ls=":",
                    label=tr("report.pq_with_factor", q=_pq_q(assessment), k=_g(assessment.factor))); n += 1
        ax.set_ylim(bottom=0)
        ax.set_ylabel(r"$U$ in V")
        idx = assessment.data.harmonics.index
        idx = idx.tz_localize(None) if idx.tz is not None else idx
        if len(idx) > 1:
            # ticks at an even spacing from the first midnight or full hour (a locator by calendar days starts again
            # at each month and the labels collide)
            t0, t1 = idx.min(), idx.max()
            span_h = (t1 - t0) / pd.Timedelta(hours=1)
            if span_h > 120 * 24.0:                                  # months: the first day of every k-th month
                months = pd.date_range(t0.floor("D"), t1, freq="MS")
                ticks = months[::max(1, int(math.ceil(len(months) / 8.0)))]
                fmt = "%m.%Y" if _comma() else "%b %Y"
            else:
                if span_h > 72.0:                                    # days
                    step = pd.Timedelta(days=max(1, int(math.ceil(span_h / 24.0 / 8.0))))
                    fmt = "%d.%m." if _comma() else "%d %b"
                    grid = pd.Timedelta(days=1)
                else:                                                # hours
                    step = pd.Timedelta(hours=next((h for h in (1, 2, 3, 4, 6, 12) if span_h / h <= 6.0), 24))
                    fmt = "%d.%m.\n%H:%M" if _comma() else "%d %b\n%H:%M"
                    grid = step
                midnight = t0.floor("D")
                first = midnight + grid * int(math.ceil((t0 - midnight) / grid))
                ticks = pd.date_range(first, t1, freq=step)
            if len(ticks):
                ax.set_xticks(list(ticks.to_pydatetime()))
            ax.xaxis.set_major_formatter(mdates.DateFormatter(fmt))
        _legend(ax, n, ncol=min(n, 5))
        if ax.get_legend() is not None:
            for h in ax.get_legend().get_lines()[:len(ords)]:
                h.set_linewidth(1.6)
    return _finish(fig, axes)


def pq_duration(assessment, order: Optional[int] = None, levels: Optional[Mapping[str, float]] = None, width="full",
                height_cm: float = 6.5):
    """Duration curve of one order of a power-quality measurement: the 10-min values of the largest phase sorted
    from the largest, over the share of time they are exceeded, with the quantile, the quantile times the factor and
    the pre-fault voltages of fault tests.

    Parameters
    ----------
    assessment : PQAssessment
        Result of [`groundfault.pq.assess`][groundfault.pq.assess].
    order : int, optional
        Order (default: the largest quantile).
    levels : mapping of str to float, optional
        ``{test: U}`` – pre-fault voltage of the largest phase of fault tests in V (points on the curve).
    width : {"full", "half"} or float
        Figure width (see [`figure`][groundfault.report.figure]).
    height_cm : float
        Figure height in cm.

    Returns
    -------
    matplotlib.figure.Figure
    """
    m = _pq_orders(assessment, None if order is None else [order], 1)[0]
    s = np.sort(assessment.largest(m).dropna().to_numpy())[::-1]
    share = 100.0 * (np.arange(1, s.size + 1) - 0.5) / s.size
    q = assessment.u95()[m]
    fig, axes = figure(width, height_cm); ax = axes[0, 0]
    n = 0
    with _rc():
        ax.plot(share, s, color="#111111", lw=1.1, label=tr("report.pq_largest_phase")); n += 1
        x_q = 100.0 * (1.0 - assessment.quantile)
        ax.axhline(q, color=COLORS[7], lw=0.9, ls="--", label=tr("report.pq_quantile", q=_pq_q(assessment))); n += 1
        ax.axvline(x_q, color=COLORS[7], lw=0.6, ls="--")
        if np.isfinite(assessment.factor) and assessment.factor != 1.0:
            ax.axhline(q * assessment.factor, color=COLORS[7], lw=0.9, ls=":",
                       label=tr("report.pq_with_factor", q=_pq_q(assessment), k=_g(assessment.factor))); n += 1
        for i, (name, u) in enumerate((levels or {}).items()):
            x = 100.0 * float((s >= u).mean())
            ax.plot([x], [u], marker="o", ms=4, color=COLORS[i % 7], ls="none", label=str(name)); n += 1
        ax.set_xlim(0, 100); ax.set_ylim(bottom=0)
        ax.set_xlabel(tr("report.pq_time_exceeded"))
        ax.set_ylabel(r"$U_{%d}$ in V" % m)
        _legend(ax, n, ncol=min(n, 4))
    return _finish(fig, axes)


def pq_profile(assessment, order: Optional[int] = None, width="full", height_cm: float = 6.5):
    """Weekly profile of one order of a power-quality measurement: median of the 10-min values of the largest phase
    per hour of the day, Monday to Friday, Saturday and Sunday, with the quantile of the whole measurement.

    Parameters
    ----------
    assessment : PQAssessment
        Result of [`groundfault.pq.assess`][groundfault.pq.assess].
    order : int, optional
        Order (default: the largest quantile).
    width : {"full", "half"} or float
        Figure width (see [`figure`][groundfault.report.figure]).
    height_cm : float
        Figure height in cm.

    Returns
    -------
    matplotlib.figure.Figure
    """
    m = _pq_orders(assessment, None if order is None else [order], 1)[0]
    s = assessment.largest(m).dropna()
    t = s.index - pd.Timedelta(seconds=float(assessment.data.interval_s))       # beginning of the interval
    frame = pd.DataFrame({"u": s.to_numpy(), "day": t.dayofweek, "hour": t.hour + t.minute / 60.0})
    fig, axes = figure(width, height_cm); ax = axes[0, 0]
    n = 0
    with _rc():
        for days, label, color in (((0, 1, 2, 3, 4), tr("report.pq_workdays"), COLORS[0]),
                                   ((5,), tr("report.pq_saturday"), COLORS[1]),
                                   ((6,), tr("report.pq_sunday"), COLORS[2])):
            f = frame[frame["day"].isin(days)]
            if not len(f):
                continue
            g = f.groupby(np.floor(f["hour"]))["u"]
            med = g.median()
            ax.plot(med.index + 0.5, med.to_numpy(), color=color, lw=1.1, label=label, marker="o", ms=2); n += 1
            if len(days) > 1:
                ax.fill_between(med.index + 0.5, g.quantile(0.05).to_numpy(), g.quantile(0.95).to_numpy(), color=color,
                                alpha=0.15, lw=0)
        ax.axhline(assessment.u95()[m], color=COLORS[7], lw=0.9, ls="--",
                   label=tr("report.pq_quantile", q=_pq_q(assessment))); n += 1
        ax.set_xlim(0, 24); ax.set_xticks(range(0, 25, 3)); ax.set_ylim(bottom=0)
        ax.set_xlabel(tr("report.pq_hour"))
        ax.set_ylabel(r"$U_{%d}$ in V" % m)
        _legend(ax, n, ncol=min(n, 4))
    return _finish(fig, axes)


def touch_voltage_cases(cases: pd.DataFrame, U_Tp: Optional[float] = None, width="full", height_cm: float = 6.0):
    """Bars of the touch voltage U_T of the assessment cases of a station, with the permissible touch voltage.

    Parameters
    ----------
    cases : pandas.DataFrame
        Table of [`groundfault.document.touch_voltage_cases`][groundfault.document.touch_voltage_cases] (columns
        ``label`` and ``U_T_V``; rows without U_T are left out).
    U_Tp : float, optional
        Permissible touch voltage in V (a horizontal line; default: the column ``U_Tp_V`` if it has a value).
    width : {"full", "half"} or float
        Figure width (see [`figure`][groundfault.report.figure]).
    height_cm : float
        Figure height in cm.

    Returns
    -------
    matplotlib.figure.Figure
    """
    df = cases[np.isfinite(pd.to_numeric(cases["U_T_V"], errors="coerce"))]
    if U_Tp is None and "U_Tp_V" in cases and pd.to_numeric(cases["U_Tp_V"], errors="coerce").notna().any():
        U_Tp = float(pd.to_numeric(cases["U_Tp_V"], errors="coerce").dropna().iloc[0])
    fig, axes = figure(width, height_cm)
    ax = axes[0, 0]
    with _rc():
        x = np.arange(len(df))
        v = pd.to_numeric(df["U_T_V"], errors="coerce").to_numpy(dtype=float)
        ax.bar(x, v, width=0.6, color=[COLORS[i % 8] for i in range(len(df))], edgecolor="none")
        for xi, vi in zip(x, v):
            ax.annotate(number(vi, 2 if vi < 10 else 3), (xi, vi), xytext=(0, 2), textcoords="offset points",
                        ha="center", va="bottom", fontsize=STYLE.font_size_small)
        ymax = max(list(v) + [U_Tp or 0.0]) * 1.15 if len(v) else 1.0
        if U_Tp is not None and np.isfinite(U_Tp):
            ax.axhline(U_Tp, color="#c00000", lw=0.9, ls="--",
                       label=tr("report.permissible_touch_voltage", value=number(U_Tp, 3)))
            _legend(ax, 1)
        ax.set_xticks(x)                                  # U_T → $U_\mathrm{T}$
        ax.set_xticklabels([re.sub(r"(?<![A-Za-z0-9])([A-Za-z])_([A-Za-z0-9]+)", r"$\1_\\mathrm{\2}$", str(t))
                            for t in df["label"]])
        ax.set_ylabel(tr("report.touch_voltage_axis"))
        ax.set_ylim(0, ymax or 1.0)
        ax.grid(axis="x", visible=False)
    return _finish(fig, axes, panel_letters=False)


def v_curve(t, ires, result=None, i_coil=None, width="full", height_cm: float = 6.0, quantity: str = "I_RES"):
    """Resonance curve (V curve): residual current or its fundamental over time or over the coil current, with the
    minimum marked.

    Parameters
    ----------
    t : array_like
        Time in s.
    ires : array_like
        Residual current (``quantity="I_RES"``) or its fundamental component (``"I_G"``) in A.
    result : VCurve, optional
        Result of [`groundfault.rms.v_curve`][groundfault.rms.v_curve] → minimum marked and labelled.
    i_coil : array_like, optional
        Coil current in A → x-axis coil current instead of time.
    width : {"full", "half"} or float
        Figure width (see [`figure`][groundfault.report.figure]).
    height_cm : float
        Figure height in cm.
    quantity : {"I_RES", "I_G"}
        Quantity of ``ires`` (y-axis label).

    Returns
    -------
    matplotlib.figure.Figure
    """
    chars = {"I_RES": r"$I_\mathrm{RES}$", "I_G": r"$I_\mathrm{G}$"}.get(quantity, quantity)
    fig, axes = figure(width, height_cm); ax = axes[0, 0]
    x = np.asarray(i_coil) if i_coil is not None else np.asarray(t)
    with _rc():
        ax.plot(x, ires, color=ROLE_COLORS["i_res"], lw=STYLE.line_width, marker="." if i_coil is not None else None,
                ls="none" if i_coil is not None else "-")
        if result is not None:
            xm = result.t_min_s if i_coil is None else x[int(np.nanargmin(ires))]
            ax.plot([xm], [result.I_min_measured], marker="o", ms=5, mfc="white", mec="#111111", ls="none")
            # below the minimum there is room (the curve lies above it); at the right edge aligned to the left
            xs = np.asarray(x, float)
            rel = (xm - np.nanmin(xs)) / max(np.nanmax(xs) - np.nanmin(xs), 1e-12)
            ax.annotate(r"$I_\mathrm{min}$ = %s A" % number(result.I_min_measured, 3), (xm, result.I_min_measured),
                        textcoords="offset points", xytext=(-6 if rel > 0.6 else 6, -8),
                        ha="right" if rel > 0.6 else "left", va="top", fontsize=STYLE.font_size_small)
        ax.set_xlabel("%s in A" % _coil_symbol() if i_coil is not None else "$t$ in s")
        ax.set_ylabel(chars + " in A"); ax.set_ylim(0, 1.1 * float(np.nanmax(np.asarray(ires, float))))
    return _finish(fig, axes)


def station_voltages(res, width="full", height_cm: float = 6.0):
    """Earth potential rise/touch voltages per harmonic and r·Z_E in V per 100 A (two panels).

    Parameters
    ----------
    res : Result
        Evaluation result with the station evaluation (table ``station_voltages``).
    width : {"full", "half"} or float
        Figure width (see [`figure`][groundfault.report.figure]).
    height_cm : float
        Figure height in cm.

    Returns
    -------
    matplotlib.figure.Figure

    Raises
    ------
    ValueError
        If the result has no station evaluation.
    """
    df = res.tables.get("station_voltages")
    if df is None or df.empty:
        raise ValueError(tr("report.no_station"))
    ords = [int(c[1:-2]) for c in df.columns if c.startswith("H") and c.endswith("_V")]
    f1 = res.steady_state.f1["f1_used"]
    fig, axes = figure(width, height_cm, ncols=2)
    x = np.arange(len(ords) + 1); n = len(df); b = 0.8 / max(n, 1)
    with _rc():
        for i, (_, r) in enumerate(df.iterrows()):
            v = [r["RMS_V"]] + [r["H%d_V" % m] for m in ords]
            axes[0, 0].bar(x - 0.4 + b * (i + 0.5), v, width=b * 0.92, color=COLORS[i % 8], label=r["channel"])
            z = [r.get("V_per_100A_RMS", np.nan)] + [100 * r.get("rZE_H%d_Ohm" % m, np.nan) for m in ords]
            axes[0, 1].bar(x - 0.4 + b * (i + 0.5), z, width=b * 0.92, color=COLORS[i % 8], label=r["channel"])
        for ax in axes[0]:
            ax.set_xticks(x); ax.set_xticklabels([tr("report.rms_short")] + ["%g" % round(m * f1) for m in ords])
            ax.set_xlabel("$f$ in Hz"); ax.grid(axis="x", visible=False)
        axes[0, 0].set_ylabel(r"$U$ in V")
        axes[0, 1].set_ylabel(tr("report.per_100_a", symbol=r"$r\,Z_\mathrm{E}$"))
        _legend(axes[0, 0], n, loc="lower left", bbox_to_anchor=(0.0, 1.0))
    return _finish(fig, axes)


def transformer_power(res, orders: Sequence[int] = (5, 7, 11, 13), width="full", height_cm: float = 6.5):
    """Active and reactive power per order at the transformer (before, during and after the fault) – two panels.

    Parameters
    ----------
    res : Result
        Evaluation result with the transformer evaluation (table ``transformer_power``).
    orders : sequence of int
        Harmonic orders.
    width : {"full", "half"} or float
        Figure width (see [`figure`][groundfault.report.figure]).
    height_cm : float
        Figure height in cm.

    Returns
    -------
    matplotlib.figure.Figure

    Raises
    ------
    ValueError
        If the result has no transformer evaluation.
    """
    df = res.tables.get("transformer_power")
    if df is None or df.empty:
        raise ValueError(tr("report.no_transformer"))
    fig, axes = figure(width, height_cm, ncols=2)
    bl = [b for b in ("pre_fault", "fault", "post_fault") if b in set(df["interval"])]
    x = np.arange(len(orders)); bw = 0.8 / len(bl)
    colors = {"pre_fault": COLORS[0], "fault": COLORS[7], "post_fault": COLORS[1]}
    with _rc():
        for i, b in enumerate(bl):
            sub = df[df["interval"] == b].set_index("order")
            for ax, sp in zip(axes[0], ("P_kW", "Q_kvar")):
                v = [sub.at[m, sp] if m in sub.index else np.nan for m in orders]
                ax.bar(x - 0.4 + bw * (i + 0.5), v, width=bw * 0.92, color=colors[b], label=tr(_INTERVAL_LEGEND[b]))
        f1 = res.steady_state.f1["f1_used"]
        for ax in axes[0]:
            ax.axhline(0, color="#333333", lw=0.6)
            ax.set_xticks(x); ax.set_xticklabels(["%g" % round(m * f1) for m in orders])
            ax.set_xlabel("$f$ in Hz"); ax.grid(axis="x", visible=False)
        axes[0, 0].set_ylabel(r"$P$ in kW"); axes[0, 1].set_ylabel(r"$Q$ in kvar")
        if len(bl) >= 2:                # above both panels: a legend of panel (a) would push panel (b) aside
            fig.legend(*axes[0, 0].get_legend_handles_labels(), loc="outside upper left", ncol=len(bl),
                       borderaxespad=0.15)
    return _finish(fig, axes)


def transformer_impedance(res, orders: Sequence[int] = (5, 7, 11, 13), width="full", height_cm: float = 11.0):
    """Impedance at the transformer node – two panels, one above the other.

    (a) apparent impedance |U_m/I_m| of the natural sequence of each order (5, 11 negative, 7, 13 positive sequence)
    before, during and after the fault (logarithmic axis); (b) upstream impedance |Z_up| (transformer and HV network
    seen from the MV busbar) from the fault-related change over the frequency, positive and negative sequence (small
    open markers: not significant; the axis is scaled to the significant orders), with the line m·X_up,1 and – with HV
    voltages – the shares of the transformer and the HV network.

    Parameters
    ----------
    res : Result
        Evaluation result with the transformer node (tables ``transformer_node`` and ``transformer_impedance``,
        [`groundfault.transformer.node_analysis`][groundfault.transformer.node_analysis]).
    orders : sequence of int
        Orders of panel (a).
    width : {"full", "half"} or float
        Figure width (see [`figure`][groundfault.report.figure]).
    height_cm : float
        Figure height in cm.

    Returns
    -------
    matplotlib.figure.Figure

    Raises
    ------
    ValueError
        If the result has no transformer node or no upstream impedance.
    """
    node = res.tables.get("transformer_node")
    zi = res.tables.get("transformer_impedance")
    if node is None or node.empty:
        raise ValueError(tr("report.no_transformer"))
    if zi is None or zi.empty:
        raise ValueError(tr("report.no_upstream"))
    fig, axes = figure(width, height_cm, nrows=2)
    natural = {1: "pos", 2: "neg", 0: "zero"}
    bl = [b for b in ("pre_fault", "fault", "post_fault") if b in set(node["interval"])]
    x = np.arange(len(orders)); bw = 0.8 / len(bl)
    colors = {"pre_fault": COLORS[0], "fault": COLORS[7], "post_fault": COLORS[1]}
    f1 = res.steady_state.f1["f1_used"]
    up = tr("report.sub_upstream")
    with _rc():
        ax = axes[0, 0]
        for i, b in enumerate(bl):
            sub = node[node["interval"] == b].set_index(["order", "sequence"])
            v = [sub.at[(m, natural[m % 3]), "Z_Ohm"] if (m, natural[m % 3]) in sub.index else np.nan
                 for m in orders]
            ax.bar(x - 0.4 + bw * (i + 0.5), v, width=bw * 0.92, color=colors[b], label=tr(_INTERVAL_LEGEND[b]))
        ax.set_xticks(x); ax.set_xticklabels(["%g" % round(m * f1) for m in orders])
        ax.set_xlabel("$f$ in Hz"); ax.grid(axis="x", visible=False)
        _log_y(ax)
        ax.set_ylabel(r"$|U_m/I_m|$ in Ω")
        _legend(ax, len(bl), loc="lower left", bbox_to_anchor=(0.0, 1.0))
        ax = axes[1, 0]
        n_series = 0
        z = zi[zi["order"] >= 2]
        for seq, mk, col, key in (("pos", "o", COLORS[0], "report.positive_sequence"),
                                  ("neg", "s", COLORS[1], "report.negative_sequence")):
            d = z[z["sequence"] == seq]
            if d.empty:
                continue
            sig = d["significant"].astype(bool)
            ax.plot(d.loc[sig, "f_Hz"], d.loc[sig, "Z_up_Ohm"], mk, color=col, ms=4.5, label=tr(key))
            ax.plot(d.loc[~sig, "f_Hz"], d.loc[~sig, "Z_up_Ohm"], mk, color=col, ms=3, mfc="none", alpha=0.35)
            n_series += 1
        x1 = (res.values.get("transformer") or {}).get("X1_up_Ohm")
        if x1 is not None and np.isfinite(x1) and not z.empty:
            ff = np.array([z["f_Hz"].min(), z["f_Hz"].max()])
            ax.plot(ff, x1 * ff / f1, "--", color=GREY, lw=0.8,
                    label="$m\\cdot X_{\\mathrm{%s},1}$" % up)
            n_series += 1
        if "Z_transformer_Ohm" in z:
            d = z[z["significant"].astype(bool)]
            ax.plot(d["f_Hz"], d["Z_transformer_Ohm"], "^", color=COLORS[2], ms=4, label=tr("report.transformer_share"))
            ax.plot(d["f_Hz"], d["Z_hv_Ohm"], "v", color=COLORS[6], ms=4, label=tr("report.hv_share"))
            n_series += 2
        ax.set_xlabel("$f$ in Hz")
        ax.set_ylabel("$|Z_\\mathrm{%s}|$ in Ω" % up)
        top = z.loc[z["significant"].astype(bool), "Z_up_Ohm"].max()
        if np.isfinite(top) and top > 0:              # scale of the significant orders (the others may scatter)
            ax.set_ylim(0, 1.6 * top)
        else:
            ax.set_ylim(bottom=0)
        _legend(ax, n_series, ncol=n_series, loc="lower left", bbox_to_anchor=(0.0, 1.0))
    return _finish(fig, axes)


# ====================================================================== station per frequency, earthing measurement
_STATION_SINGLE = ("Z_inception_Ohm", "Z_clearance_Ohm", "Z_voltage_Ohm")
#: values of the column ``source`` of :func:`groundfault.earthing.compare`: rows of this evaluation, earthing
#: measurements and fault tests from groundmeas
_SOURCE_FAULT = SOURCE_FAULT_TEST
_SOURCE_INJECTION = SOURCE_INJECTION
_SOURCE_GROUNDMEAS_FAULT = SOURCE_GROUNDMEAS_FAULT
#: values of the column ``quantity`` of :func:`groundfault.earthing.compare` in panel order → message key of the
#: panel label
_COMPARISON_QUANTITIES = {QUANTITY_EPR: "report.epr", QUANTITY_TOUCH: "report.touch_voltage"}
_TOUCH_QUANTITY = QUANTITY_TOUCH
#: prefixes of the column ``injection_frequency`` of :func:`groundfault.earthing.reduction_factor` for values of the
#: earthing measurement interpolated between two measured frequencies (English code and German text)
_INTERPOLATED = ("interpolated", "interpoliert")


def _station_table(res, name: str = "station_frequency") -> pd.DataFrame:
    """Table ``name`` of a result; ValueError if it is missing or empty."""
    df = (getattr(res, "tables", None) or {}).get(name)
    if df is None or df.empty:
        raise ValueError(tr("report.no_station_frequency", table=name))      # table key, same in both languages
    return df


def _interpolated(g: pd.DataFrame) -> np.ndarray:
    """Mask of the rows whose value of the earthing measurement is interpolated (column ``injection_frequency``)."""
    return g["injection_frequency"].astype(str).str.startswith(_INTERPOLATED).to_numpy()


def _cap(s: str) -> str:
    """Text with an upper-case first letter (panel labels)."""
    return s[:1].upper() + s[1:]


def station_colors(channels: Iterable[str]) -> Dict[str, str]:
    """Fixed colour per probe in the order of the table ``station_frequency``.

    The colour follows the probe in all station figures and in the comparison with the earthing measurement; the probe
    maximum (``max(UE)``) always has [`COLOR_MAX`][groundfault.report.COLOR_MAX] and takes no place in the colour
    sequence.

    Parameters
    ----------
    channels : iterable of str
        Channel names (duplicates allowed, the first occurrence counts).

    Returns
    -------
    dict of str to str
        Colour per channel.
    """
    ks = list(dict.fromkeys(channels))
    colors = {k: COLORS[i % len(COLORS)] for i, k in enumerate(k for k in ks if k != MAX_PROBE)}
    if MAX_PROBE in ks:
        colors[MAX_PROBE] = COLOR_MAX
    return colors


def _distance_text(a) -> str:
    """Probe distance as text ("70 m"); empty for missing or non-positive values."""
    try:
        a = float(a)
    except (TypeError, ValueError):
        return ""
    return ("%s m" % _g(a)) if np.isfinite(a) and a > 0 else ""


def _g(v: float, comma: Optional[bool] = None) -> str:
    """Short number without superfluous zeros (axes, legends): 0,02 · 1,5 · 100 with a decimal comma (and
    typographic minus) or 0.02 · 1.5 · 100 with a decimal point (``comma`` default: ``_comma``)."""
    s = "%g" % float(v)
    if comma is None:
        comma = _comma()
    return s.replace(".", ",").replace("-", "\u2212") if comma else s


def probe_label(channel: str, role: str, distance=None) -> str:
    """Legend text of a probe: ``"UE100 – earth potential rise, 100 m"`` (German: ``"UE100 – Erdungsspannung,
    100 m"``).

    Parameters
    ----------
    channel : str
        Channel (short name).
    role : str
        ``"u_t"`` (touch voltage) or another role (earth potential rise).
    distance : float, optional
        Probe distance in m (left out if missing).

    Returns
    -------
    str
    """
    kind = tr("report.touch_voltage") if role == "u_t" else tr("report.epr")
    a = _distance_text(distance)
    return "%s – %s%s" % (channel, kind, ", " + a if a else "")


def _log_y(ax):
    """Logarithmic y-axis with the decimal separator of the language; ticks at 1-2-5 (span ≤ 2 decades), 1-3 (≤ 3)
    or only at the decades – call after the data have been drawn."""
    ax.set_yscale("log")
    lo, hi = ax.get_ylim()
    dec = math.log10(hi / lo) if lo > 0 and hi > lo else 1.0
    subs = (1.0, 2.0, 5.0) if dec <= 2.2 else (1.0, 3.0) if dec <= 3.2 else (1.0,)
    comma = _comma()
    ax.yaxis.set_major_locator(ticker.LogLocator(base=10.0, subs=subs))
    ax.yaxis.set_major_formatter(ticker.FuncFormatter(lambda v, p: _g(v, comma)))
    ax.yaxis.set_minor_locator(ticker.LogLocator(base=10.0, subs=np.arange(2, 10) * 1.0))
    ax.yaxis.set_minor_formatter(ticker.NullFormatter())


def _span(g: pd.DataFrame, column: str = "Z_fault_Ohm") -> Tuple[np.ndarray, np.ndarray, np.ndarray]:
    """Value and smallest/largest value of the single methods (fault inception, fault clearance, voltage
    reference)."""
    y = pd.to_numeric(g[column], errors="coerce").to_numpy(float)
    sp = [s for s in _STATION_SINGLE if s in g]
    if not sp:
        return y, y.copy(), y.copy()
    on = g[sp].apply(pd.to_numeric, errors="coerce").to_numpy(float)
    ok = np.isfinite(on).any(axis=1)
    lo, hi = y.copy(), y.copy()
    lo[ok] = np.nanmin(on[ok], axis=1); hi[ok] = np.nanmax(on[ok], axis=1)
    return y, np.fmin(lo, y), np.fmax(hi, y)


def _uncertain_mask(g: pd.DataFrame) -> np.ndarray:
    """Mask of the rows marked as ``uncertain`` (column missing: none)."""
    if "uncertain" not in g:
        return np.zeros(len(g), bool)
    from .earthing import _yes
    return np.array([_yes(x) for x in g["uncertain"]], bool)


def _mark_uncertain(ax, f, y, color, ms: float = 3.2) -> None:
    """Draw uncertain values as ⊗ (circle with a cross in the colour of the probe) over the filled dot."""
    f = np.asarray(f, float); y = np.asarray(y, float)
    if not len(f):
        return
    ax.plot(f, y, ls="none", marker="o", ms=ms + 0.8, mfc="white", mec=color, mew=0.8, zorder=4)
    ax.plot(f, y, ls="none", marker="x", ms=ms * 0.8, mec=color, mew=0.8, zorder=5)


def _uncertain_handle():
    """Legend handle of the ⊗ marker."""
    from matplotlib.lines import Line2D
    return (Line2D([], [], ls="none", marker="o", ms=4.0, mfc="white", mec="#333333", mew=0.8),
            Line2D([], [], ls="none", marker="x", ms=2.6, mec="#333333", mew=0.8))


def _style_legend(ax, entries, n_channel: int):
    """Legend: probes (colour) and kind of display (filled/open/⊗) – above the axes."""
    from matplotlib.lines import Line2D
    handles, texts = ax.get_legend_handles_labels()
    for kind in entries:
        if kind == "fault":
            handles.append(Line2D([], [], color="#333333", marker="o", ms=3.5, lw=STYLE.line_width))
            texts.append(tr("report.fault_component_span"))
        elif kind == "uncertain":
            handles.append(_uncertain_handle())
            texts.append(tr("report.uncertain_background_methods"))
        elif kind == "direct":
            handles.append(Line2D([], [], color="#333333", marker="o", ms=3.5, mfc="white", ls="none", mew=0.8))
            texts.append(tr("report.measured_directly"))
    ax.legend(handles, texts, loc="lower left", bbox_to_anchor=(0.0, 1.0), ncol=2 if len(texts) > 3 else len(texts),
              borderaxespad=0.15)


def station_frequency(res, width="full", height_cm: float = 8.0, log: bool = True, direct: bool = True,
                      f_max: Optional[float] = None):
    """Earth potential rise and touch voltage per ampere of fault current over the frequency (table
    ``station_frequency``).

    Filled with line: fault-related component (superposition) – the vertical bar spans from the smallest to the
    largest value of the single methods (fault inception, fault clearance, voltage reference); ⊗: marked as
    ``uncertain`` (methods disagree or the background is large and fluctuating); open: direct quotient U/I_F during
    the fault (contains the background before the fault). One y-axis (V/A = Ω), logarithmic for all orders.

    Parameters
    ----------
    res : Result or ResultFile
        Result with the table ``station_frequency``.
    width : {"full", "half"} or float
        Figure width (see [`figure`][groundfault.report.figure]).
    height_cm : float
        Figure height in cm.
    log : bool
        Logarithmic y-axis.
    direct : bool
        Also draw the direct quotient.
    f_max : float, optional
        Highest frequency in Hz.

    Returns
    -------
    matplotlib.figure.Figure

    Raises
    ------
    ValueError
        If the table ``station_frequency`` is missing or empty.
    """
    df = _station_table(res)
    if f_max:
        df = df[df["f_Hz"] <= f_max]
    colors = station_colors(df["channel"])
    fig, axes = figure(width, height_cm); ax = axes[0, 0]
    with _rc():
        for channel, g in df.groupby("channel", sort=False):
            g = g.sort_values("f_Hz"); c = colors[channel]; r0 = g.iloc[0]
            if "Z_fault_Ohm" in g and g["Z_fault_Ohm"].notna().any():
                y, lo, hi = _span(g)
                ax.errorbar(g["f_Hz"], y, yerr=np.vstack([y - lo, hi - y]), color=c, marker="o", ms=3.2,
                            lw=STYLE.line_width, elinewidth=0.8, capsize=0,
                            label=probe_label(channel, r0["role"], r0.get("distance_m")))
                u = _uncertain_mask(g)
                _mark_uncertain(ax, g["f_Hz"].to_numpy(float)[u], y[u], c)
                labelled = True
            else:
                labelled = False
            if direct or not labelled:
                ax.plot(g["f_Hz"], g["Z_direct_Ohm"], ls="none", marker="o", ms=3.2, mfc="white", mec=c, mew=0.8,
                        label=None if labelled else probe_label(channel, r0["role"], r0.get("distance_m")))
        ax.set_xlabel("$f$ in Hz"); ax.set_ylabel(r"$U/I_\mathrm{F}$ in V/A")
        ax.set_xlim(left=0)
        if log:
            _log_y(ax)
        _style_legend(ax, (["fault"] if "Z_fault_Ohm" in df else []) +
                      (["uncertain"] if _uncertain_mask(df).any() else []) + (["direct"] if direct else []),
                      df["channel"].nunique())
    return _finish(fig, axes)


def station_background(res, width="full", height_per_panel_cm: float = 3.6, f_max: Optional[float] = None):
    """One panel per probe: voltage during the fault (measured), background before the fault and fault-related
    component per frequency (scatter plot, logarithmic) – shows why the direct quotient differs from the
    fault-related one.

    Parameters
    ----------
    res : Result or ResultFile
        Result with the table ``station_frequency``.
    width : {"full", "half"} or float
        Figure width (see [`figure`][groundfault.report.figure]).
    height_per_panel_cm : float
        Height per panel in cm.
    f_max : float, optional
        Highest frequency in Hz.

    Returns
    -------
    matplotlib.figure.Figure

    Raises
    ------
    ValueError
        If the table ``station_frequency`` is missing or empty.
    """
    df = _station_table(res)
    if f_max:
        df = df[df["f_Hz"] <= f_max]
    chan = list(dict.fromkeys(df["channel"]))
    fig, axes = figure(width, height_per_panel_cm * len(chan) + 0.9, nrows=len(chan), sharex=True)
    kinds = [("U_during_fault_V", tr("report.measured_during_fault"), dict(marker="o", color=COLORS[0])),
             ("U_background_V", tr("report.background_before_fault"), dict(marker="s", color=GREY, mfc="white")),
             ("U_fault_V", tr("report.fault_component"), dict(marker="D", color=COLORS[1], ms=3.0))]
    with _rc():
        for ax, k in zip(axes[:, 0], chan):
            g = df[df["channel"] == k].sort_values("f_Hz"); r0 = g.iloc[0]
            for sp, text, st in kinds:
                if sp in g and g[sp].notna().any():
                    kw = dict(ls="none", ms=3.4, mew=0.8); kw.update(st)
                    ax.plot(g["f_Hz"], g[sp], label=text, **kw)
            _log_y(ax)
            ax.set_ylabel(r"$U$ in V")
            ax.text(0.99, 0.97, probe_label(k, r0["role"], r0.get("distance_m")), transform=ax.transAxes,
                    ha="right", va="top", fontsize=STYLE.font_size_small, color="#333333")
        axes[-1, 0].set_xlabel("$f$ in Hz"); axes[-1, 0].set_xlim(left=0)
        _legend(axes[0, 0], len(kinds), loc="lower left", bbox_to_anchor=(0.0, 1.0))
    return _finish(fig, axes)


def _text_in_free_corner(ax, text: str, corners=("lr", "ur", "ll", "ul"), width: float = 0.32, height: float = 0.13,
                         **kw):
    """Put a label into the first corner of the panel without data points (``lr`` lower right, ``ur`` upper right,
    ``ll`` lower left, ``ul`` upper left) – otherwise lower right on a white background. Call after the axis limits
    have been fixed."""
    ax.get_xlim(); ax.get_ylim()                                       # finish the autoscaling
    per_axis = ax.transData + ax.transAxes.inverted()
    points = []
    for ln in ax.get_lines():
        xy = np.column_stack([np.asarray(ln.get_xdata(), float), np.asarray(ln.get_ydata(), float)])
        xy = xy[np.isfinite(xy).all(axis=1)]
        if len(xy):
            points.append(per_axis.transform(xy))
    P = np.vstack(points) if points else np.zeros((0, 2))
    position = {"lr": (0.99, 0.03, "right", "bottom"), "ur": (0.99, 0.97, "right", "top"),
                "ll": (0.01, 0.03, "left", "bottom"), "ul": (0.01, 0.97, "left", "top")}
    for e in corners:
        x, y, ha, va = position[e]
        x0, x1 = (x - width, x) if ha == "right" else (x, x + width)
        y0, y1 = (y, y + height) if va == "bottom" else (y - height, y)
        if not ((P[:, 0] >= x0) & (P[:, 0] <= x1) & (P[:, 1] >= y0) & (P[:, 1] <= y1)).any():
            return ax.text(x, y, text, transform=ax.transAxes, ha=ha, va=va, **kw)
    x, y, ha, va = position[corners[0]]
    return ax.text(x, y, text, transform=ax.transAxes, ha=ha, va=va,
                   bbox=dict(facecolor="white", alpha=0.85, edgecolor="none", pad=1.0), **kw)


def earthing_comparison(long_df: pd.DataFrame, ratio: Optional[pd.DataFrame] = None, width="full",
                        height_per_panel_cm: float = 5.0, f_max: Union[str, float, None] = "auto",
                        direct: bool = False):
    """Earthing measurement (groundmeas) and real earth fault per frequency – result of
    [`groundfault.earthing.compare`][groundfault.earthing.compare].

    One panel per quantity (earth potential rise, touch voltage; V/A = Ω, logarithmic): injection measurement per
    probe distance (grey, squares), earth fault of this evaluation per probe (colour, bar = span of the methods,
    ⊗ = uncertain), fault-test values from groundmeas (×). Below, the ratio fault test/earthing measurement
    (fault-related) – for the earth potential rise at a remote probe the effective reduction factor; black the probe
    maximum (highest voltage of all probes against the farthest probe of the earthing measurement).

    Parameters
    ----------
    long_df : pandas.DataFrame
        Long table of [`groundfault.earthing.compare`][groundfault.earthing.compare] (one row per source, probe and
        frequency).
    ratio : pandas.DataFrame, optional
        Ratio table of [`groundfault.earthing.compare`][groundfault.earthing.compare] → bottom panel (rows with
        ``distance_match == "different"`` are left out).
    width : {"full", "half"} or float
        Figure width (see [`figure`][groundfault.report.figure]).
    height_per_panel_cm : float
        Height per panel in cm.
    f_max : "auto", float or None
        Highest frequency in Hz; ``"auto"``: 15 % above the highest frequency of the earthing measurement.
    direct : bool
        Also draw the direct quotients of the fault test (open circles).

    Returns
    -------
    matplotlib.figure.Figure

    Raises
    ------
    ValueError
        If there are no values to compare.
    """
    long_df = long_df.copy()
    gf = long_df[long_df["source"] == _SOURCE_FAULT]
    inj = long_df[long_df["source"] == _SOURCE_INJECTION]
    gm = long_df[long_df["source"] == _SOURCE_GROUNDMEAS_FAULT]
    if f_max == "auto":
        f_max = 1.15 * float(inj["f_Hz"].max()) if not inj.empty else None
    if f_max:
        gf, inj, gm = (d[d["f_Hz"] <= f_max] for d in (gf, inj, gm))
    groups = [g for g in _COMPARISON_QUANTITIES if g in set(gf["quantity"]) | set(inj["quantity"])]
    if not groups:
        raise ValueError(tr("report.no_values_to_compare"))
    v = ratio if ratio is not None and not ratio.empty else None
    if v is not None:
        v = v[v["distance_match"] != "different"] if "distance_match" in v else v
        if f_max:
            v = v[v["f_Hz"] <= f_max]
        v = v if not v.empty else None
    n = len(groups) + (1 if v is not None else 0)
    fig, axes = figure(width, height_per_panel_cm * n + 1.0, nrows=n, sharex=True)
    colors = station_colors(gf["channel"].dropna())
    ticks = ["-", "--", ":", "-."]
    labels = []
    with _rc():
        for ax, grp in zip(axes[:, 0], groups):
            sub = inj[inj["quantity"] == grp]
            dist = sorted(sub["distance_m"].dropna().unique(), reverse=True)
            for j, a in enumerate(dist):
                for mid, s in sub[sub["distance_m"] == a].groupby("measurement_id"):
                    s = s.sort_values("f_Hz")
                    date = ", " + s["date"].iloc[0] if "date" in s and s["date"].iloc[0] else ""
                    label = tr("report.injection_at_distance", distance=_distance_text(a), id=int(mid), date=date)
                    if grp != _TOUCH_QUANTITY:
                        label = tr("report.injection_ze", label=label, symbol="$Z_\\mathrm{E}$")
                    ax.plot(s["f_Hz"], s["Z_Ohm"], color="#4d4d4d", ls=ticks[j % 4], marker="s", ms=3.0,
                            lw=STYLE.line_width, label=label)
            e = gf[(gf["quantity"] == grp) & (gf["type"] == "fault")]
            has_uncertain = False
            for channel, g in e.groupby("channel", sort=False):
                g = g.sort_values("f_Hz"); y = g["Z_Ohm"].to_numpy(float)
                lo = np.fmin(pd.to_numeric(g["Z_min_Ohm"], errors="coerce").to_numpy(float), y)
                hi = np.fmax(pd.to_numeric(g["Z_max_Ohm"], errors="coerce").to_numpy(float), y)
                lo = np.where(np.isfinite(lo), lo, y); hi = np.where(np.isfinite(hi), hi, y)
                a = _distance_text(g["distance_m"].iloc[0])
                label = (tr("report.fault_probe", channel=channel, distance=" (%s)" % a if a else "")
                         if grp == _TOUCH_QUANTITY else
                         tr("report.fault_probe_symbol", channel=channel, distance=" (%s)" % a if a else "",
                            symbol="$r\\,Z_\\mathrm{E}$"))
                ax.errorbar(g["f_Hz"], y, yerr=np.vstack([y - lo, hi - y]), color=colors[channel], marker="o", ms=3.2,
                            lw=STYLE.line_width, elinewidth=0.8, capsize=0, label=label)
                u = _uncertain_mask(g)
                _mark_uncertain(ax, g["f_Hz"].to_numpy(float)[u], y[u], colors[channel])
                has_uncertain = has_uncertain or bool(u.any())
            if direct:
                d = gf[(gf["quantity"] == grp) & (gf["type"] == "direct")]
                for channel, g in d.groupby("channel", sort=False):
                    g = g.sort_values("f_Hz")
                    ax.plot(g["f_Hz"], g["Z_Ohm"], ls="none", marker="o", ms=3.2, mfc="white",
                            mec=colors.get(channel, GREY), mew=0.8)
            gg = gm[gm["quantity"] == grp]
            for mid, s in gg.groupby("measurement_id"):
                a = _distance_text(s["distance_m"].iloc[0])
                ax.plot(s["f_Hz"], s["Z_Ohm"], ls="none", marker="x", ms=4.5, mew=1.0, color="#111111",
                        label=tr("report.fault_groundmeas", id=int(mid), distance=", " + a if a else ""))
            _log_y(ax)
            ax.set_ylabel(r"$U_\mathrm{T}/I$ in V/A" if grp == _TOUCH_QUANTITY
                          else r"$r\,Z_\mathrm{E}$, $Z_\mathrm{E}$ in V/A")
            labels.append((ax, _cap(tr(_COMPARISON_QUANTITIES[grp]))))
            handles, texts = ax.get_legend_handles_labels()
            if has_uncertain:
                handles.append(_uncertain_handle()); texts.append(tr("report.uncertain_background_methods"))
            ax.legend(handles, texts, loc="lower left", bbox_to_anchor=(0.0, 1.0), ncol=2, borderaxespad=0.15)
        if v is not None:
            ax = axes[-1, 0]
            for channel, g in v.groupby("channel", sort=False):
                g = g.sort_values("f_Hz")
                c = COLOR_MAX if channel == MAX_PROBE else colors.get(channel, GREY)
                name = (tr("report.probe_maximum", probes=" (%s)" % ", ".join(dict.fromkeys(g["probe"].astype(str))))
                        if channel == MAX_PROBE and "probe" in g else channel)
                ax.plot(g["f_Hz"], g["ratio_fault"], color=c, marker="o",
                        ms=3.2, lw=STYLE.line_width, label=name)
                u = _uncertain_mask(g)
                _mark_uncertain(ax, g["f_Hz"].to_numpy(float)[u],
                                g["ratio_fault"].to_numpy(float)[u], c)
            if float(np.nanmax(v["ratio_fault"])) > 0.6:      # reference line at 1 only if it is close
                ax.axhline(1.0, color=GREY, lw=0.6, ls="--")
            ax.set_ylim(bottom=0)
            ax.set_ylabel(tr("report.fault_to_injection"))
            ax.legend(loc="lower left", bbox_to_anchor=(0.0, 1.0), ncol=min(5, v["channel"].nunique()),
                      borderaxespad=0.15)
        axes[-1, 0].set_xlabel("$f$ in Hz"); axes[-1, 0].set_xlim(left=0)
        for ax, label in labels:                  # only now: the axis limits are fixed
            _text_in_free_corner(ax, label, fontsize=STYLE.font_size_small, color="#333333")
    return _finish(fig, axes)


def _remote_probes(df: pd.DataFrame) -> pd.DataFrame:
    """Only the rows with a comparable probe distance ("equal"/"similar") – if there are none, all rows."""
    if "distance_match" not in df:
        return df
    remote = df[df["distance_match"].isin(["equal", "similar"])]
    return remote if not remote.empty else df


def _probe_selection(df: pd.DataFrame, probes: str = "maximum") -> pd.DataFrame:
    """Rows for the figures of the earthing measurement: ``"maximum"`` the probe maximum (``channel`` =
    ``max(UE)``; without such rows like ``"remote"``), ``"remote"`` the probes with a comparable distance,
    ``"all"`` all rows."""
    if probes == "all" or df.empty or "channel" not in df:
        return df
    is_max = df["channel"] == MAX_PROBE
    if probes == "maximum" and is_max.any():
        return df[is_max]
    return _remote_probes(df[~is_max]) if (~is_max).any() else df


def _pair_text(channel: str, g: pd.DataFrame) -> str:
    """Legend text "Probe maximum (UE_10m) ↔ Earthing measurement ID 3 (40 m)" or "UE_40m (40 m) ↔ …"."""
    d = _distance_text(g["distance_injection_m"].iloc[0]) if "distance_injection_m" in g else ""
    right = tr("report.injection_id", id=int(g["injection_id"].iloc[0]), distance=" (%s)" % d if d else "")
    if channel == MAX_PROBE:
        s = ", ".join(dict.fromkeys(str(x) for x in g["probe"].dropna())) if "probe" in g else ""
        return "%s ↔ %s" % (tr("report.probe_maximum", probes=" (%s)" % s if s else ""), right)
    a = _distance_text(g["distance_fault_m"].iloc[0]) if "distance_fault_m" in g else ""
    return "%s%s ↔ %s" % (channel, " (%s)" % a if a else "", right)


def typical_band(r_typical: Optional[Tuple[float, float]] = None,
                 effective: Optional[pd.DataFrame] = None) -> Optional[Tuple[float, float]]:
    """Typical range of the reduction factor at 50 Hz that the figures draw as a grey bar.

    Parameters
    ----------
    r_typical : (float, float), optional
        Range (low, high) given by the user.
    effective : pandas.DataFrame, optional
        Effective table of [`groundfault.earthing.reduction_factor`][groundfault.earthing.reduction_factor]: its
        columns ``r_typical_min``/``r_typical_max`` hold the range used for the classification.

    Returns
    -------
    (float, float) or None
        ``r_typical``, otherwise the range of the effective table; None if neither gives one (then no bar – groundfault
        has no default range).
    """
    if r_typical is not None:
        return float(r_typical[0]), float(r_typical[1])
    if effective is not None and len(effective) and {"r_typical_min", "r_typical_max"} <= set(effective.columns):
        low = pd.to_numeric(effective["r_typical_min"], errors="coerce").dropna()
        high = pd.to_numeric(effective["r_typical_max"], errors="coerce").dropna()
        if len(low) and len(high):
            return float(low.iloc[0]), float(high.iloc[0])
    return None


def reduction_factor(spectrum: pd.DataFrame, effective: Optional[pd.DataFrame] = None, width="full",
                     height_cm: Optional[float] = None, probes: str = "maximum", shield_current: bool = True,
                     impedances: bool = True, r_typical: Optional[Tuple[float, float]] = None,
                     remote_probes_only: Optional[bool] = None, direct: bool = False, trend: bool = True):
    """Reduction factor over the frequency (magnitude) – result of
    [`groundfault.earthing.reduction_factor`][groundfault.earthing.reduction_factor].

    Top: r = r·Z_E/Z_E (fault test ÷ earthing measurement) over f, drawn like the impedances: ● r per order
    (earthing measurement measured at this frequency), ○ earthing measurement interpolated, ⊗ uncertain fault-test
    value, bar = span of the superposition methods (``r_min`` … ``r_max``); dotted the trend r ∝ f^b fitted at the
    reliable orders (``trend``, from ``effective``); dashed in the same colour the RMS value r_eff over all
    harmonics; grey bar at 50 Hz: the typical range given by the user (``r_typical``); ◆ grey: cross-check from the
    shield current (r = |I_F − I_S|/|I_F|, only meaningful if all return conductors are measured). With ``direct``
    also r from the values as measured (open circles, with background). Logarithmic if the values differ by more than
    a factor of 20. Bottom (``impedances``, V/A, logarithmic): the two compared quantities – fault test
    r·Z_E = U_E/I_F (●), fault test with the earth current Z_E = U_E/I_E (◇, if the shield current is measured) and
    earthing measurement Z_E (■ grey, □ interpolated); they differ by the reduction factor.

    Parameters
    ----------
    spectrum : pandas.DataFrame
        Reduction factor per order (first table of
        [`groundfault.earthing.reduction_factor`][groundfault.earthing.reduction_factor]).
    effective : pandas.DataFrame, optional
        RMS values r_eff (second table of
        [`groundfault.earthing.reduction_factor`][groundfault.earthing.reduction_factor]).
    width : {"full", "half"} or float
        Figure width (see [`figure`][groundfault.report.figure]).
    height_cm : float, optional
        Figure height in cm (default: 11.5 with, 6.5 without the impedance panel).
    probes : {"maximum", "remote", "all"}
        ``"maximum"`` the probe maximum (default), ``"remote"`` the probes with a comparable distance, ``"all"``
        all probes.
    shield_current : bool
        Show the cross-check from the shield current.
    impedances : bool
        Add the panel with the compared impedances.
    r_typical : (float, float), optional
        Typical range of the reduction factor at 50 Hz drawn as a grey bar; default: the range stored in ``effective``
        (``r_typical`` of [`groundfault.earthing.reduction_factor`][groundfault.earthing.reduction_factor]), without
        one no bar.
    remote_probes_only : bool, optional
        Earlier interface: True → ``probes="remote"``, False → ``probes="all"``.
    direct : bool
        Also draw r from the values as measured (``r_direct``, with background).
    trend : bool
        Draw the trend r ∝ f^b (needs ``effective``).

    Returns
    -------
    matplotlib.figure.Figure

    Raises
    ------
    ValueError
        If there is no reduction factor (no earthing measurement at the frequencies of the fault test).
    """
    from .earthing import power_fit
    band = typical_band(r_typical, effective)
    if remote_probes_only is not None:                    # earlier interface
        probes = "remote" if remote_probes_only else "all"
    sp = spectrum[np.isfinite(pd.to_numeric(spectrum["r"], errors="coerce"))] if not spectrum.empty else spectrum
    if sp.empty:
        raise ValueError(tr("report.no_reduction_factor"))
    sp = _probe_selection(sp, probes)
    f_max = 1.08 * float(sp["f_Hz"].max())
    colors = station_colors(spectrum["channel"])
    fig, axes = figure(width, height_cm or (11.5 if impedances else 6.5), nrows=2 if impedances else 1, sharex=True)
    ax = axes[0, 0]
    from matplotlib.lines import Line2D
    extra = []
    y_max = band[1] if band else 0.0
    all_y = list(band or ())
    with _rc():
        f1 = sp.loc[sp["order"] == 1, "f_Hz"]
        f1 = float(f1.iloc[0]) if len(f1) else 50.0
        if band:
            ax.vlines(f1, band[0], band[1], colors=COLOR_TYPICAL, lw=7.0, zorder=1)
        has_span = has_direct = False
        for (channel, mid), g in sp.groupby(["channel", "injection_id"], sort=False):
            g = g.sort_values("f_Hz"); c = colors.get(channel, GREY)
            y = g["r"].to_numpy(float); f = g["f_Hz"].to_numpy(float)
            interp = _interpolated(g)
            ax.plot(f, y, color=c, lw=STYLE.line_width, marker=None, label=_pair_text(channel, g))
            if {"r_min", "r_max"} <= set(g.columns):              # span of the superposition methods
                lo = np.fmin(pd.to_numeric(g["r_min"], errors="coerce").to_numpy(float), y)
                hi = np.fmax(pd.to_numeric(g["r_max"], errors="coerce").to_numpy(float), y)
                lo = np.where(np.isfinite(lo), lo, y); hi = np.where(np.isfinite(hi), hi, y)
                if np.any(hi - lo > 0.01 * np.abs(y)):                # visible span
                    ax.errorbar(f, y, yerr=np.vstack([y - lo, hi - y]), fmt="none", ecolor=c, elinewidth=0.8,
                                capsize=0, zorder=2)
                    has_span = True
                    all_y += list(lo) + list(hi)
                    y_max = max(y_max, float(np.nanmax(hi)))
            ax.plot(f[~interp], y[~interp], ls="none", marker="o", ms=3.4, color=c, zorder=3)
            ax.plot(f[interp], y[interp], ls="none", marker="o", ms=3.4, mfc="white", mec=c, mew=0.8, zorder=3)
            u = _uncertain_mask(g)
            _mark_uncertain(ax, f[u], y[u], c)
            if direct and "r_direct" in g:
                yd = pd.to_numeric(g["r_direct"], errors="coerce").to_numpy(float)
                ax.plot(f, yd, ls="none", marker="o", ms=3.2, mfc="white", mec=c, mew=0.8, alpha=0.8, zorder=2)
                has_direct = bool(np.isfinite(yd).any()) or has_direct
                all_y += list(yd)
                y_max = max(y_max, float(np.nanmax(yd)) if np.isfinite(yd).any() else 0.0)
            y_max = max(y_max, float(np.nanmax(y)))
            all_y += list(y)
            if effective is not None and not effective.empty:
                e = effective[(effective["channel"] == channel) & (effective["injection_id"] == mid)]
                if not e.empty and np.isfinite(e["r_eff"].iloc[0]):
                    r_eff = float(e["r_eff"].iloc[0])
                    ax.axhline(r_eff, color=c, lw=0.9, ls="--")
                    ax.text(f_max, r_eff, " $r_\\mathrm{eff}$ = %s" % number(r_eff, 2), color=c, va="bottom",
                            ha="right", fontsize=STYLE.font_size_small)
                    y_max = max(y_max, r_eff)
                orders_ = str(e["trend_orders"].iloc[0]) if not e.empty and "trend_orders" in e else ""
                chosen = [int(x) for x in orders_.replace(" ", "").split(",") if x.isdigit()]
                if trend and len(chosen) >= 2:                   # r ∝ f^b at the reliable orders
                    gt = g[g["order"].isin(chosen)]
                    c_, b_ = power_fit(gt["f_Hz"], gt["r"])
                    if np.isfinite(b_):
                        ft = np.geomspace(float(gt["f_Hz"].min()), float(gt["f_Hz"].max()), 50)
                        ax.plot(ft, c_ * ft ** b_, color=c, lw=0.9, ls=":", zorder=2,
                                label=tr("report.trend_r", formula="$r \\propto f^{\\,b}$", b=number(b_, 2),
                                         orders=orders_))
            if interp.any():
                extra = [(Line2D([], [], ls="none", marker="o", ms=3.4, mfc="white", mec="#333333", mew=0.8),
                          tr("report.injection_interpolated"))]
        if shield_current and "r_shield" in spectrum and spectrum["r_shield"].notna().any():
            s_ = spectrum.drop_duplicates("order").sort_values("f_Hz")
            s_ = s_[(s_["f_Hz"] <= f_max) & s_["r_shield"].notna()]
            if not s_.empty:
                ax.plot(s_["f_Hz"], s_["r_shield"], ls="none", marker="D", ms=6.0, mfc="none", mec=GREY,
                        mew=0.9, zorder=2, label=tr("report.from_shield_current",
                                                    quotient=r"$|I_\mathrm{F}-I_\mathrm{S}|/|I_\mathrm{F}|$"))
                y_max = max(y_max, float(np.nanmax(s_["r_shield"])))
                all_y += list(s_["r_shield"])
        ax.set_xlim(0, f_max)
        w = np.asarray(all_y, float); w = w[np.isfinite(w) & (w > 0)]
        if w.size and w.max() / w.min() > 20:
            _log_y(ax)
        else:
            ax.set_ylim(0, max(0.05, 1.15 * y_max))
        ax.set_ylabel(tr("report.reduction_factor_axis", symbol="$r$"))
        handles, texts = ax.get_legend_handles_labels()
        for gr, tx in extra:
            handles.append(gr); texts.append(tx)
        if has_span:
            handles.append(Line2D([], [], color="#333333", lw=0.8, marker="|", ms=8, mew=0.8))
            texts.append(tr("report.span_methods"))
        if has_direct:
            handles.append(Line2D([], [], ls="none", marker="o", ms=3.2, mfc="white", mec="#333333", mew=0.8))
            texts.append(tr("report.r_direct"))
        if _uncertain_mask(sp).any():
            handles.append(_uncertain_handle()); texts.append(tr("report.uncertain_background_methods"))
        handles.append(Line2D([], [], color="#333333", lw=0.9, ls="--")); texts.append(tr("report.rms_over_harmonics"))
        if band:
            handles.append(Line2D([], [], color=COLOR_TYPICAL, lw=7.0))
            texts.append(tr("report.typical_at_50hz", low=_g(band[0]), high=_g(band[1])))
        ax.legend(handles, texts, loc="lower left", bbox_to_anchor=(0.0, 1.0), ncol=2, borderaxespad=0.15)
        if impedances:
            az = axes[1, 0]
            all_ = _probe_selection(spectrum[spectrum["f_Hz"] <= f_max], probes)
            all_ = all_[all_.set_index(["channel", "injection_id"]).index.isin(
                sp.set_index(["channel", "injection_id"]).index.unique())]
            for (channel, mid), g in all_.groupby(["channel", "injection_id"], sort=False):
                g = g.sort_values("f_Hz"); c = colors.get(channel, GREY)
                left = _pair_text(channel, g).split(" ↔ ")[0]
                y = pd.to_numeric(g["rZE_fault_Ohm"], errors="coerce").to_numpy(float)
                az.plot(g["f_Hz"], y, color=c, marker="o", ms=3.4, lw=STYLE.line_width,
                        label=tr("report.fault_quotient", name=left,
                                 quotient="$r\\,Z_\\mathrm{E} = U_\\mathrm{E}/I_\\mathrm{F}$"))
                u = _uncertain_mask(g)
                _mark_uncertain(az, g["f_Hz"].to_numpy(float)[u], y[u], c)
                if "Z_E_earth_current_Ohm" in g and pd.to_numeric(g["Z_E_earth_current_Ohm"],
                                                                   errors="coerce").notna().any():
                    ie = pd.to_numeric(g["Z_E_earth_current_Ohm"], errors="coerce").to_numpy(float)
                    if np.nanmax(np.abs(ie / y - 1.0)) > 0.05:            # only if visibly different
                        az.plot(g["f_Hz"], ie, color=c, ls=":", lw=0.9, marker="D", ms=3.0, mfc="white", mec=c,
                                mew=0.8, label=tr("report.fault_quotient", name=left,
                                                  quotient="$Z_\\mathrm{E} = U_\\mathrm{E}/I_\\mathrm{E}$"))
            for mid, g in sp.groupby("injection_id", sort=False):
                g = g.drop_duplicates("order").sort_values("f_Hz")
                interp = _interpolated(g)
                y = g["Z_E_injection_Ohm"].to_numpy(float); f = g["f_Hz"].to_numpy(float)
                d = _distance_text(g["distance_injection_m"].iloc[0])
                az.plot(f, y, color="#4d4d4d", ls="--", lw=STYLE.line_width, marker=None,
                        label=tr("report.injection_ze", symbol="$Z_\\mathrm{E}$",
                                 label=tr("report.injection_id", id=int(mid), distance=" (%s)" % d if d else "")))
                az.plot(f[~interp], y[~interp], ls="none", marker="s", ms=3.0, color="#4d4d4d", zorder=3)
                az.plot(f[interp], y[interp], ls="none", marker="s", ms=3.0, mfc="white", mec="#4d4d4d", mew=0.8,
                        zorder=3)
            _log_y(az)
            az.set_ylabel(r"$r\,Z_\mathrm{E}$, $Z_\mathrm{E}$ in V/A")
            handles, texts = az.get_legend_handles_labels()
            if _uncertain_mask(all_).any():
                handles.append(_uncertain_handle()); texts.append(tr("report.uncertain"))
            az.legend(handles, texts, loc="lower left", bbox_to_anchor=(0.0, 1.0), ncol=2,
                      borderaxespad=0.15)
        axes[-1, 0].set_xlabel("$f$ in Hz")
    return _finish(fig, axes)


def lv_network(lv_table: pd.DataFrame, check: Optional[pd.DataFrame] = None, width="full",
               height_per_panel_cm: float = 5.0, all_probes: bool = False):
    """Earthing impedance of the LV network and parallel-connection check (table ``station_lv_network`` and –
    with an earthing measurement – [`groundfault.earthing.check_lv_network`][groundfault.earthing.check_lv_network]).

    Top (V/A, logarithmic): Z_E,LV = ΔU_E/I_LV (●), from the fault test r·Z_E = max(U_E)/I_F (◇; with a shield
    current also Z_E = ΔU_E/I_E, △, if visibly different) and earthing measurement (■). As a parallel branch,
    Z_E,LV must lie above r·Z_E (and above Z_E); if it lies below the earthing measurement, that measurement
    probably measures too much. Bottom (A, logarithmic): fault-related current into the LV network
    (●, mean; ▲ before, ▼ after the transformer star point), fault current I_F (□) and earth current I_E (◇, with a
    shield current) – I_LV must lie below them.

    Parameters
    ----------
    lv_table : pandas.DataFrame
        Table ``station_lv_network`` of the result.
    check : pandas.DataFrame, optional
        Check against the earthing measurement (first table of
        [`groundfault.earthing.check_lv_network`][groundfault.earthing.check_lv_network]).
    width : {"full", "half"} or float
        Figure width (see [`figure`][groundfault.report.figure]).
    height_per_panel_cm : float
        Height per panel in cm.
    all_probes : bool
        All probes; otherwise only the probe maximum (``max(UE)``; without such rows the farthest probe).

    Returns
    -------
    matplotlib.figure.Figure

    Raises
    ------
    ValueError
        If there is no LV network evaluation (PEN currents).
    """
    if lv_table is None or lv_table.empty:
        raise ValueError(tr("report.no_lv_network"))
    so = lv_table.copy()
    if not all_probes:
        if (so["channel"] == MAX_PROBE).any():
            so = so[so["channel"] == MAX_PROBE]
        elif so["distance_m"].notna().any():
            so = so[so["distance_m"] == so["distance_m"].max()]
    if check is not None and not check.empty:
        check = check[check["channel"].isin(set(so["channel"]))]
    colors = station_colors(so["channel"])
    sub_lv = tr("report.sub_lv_current")
    fig, axes = figure(width, 2 * height_per_panel_cm + 1.2, nrows=2, sharex=True)
    with _rc():
        ax = axes[0, 0]
        for channel, g in so.groupby("channel", sort=False):
            g = g.sort_values("f_Hz"); c = colors.get(channel, GREY)
            if channel == MAX_PROBE:
                s = ", ".join(dict.fromkeys(str(x) for x in g["probe"].dropna())) if "probe" in g else ""
                name = tr("report.probe_maximum", probes=" (%s)" % s if s else "")
                reference = "\\mathrm{max}(U_\\mathrm{E})"
            else:
                a = _distance_text(g["distance_m"].iloc[0])
                name, reference = "%s%s" % (channel, " (%s)" % a if a else ""), "U_\\mathrm{E}"
            y = g["Z_E_LV_Ohm"].to_numpy(float)
            ax.plot(g["f_Hz"], y, color=c, marker="o", ms=3.4, lw=STYLE.line_width,
                    label="$Z_\\mathrm{E,%s}$ %s" % (tr("report.sub_lv_impedance"), name))
            u = _uncertain_mask(g)
            _mark_uncertain(ax, g["f_Hz"].to_numpy(float)[u], y[u], c)
            z_f = pd.to_numeric(g["rZE_fault_Ohm"], errors="coerce") if "rZE_fault_Ohm" in g else None
            if z_f is not None and z_f.notna().any():
                ax.plot(g["f_Hz"], z_f, color=c, ls=":", lw=0.9, marker="D", ms=3.0, mfc="white", mec=c, mew=0.8,
                        label=tr("report.fault_rze", symbol="$r\\,Z_\\mathrm{E}$",
                                 quotient="$%s/I_\\mathrm{F}$" % reference))
            with_s = str(g["total_earthing_reference"].iloc[0]).startswith("I_E")
            z_g = pd.to_numeric(g["Z_E_earth_current_Ohm"], errors="coerce")
            if with_s and (z_f is None or not z_f.notna().any() or np.nanmax(np.abs(z_g / z_f - 1.0)) > 0.05):
                ax.plot(g["f_Hz"], z_g, color=c, ls=":", lw=0.9, marker="^", ms=3.0, mfc="white", mec=c, mew=0.8,
                        label=tr("report.fault_rze", symbol="$Z_\\mathrm{E}$",
                                 quotient="$%s/I_\\mathrm{E}$" % reference))
        if check is not None and not check.empty:
            pr = check if (check["channel"] == MAX_PROBE).any() else _remote_probes(check)
            for (channel, mid), g in pr.groupby(["channel", "injection_id"], sort=False):
                g = g.sort_values("f_Hz")
                ax.plot(g["f_Hz"], g["Z_E_injection_Ohm"], color="#4d4d4d", marker="s", ms=3.0, ls="--",
                        lw=STYLE.line_width, label=tr("report.injection_ze", symbol="$Z_\\mathrm{E}$",
                                                      label=tr("report.injection_id", id=int(mid), distance=" (%s)"
                                                               % _distance_text(g["distance_injection_m"].iloc[0]))))
        _log_y(ax)
        ax.set_ylabel(r"$|Z|$ in V/A")
        handles, texts = ax.get_legend_handles_labels()
        if _uncertain_mask(so).any():
            handles.append(_uncertain_handle()); texts.append(tr("report.uncertain_background_methods"))
        ax.legend(handles, texts, loc="lower left", bbox_to_anchor=(0.0, 1.0), ncol=2, borderaxespad=0.15)
        ax = axes[1, 0]
        g = so[so["channel"] == so["channel"].iloc[0]].sort_values("f_Hz")
        ax.plot(g["f_Hz"], g["I_LV_A"], color="#111111", marker="o", ms=3.4, lw=STYLE.line_width,
                label=tr("report.lv_current", symbol="$I_\\mathrm{%s}$" % sub_lv))
        for sp_, mk, key in (("I_LV_feeders_A", "^", "report.before_star_point"),
                             ("I_LV_busbar_A", "v", "report.after_star_point")):
            if sp_ in g and g[sp_].notna().any():
                ax.plot(g["f_Hz"], g[sp_], ls="none", marker=mk, ms=3.4, mfc="white", mec="#111111", mew=0.8,
                        label=tr(key))
        ax.plot(g["f_Hz"], g["I_F_A"], ls="none", marker="s", ms=3.4, mfc="white", mec=GREY, mew=0.8,
                label=tr("report.fault_current", symbol="$I_\\mathrm{F}$"))
        if str(g["total_earthing_reference"].iloc[0]).startswith("I_E"):
            i_e = (g["I_LV_A"] / g["I_LV_to_I_E"]).where(g["I_LV_to_I_E"] > 0)
            if np.nanmax(np.abs(i_e / g["I_F_A"] - 1.0)) > 0.05:
                ax.plot(g["f_Hz"], i_e, ls="none", marker="D", ms=3.4, mfc="white", mec=GREY, mew=0.8,
                        label=tr("report.earth_current", symbol="$I_\\mathrm{E} = I_\\mathrm{F} - I_\\mathrm{S}$"))
        _log_y(ax)
        ax.set_ylabel("$I$ in A"); ax.set_xlabel("$f$ in Hz"); ax.set_xlim(left=0)
        ax.legend(loc="lower left", bbox_to_anchor=(0.0, 1.0), ncol=2, borderaxespad=0.15)
    return _finish(fig, axes)


# =========================================================================== numbers and tables
def number(x, digits: int = 3, mode: str = "significant", thousands: bool = False, lang: Optional[str] = None) -> str:
    """Number as text for report tables and figure labels.

    Rounds half away from zero (DIN 1333: 12.25 → 12.3, whereas Python rounds half to even), writes the decimal
    separator of the output language, a typographic minus (U+2212) and – optionally – a thin space (U+2009) as
    thousands separator. Integers (``int``, e.g. counts) stay integers.

    Parameters
    ----------
    x : float, int or None
        The number; ``None``, NaN and ±inf give ``"–"``, values that are not numbers are returned as ``str(x)``.
    digits : int
        Significant digits (``mode="significant"``) or decimals (``mode="fixed"``).
    mode : {"significant", "fixed"}
        Significant digits or a fixed number of decimals.
    thousands : bool
        Group the thousands with a thin space.
    lang : {"en", "de"}, optional
        Explicit language of the decimal separator (``"de"``: comma, ``"en"``: point). Default:
        ``STYLE.decimal_comma`` if set, otherwise the active output language.

    Returns
    -------
    str

    Examples
    --------
    >>> number(0.012345, 3, lang="de")
    '0,0123'
    >>> number(12.5, 2, "fixed", lang="en")
    '12.50'
    >>> number(1234.5678, 3)
    '1235'
    """
    if x is None:
        return "–"
    if isinstance(x, (int, np.integer)) and not isinstance(x, (bool, np.bool_)):
        s = "{:,d}".format(int(x)) if thousands else str(int(x))
    else:
        try:
            xf = float(x)
        except (TypeError, ValueError):
            return str(x)
        if not math.isfinite(xf):
            return "–"

        def rounding(n: int) -> Decimal:
            # repr() gives the shortest decimal representation → 12.25 becomes 12.3 (Python's "round half to
            # even" would give 12.2)
            return Decimal(repr(xf)).quantize(Decimal(1).scaleb(-n), rounding=ROUND_HALF_UP)

        if mode == "fixed":
            n = digits
            d = rounding(n)
        elif xf == 0:
            n, d = 0, Decimal(0)
        else:
            exp = int(math.floor(math.log10(abs(xf))))
            n = max(0, digits - 1 - exp)
            d = rounding(n)
            if n > 0 and abs(d) >= Decimal(10) ** (exp + 1):          # 9.995 → 10.0 (not 10.00)
                n -= 1
                d = rounding(n)
        s = "{:,.{n}f}".format(d, n=n) if thousands else "{:.{n}f}".format(d, n=n)
    s = s.replace(",", "\u2009").replace("-", "\u2212")
    return s.replace(".", ",") if _comma(lang) else s


def _localize(df, lang: Optional[str] = None):
    """[`groundfault.i18n.localize`][groundfault.i18n.localize] for display and export, with text columns of the pandas
    string dtype converted to ``object`` beforehand.

    The conversion is not needed for the translation: ``localize`` translates the code values of text columns of
    dtype ``object`` and of the pandas string dtype alike (``str``, the default for text in pandas ≥ 3, also has
    ``dtype.kind == "O"``). Translated columns get the dtype pandas infers for the translated values.
    """
    if hasattr(df, "columns") and df.columns.is_unique:
        text = [c for c in df.columns if df[c].dtype != object and pd.api.types.is_string_dtype(df[c])]
        if text:
            df = df.astype({c: object for c in text})
    return i18n.localize(df, lang)


def _column_digits(digits: Union[int, Mapping[str, int]], column, lang: Optional[str] = None) -> int:
    """Digits of ``column``: ``digits`` itself or its entry for the English or the localized column name (3)."""
    if not isinstance(digits, Mapping):
        return digits
    if column in digits:
        return digits[column]
    return digits.get(i18n.column_label(column, lang), 3)


def format_table(df: pd.DataFrame, digits: Union[int, Mapping[str, int]] = 3, mode: str = "significant",
                 thousands: bool = False, lang: Optional[str] = None) -> pd.DataFrame:
    """Table with the numbers formatted as text, in the output language (Markdown, LaTeX, DOCX, display).

    Parameters
    ----------
    df : pandas.DataFrame
        Table with English column names and code values (or a report table of
        [`report_tables`][groundfault.report.report_tables]).
    digits : int or mapping of str to int
        Digits for [`number`][groundfault.report.number]: for all columns, or per column (English or localized column
        name; columns that are not in the mapping: 3).
    mode : {"significant", "fixed"}
        See [`number`][groundfault.report.number].
    thousands : bool
        Thin space as thousands separator.
    lang : {"en", "de"}, optional
        Output language (default: the active one; see [`number`][groundfault.report.number] for the decimal separator).

    Returns
    -------
    pandas.DataFrame
        Copy of ``df``: numeric columns as text (integer columns stay integers, ``None`` → "–", numbers in mixed
        columns are formatted as well); column names and code values translated with
        [`groundfault.i18n.localize`][groundfault.i18n.localize] (German) or unchanged (English).
    """
    out = df.copy()
    for c in out.columns:
        n = _column_digits(digits, c, lang)
        if pd.api.types.is_numeric_dtype(out[c]) and not pd.api.types.is_bool_dtype(out[c]):
            whole = pd.api.types.is_integer_dtype(out[c])
            out[c] = [number(int(v) if whole else v, n, mode, thousands, lang) for v in out[c]]
        elif out[c].dtype == object:                   # mixed columns (e.g. total rows with text)
            out[c] = [number(v, n, mode, thousands, lang) if isinstance(v, (int, float, np.number))
                      and not isinstance(v, (bool, np.bool_)) else ("–" if v is None else v) for v in out[c]]
    return _localize(out, lang)


def _latex(df: pd.DataFrame) -> str:
    """Table as LaTeX ``tabular`` (booktabs rules, first column left-aligned)."""
    def esc(s):
        s = str(s)
        for a, b in (("\\", r"\textbackslash{}"), ("&", r"\&"), ("%", r"\%"), ("_", r"\_"), ("#", r"\#"),
                     ("Ω", r"$\Omega$"), ("°", r"$^\circ$"), ("\u2212", r"$-$"), ("\u2009", r"\,")):
            s = s.replace(a, b)
        return s
    columns = "l" + "r" * (len(df.columns) - 1)
    z = [r"\begin{tabular}{%s}" % columns, r"\toprule", " & ".join(esc(c) for c in df.columns) + r" \\", r"\midrule"]
    for _, r in df.iterrows():
        z.append(" & ".join(esc(v) for v in r.values) + r" \\")
    z += [r"\bottomrule", r"\end{tabular}"]
    return "\n".join(z) + "\n"


def _markdown(df: pd.DataFrame) -> str:
    """Table as Markdown (first column left-aligned, the others right-aligned)."""
    k = "| " + " | ".join(str(c) for c in df.columns) + " |"
    t = "|" + "|".join([":---"] + ["---:"] * (len(df.columns) - 1)) + "|"
    return "\n".join([k, t] + ["| " + " | ".join(str(v) for v in r.values) + " |" for _, r in df.iterrows()]) + "\n"


def _csv_ready(df: pd.DataFrame, comma: bool) -> pd.DataFrame:
    """Floats in mixed (object) columns with a decimal comma – pandas applies ``decimal=","`` to float columns only."""
    if not comma:
        return df
    out = df.copy()
    for c in out.columns:
        if out[c].dtype == object:
            out[c] = [str(v).replace(".", ",") if isinstance(v, (float, np.floating)) and math.isfinite(v) else v
                      for v in out[c]]
    return out


def export_table(df: pd.DataFrame, path_stem: str, formats: Sequence[str] = ("csv", "xlsx", "md"),
                 digits: Union[int, Mapping[str, int]] = 3, index: bool = False) -> List[str]:
    """Export a table in the output language.

    CSV and XLSX keep numbers as numbers (Excel can calculate with them; CSV with semicolons and the decimal separator
    of the language), Markdown, LaTeX and DOCX contain formatted numbers
    ([`format_table`][groundfault.report.format_table]). Column names and code values are translated with
    [`groundfault.i18n.localize`][groundfault.i18n.localize] in German and keep their English names in English; switch
    the language temporarily with ``with i18n.language("de"): …``.

    Parameters
    ----------
    df : pandas.DataFrame
        Table with English column names and code values (or a report table of
        [`report_tables`][groundfault.report.report_tables]).
    path_stem : str
        Path without extension; the folder is created if necessary.
    formats : sequence of {"csv", "xlsx", "md", "tex", "docx"}
        File formats. XLSX needs ``openpyxl``, DOCX ``python-docx`` (extra ``export``).
    digits : int or mapping of str to int
        Digits of the formatted numbers (see [`format_table`][groundfault.report.format_table]).
    index : bool
        Write the index as the first column(s).

    Returns
    -------
    list of str
        Paths of the written files.

    Raises
    ------
    ValueError
        For an unknown format.
    ImportError
        For XLSX or DOCX if the library is not installed (the message names ``pip install "groundfault[export]"``);
        nothing is written then.
    """
    for fmt, library in (("xlsx", "openpyxl"), ("docx", "python-docx")):
        if fmt in formats and not _export_library(library):
            raise ImportError(tr("report.export_library_missing", format=fmt.upper(), library=library))
    os.makedirs(os.path.dirname(os.path.abspath(path_stem)) or ".", exist_ok=True)
    if index:
        df = df.reset_index()
    shown = _localize(df)
    comma = _comma()
    sheet = tr("report.sheet_table")
    paths = []
    for fmt in formats:
        p = "%s.%s" % (path_stem, fmt)
        if fmt == "csv":
            _csv_ready(shown, comma).to_csv(p, sep=";", decimal="," if comma else ".", index=False,
                                            encoding="utf-8-sig")
        elif fmt == "xlsx":
            with pd.ExcelWriter(p, engine="openpyxl") as w:
                _excel_ready(shown).to_excel(w, index=False, sheet_name=sheet)
                _format_xlsx(w.sheets[sheet], shown)
        elif fmt == "md":
            with open(p, "w", encoding="utf-8") as f:
                f.write(_markdown(format_table(df, digits)))
        elif fmt == "tex":
            with open(p, "w", encoding="utf-8") as f:
                f.write(_latex(format_table(df, digits)))
        elif fmt == "docx":
            docx_tables({sheet: df}, p, digits)
        else:
            raise ValueError(tr("report.unknown_table_format", format=fmt))
        paths.append(p)
    return paths


def _excel_ready(df: pd.DataFrame) -> pd.DataFrame:
    """Time stamps with a time zone as text (Excel has no time zones): columns of a datetime dtype with a time zone
    and single values with ``tzinfo`` in columns of dtype object."""
    out = None
    for c in df.columns:
        s = df[c]
        if isinstance(s.dtype, pd.DatetimeTZDtype):
            out = df.copy() if out is None else out
            out[c] = s.astype(str)
        elif s.dtype == object and s.map(lambda v: getattr(v, "tzinfo", None) is not None).any():
            out = df.copy() if out is None else out
            out[c] = s.map(lambda v: str(v) if getattr(v, "tzinfo", None) is not None else v)
    return df if out is None else out


def _export_library(library: str) -> bool:
    """True if the library of an export format can be imported (``"openpyxl"``, ``"python-docx"``)."""
    try:
        importlib.import_module({"python-docx": "docx"}.get(library, library))
    except ImportError:
        return False
    return True


def _format_xlsx(ws, df: pd.DataFrame):
    """Bold header, column widths, three decimals for float columns and a frozen header row."""
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    for j, c in enumerate(df.columns, start=1):
        ws.cell(row=1, column=j).font = Font(bold=True)
        width = max(len(str(c)), *(len(number(v) if isinstance(v, (int, float, np.floating)) else str(v))
                                   for v in df[c].head(200))) if len(df) else len(str(c))
        ws.column_dimensions[get_column_letter(j)].width = min(max(8, width + 2), 60)
        if pd.api.types.is_float_dtype(df[c]):
            for i in range(2, len(df) + 2):
                ws.cell(row=i, column=j).number_format = "0.000"
    ws.freeze_panes = "A2"


def export_tables(tables: Mapping[str, pd.DataFrame], folder: str, prefix: str = "",
                  formats: Sequence[str] = ("csv", "md"), excel_workbook: bool = True,
                  docx: bool = True, digits: Union[int, Mapping[str, int]] = 3, names: Iterable[str] = ()) -> List[str]:
    """Export several tables in the output language: single files, one Excel workbook and one DOCX file.

    Column names and code values follow the output language as in [`export_table`][groundfault.report.export_table]; the
    file names do not: they are English in every language and folder layout. File and sheet names are the table names
    (keys of ``tables``); the workbook and the Word file are called ``<prefix>tables.xlsx`` and ``<prefix>tables.docx``.
    The headings in the Word file are the titles of the tables ([`table_title`][groundfault.report.table_title]).

    Parameters
    ----------
    tables : mapping of str to pandas.DataFrame
        Tables by name; empty tables and ``None`` are skipped.
    folder : str
        Target folder (created if necessary).
    prefix : str
        Prefix of all file names (e.g. the measurement).
    formats : sequence of str
        Formats of the single files (see [`export_table`][groundfault.report.export_table]).
    excel_workbook : bool
        Write all tables into one workbook (one sheet per table, sheet name: the first 31 characters of the name).
        Needs ``openpyxl``; without it the workbook is skipped with a warning.
    docx : bool
        Write all tables into one Word file. Needs ``python-docx``; without it the file is skipped with a warning.
    digits : int or mapping of str to int
        Digits of the formatted numbers (see [`format_table`][groundfault.report.format_table]).
    names : iterable of str, optional
        Names of the measurements, for the headings of the tables per measurement (see
        [`table_title`][groundfault.report.table_title]).

    Returns
    -------
    list of str
        Paths of the written files.

    Raises
    ------
    ImportError
        If ``formats`` asks for XLSX or DOCX single files and the library is missing (see
        [`export_table`][groundfault.report.export_table]).

    Warns
    -----
    UserWarning
        If the workbook or the Word file is skipped because ``openpyxl`` or ``python-docx`` is not installed (the
        message names ``pip install "groundfault[export]"``).
    """
    os.makedirs(folder, exist_ok=True)
    paths = []
    for name, df in tables.items():
        if df is None or df.empty:
            continue
        paths += export_table(df, os.path.join(folder, prefix + name), formats, digits)
    stem = os.path.join(folder, prefix + "tables")             # file names are English in every language
    if excel_workbook and not _export_library("openpyxl"):
        warnings.warn(tr("report.export_skipped", file=os.path.basename(stem + ".xlsx"), library="openpyxl"),
                      stacklevel=2)
    elif excel_workbook:
        p = stem + ".xlsx"
        with pd.ExcelWriter(p, engine="openpyxl") as w:
            for name, df in tables.items():
                if df is None or df.empty:
                    continue
                sheet = name[:31]
                shown = _excel_ready(_localize(df))
                shown.to_excel(w, index=False, sheet_name=sheet)
                _format_xlsx(w.sheets[sheet], shown)
        paths.append(p)
    if docx:
        p = stem + ".docx"
        if docx_tables(tables, p, digits, names=names):
            paths.append(p)
        else:
            warnings.warn(tr("report.export_skipped", file=os.path.basename(p), library="python-docx"), stacklevel=2)
    return paths


def docx_tables(tables: Mapping[str, pd.DataFrame], path: str, digits=3,
                font: Optional[str] = None, size_pt: float = 10.0, names: Iterable[str] = ()) -> bool:
    """Word file with all tables (formatted numbers and headers in the output language, font as in the report).

    Parameters
    ----------
    tables : mapping of str to pandas.DataFrame
        Tables by name (the name gives the heading, see [`export_tables`][groundfault.report.export_tables]); empty
        tables are skipped.
    path : str
        Path of the ``.docx`` file.
    digits : int or mapping of str to int
        Digits of the formatted numbers (see [`format_table`][groundfault.report.format_table]).
    font : str, optional
        Font (default: the first font of the report style).
    size_pt : float
        Font size in pt.
    names : iterable of str, optional
        Names of the measurements, for the headings of the tables per measurement (see
        [`table_title`][groundfault.report.table_title]).

    Returns
    -------
    bool
        False if ``python-docx`` is not installed (no file written).
    """
    try:
        import docx
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt, Cm
    except Exception:
        return False
    font = font or STYLE.fonts[0]
    d = docx.Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for r in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, r, Cm(2.54))
    st = d.styles["Normal"]; st.font.name = font; st.font.size = Pt(size_pt)
    for name, df in tables.items():
        if df is None or df.empty:
            continue
        d.add_paragraph(table_title(name, names), style="Heading 2")
        f = format_table(df, digits)
        t = d.add_table(rows=1, cols=len(f.columns))
        t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, c in enumerate(f.columns):
            cell = t.rows[0].cells[j]; cell.text = str(c)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.bold = True
        for _, r in f.iterrows():
            cells = t.add_row().cells
            for j, v in enumerate(r.values):
                cells[j].text = str(v)
                if j > 0:
                    cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        d.add_paragraph("")
    d.save(path)
    return True


#: tables with a title (message ``report.title_<name>``): the tables of :func:`report_tables`, the tables of a result
#: and those the notebooks export; a table per measurement carries its name after the table name (``epr_m1``)
_TABLE_TITLES = ("voltages", "residual_current", "split_50hz", "epr", "shield_current", "transformer",
                 "assessment_quantities", "equivalent_sources", "overview", "station_frequency", "earthing_comparison",
                 "earthing_ratio", "station_lv_network", "station_lv_network_eff", "reduction_factor",
                 "reduction_factor_eff", "lv_network_check", "lv_network_check_eff", "plausibility", "statistics",
                 "harmonics", "coil_currents", "voltage_metrics", "voltage_pre_fault", "voltage_fault",
                 "voltage_post_fault", "station_metrics", "station_voltages", "station_currents",
                 "station_currents_frequency", "station_extrapolation", "transformer_power", "transformer_node",
                 "transformer_power_change", "transformer_impedance", "transformer_divider", "transformer_summary",
                 "steady_state_segments")


def table_title(name: str, names: Iterable[str] = ()) -> str:
    """Heading of a table (Word file, notebook) in the output language.

    Every table of [`report_tables`][groundfault.report.report_tables], of a result and of the notebook exports has a
    title. A table per measurement (``epr_m4_station``, ``station_frequency_m4_station``) gets the title of its table
    followed by the name of the measurement: as given if it is one of ``names`` (the table name only holds it in the
    form of [`slug`][groundfault.report.slug]), otherwise as in the table name.

    Parameters
    ----------
    name : str
        Table name such as ``"residual_current"``, ``"epr_m4_station"`` or ``"station_frequency_m4_station"``.
    names : iterable of str, optional
        Names of the measurements, e.g. the keys of the results passed to
        [`report_tables`][groundfault.report.report_tables].

    Returns
    -------
    str
        Title in the output language, e.g. ``"Residual earth-fault current"`` or ``"Earth potential rise and touch
        voltage per frequency – M4 (Station)"``. A part of the package imported on demand registers the titles of
        its tables as messages ``report.title_<name>`` ([`groundfault.i18n.register`][groundfault.i18n.register]).
        Other names are shown as they are (German: with the German column name if there is one).

    Examples
    --------
    >>> table_title("station_frequency_m4_station", ["M3 (SUB 0 Ω)", "M4 (Station)"])
    'Earth potential rise and touch voltage per frequency – M4 (Station)'
    """
    name = str(name)
    if name in _TABLE_TITLES or (name != "of" and i18n.has_message("report.title_" + name)):   # also registered
        return tr("report.title_" + name)                                                       # by other parts
    splits = [(t, name[len(t) + 1:]) for t in _TABLE_TITLES if name.startswith(t + "_")]
    if splits:
        given = {slug(n): str(n) for n in names}
        known = [(t, given[rest]) for t, rest in splits if rest in given]
        table, subject = max(known or [(t, rest.replace("_", " ")) for t, rest in splits], key=lambda x: len(x[0]))
        return tr("report.title_of", title=tr("report.title_" + table), name=subject)
    label = i18n.column_label(name)
    if label == name and i18n.get_language() == "de":
        keys = [k for k in i18n.COLUMNS_DE if len(k) > 2 and "%" not in k
                and name.startswith(k if k.endswith("_") else k + "_")]
        if keys:
            k = max(keys, key=len)
            label = i18n.COLUMNS_DE[k] + name[len(k):]
    return _cap(label.replace("_", " "))


# =========================================================================== report tables (layout of earlier reports)
#: row labels of the sequence systems in the voltage table (message keys)
_SEQUENCE_LABELS = {"pos": "report.positive_sequence", "neg": "report.negative_sequence"}


def _num(x) -> float:
    """Float of a value of a result (None and non-numbers → NaN)."""
    try:
        return float(x) if x is not None else float("nan")
    except (TypeError, ValueError):
        return float("nan")


def report_tables(results: Mapping[str, object], series=None, Z_E: Optional[float] = None,
                  orders: Sequence[int] = (1, 3, 5, 7), comparison=None) -> Dict[str, pd.DataFrame]:
    """Tables in the layout of earlier test reports (one column per measurement).

    Parameters
    ----------
    results : mapping of str to Result or ResultFile
        Results in report order, e.g. ``{"M1 (substation, 2.5 Ω)": result, …}``; the keys become column names.
    series : SourceSeries, optional
        Equivalent sources
        ([`groundfault.equivalent_source.fit_equivalent_source`][groundfault.equivalent_source.fit_equivalent_source] or
        [`groundfault.equivalent_source.measured_source`][groundfault.equivalent_source.measured_source]) → table
        ``equivalent_sources``; with the measured source a column names the source of each order (the orders of the
        zero sequence keep the fitted one).
    Z_E : float, optional
        Earthing impedance in Ω: adds the current per harmonic at Z_E, the remainder that is not fitted and I_RES to
        ``equivalent_sources``.
    orders : sequence of int
        Harmonic orders of the residual-current table (the fundamental is always included).
    comparison : SourceSeries, optional
        A second series of the same measurements (the fitted source when ``series`` is the measured one, or the
        other way round): adds its current at Z_E; the Thevenin check U_q/U_faulted phase,pre is always that of the
        fitted source.

    Returns
    -------
    dict of str to pandas.DataFrame
        Tables without data are left out:

        * ``voltages`` – positive/negative-sequence voltage per harmonic and the largest line-to-line THD before
          and after the fault,
        * ``residual_current`` – residual earth-fault current: RMS value, harmonics and harmonic content,
        * ``split_50hz`` – active and detuning current, coil currents,
        * ``epr_<measurement>`` – largest measured earth potential rise and r·Z_E per 100 A (stations),
        * ``shield_current_<measurement>`` – shield current and reduction factor per harmonic (stations),
        * ``transformer`` – harmonic power and sequence currents at the transformer,
        * ``transformer_summary`` – transformer node: polarity of the currents, upstream impedance X_up/m (with HV
          voltages split into transformer and HV network) and change of the fundamental power during the fault,
        * ``assessment_quantities`` – measured and assessment value of each measurement,
        * ``equivalent_sources`` – source impedance and voltage per order (with ``series``).

    Notes
    -----
    The tables are display tables in the output language that is active when they are created: headers and row labels
    are messages, the canonical columns (``quantity``, ``unit``, ``measurement``, ``sequence``) and their code values
    are translated with [`groundfault.i18n.localize`][groundfault.i18n.localize] (German: ``Groesse``, ``Einheit``,
    ``Messung``, ``Sequenz``).
    """
    tab: Dict[str, pd.DataFrame] = {}
    # ---- voltages (sequence system of the natural order) before/after the fault
    rows = {}
    for nm, e in results.items():
        sp = e.values.get("voltage")
        if not sp:
            continue
        ords = sp["orders"]
        for b, key in (("pre_fault", "report.pre_fault_of"), ("post_fault", "report.post_fault_of")):
            if b not in sp:
                continue
            column = tr(key, name=nm)
            for m, seq in ((1, "pos"), (5, "neg"), (7, "pos"), (11, "neg"), (13, "pos")):
                if m in ords:
                    rows.setdefault("%d Hz (%s)" % (round(m * e.settings.f_nominal), tr(_SEQUENCE_LABELS[seq])),
                                    {})[column] = sp[b]["U_" + seq][ords.index(m)]
            k = e.tables["voltage_metrics"]; kk = k[k["interval"] == b]
            if not kk.empty:
                rows.setdefault(tr("report.thd_max_loop"), {})[column] = float(
                    kk[["UTHD12_pct", "UTHD23_pct", "UTHD31_pct"]].max(axis=1).iloc[0])
    if rows:
        tab["voltages"] = pd.DataFrame(rows).T.rename_axis(tr("report.quantity_in", unit="V")).reset_index()
    # ---- residual earth-fault current
    z = {}
    for nm, e in results.items():
        if "IRES" not in e.values:
            continue
        hz = _main_channel_harmonics(e)
        z.setdefault(tr("report.rms"), {})[nm] = e.values["IRES"]
        for m in sorted(set(orders) | {1}):
            if m in hz:
                z.setdefault("%d Hz" % round(m * e.settings.f_nominal), {})[nm] = hz[m]
        z.setdefault(tr("report.harmonic_content"), {})[nm] = e.values.get("IH_total", np.nan)
    if z:
        tab["residual_current"] = pd.DataFrame(z).T.rename_axis(tr("report.quantity_in", unit="A")).reset_index()
    # ---- split of the 50 Hz current
    z = {}
    for nm, e in results.items():
        if "Iw" in e.values:
            z.setdefault(tr("report.active_current"), {})[nm] = e.values["Iw"]
            z.setdefault(tr("report.detuning_current"), {})[nm] = e.values["Iv"]
        sp = e.tables.get("coil_currents")
        if sp is not None and not sp.empty:
            for _, r in sp.iterrows():
                z.setdefault(tr("report.coil_current", channel=r["channel"]), {})[nm] = r["H1_A"]
    if z:
        tab["split_50hz"] = pd.DataFrame(z).T.rename_axis(tr("report.quantity_in", unit="A")).reset_index()
    # ---- earth potential rise / shield current (stations)
    for nm, e in results.items():
        df = e.tables.get("station_voltages")
        if df is None or df.empty:
            continue
        ue_df = df[df["role"] == "u_e"] if (df["role"] == "u_e").any() else df
        r = ue_df.loc[ue_df["RMS_V"].idxmax()]      # "max. U_E measured" (as in earlier test reports)
        ords = [int(c[1:-2]) for c in df.columns if c.startswith("H") and c.endswith("_V")]
        zz = [dict(quantity=tr("report.rms"), UE_V=r["RMS_V"], rZE_V_per_100A=r.get("V_per_100A_RMS", np.nan))]
        for m in ords:
            zz.append(dict(quantity="%d Hz" % round(m * e.settings.f_nominal), UE_V=r["H%d_V" % m],
                           rZE_V_per_100A=100 * r.get("rZE_H%d_Ohm" % m, np.nan)))
        tab["epr_" + slug(nm)] = pd.DataFrame(zz).rename(
            columns={"UE_V": "U_E max. in V (%s)" % r["channel"], "rZE_V_per_100A": tr("report.rze_per_100_a")})
        ds = e.tables.get("station_currents")
        if ds is not None and not ds.empty:
            for _, s in ds[ds["role"] == "i_shield"].iterrows():
                zz = [dict(quantity=tr("report.rms"), I_S_A=s["RMS_A"])]
                for m in ords:
                    zz.append(dict(quantity="%d Hz" % round(m * e.settings.f_nominal), I_S_A=s.get("H%d_A" % m, np.nan),
                                   r_magnitude=s.get("r_H%d_magnitude" % m, np.nan),
                                   r_complex=s.get("r_H%d_complex" % m, np.nan)))
                tab["shield_current_" + slug(nm)] = pd.DataFrame(zz).rename(
                    columns={"I_S_A": "I_S in A", "r_magnitude": "1 − I_S/I_F", "r_complex": "|I_F − I_S|/|I_F|"})
    # ---- transformer
    z = []
    for nm, e in results.items():
        df = e.tables.get("transformer_power")
        if df is None or df.empty:
            continue
        for b, key in (("pre_fault", "report.label_pre_fault"), ("fault", "report.label_fault"),
                       ("post_fault", "report.label_post_fault")):
            for m in (5, 7):
                s = df[(df["interval"] == b) & (df["order"] == m)]
                if s.empty:
                    continue
                s = s.iloc[0]
                z.append({"measurement": "%s %s" % (nm, tr(key)), "f in Hz": round(m * e.settings.f_nominal),
                          "P in kW": s["P_kW"], "Q in kvar": s["Q_kvar"],
                          tr("report.i_negative_sequence"): s["I_neg_A"],
                          tr("report.i_positive_sequence"): s["I_pos_A"]})
    if z:
        tab["transformer"] = pd.DataFrame(z)
    # ---- transformer node: upstream impedance, polarity and power change at the fundamental per measurement
    z = {}
    for nm, e in results.items():
        v = e.values.get("transformer")
        if not v:
            continue
        z.setdefault(tr("report.tr_polarity"), {})[nm] = tr("report.tr_consumer" if (v.get("sign") or 1) > 0
                                                            else "report.tr_generator")
        z.setdefault(tr("report.tr_x1_up"), {})[nm] = _num(v.get("X1_up_Ohm"))
        z.setdefault(tr("report.tr_x1_count"), {})[nm] = v.get("X1_up_count") or 0
        if v.get("X1_transformer_Ohm") is not None:
            z.setdefault(tr("report.tr_x1_transformer"), {})[nm] = _num(v.get("X1_transformer_Ohm"))
            z.setdefault(tr("report.tr_x1_hv"), {})[nm] = _num(v.get("X1_hv_Ohm"))
        if v.get("dP1_fault_kW") is not None:
            z.setdefault(tr("report.tr_dp1"), {})[nm] = _num(v.get("dP1_fault_kW"))
            z.setdefault(tr("report.tr_dq1"), {})[nm] = _num(v.get("dQ1_fault_kvar"))
    if z:
        tab["transformer_summary"] = pd.DataFrame(z).T.rename_axis(tr("report.tr_quantity")).reset_index()
    # ---- assessment quantities of all measurements side by side
    z = {}
    for nm, e in results.items():
        df = e.tables.get("assessment_quantities")
        if df is None or df.empty:
            continue
        for _, r in df.iterrows():
            z.setdefault((r["quantity"], r["unit"]), {})[tr("report.measured_value_of", name=nm)] = r["measured_value"]
            z.setdefault((r["quantity"], r["unit"]), {})[tr("report.assessment_value_of", name=nm)] = r["assessment"]
    if z:
        d = pd.DataFrame(z).T
        d.index = pd.MultiIndex.from_tuples(d.index, names=["quantity", "unit"])
        tab["assessment_quantities"] = d.reset_index()
    # ---- equivalent sources
    if series is not None:
        t = series.source_table()
        cols = ["f_Hz", "sequence", "Rq_Ohm", "Xq_Ohm", "Uq_V", "residual_pct"]
        f_column = "f in Hz"
        e = t[cols].rename(columns={"f_Hz": f_column, "Rq_Ohm": "R_q in Ω", "Xq_Ohm": "X_q in Ω", "Uq_V": "U_q in V",
                                    "residual_pct": tr("report.residual_in_pct")})
        measured = (series.settings or {}).get("source") == "measured"
        if measured:                                       # source of each order (zero sequence: fitted)
            e.insert(2, tr("report.source_column"),
                     [tr("report.source_value_measured" if x == "measured" else "report.source_value_fitted")
                      for x in t["method"]])
        fitted = comparison if measured else series        # Thevenin check (expected ≈ 1∠0°) of the fitted source
        if fitted is not None:
            tf = fitted.source_table().set_index("order")
            if "Uq_ref_faulted_phase" in tf:
                label = tr("report.uq_fitted_to_faulted_phase" if measured else "report.uq_to_faulted_phase")
                e[label] = [tf["Uq_ref_faulted_phase"].get(m, np.nan) for m in t["order"]]
                e[tr("report.angle_in_deg")] = [tf["Uq_angle_ref_faulted_phase_deg"].get(m, np.nan)
                                                for m in t["order"]]
        if Z_E is not None:
            from .equivalent_source import ires_at_impedance
            r = ires_at_impedance(series, Z_E)
            rc = ires_at_impedance(comparison, Z_E) if comparison is not None else None
            if rc is None:
                i_column, c_column = tr("report.current_at_ze", impedance=number(Z_E)), None
            else:
                i_column, c_column = (tr("report.current_at_ze_source", impedance=number(Z_E), source=_source_label(x))
                                      for x in (series, comparison))
            e[i_column] = [r["parts"][m] for m in t["order"]]
            if rc is not None:
                e[c_column] = [rc["parts"].get(m, np.nan) for m in t["order"]]
            if r.get("IH_rest", 0) > 0:
                line = {c: np.nan for c in e.columns}
                line.update({f_column: tr("report.remainder"), "sequence": "–", i_column: r["IH_rest"]})
                if rc is not None:
                    line[c_column] = rc["IH_rest"]
                e = pd.concat([e.astype(object), pd.DataFrame([line])], ignore_index=True)
            total_ = {c: np.nan for c in e.columns}
            total_.update({f_column: "I_RES", "sequence": "", i_column: r["IRES"]})
            if rc is not None:
                total_[c_column] = rc["IRES"]
            e = pd.concat([e.astype(object), pd.DataFrame([total_])], ignore_index=True)
        tab["equivalent_sources"] = e
    return {name: _localize(df) for name, df in tab.items()}


def _main_channel_harmonics(e, role: str = "i_res") -> Dict[int, float]:
    """{order: magnitude} of the main channel of a role from the table ``harmonics`` (max window)."""
    h = e.tables.get("harmonics")
    if h is None or h.empty:
        return {}
    channel = e.values.get("main_channel_%s" % role)
    sub = h[(h["role"] == role) & ((h["channel"] == channel) if channel else True)]
    if channel is None and len(sub):
        sub = sub[sub["channel"] == sub["channel"].iloc[0]]
    return {int(r["order"]): float(r["magnitude"]) for _, r in sub.iterrows()}


def slug(name: str) -> str:
    """Name as part of file and table names: letters and digits, everything else ``_``, lower case.

    The campaign notebook and [`report_tables`][groundfault.report.report_tables] name per-measurement tables and
    figures with it, so that all names of a campaign export follow one rule.

    Examples
    --------
    >>> slug("M1 (2,5 Ω)")
    'm1_2_5'
    >>> slug("Measurement 4 (Station)")
    'measurement_4_station'
    """
    return re.sub(r"[^A-Za-z0-9]+", "_", str(name)).strip("_").lower()


# =========================================================================== complete export
def all_figures(res, folder: str, rec=None, prefix: str = "", formats: Optional[Sequence[str]] = None) -> List[str]:
    """Save the standard figures of a measurement (depending on the available channels) in ``folder``.

    Parameters
    ----------
    res : Result
        Evaluation result.
    folder : str
        Target folder.
    rec : Recording, optional
        The recording → also instantaneous values and RMS trend.
    prefix : str
        Prefix of the file names (default: the name of the measurement and ``_``).
    formats : sequence of str, optional
        File formats (default: ``STYLE.formats``).

    Returns
    -------
    list of str
        Paths of the written files.
    """
    paths = []
    prefix = prefix or (slug(res.name) + "_")

    def save_(fig, name):
        paths.extend(save(fig, os.path.join(folder, prefix + name), formats))
        plt.close(fig)

    roles = res.roles
    if rec is not None:
        t0 = res.t_max_window - 0.04
        save_(instantaneous(rec, roles, t0, 0.08), "instantaneous")
        chan = [r for r in ("i_res", "i_coil", "i_shield", "u_e") if roles.has(r)]
        if chan:
            save_(rms_trend(res, rec, channels=chan[:3]), "rms_trend")
    if roles.has("i_res"):
        save_(spectrum(res), "spectrum_fault_current")
        save_(harmonics_trend(res), "harmonics_trend")
        if "U0" in res.values:
            save_(phasor_diagram(res), "phasor_diagram")
    if "voltage" in res.values:
        save_(voltage_spectrum(res, "U_neg", (5, 11)), "voltage_negative_sequence")
        save_(voltage_spectrum(res, "U_pos", (7, 13)), "voltage_positive_sequence")
    if res.tables.get("station_voltages") is not None and not res.tables["station_voltages"].empty:
        save_(station_voltages(res), "station_voltages")
    if res.tables.get("station_frequency") is not None and not res.tables["station_frequency"].empty:
        save_(station_frequency(res), "station_frequency")
        save_(station_background(res), "station_background")
    if res.tables.get("station_lv_network") is not None and not res.tables["station_lv_network"].empty:
        save_(lv_network(res.tables["station_lv_network"]), "station_lv_network")
    if res.tables.get("transformer_power") is not None and not res.tables["transformer_power"].empty:
        save_(transformer_power(res), "transformer_power")
    if res.tables.get("transformer_impedance") is not None and not res.tables["transformer_impedance"].empty:
        save_(transformer_impedance(res), "transformer_impedance")
    return paths
