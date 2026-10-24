# -*- coding: utf-8 -*-
"""Word report of a campaign of staged earth-fault tests.

[`campaign_report`][groundfault.document.campaign_report] writes the evaluation of a campaign as a Word document
(``.docx``) in the structure of a test report:

1. **Set-up and procedure** – object of investigation, the tests (place, additional impedance, time and duration of
   every earth fault), the measured quantities, photo placeholders, the procedure.
2. **Results** – measurement uncertainty, processing of the recordings, residual earth-fault current (waveforms,
   spectrum, active and reactive part, coil currents), earth potential rise at stations (per frequency, earthing
   measurement, reduction factor), network voltages (power-quality measurement, voltages before and after each test).
3. **Assessment** – the residual current with the measurement-uncertainty margin and scaled order by order to the
   power-quality measurement, the equivalent-source method (active two-terminal network per order) and what it means
   for the forecast, the touch voltage at stations, the uncertainties, the summary.
4. **Appendices** – data of every recording, plausibility checks, symbols.

Sections without data are left out or say so (no station: no station sections; no power-quality measurement: a
sentence in its place). Every number is taken from the results and the objects of the campaign evaluation – the report
repeats no value by hand; the assessment values carry the margins m_I and m_U of the user in the single tests and at
the earthing impedance alike. Texts are German or English (the output language of [`groundfault.i18n`][groundfault.i18n]).

The package contains no names of networks, sites, organisations or persons and no limits or factors: they come from
[`ReportInfo`][groundfault.document.ReportInfo] and the arguments of the call (the limits and factors of the campaign
notebook). Photos taken on site are inserted where you name them; otherwise a framed placeholder with its caption
marks the place. Texts only you can write (the recommendation) are marked as placeholders. A Word file of your own
(``template=``) gives header, footer, logo and styles – e.g. your organisation's report template, made with
[`make_template`][groundfault.document.make_template] from an earlier report.

Requires ``python-docx`` (extra ``export``). The figures are those of [`groundfault.report`][groundfault.report];
with the raw recordings at hand (campaign folder) the waveforms and the trend of the fault current are added.
"""
from __future__ import annotations

import datetime as _dt
import functools
import math
import os
import pathlib
import re
import shutil
import subprocess
import tempfile
import warnings
from dataclasses import dataclass, field, fields
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence, Tuple, Union

import numpy as np
import pandas as pd

from . import __version__, i18n, report
from . import formulas as fml
from .i18n import tr

__all__ = ["ReportInfo", "PHOTO_SLOTS", "campaign_report", "touch_voltage_cases", "make_template",
           "to_pdf", "docx_available", "omml", "REFERENCE_TOLERANCE_HZ"]

#: photo slots of the report and the section they belong to (``photos={slot: path}``; None = placeholder)
PHOTO_SLOTS = ("setup_impedance", "recorder", "station")

i18n.register({
    # ------------------------------------------------------------------ general
    "document.no_docx": {
        "en": "The Word report needs python-docx: pip install python-docx (or groundfault[export]).",
        "de": "Der Word-Bericht braucht python-docx: pip install python-docx (oder groundfault[export])."},
    "document.no_results": {"en": "No results – nothing to report.", "de": "Keine Ergebnisse – nichts zu berichten."},
    "document.info_unknown": {
        "en": "unknown entries in the report information: {keys} (known: {known})",
        "de": "unbekannte Angaben zum Bericht: {keys} (bekannt: {known})"},
    "document.photo_unknown": {
        "en": "unknown photo slot {slot!r} (known: {known}; further photos: \"extra\" with a list)",
        "de": "unbekannter Fotoplatz {slot!r} (bekannt: {known}; weitere Fotos: „extra“ mit einer Liste)"},
    "document.photo_missing": {
        "en": "photo not found: {path} – placeholder used",
        "de": "Foto nicht gefunden: {path} – Platzhalter verwendet"},
    "document.template_missing": {"en": "template not found: {path}", "de": "Vorlage nicht gefunden: {path}"},
    "document.raw_missing": {
        "en": "{name}: recording {file} not found – no waveforms",
        "de": "{name}: Aufzeichnung {file} nicht gefunden – keine Momentanwerte"},
    "document.raw_failed": {
        "en": "{name}: recording not evaluated again ({error}) – no waveforms",
        "de": "{name}: Aufzeichnung nicht erneut ausgewertet ({error}) – keine Momentanwerte"},
    "document.figure_failed": {
        "en": "figure {name} left out: {error}", "de": "Abbildung {name} ausgelassen: {error}"},
    "document.written": {"en": "Report written: {path}", "de": "Bericht geschrieben: {path}"},
    "document.kept_edited": {
        "en": "The existing report had been edited after it was written – kept as {name}.",
        "de": "Der vorhandene Bericht war nach dem Erstellen bearbeitet – aufbewahrt als {name}."},
    "document.factor_large": {
        "en": "{name}: scaling factor up to {factor} (orders {orders}) – check the voltage levels and the THD of the test",
        "de": "{name}: Skalierungsfaktor bis {factor} (Ordnungen {orders}) – Spannungspegel und THD des Versuchs prüfen"},
    "document.pdf_none": {
        "en": "No PDF: LibreOffice (soffice) not found – save the report as PDF in Word.",
        "de": "Kein PDF: LibreOffice (soffice) nicht gefunden – den Bericht in Word als PDF speichern."},
    "document.pdf_failed": {"en": "No PDF: {error}", "de": "Kein PDF: {error}"},
    # ------------------------------------------------------------------ labels
    "document.figure": {"en": "Figure", "de": "Abbildung"},
    "document.table": {"en": "Table", "de": "Tabelle"},
    "document.toc": {"en": "Contents", "de": "Inhalt"},
    "document.toc_update": {
        "en": "(Contents – update in Word: right click → Update field)",
        "de": "(Inhaltsverzeichnis – in Word aktualisieren: Rechtsklick → Felder aktualisieren)"},
    "document.appendix": {"en": "Appendix {letter} – {title}", "de": "Anhang {letter} – {title}"},
    "document.placeholder_photo": {"en": "[Insert photo: {caption}]", "de": "[Foto einfügen: {caption}]"},
    "document.placeholder_text": {"en": "[To be completed: {what}]", "de": "[Bitte ergänzen: {what}]"},
    "document.yes": {"en": "yes", "de": "ja"},
    "document.no": {"en": "no", "de": "nein"},
    "document.met": {"en": "met", "de": "eingehalten"},
    "document.exceeded": {"en": "exceeded", "de": "überschritten"},
    "document.valid": {"en": "valid", "de": "gültig"},
    "document.not_valid": {"en": "not valid", "de": "nicht gültig"},
    "document.inside": {"en": "inside", "de": "innerhalb"},
    "document.outside": {"en": "outside (extrapolation)", "de": "außerhalb (Extrapolation)"},
    "document.in_assessment": {"en": "in the assessment value", "de": "im Bewertungswert"},
    "document.informative": {"en": "for information", "de": "informativ"},
    # ------------------------------------------------------------------ title page
    "document.default_title": {"en": "Evaluation of staged earth-fault tests", "de": "Auswertung Erdschlussmessung"},
    "document.carried_out": {"en": "Carried out on {date} at {places}", "de": "Durchführung am {date}: {places}"},
    "document.further_info": {"en": "Further information", "de": "Weitere Informationen"},
    "document.date_label": {"en": "Date:", "de": "Datum:"},
    "document.authors_label": {"en": "Prepared by:", "de": "Bearbeitung:"},
    "document.version_label": {"en": "Version:", "de": "Stand:"},
    "document.software_line": {
        "en": "Evaluated with groundfault {version}", "de": "Ausgewertet mit groundfault {version}"},
    # ------------------------------------------------------------------ key results
    "document.key_results": {"en": "Summary of results", "de": "Kurzfassung"},
    "document.key_results_intro": {
        "en": "Characteristic values of the tests and of the assessment (details in the section given). Limits and "
              "factors are those of the network operator as set for this evaluation.",
        "de": "Kennwerte der Versuche und der Bewertung (Einzelheiten im angegebenen Abschnitt). Grenzwerte und "
              "Faktoren sind die für diese Auswertung gesetzten Vorgaben des Netzbetreibers."},
    "document.kr_quantity": {"en": "Quantity", "de": "Größe"},
    "document.kr_value": {"en": "Value", "de": "Wert"},
    "document.kr_reference": {"en": "Limit / reference", "de": "Grenzwert / Bezug"},
    "document.kr_result": {"en": "Result", "de": "Ergebnis"},
    "document.kr_section": {"en": "Section", "de": "Abschnitt"},
    "document.kr_ice": {"en": "Capacitive earth-fault current I_[CE]", "de": "Kapazitiver Erdschlussstrom I_[CE]"},
    "document.kr_damping": {"en": "Damping d = I_[w]/I_[CE]", "de": "Dämpfung d = I_[w]/I_[CE]"},
    "document.kr_detuning": {"en": "Detuning v = I_[v]/I_[CE]", "de": "Verstimmung v = I_[v]/I_[CE]"},
    "document.kr_ires": {"en": "Residual current I_[RES] measured", "de": "Erdschlussreststrom I_[RES] gemessen"},
    "document.kr_ires_margin": {
        "en": "… with measurement-uncertainty margin", "de": "… mit Messunsicherheitszuschlag"},
    "document.kr_ires_scaled": {
        "en": "I_[RES,ass] scaled to the PQ measurement", "de": "I_[RES,Bew] skaliert auf die PQ-Messung"},
    "document.kr_ires_ze": {
        "en": "I_[RES,ass] at Z_[E] = {Z} Ω (equivalent sources)",
        "de": "I_[RES,Bew] an Z_[E] = {Z} Ω (Ersatzquellen)"},
    "document.kr_ut": {"en": "Touch voltage U_[T] ({station}), largest approach: {case}",
                       "de": "Berührungsspannung U_[T] ({station}), größter Ansatz: {case}"},
    "document.kr_decisive": {"en": " – decisive", "de": " – maßgeblich"},
    "document.probe": {"en": "probe", "de": "Sonde"},
    "document.kr_pq": {"en": "Power-quality measurement", "de": "PQ-Messung"},
    "document.kr_pq_value": {
        "en": "{days} days gapless, {q} % quantile, factor {k}", "de": "{days} Tage lückenlos, {q}-%-Quantil, Faktor {k}"},
    "document.kr_pq_value_periods": {
        "en": "{covered} days in {n} periods ({days} days gapless), {q} % quantile, factor {k}",
        "de": "{covered} Tage in {n} Abschnitten ({days} Tage lückenlos), {q}-%-Quantil, Faktor {k}"},
    "document.kr_tests": {"en": "Earth-fault tests", "de": "Erdschlussversuche"},
    "document.kr_tests_value": {
        "en": "{n_sub} at the substation{stations}", "de": "{n_sub} im Umspannwerk{stations}"},
    "document.kr_tests_station": {"en": ", 1 at a station", "de": ", 1 an einer Station"},
    "document.kr_tests_stations": {"en": ", {n} at stations", "de": ", {n} an Stationen"},
    # ------------------------------------------------------------------ chapter 1
    "document.ch_setup": {"en": "Set-up and procedure", "de": "Aufbau und Durchführung"},
    "document.sec_object": {
        "en": "Object of investigation and measurement set-up", "de": "Untersuchungsgegenstand und Messaufbau"},
    "document.object_intro": {
        "en": "Measuring points of the earth-fault tests: {places}.", "de": "Messorte der Erdschlussmessungen: {places}."},
    "document.normal_state": {
        "en": "The tests were made in the normal switching state of the network; no special switching states were set "
              "up.",
        "de": "Gemessen wurde im Normalschaltzustand des Netzes; Sonderschaltzustände wurden nicht hergestellt."},
    "document.coil_unchanged": {
        "en": "The position of the arc-suppression coils was not changed deliberately during the tests, i.e. no "
              "resonance curve (V curve) was recorded.",
        "de": "Während der Messung wurde die Stellung der Erdschlusslöschspulen nicht gezielt verändert, d. h. keine "
              "Resonanzkurve (V-Kurve) aufgenommen."},
    "document.auto_tuning": {
        "en": "Before each fault the coils were tuned by the automatic coil controller, as in normal operation.",
        "de": "Die Spulenstellung wurde vor dem Erdschluss wie im Netzbetrieb durch die automatische Abstimmung des "
              "Spulenreglers bestimmt."},
    "document.n_tests": {"en": "In total, {n} earth faults were applied:", "de": "Insgesamt wurden {n} Erdschlüsse erzeugt:"},
    "document.tab_tests": {"en": "Staged earth-fault tests", "de": "Durchgeführte Erdschlussversuche"},
    "document.col_test": {"en": "Test", "de": "Bezeichnung"},
    "document.col_place": {"en": "Fault location", "de": "Erdschlussort"},
    "document.col_z": {"en": "Additional impedance in Ω", "de": "Zusatzimpedanz in Ω"},
    "document.col_start": {"en": "Fault inception ({tz})", "de": "Fehlereintritt ({tz})"},
    "document.col_duration": {"en": "Duration in s", "de": "Dauer in s"},
    "document.purpose_default": {
        "en": "The different additional impedances show the influence of the earthing or fault impedance on the "
              "residual earth-fault current{station}.",
        "de": "Mit den verschiedenen Zusatzimpedanzen wird der Einfluss der Erdungs- bzw. Fehlerimpedanz auf den "
              "Erdschlussreststrom bestimmt{station}."},
    "document.purpose_station": {
        "en": "; the test at a station shows the conditions at an installation in the network",
        "de": "; die Messung an einer Station zeigt die Verhältnisse an einer Anlage im Netz"},
    "document.photo_setup_impedance": {
        "en": "Connection of the additional impedance in the fault circuit",
        "de": "Einbindung der Zusatzimpedanz in den Erdschlusskreis"},
    "document.photo_recorder": {
        "en": "Recording set-up with transient recorder and current clamps in the switchgear",
        "de": "Messwertaufzeichnung mit Transientenrekorder und Stromzangen in der Schaltanlage"},
    "document.photo_station": {
        "en": "Measurement set-up – {station} (probes, fault current, shield current)",
        "de": "Messaufbau – {station} (Sonden, Fehlerstrom, Schirmstrom)"},
    "document.recorder_text": {
        "en": "The tests were recorded with {instrument} at a sampling rate of {fs} kHz. The measured quantities are "
              "listed in the following table.",
        "de": "Aufgezeichnet wurde mit {instrument} bei einer Abtastrate von {fs} kHz. Die Messgrößen zeigt die "
              "folgende Tabelle."},
    "document.instrument_default": {"en": "a transient recorder", "de": "einem Transientenrekorder"},
    "document.tab_quantities": {"en": "Measured quantities", "de": "Messgrößen der Erdschlussmessung"},
    "document.col_quantity": {"en": "Quantity", "de": "Netzgröße"},
    "document.col_channels": {"en": "Channels", "de": "Kanäle"},
    "document.col_location": {"en": "Measuring point", "de": "Messort"},
    "document.col_transducer": {"en": "Transducer", "de": "Messwandler"},
    "document.sec_procedure": {"en": "Procedure", "de": "Durchführung"},
    "document.procedure_intro": {
        "en": "Each earth-fault test consisted of the following steps:",
        "de": "Die Erdschlussmessung gliedert sich in die folgenden Schritte:"},
    "document.step_tuning": {"en": "coils tuned to the set detuning", "de": "Spulenregler in Sollverstimmung"},
    "document.step_on": {
        "en": "earth fault switched on by remote control{duration}",
        "de": "Einschalten des Erdschlusses über die Fernsteuerung{duration}"},
    "document.step_on_duration": {"en": " for {d} s", "de": " für {d} s"},
    "document.step_on_durations": {"en": " for {lo} … {hi} s", "de": " für {lo} … {hi} s"},
    "document.step_off": {"en": "earth fault switched off by remote control", "de": "Ausschalten des Erdschlusses über die Fernsteuerung"},
    "document.step_impedance": {"en": "additional impedance changed, if planned", "de": "ggf. Anpassen der Zusatzimpedanz"},
    "document.corrections_intro": {
        "en": "Corrections of the measurement chain found after the tests (applied to every evaluation):",
        "de": "Nach der Messung festgestellte Korrekturen der Messkette (in jeder Auswertung angewendet):"},
    "document.tab_corrections": {"en": "Corrections of the measurement chain", "de": "Korrekturen der Messkette"},
    "document.factor_total": {"en": "{f} (total {t})", "de": "{f} (zusammen {t})"},
    "document.polarity_rotated": {
        "en": "In addition, the phasor of the fault current was turned by 180° in {tests}: the active current was "
              "negative, i.e. the current measurement reversed (appendix B).",
        "de": "Zusätzlich wurde in {tests} der Zeiger des Fehlerstroms um 180° gedreht: Der Wirkstrom war negativ, die "
              "Strommessung also verpolt (Anhang B)."},
    "document.col_files": {"en": "Recordings", "de": "Aufzeichnungen"},
    "document.col_channel": {"en": "Channel", "de": "Kanal"},
    "document.col_factor": {"en": "Factor", "de": "Faktor"},
    "document.col_remark": {"en": "Reason", "de": "Begründung"},
    # ------------------------------------------------------------------ chapter 2
    "document.ch_results": {"en": "Results", "de": "Ergebnisse"},
    "document.sec_general": {"en": "General", "de": "Allgemeines zur Messung"},
    "document.sec_uncertainty": {"en": "Measurement uncertainty", "de": "Messunsicherheiten"},
    "document.uncertainty_components": {
        "en": "The measurement uncertainty is estimated from the accuracies[^From the manuals and type "
              "specifications of the devices.^] of the devices of each measurement chain:",
        "de": "Die Messunsicherheit lässt sich aus den Fehlergrenzen[^Nach Handbuch bzw. Typenspezifikation der "
              "Geräte.^] der Geräte der jeweiligen Messkette abschätzen:"},
    "document.tab_uncertainty_components": {"en": "Accuracy of the measurement chain", "de": "Fehlergrenzen der Messkette"},
    "document.col_component": {"en": "Component", "de": "Komponente"},
    "document.col_accuracy": {"en": "Accuracy in %", "de": "Fehlergrenze in %"},
    "document.uncertainty_sum": {
        "en": "Sum of the limits of all components {lin} %, root sum of squares {rss} %. The limits apply at the rated "
              "frequency; for the harmonics the frequency response of the transducers and clamps limits the accuracy "
              "above all.",
        "de": "Summe der Fehlergrenzen aller Komponenten {lin} %, quadratische Summe {rss} %. Die Fehlergrenzen gelten "
              "bei Nennfrequenz; für die Oberschwingungen begrenzt vor allem der Frequenzgang der Wandler und Zangen "
              "die Genauigkeit."},
    "document.uncertainty_margin": {
        "en": "For the assessment, a margin of m_[I] = {mi} % is added to the currents{mu}.",
        "de": "Für die Bewertung wird auf die Ströme ein Zuschlag von m_[I] = {mi} % angesetzt{mu}."},
    "document.uncertainty_margin_u": {
        "en": " and the measured harmonic voltages are reduced by m_[U] = {mu} % (a larger scaling factor)",
        "de": " und die gemessenen Oberschwingungsspannungen werden um m_[U] = {mu} % verringert (größerer Skalierungsfaktor)"},
    "document.uncertainty_none": {
        "en": "No margin for the measurement uncertainty is set (m_[I] = 0).", "de": "Ein Zuschlag für die Messunsicherheit ist nicht gesetzt (m_[I] = 0)."},
    "document.sampling_text": {
        "en": "The sampling rate of {fs} kHz (Nyquist frequency {ny} kHz) covers the harmonics up to the {order}th "
              "order ({f} Hz) with a large margin{df}.",
        "de": "Die Abtastrate von {fs} kHz (Nyquist-Frequenz {ny} kHz) erfasst die Harmonischen bis zur {order}. "
              "Ordnung ({f} Hz) mit großem Abstand{df}."},
    "document.sampling_df": {
        "en": "; the frequency resolution of {df} Hz follows from the window length",
        "de": "; die Frequenzauflösung von {df} Hz folgt aus der Fensterlänge"},
    "document.sec_processing": {"en": "Processing of the recordings", "de": "Messwert-Aufbereitung"},
    "document.processing_intro": {
        "en": "The recordings were processed with the Python package groundfault {version}[^Open-source software "
              "(MIT licence): {url}^]:",
        "de": "Die Aufzeichnungen wurden mit dem Python-Paket groundfault {version}[^Open-Source-Software "
              "(MIT-Lizenz): {url}^] aufbereitet:"},
    "document.proc_read": {
        "en": "reading the recordings (COMTRADE[^Format of transient recordings, IEEE C37.111 and IEC 60255-24.^]){corr}",
        "de": "Einlesen der Rohdaten (COMTRADE[^Datenformat für Störschriebe nach IEEE C37.111 bzw. "
              "IEC 60255-24.^]){corr}"},
    "document.proc_read_corr": {"en": " with the corrections of the measurement chain", "de": " mit den Korrekturen der Messkette"},
    "document.proc_dc": {
        "en": "removal of the DC component per window ({method})", "de": "Entfernen des Gleichanteils je Fenster ({method})"},
    "document.proc_f1": {
        "en": "fundamental frequency from the line-to-line voltage ({f1} Hz)",
        "de": "Grundfrequenz aus der Leiter-Leiter-Spannung ({f1} Hz)"},
    "document.proc_f1_range": {
        "en": "fundamental frequency from the line-to-line voltage ({lo} … {hi} Hz, per recording in appendix A)",
        "de": "Grundfrequenz aus der Leiter-Leiter-Spannung ({lo} … {hi} Hz, je Aufzeichnung in Anhang A)"},
    "document.dc_median": {"en": "median", "de": "Median"},
    "document.dc_mean": {"en": "mean", "de": "Mittelwert"},
    "document.window_hann": {"en": "Hann", "de": "Hann"},
    "document.window_hanning": {"en": "Hann", "de": "Hann"},
    "document.window_rect": {"en": "rectangular", "de": "Rechteck"},
    "document.window_rectangular": {"en": "rectangular", "de": "Rechteck"},
    "document.window_boxcar": {"en": "rectangular", "de": "Rechteck"},
    "document.proc_dft": {
        "en": "harmonics with a DFT in windows of {T} ms ({n} periods, {window} window) up to the {order}th order",
        "de": "Harmonische mit einer DFT in Fenstern von {T} ms ({n} Perioden, {window}-Fenster) bis zur {order}. Ordnung"},
    "document.proc_intervals": {
        "en": "intervals before, during and after the fault set per recording{fig}; in the steady state the "
              "window with the largest residual current is assessed",
        "de": "Bereiche vor, im und nach dem Fehler je Aufzeichnung festgelegt{fig}; im stationären Bereich "
              "wird das Fenster mit dem größten Erdschlussreststrom bewertet"},
    "document.proc_intervals_fig": {
        "en": " ({name} in the following figure)", "de": " ({name} in der folgenden Abbildung)"},
    "document.proc_phasors": {
        "en": "phasors referred to the line-to-line voltage U_[L1L2]; symmetrical components of the voltages",
        "de": "Zeiger bezogen auf die Leiter-Leiter-Spannung U_[L1L2]; symmetrische Komponenten der Spannungen"},
    "document.fig_rms_trend": {
        "en": "RMS values of {name} with the intervals of the evaluation", "de": "Effektivwerte von {name} mit den Bereichen der Auswertung"},
    "document.sec_fault_current": {"en": "Residual earth-fault current", "de": "Fehlerstrom"},
    "document.sec_waveforms": {"en": "Waveforms", "de": "Qualitative Fehlerstromverläufe"},
    "document.waveforms_intro": {
        "en": "Five periods of the fault current in the assessed window of each test show the influence of the "
              "additional impedance on the waveform.",
        "de": "Je Versuch fünf Perioden des Fehlerstroms im bewerteten Fenster zeigen den Einfluss der Zusatzimpedanz "
              "auf den Kurvenverlauf."},
    "document.fig_waveform": {"en": "Fault current – {name}", "de": "Fehlerstromverlauf – {name}"},
    "document.waveforms_dc_free": {
        "en": "The waveforms are shown without the DC component, as in the evaluation{items}.",
        "de": "Die Verläufe sind ohne Gleichanteil dargestellt, wie in der Auswertung{items}."},
    "document.waveforms_dc_items": {"en": "; removed in the excerpt: {items}", "de": "; im Ausschnitt entfernt: {items}"},
    "document.waveforms_missing": {
        "en": "The raw recordings were not available for this report; the waveforms are missing.",
        "de": "Die Rohdaten lagen für diesen Bericht nicht vor; die Kurvenverläufe fehlen."},
    "document.sec_frequency": {"en": "Values in the frequency domain", "de": "Messwerte im Frequenzbereich"},
    "document.frequency_intro": {
        "en": "Residual earth-fault current of every test: RMS value, fundamental and all harmonic orders that reach "
              "{threshold} A in at least one test (up to the {order}th order, {f} Hz). The harmonic content I_[H] is "
              "everything above the fundamental – the orders ≥ 2 and the content between the orders – without the "
              "low-frequency part below 50 Hz; THD_[I] = I_[H]/I_[G].",
        "de": "Erdschlussreststrom aller Versuche: Effektivwert, Grundschwingung und alle Ordnungen, die in mindestens "
              "einem Versuch {threshold} A erreichen (bis zur {order}. Ordnung, {f} Hz). Der Oberschwingungsanteil "
              "I_[H] umfasst alles oberhalb der Grundschwingung – die Ordnungen ≥ 2 und die Anteile zwischen den "
              "Ordnungen –, ohne den niederfrequenten Anteil unter 50 Hz; THD_[I] = I_[H]/I_[G]."},
    "document.tab_ires": {"en": "Residual earth-fault current per order in A", "de": "Harmonische Komponenten des Fehlerstroms in A"},
    "document.row_rms": {"en": "RMS value I_[RES]", "de": "Effektivwert I_[RES]"},
    "document.row_harmonic_content": {"en": "Harmonic content I_[H]", "de": "Oberschwingungsanteil I_[H]"},
    "document.row_low_frequency": {"en": "Low-frequency part (< 50 Hz)", "de": "Niederfrequenter Anteil (< 50 Hz)"},
    "document.row_thd": {"en": "THD_[I] in %", "de": "THD_[I] in %"},
    "document.row_order": {"en": "{f} Hz ({m})", "de": "{f} Hz ({m}.)"},
    "document.category_pos_neg": {"en": "pos./neg.", "de": "Mit/Gegen"},
    "document.category_zero": {"en": "zero", "de": "Null"},
    "document.category_even": {"en": "even", "de": "gerade"},
    "document.category_fundamental": {"en": "fundamental", "de": "Grund"},
    "document.col_category": {"en": "System", "de": "System"},
    "document.fig_spectrum": {"en": "Harmonic components of the residual current", "de": "Harmonische Komponenten des Erdschlussreststroms"},
    "document.split_intro": {
        "en": "The fundamental consists of an active and a reactive (detuning) part, split with the angle between "
              "the fundamental of the residual current and the zero-sequence voltage:",
        "de": "Die Grundschwingungskomponente setzt sich aus einem Wirk- und einem Blindanteil (Verstimmung) "
              "zusammen; die Aufteilung erfolgt mit dem Winkel zwischen der Grundschwingung des Fehlerstroms und der "
              "Nullspannung:"},
    "document.eq_iw": {"en": "I_[w] = I_[G]·cos φ(U_[0], I_[G]),   I_[v] = I_[G]·sin φ(U_[0], I_[G])",
                       "de": "I_[w] = I_[G]·cos φ(U_[0], I_[G]),   I_[v] = I_[G]·sin φ(U_[0], I_[G])"},
    "document.tab_split": {"en": "Active and reactive part of the fundamental, coil currents", "de": "Aufteilung der Grundschwingung in Wirk- und Blindanteil, Spulenströme"},
    "document.row_ig": {"en": "Fundamental I_[G] in A", "de": "Grundschwingung I_[G] in A"},
    "document.row_iw": {"en": "Active current I_[w] in A", "de": "Wirkstrom I_[w] in A"},
    "document.row_iv": {"en": "Reactive current I_[v] in A", "de": "Blindstrom I_[v] in A"},
    "document.row_coil": {"en": "Coil current {channel} in A", "de": "Spulenstrom {channel} in A"},
    "document.row_coil_sum": {"en": "Sum of the coil currents in A", "de": "Summe der Spulenströme in A"},
    "document.row_damping": {"en": "Damping d in %", "de": "Dämpfung d in %"},
    "document.row_detuning": {"en": "Detuning v in %", "de": "Verstimmung v in %"},
    "document.ice_text": {
        "en": "The capacitive earth-fault current I_[CE] is taken as {source}; d = I_[w]/I_[CE], v = I_[v]/I_[CE]{others}.",
        "de": "Als kapazitiver Erdschlussstrom I_[CE] gilt {source}; d = I_[w]/I_[CE], v = I_[v]/I_[CE]{others}."},
    "document.ice_from_display": {"en": "the display of the coil controller", "de": "die Anzeige des Spulenreglers"},
    "document.ice_from_controller": {"en": "the value computed from the display of the coil controller, I_[L]/(1 + v)",
                                     "de": "der aus der Anzeige des Spulenreglers berechnete Wert I_[L]/(1 + v)"},
    "document.sum_ice_controller": {
        "en": "computed from the display of the coil controller {display} A, the measured coil currents add up to {coils} A",
        "de": "aus der Anzeige des Spulenreglers berechnet {display} A, Summe der gemessenen Spulenströme {coils} A"},
    "document.ice_from_coils": {"en": "the sum of the measured coil currents", "de": "die Summe der gemessenen Spulenströme"},
    "document.ice_from_coils_long": {
        "en": "approximately the sum of the measured coil currents (median of the tests; exact at resonance only – "
              "the coil current differs from I_[CE] by the reactive current I_[v], here {v} %)",
        "de": "näherungsweise die Summe der gemessenen Spulenströme (Median der Versuche; genau nur bei Resonanz – der "
              "Spulenstrom weicht um den Blindstrom I_[v] von I_[CE] ab, hier {v} %)"},
    "document.ice_others": {"en": " (also for {names}, without coil currents)",
                            "de": " (auch für {names}, ohne Spulenströme)"},
    # ---- station
    "document.sec_station": {"en": "Earth potential rise – {station}", "de": "Erdungsspannung – {station}"},
    "document.sec_station_waveforms": {"en": "Waveforms", "de": "Qualitative Verläufe"},
    "document.station_waveforms_intro": {
        "en": "Excerpt of the measured values without the DC component: {items}.",
        "de": "Ausschnitt der Messwerte ohne Gleichanteil: {items}."},
    "document.wf_u_e": {"en": "earth potential rise", "de": "Erdungsspannung"},
    "document.wf_i_res": {"en": "fault current", "de": "Fehlerstrom"},
    "document.wf_i_shield": {"en": "shield current", "de": "Schirmstrom"},
    "document.fig_station_waveform": {"en": "Earth potential rise and currents – {station}{panels}",
                                      "de": "Erdungsspannung und Ströme – {station}{panels}"},
    "document.sec_station_frequency": {"en": "Values in the frequency domain", "de": "Messwerte im Frequenzbereich"},
    "document.station_probes": {
        "en": "The earth potential rise U_[E] was measured against the earth probe {probes}.",
        "de": "Die Erdungsspannung U_[E] wurde gegen den Erdungsspieß {probes} gemessen."},
    "document.station_probes_many": {
        "en": "The earth potential rise U_[E] was measured against the earth probes {probes}; the table shows the one "
              "with the largest value.",
        "de": "Die Erdungsspannung U_[E] wurde gegen die Erdungsspieße {probes} gemessen; die Tabelle zeigt den mit dem "
              "größten Wert."},
    "document.tab_epr": {
        "en": "Earth potential rise per order ({probe}): value during the fault and r·Z_[E] per 100 A",
        "de": "Harmonische Komponenten der Erdungsspannung ({probe}): Messwert im Fehler und r·Z_[E] hochgerechnet auf "
              "100 A"},
    "document.col_ue_max": {"en": "U_[E] during the fault in V", "de": "U_[E] im Fehler in V"},
    "document.col_per_100": {"en": "r·Z_[E]·100 A direct in V", "de": "r·Z_[E]·100 A direkt in V"},
    "document.col_per_100_fault": {"en": "r·Z_[E]·100 A fault-related in V", "de": "r·Z_[E]·100 A fehlerbedingt in V"},
    "document.col_if": {"en": "I_[F] in A", "de": "I_[F] in A"},
    "document.row_rms_plain": {"en": "RMS value", "de": "Effektivwert"},
    "document.epr_note": {
        "en": "Direct: voltage during the fault per ampere of fault current, including the voltage present before "
              "the fault; fault-related: difference of the phasors during and before the fault (superposition) – "
              "the value of the assessment (case 3 of the touch voltage){uncertain}.",
        "de": "Direkt: Spannung im Fehler je Ampere Fehlerstrom, einschließlich der schon vor dem Fehler anliegenden "
              "Spannung; fehlerbedingt: Differenz der Zeiger im Fehler und davor (Überlagerungsverfahren) – der Wert "
              "der Bewertung (Fall 3 der Berührungsspannung){uncertain}."},
    "document.epr_uncertain": {
        "en": "; * uncertain (the superposition methods disagree or the voltage before the fault is large)",
        "de": "; * unsicher (die Überlagerungsverfahren weichen voneinander ab oder die Spannung vor dem Fehler ist "
              "groß)"},
    "document.probe_inverted": {
        "en": "The voltage against the probe {ut_probe} ({ut} V) is larger than the earth potential rise against "
              "{ue_probe} ({ue} V), whereas U_[T] ≤ U_[E] is expected. Either {ue_probe} does not capture the full "
              "earth potential rise – then U_[E] and the cases 3, 4 and 4b are too small – or {ut_probe} does not "
              "measure a touch voltage; the reduction factor therefore uses the probe maximum.",
        "de": "Die Spannung gegen den Spieß {ut_probe} ({ut} V) ist größer als die Erdungsspannung gegen {ue_probe} "
              "({ue} V), erwartet ist U_[T] ≤ U_[E]. Entweder erfasst {ue_probe} nicht die volle Erdungsspannung – dann "
              "sind U_[E] und die Fälle 3, 4 und 4b zu klein – oder {ut_probe} misst keine Berührungsspannung; der "
              "Reduktionsfaktor verwendet deshalb das Sondenmaximum."},
    "document.fig_station_frequency": {
        "en": "Voltage per ampere of fault current over the frequency – {station}",
        "de": "Spannung je Ampere Fehlerstrom über der Frequenz – {station}"},
    "document.station_comparison_intro": {
        "en": "The earthing measurement of the station[^Current injection{details}; groundmeas measurement {ids}.^] "
              "gives the earthing impedance Z_[E] per frequency. The earth-fault test gives r·Z_[E] as the probe "
              "maximum (per frequency the largest voltage of all probes{probe}) in its fault-related part; the ratio of "
              "the two is the reduction factor r.",
        "de": "Die Erdungsmessung der Station[^Stromeinspeisung{details}; groundmeas-Messung {ids}.^] liefert die "
              "Erdungsimpedanz Z_[E] je Frequenz. Die Erdschlussmessung liefert r·Z_[E] als Sondenmaximum (je Frequenz "
              "die größte Spannung aller Spieße{probe}) im fehlerbedingten Anteil; das Verhältnis beider ist der "
              "Reduktionsfaktor r."},
    "document.earthing_details": {"en": ", {date}", "de": ", {date}"},
    "document.earthing_distance": {"en": ", probe at {d} m", "de": ", Sonde in {d} m"},
    "document.probe_maximum_here": {"en": ", here {probe}", "de": ", hier {probe}"},
    "document.tab_earthing": {
        "en": "Earthing measurement and earth-fault test per frequency", "de": "Erdungsmessung und Erdschlussmessung je Frequenz"},
    "document.col_ze_injection": {"en": "Z_[E] earthing measurement in Ω", "de": "Z_[E] Erdungsmessung in Ω"},
    "document.col_rze_fault": {"en": "r·Z_[E] earth-fault test (probe maximum) in Ω",
                               "de": "r·Z_[E] Erdschlussmessung (Sondenmaximum) in Ω"},
    "document.col_r": {"en": "r", "de": "r"},
    "document.col_f": {"en": "f in Hz", "de": "f in Hz"},
    "document.fig_earthing": {
        "en": "Earth-fault test and earthing measurement – {station}", "de": "Gegenüberstellung Erdschlussmessung und Erdungsmessung – {station}"},
    "document.reduction_text": {
        "en": "Effective reduction factor (RMS value over the orders, r_[eff] = √Σ|ΔU_[E,m]|² / √Σ(Z_[E,m]·I_[F,m])²) "
              "r = {r} (at 50 Hz {r50}{classification}); share of the shield current in the fault current {shield} %.",
        "de": "Wirksamer Reduktionsfaktor (Effektivwert über die Ordnungen, r_[eff] = √Σ|ΔU_[E,m]|² / "
              "√Σ(Z_[E,m]·I_[F,m])²) r = {r} (bei 50 Hz {r50}{classification}); Anteil des Schirmstroms am Fehlerstrom "
              "{shield} %."},
    "document.reduction_class": {"en": ", {text}", "de": ", {text}"},
    "document.reduction_range": {"en": " {low} … {high}", "de": " {low} … {high}"},
    "document.earthing_hints": {"en": "Notes on the earthing:", "de": "Hinweise zur Erdung:"},
    "document.r_class_below": {"en": "below the usual range", "de": "unter dem üblichen Bereich"},
    "document.r_class_typical": {"en": "within the usual range", "de": "im üblichen Bereich"},
    "document.r_class_above": {"en": "above the usual range", "de": "über dem üblichen Bereich"},
    "document.r_class_implausible": {"en": "above 1 – not plausible", "de": "über 1 – nicht plausibel"},
    "document.fig_reduction": {"en": "Reduction factor per order – {station}", "de": "Reduktionsfaktor je Ordnung – {station}"},
    "document.lv_text": {
        "en": "Current into the LV network I_[ON] = {ion} A ({share} % of the fault current), earthing impedance of the "
              "LV network Z_[E,ON] = {zon} Ω.",
        "de": "Strom ins Ortsnetz I_[ON] = {ion} A ({share} % des Fehlerstroms), Erdungsimpedanz des Ortsnetzes "
              "Z_[E,ON] = {zon} Ω."},
    "document.fig_lv": {"en": "LV network – {station}", "de": "Ortsnetz – {station}"},
    "document.no_earthing": {
        "en": "No earthing measurement of the station was available (groundmeas); the comparison is missing.",
        "de": "Eine Erdungsmessung der Station lag nicht vor (groundmeas); der Vergleich fehlt."},
    # ---- voltages
    "document.sec_voltages": {"en": "Network voltages", "de": "Netzspannungen"},
    "document.sec_pq": {"en": "Power-quality measurement", "de": "Power-Quality-Messung"},
    "document.pq_intro": {
        "en": "In preparation of the tests, the voltage quality of the busbar was recorded{device}: {start} to {end} "
              "({span} days, {n} 10-min values, {complete} % complete). The longest stretch of gapless values covers "
              "{gapless} days.",
        "de": "In Vorbereitung der Erdschlussmessung wurde die Spannungsqualität der Sammelschiene aufgezeichnet"
              "{device}: {start} bis {end} ({span} Tage, {n} 10-min-Werte, {complete} % vollständig). Die längste "
              "Strecke lückenloser Werte umfasst {gapless} Tage."},
    "document.pq_device": {"en": " ({device})", "de": " ({device})"},
    "document.pq_intro_periods": {
        "en": "In preparation of the tests, the voltage quality of the busbar was recorded in {k} periods{device}: "
              "{items} – together {covered} days with {n} 10-min values. The longest stretch of gapless values covers "
              "{gapless} days.",
        "de": "In Vorbereitung der Erdschlussmessung wurde die Spannungsqualität der Sammelschiene in {k} Abschnitten "
              "aufgezeichnet{device}: {items} – zusammen {covered} Tage mit {n} 10-min-Werten. Die längste Strecke "
              "lückenloser Werte umfasst {gapless} Tage."},
    "document.pq_period_item": {"en": "{start} to {end} ({days} days)", "de": "{start} bis {end} ({days} Tage)"},
    "document.pq_periods": {
        "en": "Per period the quantile of the {order}th order is {values} V (whole measurement {whole} V) – the "
              "highest period lies {spread} % above the lowest{assessment}.",
        "de": "Je Abschnitt beträgt das Quantil der {order}. Ordnung {values} V (ganze Messung {whole} V) – der "
              "höchste Abschnitt liegt {spread} % über dem niedrigsten{assessment}."},
    "document.pq_periods_above": {
        "en": "; the assessment voltage (quantile × k_[s] = {ua} V) lies {pct} % above the highest period",
        "de": "; die Bewertungsspannung (Quantil × k_[s] = {ua} V) liegt {pct} % über dem höchsten Abschnitt"},
    "document.pq_periods_below": {
        "en": "; the assessment voltage (quantile × k_[s] = {ua} V) lies {pct} % below the highest period",
        "de": "; die Bewertungsspannung (Quantil × k_[s] = {ua} V) liegt {pct} % unter dem höchsten Abschnitt"},
    "document.pq_method": {
        "en": "Per 10-min interval[^Intervals aligned to the clock; shorter values of the export are combined per "
              "interval into their root mean square (as in IEC 61000-4-30).^] the largest of the three phases "
              "(harmonic voltage line-to-earth per order) and the largest loop (THD line-to-line) are taken; the "
              "assessment uses their {q} % quantile. Intervals with an earth fault or a dip (unbalanced fundamental) "
              "are left out{excluded}.",
        "de": "Je 10-min-Intervall[^Uhrzeitsynchrone Intervalle; aus kürzeren Werten des Exports wird je Intervall "
              "der quadratische Mittelwert gebildet (wie in IEC 61000-4-30).^] zählt der größte der drei Leiter "
              "(Oberschwingungsspannung Leiter-Erde je Ordnung) bzw. die größte Schleife (THD Leiter-Leiter); bewertet "
              "wird ihr {q}-%-Quantil. Intervalle mit Erdschluss oder Spannungseinbruch (unsymmetrische "
              "Grundschwingung) bleiben unberücksichtigt{excluded}."},
    "document.pq_excluded": {"en": " ({n} intervals)", "de": " ({n} Intervalle)"},
    "document.pq_excluded_none": {"en": " (none in this measurement)", "de": " (in dieser Messung keine)"},
    "document.pq_factor": {
        "en": "Factor on the quantile by the duration of the gapless values: k_[s] = {k} ({hint}).",
        "de": "Faktor auf das Quantil nach der Dauer der lückenlosen Werte: k_[s] = {k} ({hint})."},
    "document.pq_factor_rule": {
        "en": "{days} days gapless; rule: {rule}", "de": "{days} Tage lückenlos; Regel: {rule}"},
    "document.pq_rule_item": {"en": "≥ {days} days {k}", "de": "≥ {days} Tage {k}"},
    "document.pq_invalid": {
        "en": "The measurement is not valid for the assessment (gapless values shorter than every duration of the "
              "rule) – nothing is scaled with it.",
        "de": "Die Messung ist für die Bewertung nicht gültig (lückenlose Werte kürzer als jede Dauer der Regel) – mit "
              "ihr wird nichts skaliert."},
    "document.pq_dominant": {
        "en": "The {order}th order ({f} Hz) dominates the harmonic voltage – in the median interval it is {share} % "
              "of U_[H] = √(Σ U_[m]²); its quantile is {u} V (maximum {umax} V).",
        "de": "Die {order}. Ordnung ({f} Hz) dominiert die Oberschwingungsspannung – im Median der Intervalle {share} % "
              "von U_[H] = √(Σ U_[m]²); ihr Quantil beträgt {u} V (Maximum {umax} V)."},
    "document.tab_pq": {
        "en": "Quantiles of the power-quality measurement (largest phase per interval)",
        "de": "Quantile der Power-Quality-Messung (größter Leiter je Intervall)"},
    "document.col_order": {"en": "Order", "de": "Ordnung"},
    "document.col_u_quantile": {"en": "Quantile in V", "de": "Quantil in V"},
    "document.col_u_factored": {"en": "× factor in V", "de": "× Faktor in V"},
    "document.col_u_median": {"en": "Median in V", "de": "Median in V"},
    "document.col_u_max": {"en": "Maximum in V", "de": "Maximum in V"},
    "document.row_thd_ll": {"en": "THD line-to-line in %", "de": "THD Leiter-Leiter in %"},
    "document.pq_thd": {
        "en": "THD line-to-line (largest loop): quantile {q} %, times the factor {k} %.",
        "de": "THD Leiter-Leiter (größte Schleife): Quantil {q} %, mal Faktor {k} %."},
    "document.fig_pq_series": {
        "en": "10-min values of the power-quality measurement with the quantile and the quantile times the factor",
        "de": "10-min-Werte der Power-Quality-Messung mit dem Quantil und dem Quantil mal Faktor"},
    "document.fig_pq_duration": {
        "en": "Duration curve of the {order}th order with the pre-fault voltages of the tests",
        "de": "Dauerlinie der {order}. Ordnung mit den Vorfehlerspannungen der Versuche"},
    "document.fig_pq_profile": {"en": "Weekly profile of the {order}th order", "de": "Wochengang der {order}. Ordnung"},
    "document.pq_comparability": {
        "en": "Where the pre-fault voltages of the tests lie in the distribution of the measurement – over the whole "
              "measurement and in the same time window of the week. The scaling assumes the same network and switching "
              "state in the power-quality measurement and in the tests.",
        "de": "Lage der Vorfehlerspannungen der Versuche in der Verteilung der Messung – über die ganze Messung und im "
              "selben Zeitfenster der Woche. Die Skalierung setzt dasselbe Netz und denselben Schaltzustand in "
              "PQ-Messung und Versuchen voraus."},
    "document.tab_pq_tests": {"en": "Tests in the distribution of the power-quality measurement", "de": "Erdschlussversuche in der Verteilung der PQ-Messung"},
    "document.pq_tests_note": {
        "en": "U pre-fault: the largest of the three line-to-earth voltages of the order before the fault – the "
              "quantity of the power-quality measurement.{others}",
        "de": "U vor dem Fehler: die größte der drei Leiter-Erde-Spannungen der Ordnung vor dem Fehler – die Größe der "
              "PQ-Messung.{others}"},
    "document.pq_tests_others": {
        "en": " Not listed: {names} (voltages not measured at the busbar).",
        "de": " Nicht aufgeführt: {names} (Spannungen nicht an der Sammelschiene gemessen)."},
    "document.col_u_test": {"en": "U pre-fault in V", "de": "U vor dem Fehler in V"},
    "document.col_percentile": {"en": "Percentile (all)", "de": "Perzentil (gesamt)"},
    "document.col_window": {"en": "Time window", "de": "Zeitfenster"},
    "document.col_percentile_window": {"en": "Percentile (window)", "de": "Perzentil (Fenster)"},
    "document.col_factor_test": {"en": "Quantile/U_[test]", "de": "Quantil/U_[Versuch]"},
    "document.no_pq": {
        "en": "No power-quality measurement of the busbar was available for this evaluation.",
        "de": "Eine Power-Quality-Messung der Sammelschiene lag für diese Auswertung nicht vor."},
    "document.no_pq_given": {
        "en": "No file of the power-quality measurement of the busbar was available for this evaluation; the scaling "
              "takes the quantiles given in the evaluation: {values}, factor k_[s] = {k}.",
        "de": "Eine Datei der Power-Quality-Messung der Sammelschiene lag für diese Auswertung nicht vor; skaliert "
              "wird mit den in der Auswertung angegebenen Quantilen: {values}, Faktor k_[s] = {k}."},
    "document.given_u95": {"en": "U_[95,{m}] = {u} V", "de": "U_[95,{m}] = {u} V"},
    "document.given_thd": {"en": "THD_[95] = {thd} %", "de": "THD_[95] = {thd} %"},
    "document.given_ll": {"en": "line-to-line", "de": "Leiter-Leiter"},
    "document.given_seq": {"en": "sequence voltage", "de": "Sequenzspannung"},
    "document.sec_test_voltages": {"en": "Voltages of the earth-fault tests", "de": "Erdschlussmessung"},
    "document.test_voltages_intro": {
        "en": "The line-to-earth voltages were transformed into symmetrical components per order; shown is the "
              "sequence system of the natural order (positive for 1, 7, 13 …, negative for 5, 11, 17 …) before and "
              "after each test – the orders that reach 1 V – and the largest line-to-line THD. During the fault the "
              "components are coupled and are not shown.",
        "de": "Die Leiter-Erde-Spannungen wurden je Ordnung in symmetrische Komponenten transformiert; dargestellt ist "
              "das Sequenzsystem der natürlichen Folge (Mitsystem für 1, 7, 13 …, Gegensystem für 5, 11, 17 …) vor und "
              "nach jedem Versuch – die Ordnungen, die 1 V erreichen – sowie der größte Leiter-Leiter-THD. Während des "
              "Fehlers sind die Komponenten verkoppelt und werden nicht dargestellt."},
    "document.tab_voltages": {"en": "Voltages before and after the tests", "de": "Spannungen vor und nach den Versuchen"},
    "document.col_before": {"en": "before", "de": "vor"},
    "document.col_after": {"en": "after", "de": "nach"},
    "document.positive": {"en": "pos.", "de": "Mit"},
    "document.negative": {"en": "neg.", "de": "Gegen"},
    "document.zero": {"en": "zero", "de": "Null"},
    # ------------------------------------------------------------------ chapter 3
    "document.ch_assessment": {"en": "Assessment", "de": "Bewertung"},
    "document.sec_ires": {"en": "Processing of the residual earth-fault current", "de": "Aufbereitung Erdschlussreststrom"},
    "document.ires_intro": {
        "en": "For the assessment the residual current of the test with the largest assessment value, {name}, is "
              "processed in three steps.",
        "de": "Für die abschließende Bewertung wird der Erdschlussreststrom des Versuchs mit dem größten "
              "Bewertungswert, {name}, in drei Schritten aufbereitet."},
    "document.step1": {
        "en": "Step 1: margin for the measurement uncertainty, I_[ass] = I·(1 + m_[I]) with m_[I] = {mi} %.",
        "de": "Schritt 1: Zuschlag der Messunsicherheit, I_[Bew] = I·(1 + m_[I]) mit m_[I] = {mi} %."},
    "document.tab_step1": {"en": "Residual current with the measurement-uncertainty margin", "de": "Berücksichtigung der Messunsicherheit im Fehlerstrom"},
    "document.col_measured": {"en": "Measured in A", "de": "Originalmessung in A"},
    "document.col_with_margin": {"en": "With margin in A", "de": "Inkl. Messunsicherheit in A"},
    "document.step2": {
        "en": "Step 2: scaling of the harmonic currents to the voltage of the power-quality measurement – order by "
              "order. Each order of the positive or negative sequence is scaled with its own voltage, "
              "k_[m] = k_[s]·U_[95,m]/U_[m] (k_[s] = {k}): U_[95,m] is the quantile of the largest phase of the PQ "
              "measurement, U_[m] the magnitude of the positive- or negative-sequence component of the order in the "
              "test (mean of before and after the fault{u_ref}){safe}. An order without its own "
              "value is scaled with the THD factor. The orders of the zero sequence (3, 9, 15 …) keep their measured "
              "value – assumption: their current arises with the displacement in the earth fault and is not driven by "
              "the voltage before the fault; even orders are unusual for the voltage of a network and stay as "
              "measured as well.",
        "de": "Schritt 2: Skalierung der Oberschwingungsströme auf die Spannung der Power-Quality-Messung – je "
              "Ordnung. Jede Ordnung des Mit- oder Gegensystems wird mit ihrer eigenen Spannung skaliert, "
              "k_[m] = k_[s]·U_[95,m]/U_[m] (k_[s] = {k}): U_[95,m] ist das Quantil des größten Leiters der "
              "PQ-Messung, U_[m] der Betrag der Mit- bzw. Gegensystemkomponente der Ordnung im Versuch (Mittel aus vor "
              "und nach dem Fehler{u_ref}){safe}. Eine Ordnung ohne eigenen Wert "
              "wird mit dem THD-Faktor skaliert. Die Ordnungen des Nullsystems (3, 9, 15 …) behalten ihren Messwert – "
              "Annahme: Ihr Strom entsteht mit der Verlagerung im Erdschluss und wird nicht von der Spannung vor dem "
              "Fehler getrieben; gerade Ordnungen sind für die Spannung eines Netzes ungewöhnlich und bleiben "
              "ebenfalls wie gemessen."},
    "document.step2_thd_only": {
        "en": "Step 2: scaling of the harmonic currents to the voltage of the power-quality measurement. It gives the "
              "THD only, no voltage per order: each order of the positive or negative sequence is scaled with the THD "
              "factor k_[THD] (k_[s] = {k}). The orders of the zero sequence (3, 9, 15 …) keep their measured value – "
              "assumption: their current arises with the displacement in the earth fault and is not driven by the "
              "voltage before the fault; even orders are unusual for the voltage of a network and stay as measured as "
              "well.",
        "de": "Schritt 2: Skalierung der Oberschwingungsströme auf die Spannung der Power-Quality-Messung. Sie gibt nur "
              "den THD an, keine Spannung je Ordnung: Jede Ordnung des Mit- oder Gegensystems wird mit dem THD-Faktor "
              "k_[THD] skaliert (k_[s] = {k}). Die Ordnungen des Nullsystems (3, 9, 15 …) behalten ihren Messwert – "
              "Annahme: Ihr Strom entsteht mit der Verlagerung im Erdschluss und wird nicht von der Spannung vor dem "
              "Fehler getrieben; gerade Ordnungen sind für die Spannung eines Netzes ungewöhnlich und bleiben "
              "ebenfalls wie gemessen."},
    "document.step2_thd_value": {
        "en": "The THD factor is k_[THD] = k_[s]·THD_[95]/THD_[test] = {k}·{thd95} %/{thd} % = {f}: THD_[95] is the "
              "quantile of the largest line-to-line loop of the PQ measurement, THD_[test] the largest line-to-line "
              "THD of the test{u_ref}.",
        "de": "Der THD-Faktor ist k_[THD] = k_[s]·THD_[95]/THD_[Versuch] = {k}·{thd95} %/{thd} % = {f}: THD_[95] ist "
              "das Quantil der größten Leiter-Leiter-Schleife der PQ-Messung, THD_[Versuch] der größte "
              "Leiter-Leiter-THD des Versuchs{u_ref}."},
    "document.ref_table_paren": {"en": " (table {n})", "de": " (Tabelle {n})"},
    "document.ref_table": {"en": ", table {n}", "de": ", Tabelle {n}"},
    "document.step2_safe": {
        "en": "[^The component is never larger than the largest phase, so k_[m] is on the safe side compared with "
              "the ratio of the largest phases (quantile/U_[test], table {n}).^]",
        "de": "[^Die Komponente ist nie größer als der größte Leiter; k_[m] liegt daher gegenüber dem Verhältnis der "
              "größten Leiter (Quantil/U_[Versuch], Tabelle {n}) auf der sicheren Seite.^]"},
    "document.step2_condition": {
        "en": "The scaling with the voltage is admissible if the network (distribution of inductances and "
              "capacitances, switching state) is the same in the power-quality measurement and in the tests; the "
              "series resonance of network capacitance and transformer inductance can amplify or damp the voltage.",
        "de": "Die Skalierung mit der Oberschwingungsspannung ist zulässig, wenn sich die Netzsituation (Verteilung "
              "von Induktivitäten und Kapazitäten, Schaltzustand) zwischen PQ-Messung und Erdschlussmessung nicht "
              "unterscheidet; der Reihenschwingkreis aus Netzkapazität und Streuinduktivität des Transformators kann "
              "die Spannung verstärken oder dämpfen."},
    "document.tab_step2": {"en": "Scaling per order – {name}", "de": "Skalierung je Ordnung – {name}"},
    "document.col_i_measured": {"en": "I_[m] in A", "de": "I_[m] in A"},
    "document.col_level": {"en": "U_[m] test in V", "de": "U_[m] Versuch in V"},
    "document.col_u95": {"en": "U_[95,m] in V", "de": "U_[95,m] in V"},
    "document.col_k": {"en": "k_[m]", "de": "k_[m]"},
    "document.col_rule": {"en": "Rule", "de": "Regel"},
    "document.col_i_scaled": {"en": "I_[m,scaled] in A", "de": "I_[m,skaliert] in A"},
    "document.rule_own voltage": {"en": "own voltage", "de": "eigene Spannung"},
    "document.rule_THD": {"en": "THD factor", "de": "THD-Faktor"},
    "document.rule_as measured": {"en": "as measured", "de": "wie gemessen"},
    "document.rule_no information": {"en": "no information", "de": "keine Angabe"},
    "document.step2_rest": {
        "en": "Orders not listed (below {threshold} A): {rest} A together, scaled with the same rules{thd}. Not "
              "scaled: zero sequence {zero} A, even orders {even} A, content between the orders {other} A.",
        "de": "Nicht aufgeführte Ordnungen (unter {threshold} A): zusammen {rest} A, nach denselben Regeln "
              "skaliert{thd}. Nicht skaliert: Nullsystem {zero} A, gerade Ordnungen {even} A, Anteile zwischen den "
              "Ordnungen {other} A."},
    "document.step2_thd_orders": {"en": " (with the THD factor: orders {orders})",
                                  "de": " (mit dem THD-Faktor: Ordnungen {orders})"},
    "document.step2_none": {
        "en": "Step 2 is left out: no valid power-quality measurement and no quantiles were given.",
        "de": "Schritt 2 entfällt: keine gültige PQ-Messung und keine Quantile angegeben."},
    "document.eq_ih_scaled": {
        "en": "I_[H,scaled] = √(Σ (k_[m]·I_[m])² + I_[zero]² + I_[even]² + I_[other]²) = {ih} A",
        "de": "I_[H,skaliert] = √(Σ (k_[m]·I_[m])² + I_[Null]² + I_[gerade]² + I_[sonst]²) = {ih} A"},
    "document.eq_ih_assessed": {
        "en": "I_[H,ass] = (1 + m_[I])·I_[H,scaled] = {ih} A",
        "de": "I_[H,Bew] = (1 + m_[I])·I_[H,skaliert] = {ih} A"},
    "document.eq_ih_assessed_mu": {
        "en": "I_[H,ass] = (1 + m_[I])·√(Σ (k_[m]·I_[m]/(1 − m_[U]))² + I_[zero]² + I_[even]² + I_[other]²) = {ih} A",
        "de": "I_[H,Bew] = (1 + m_[I])·√(Σ (k_[m]·I_[m]/(1 − m_[U]))² + I_[Null]² + I_[gerade]² + I_[sonst]²) = {ih} A"},
    "document.step3": {
        "en": "Step 3: residual current with the scaled harmonic content; assessment values with the margin m_[I]:",
        "de": "Schritt 3: Hochrechnung des Erdschlussreststroms mit dem skalierten Oberschwingungsanteil; "
              "Bewertungswerte mit dem Zuschlag m_[I]:"},
    "document.step3_mu": {
        "en": "Step 3: residual current with the scaled harmonic content; assessment values with the margins m_[I] "
              "and – for the orders scaled with a measured voltage, which may be measured too high by m_[U] – "
              "m_[U] = {mu} %:",
        "de": "Schritt 3: Hochrechnung des Erdschlussreststroms mit dem skalierten Oberschwingungsanteil; "
              "Bewertungswerte mit den Zuschlägen m_[I] und – für die mit einer gemessenen Spannung skalierten "
              "Ordnungen, deren Spannung um m_[U] zu hoch gemessen sein kann – m_[U] = {mu} %:"},
    "document.eq_ires_scaled": {
        "en": "I_[RES,max] = √(I_[G]² + I_[H,scaled]²) = √({ig}² + {ih}²) A = {ires} A (without the margins)",
        "de": "I_[RES,max] = √(I_[G]² + I_[H,skaliert]²) = √({ig}² + {ih}²) A = {ires} A (ohne Zuschläge)"},
    "document.eq_ires_max": {
        "en": "I_[RES,ass] = √(I_[G,ass]² + I_[H,ass]²) = √({ig}² + {ih}²) A = {ires} A with I_[G,ass] = (1 + m_[I])·I_[G]",
        "de": "I_[RES,Bew] = √(I_[G,Bew]² + I_[H,Bew]²) = √({ig}² + {ih}²) A = {ires} A mit I_[G,Bew] = (1 + m_[I])·I_[G]"},
    "document.step3_unscaled": {
        "en": "Without a scaling the assessment value is the measured RMS value with the margin: I_[RES,ass] = {ires} A.",
        "de": "Ohne Skalierung ist der Bewertungswert der gemessene Effektivwert mit Zuschlag: I_[RES,Bew] = {ires} A."},
    "document.tab_all_tests": {
        "en": "Residual current of every test: measured, with margin and scaled",
        "de": "Erdschlussreststrom aller Versuche: gemessen, mit Zuschlag und skaliert"},
    "document.col_ires_measured": {"en": "I_[RES] in A", "de": "I_[RES] in A"},
    "document.col_ires_margin": {"en": "with margin in A", "de": "mit Zuschlag in A"},
    "document.col_ih_scaled": {"en": "I_[H,scaled] in A", "de": "I_[H,skaliert] in A"},
    "document.col_ires_scaled": {"en": "I_[RES,max] in A", "de": "I_[RES,max] in A"},
    "document.col_ires_scaled_a": {"en": "I_[RES,ass] in A", "de": "I_[RES,Bew] in A"},
    "document.station_scaled_note": {
        "en": "{names}: scaled with the voltages measured at the station (the power-quality measurement holds for the "
              "busbar) – for information.",
        "de": "{names}: mit den an der Station gemessenen Spannungen skaliert (die PQ-Messung gilt für die "
              "Sammelschiene) – informativ."},
    "document.single_factor_note": {
        "en": "The single evaluations were saved with other PQ settings (factor {stored}); this report scales every "
              "test with the PQ measurement of the campaign (factor {used}).",
        "de": "Die Einzelauswertungen sind mit anderen PQ-Einstellungen gespeichert (Faktor {stored}); dieser Bericht "
              "skaliert jeden Versuch mit der PQ-Messung der Kampagne (Faktor {used})."},
    # ---- equivalent sources
    "document.sec_es": {"en": "Equivalent-source method and forecast", "de": "Ersatzquellenverfahren und Prognose"},
    "document.es_intro": {
        "en": "Per harmonic order the network seen from the fault is an active two-terminal network: a source "
              "voltage U_[q,m] behind the source impedance Z_[q,m]. The current of order m at an earthing or fault "
              "impedance Z follows from",
        "de": "Je Oberschwingungsordnung wirkt das Netz von der Fehlerstelle aus als aktiver Zweipol: eine "
              "Quellenspannung U_[q,m] hinter der Quellenimpedanz Z_[q,m]. Der Strom der Ordnung m bei einer Erdungs- "
              "bzw. Fehlerimpedanz Z folgt aus"},
    "document.eq_es": {"en": "I_[m](Z) = U_[q,m] / |Z_[q,m] + Z|", "de": "I_[m](Z) = U_[q,m] / |Z_[q,m] + Z|"},
    "document.es_source_measured": {
        "en": "With the measured source, U_[q,m] is the voltage of the order before the fault – the component of the "
              "natural sequence at the faulted phase, mean of before and after the fault, median of the tests – and "
              "Z_[q,m] = U_[q,m]/I_[m,k] − Z_[k] follows from the tests k with their additional impedances ({zs} Ω; "
              "{how}). For the orders of the zero sequence the voltage before the fault is not their source "
              "(assumption: they arise in the fault); they keep the source fitted to the currents.",
        "de": "Mit der gemessenen Quelle ist U_[q,m] die Spannung der Ordnung vor dem Fehler – die Komponente der "
              "natürlichen Folge am Fehlerleiter, Mittel aus vor und nach dem Fehler, Median der Versuche – und "
              "Z_[q,m] = U_[q,m]/I_[m,k] − Z_[k] folgt aus den Versuchen k mit ihren Zusatzimpedanzen ({zs} Ω; {how}). "
              "Für die Ordnungen des Nullsystems ist die Spannung vor dem Fehler nicht die Quelle (Annahme: sie "
              "entstehen im Fehler); sie behalten die aus den Strömen geschätzte Quelle."},
    "document.zq_median": {"en": "median over the tests", "de": "Median über die Versuche"},
    "document.zq_smallest": {"en": "from the test with the smallest additional impedance",
                             "de": "aus dem Versuch mit der kleinsten Zusatzimpedanz"},
    "document.es_source_fitted": {
        "en": "U_[q,m] and Z_[q,m] are fitted to the currents of the tests with their additional impedances ({zs} Ω).",
        "de": "U_[q,m] und Z_[q,m] sind an die Ströme der Versuche mit ihren Zusatzimpedanzen ({zs} Ω) angepasst."},
    "document.tab_es": {"en": "Equivalent sources per order and current at Z_[E] = {Z} Ω", "de": "Ersatzquellen je Ordnung und Strom an Z_[E] = {Z} Ω"},
    "document.col_sequence": {"en": "Sequence", "de": "System"},
    "document.col_source": {"en": "Source", "de": "Quelle"},
    "document.col_uq": {"en": "U_[q] in V", "de": "U_[q] in V"},
    "document.col_rq": {"en": "R_[q] in Ω", "de": "R_[q] in Ω"},
    "document.col_xq": {"en": "X_[q] in Ω", "de": "X_[q] in Ω"},
    "document.col_residual": {"en": "Residual in %", "de": "Residuum in %"},
    "document.col_loo": {"en": "Leave-one-out max. in %", "de": "Leave-one-out max. in %"},
    "document.col_i_ze": {"en": "I at Z_[E] in A", "de": "I an Z_[E] in A"},
    "document.col_i_ze_scaled": {"en": "scaled in A", "de": "skaliert in A"},
    "document.source_measured": {"en": "measured", "de": "gemessen"},
    "document.source_fitted": {"en": "fitted", "de": "geschätzt"},
    "document.fig_source_fit": {
        "en": "Current per order over the additional impedance: model and tests (normalised to the source level)",
        "de": "Strom je Ordnung über der Zusatzimpedanz: Modell und Versuche (auf den Quellenpegel normiert)"},
    "document.es_result": {
        "en": "At the earthing impedance given for the assessment, Z_[E] = {Z} Ω{basis} (ohmic), the residual current is "
              "I_[RES] = {ires} A (I_[G] = {ig} A, I_[H] = {ih} A){assessed}. Z_[E] lies {inside} the range of the "
              "tests ({zmin} … {zmax} Ω).",
        "de": "An der für die Bewertung vorgegebenen Erdungsimpedanz Z_[E] = {Z} Ω{basis} (ohmsch) beträgt der "
              "Erdschlussreststrom I_[RES] = {ires} A (I_[G] = {ig} A, I_[H] = {ih} A){assessed}. Z_[E] liegt {inside} "
              "des Bereichs der Versuche ({zmin} … {zmax} Ω)."},
    "document.es_result_assessed": {"en": ", with margin {a} A", "de": ", mit Zuschlag {a} A"},
    "document.es_scaled": {
        "en": "Scaled order by order to the power-quality measurement ({how}; {kept} as measured): "
              "I_[H,scaled] = {ih} A, I_[RES] = {ires} A, assessed I_[RES,ass] = {a} A.",
        "de": "Je Ordnung skaliert auf die PQ-Messung ({how}; {kept} wie gemessen): I_[H,skaliert] = {ih} A, "
              "I_[RES] = {ires} A, Bewertung I_[RES,Bew] = {a} A."},
    "document.es_how_own": {"en": "U_[q,m] set to k_[s]·U_[95,m]", "de": "U_[q,m] auf k_[s]·U_[95,m] gesetzt"},
    "document.es_how_thd": {
        "en": "U_[q,m] multiplied by the THD factor k_[THD] = {f}, there is no quantile per order",
        "de": "U_[q,m] mit dem THD-Faktor k_[THD] = {f} multipliziert, ein Quantil je Ordnung gibt es nicht"},
    "document.es_how_mixed": {
        "en": "U_[q,m] set to k_[s]·U_[95,m], an order without its own quantile multiplied by the THD factor "
              "k_[THD] = {f}",
        "de": "U_[q,m] auf k_[s]·U_[95,m] gesetzt, eine Ordnung ohne eigenes Quantil mit dem THD-Faktor k_[THD] = {f} "
              "multipliziert"},
    "document.es_how_own_unscaled": {
        "en": "U_[q,m] set to k_[s]·U_[95,m], an order without its own quantile as measured",
        "de": "U_[q,m] auf k_[s]·U_[95,m] gesetzt, eine Ordnung ohne eigenes Quantil wie gemessen"},
    "document.es_how_none": {"en": "no order with a quantile", "de": "keine Ordnung mit einem Quantil"},
    "document.es_kept_between": {"en": "zero sequence, even orders and the content between the orders",
                                 "de": "Nullsystem, gerade Ordnungen und Anteile zwischen den Ordnungen"},
    "document.es_kept_rest": {"en": "zero sequence, even orders and the rest outside the orders of the model",
                              "de": "Nullsystem, gerade Ordnungen und der Rest außerhalb der Ordnungen des Modells"},
    "document.es_mc": {
        "en": "Monte Carlo of the harmonic content[^Random errors (1σ) of the currents {sm} % and {sp}°, of the "
              "additional impedances {sz} %{level}{model}; I_[G] fixed; {n} runs.^]: I_[H](Z_[E]) P5 … P95 = {h5} … "
              "{h95} A, hence I_[RES](Z_[E]) = {p5} … {p95} A (not scaled).",
        "de": "Monte Carlo des Oberschwingungsanteils[^Zufallsfehler (1σ) der Ströme {sm} % und {sp}°, der "
              "Zusatzimpedanzen {sz} %{level}{model}; I_[G] fest; {n} Läufe.^]: I_[H](Z_[E]) P5 … P95 = {h5} … "
              "{h95} A, damit I_[RES](Z_[E]) = {p5} … {p95} A (nicht skaliert)."},
    "document.es_mc_level": {"en": ", of the level normalisation {sl} %", "de": ", der Pegelnormierung {sl} %"},
    "document.es_mc_model": {"en": ", plus the residual of the fit per order",
                             "de": ", dazu das Residuum der Anpassung je Ordnung"},
    "document.es_loo": {
        "en": "Cross-validation[^Leave-one-out: each test is forecast from the other tests.^]: largest deviation "
              "{dev} % ({order}th order){limit}.",
        "de": "Kreuzvalidierung[^Leave-one-out: Jeder Versuch wird aus den übrigen Versuchen vorhergesagt.^]: größte "
              "Abweichung {dev} % ({order}. Ordnung){limit}."},
    "document.es_loo_limit": {
        "en": "; {bad} of {n} values outside {lim} – a hint that the network or the sources changed between the tests",
        "de": "; {bad} von {n} Werten außerhalb {lim} – ein Hinweis auf eine Änderung von Netz oder Quellen zwischen "
              "den Versuchen"},
    "document.loo_limit_unknown": {"en": "the permissible deviation", "de": "der zulässigen Abweichung"},
    "document.or": {"en": "or", "de": "bzw."},
    "document.fig_ires_ze": {
        "en": "Residual current over the earthing impedance with the limits", "de": "Erdschlussreststrom über der Erdungsimpedanz mit den Grenzwerten"},
    "document.tab_methods": {
        "en": "Residual current at Z_[E] = {Z} Ω with the power-quality measurement – methods",
        "de": "Erdschlussreststrom an Z_[E] = {Z} Ω mit der PQ-Messung – Verfahren"},
    "document.col_method": {"en": "Method", "de": "Verfahren"},
    "document.col_ih": {"en": "I_[H] in A", "de": "I_[H] in A"},
    "document.col_ires": {"en": "I_[RES] in A", "de": "I_[RES] in A"},
    "document.col_ires_assessed": {"en": "I_[RES,ass] in A", "de": "I_[RES,Bew] in A"},
    "document.forecast_title": {"en": "What the equivalent sources mean for the forecast", "de": "Bedeutung für die Prognose"},
    "document.forecast_1": {
        "en": "Earthing impedance: for an ohmic earthing or fault impedance the residual current can be estimated for "
              "every value, not only at the impedances of the tests{fig}; the harmonic part then falls with growing "
              "resistance.{reactive}",
        "de": "Erdungsimpedanz: Für eine ohmsche Erdungs- bzw. Fehlerimpedanz ist der Reststrom für jeden Wert "
              "abschätzbar, nicht nur an den Impedanzen der Versuche{fig}; der Oberschwingungsanteil sinkt dann mit "
              "wachsendem Widerstand.{reactive}"},
    "document.ref_figure": {"en": " (figure {n})", "de": " (Abbildung {n})"},
    "document.forecast_reactive": {
        "en": " The source impedances are capacitive (X_[q] < 0): an inductive earthing impedance can increase the "
              "current of an order (series resonance) – the model takes the impedance as ohmic.",
        "de": " Die Quellenimpedanzen sind kapazitiv (X_[q] < 0): Eine induktive Erdungsimpedanz kann den Strom einer "
              "Ordnung erhöhen (Reihenresonanz) – das Modell rechnet mit einer ohmschen Impedanz."},
    "document.forecast_2": {
        "en": "Network state: scaled to the quantiles of the power-quality measurement, the value holds for the "
              "unfavourable state of the measured period; the model is valid for the network and switching state of "
              "the tests – reliable within the range of the additional impedances, beyond it an extrapolation.",
        "de": "Netzzustand: Skaliert auf die Quantile der PQ-Messung gilt der Wert für den ungünstigen Zustand des "
              "gemessenen Zeitraums; das Modell gilt für Netz und Schaltzustand der Versuche – im Bereich der "
              "Zusatzimpedanzen belastbar, darüber hinaus eine Extrapolation."},
    "document.forecast_2_unscaled": {
        "en": "Network state: without a scaling to a power-quality measurement, the value holds for the harmonic "
              "voltage on the day of the tests; the model is valid for the network and switching state of the tests – "
              "reliable within the range of the additional impedances, beyond it an extrapolation.",
        "de": "Netzzustand: Ohne Skalierung auf eine PQ-Messung gilt der Wert für die Oberschwingungsspannung am "
              "Versuchstag; das Modell gilt für Netz und Schaltzustand der Versuche – im Bereich der "
              "Zusatzimpedanzen belastbar, darüber hinaus eine Extrapolation."},
    "document.forecast_3": {
        "en": "Growth of the network: the active part of the fundamental grows with I_[CE] (I_[w] = d·I_[CE]); with "
              "the present damping d = {d} % the active current alone reaches {limit} A at I_[CE] ≈ {ice} A (now "
              "{ice_now} A){ig}. The harmonics do not grow in proportion (resonance) – for them the network model of "
              "the forecast is needed.",
        "de": "Netzausbau: Der Wirkanteil der Grundschwingung wächst mit I_[CE] (I_[w] = d·I_[CE]); mit der heutigen "
              "Dämpfung d = {d} % erreicht der Wirkstrom allein {limit} A bei I_[CE] ≈ {ice} A (heute {ice_now} A){ig}. "
              "Die Oberschwingungen wachsen nicht proportional (Resonanz) – für sie ist das Netzmodell der Prognose "
              "nötig."},
    "document.forecast_3_ig": {
        "en": ", with the detuning of the tests (v = {v} %) the fundamental at I_[CE] ≈ {ice} A",
        "de": ", mit der Verstimmung der Versuche (v = {v} %) die Grundschwingung bei I_[CE] ≈ {ice} A"},
    "document.no_es": {
        "en": "The equivalent-source method needs at least one test at the substation with its additional impedance.",
        "de": "Das Ersatzquellenverfahren braucht mindestens einen Versuch im Umspannwerk mit seiner Zusatzimpedanz."},
    "document.no_es_given": {
        "en": "No equivalent sources were determined for this evaluation.",
        "de": "Für diese Auswertung wurden keine Ersatzquellen bestimmt."},
    # ---- station assessment
    "document.sec_station_assessment": {"en": "Assessment of the earthing – {station}", "de": "Aufbereitung Erdung – {station}"},
    "document.ut_condition": {
        "en": "The touch voltage must not exceed the permissible touch voltage, U_[T] ≤ U_[Tp], with",
        "de": "Die Berührungsspannung darf die zulässige Berührungsspannung nicht überschreiten, U_[T] ≤ U_[Tp], mit"},
    "document.ut_cases_note": {
        "en": "Cases 1 and 2 take the limit current of the network ({limit} A) and Z_[E] of the earthing measurement at "
              "the reference frequency {f} Hz set for the evaluation ({z} Ω{how}{z50}).",
        "de": "Die Fälle 1 und 2 setzen den Grenzstrom des Netzgebiets ({limit} A) und Z_[E] der Erdungsmessung bei der "
              "für die Auswertung gesetzten Referenzfrequenz {f} Hz an ({z} Ω{how}{z50})."},
    "document.ut_cases_z50": {"en": "; at 50 Hz {z} Ω", "de": "; bei 50 Hz {z} Ω"},
    "document.ut_ref_interpolated": {"en": ", interpolated linearly between {lo} and {hi} Hz",
                                     "de": ", linear interpoliert zwischen {lo} und {hi} Hz"},
    "document.ut_ref_nearest": {"en": " at {used} Hz – the earthing measurement has no value at {req} Hz",
                                "de": " bei {used} Hz – die Erdungsmessung hat bei {req} Hz keinen Wert"},
    "document.eq_ut": {"en": "U_[T] = r·I_[RES]·Z_[E] / F = U_[E] / F", "de": "U_[T] = r·I_[RES]·Z_[E] / F = U_[E] / F"},
    "document.ut_parameters": {
        "en": "Parameters of the network operator: U_[Tp] = {utp} V, F = {f}{r}.",
        "de": "Vorgaben des Netzbetreibers: U_[Tp] = {utp} V, F = {f}{r}."},
    "document.ut_r_flat": {"en": ", flat reduction factor r = {r}", "de": ", pauschaler Reduktionsfaktor r = {r}"},
    "document.ut_case2_missing": {
        "en": "Case 2 (limit current with a flat reduction factor r) is left out – no flat r is given.",
        "de": "Fall 2 (Grenzstrom mit pauschalem Reduktionsfaktor r) entfällt – kein pauschales r vorgegeben."},
    "document.ut_missing": {
        "en": "U_[Tp] or F is not set – the touch voltage is not assessed.",
        "de": "U_[Tp] oder F ist nicht gesetzt – die Berührungsspannung wird nicht bewertet."},
    "document.tab_ut": {"en": "Approaches to the touch voltage – {station}", "de": "Ansätze zur Beurteilung der Berührungsspannung – {station}"},
    "document.col_case": {"en": "Case", "de": "Fall"},
    "document.col_i": {"en": "I in A", "de": "I in A"},
    "document.col_ze": {"en": "Z_[E] in Ω", "de": "Z_[E] in Ω"},
    "document.col_ue": {"en": "U_[E] in V", "de": "U_[E] in V"},
    "document.col_ut": {"en": "U_[T] in V", "de": "U_[T] in V"},
    "document.col_share": {"en": "of U_[Tp] in %", "de": "von U_[Tp] in %"},
    "document.case_1": {
        "en": "1 limit current, r = 1, Z_[E] at {f} Hz", "de": "1 Grenzstrom, r = 1, Z_[E] bei {f} Hz"},
    "document.case_2": {
        "en": "2 limit current, flat r, Z_[E] at {f} Hz", "de": "2 Grenzstrom, pauschales r, Z_[E] bei {f} Hz"},
    "document.case_3": {
        "en": "3 per frequency: r·Z_[E] measured, assessed current per order",
        "de": "3 je Frequenz: r·Z_[E] gemessen, Bewertungsstrom je Ordnung"},
    "document.case_4": {"en": "4 measured earth potential rise", "de": "4 gemessene Erdungsspannung"},
    "document.case_4b": {"en": "4b measured, extrapolated to the limit current", "de": "4b gemessen, hochgerechnet auf den Grenzstrom"},
    "document.case_ut_measured": {"en": "measured touch voltage (probe)", "de": "gemessene Berührungsspannung (Sonde)"},
    "document.case_short_1": {"en": "Case 1", "de": "Fall 1"},
    "document.case_short_2": {"en": "Case 2", "de": "Fall 2"},
    "document.case_short_3": {"en": "Case 3", "de": "Fall 3"},
    "document.case_short_4": {"en": "Case 4", "de": "Fall 4"},
    "document.case_short_4b": {"en": "Case 4b", "de": "Fall 4b"},
    "document.case_short_ut": {"en": "U_T probe", "de": "U_T Sonde"},
    "document.ut_case3_note": {
        "en": "Case 3 uses the assessment currents per order of {name} – the largest assessment value, conservative; "
              "I = √(Σ I_[m]²) without the content between the orders – and r·Z_[E] of the earth-fault test per order "
              "(probe {probe}, fault-related; orders {orders}); orders without a value of their own take the largest "
              "measured value of r·Z_[E] ({covered} % of Σ I_[m]² lie at measured orders).",
        "de": "Fall 3 verwendet die Bewertungsströme je Ordnung von {name} – dem größten Bewertungswert, konservativ; "
              "I = √(Σ I_[m]²) ohne die Anteile zwischen den Ordnungen – und r·Z_[E] der Erdschlussmessung je Ordnung "
              "(Spieß {probe}, fehlerbedingt; Ordnungen {orders}); Ordnungen ohne eigenen Wert erhalten den größten "
              "gemessenen Wert von r·Z_[E] ({covered} % von Σ I_[m]² liegen bei gemessenen Ordnungen)."},
    "document.fig_ut": {
        "en": "Approaches to the touch voltage – {station}", "de": "Gegenüberstellung der Ansätze zur Beurteilung der Berührungsspannung – {station}"},
    "document.ze_permissible": {
        "en": "Permissible earthing impedance for I_[RES,ass] = {i} A: Z_[E,perm] = F·U_[Tp]/(r·I_[RES,ass]) = {z1} Ω "
              "with r = 1{zr}{cmp}.",
        "de": "Zulässige Erdungsimpedanz für I_[RES,Bew] = {i} A: Z_[E,zul] = F·U_[Tp]/(r·I_[RES,Bew]) = {z1} Ω "
              "mit r = 1{zr}{cmp}."},
    "document.ze_permissible_r": {"en": ", {z} Ω with the effective r = {r} of the test",
                                  "de": ", {z} Ω mit dem wirksamen r = {r} der Messung"},
    "document.ze_permissible_cmp": {
        "en": " – to be compared with Z_[E] of the earthing measurement at the reference frequency ({f} Hz: {z} Ω)",
        "de": " – zu vergleichen mit Z_[E] der Erdungsmessung bei der Referenzfrequenz ({f} Hz: {z} Ω)"},
    # ---- uncertainties
    "document.sec_uncertainties": {"en": "Uncertainties", "de": "Unsicherheiten"},
    "document.uncertainties_intro": {
        "en": "Contributions to the uncertainty of the assessment and how they are treated:",
        "de": "Beiträge zur Unsicherheit der Bewertung und wie sie berücksichtigt sind:"},
    "document.tab_uncertainties": {"en": "Uncertainties of the assessment", "de": "Unsicherheiten der Bewertung"},
    "document.col_contribution": {"en": "Contribution", "de": "Beitrag"},
    "document.col_size": {"en": "Size", "de": "Größe"},
    "document.col_treatment": {"en": "Treatment", "de": "Berücksichtigung"},
    "document.u_chain": {"en": "Measurement chain", "de": "Messkette"},
    "document.u_chain_value": {"en": "sum {lin} %, RSS {rss} %", "de": "Summe {lin} %, quadr. {rss} %"},
    "document.u_margin_value": {"en": "margin m_[I] = {mi} %", "de": "Zuschlag m_[I] = {mi} %"},
    "document.u_margin_value_mu": {"en": "margins m_[I] = {mi} %, m_[U] = {mu} %",
                                   "de": "Zuschläge m_[I] = {mi} %, m_[U] = {mu} %"},
    "document.u_stationarity": {"en": "Steady state of the fault, {name}", "de": "Stationarität im Fehler, {name}"},
    "document.u_stationarity_value": {
        "en": "I_[RES] over {n} windows: {min} … {max} A (σ = {std} %)", "de": "I_[RES] über {n} Fenster: {min} … {max} A (σ = {std} %)"},
    "document.u_stationarity_treat": {"en": "largest window assessed", "de": "größtes Fenster bewertet"},
    "document.u_stationarity_others": {"en": "Steady state, other tests", "de": "Stationarität, übrige Versuche"},
    "document.u_stationarity_item": {"en": "{name}: {min} … {max} A (σ = {std} %)",
                                     "de": "{name}: {min} … {max} A (σ = {std} %)"},
    "document.u_stationarity_others_treat": {"en": "for information (findings in appendix B)",
                                             "de": "informativ (Befunde in Anhang B)"},
    "document.u_tests": {"en": "Differences between the tests", "de": "Unterschiede zwischen den Versuchen"},
    "document.u_tests_value": {"en": "I_[G] {min} … {max} A ({span} %)", "de": "I_[G] {min} … {max} A ({span} %)"},
    "document.u_tests_treat": {"en": "I_[G]: {choice}", "de": "I_[G]: {choice}"},
    "document.ig_max": {"en": "largest of the tests", "de": "größter der Versuche"},
    "document.ig_median": {"en": "median of the tests", "de": "Median der Versuche"},
    "document.u_model": {"en": "Equivalent-source model", "de": "Ersatzquellenmodell"},
    "document.u_model_value": {"en": "Monte Carlo (I_[G] fixed) P5 … P95 {p5} … {p95} A",
                               "de": "Monte Carlo (I_[G] fest) P5 … P95 {p5} … {p95} A"},
    "document.u_loo_value": {"en": "leave-one-out up to {dev} %", "de": "Leave-one-out bis {dev} %"},
    "document.u_pq": {"en": "Sample of the PQ measurement", "de": "Stichprobe der PQ-Messung"},
    "document.u_pq_value": {
        "en": "bootstrap over days: I_[RES](Z_[E]) scaled P5 … P95 = {p5} … {p95} A",
        "de": "Bootstrap über Tage: I_[RES](Z_[E]) skaliert P5 … P95 = {p5} … {p95} A"},
    "document.u_pq_value_order": {
        "en": "quantile of the {order}th order {low} … {high} %", "de": "Quantil der {order}. Ordnung {low} … {high} %"},
    "document.u_pq_treat": {"en": "factor k_[s] = {k} on the quantile", "de": "Faktor k_[s] = {k} auf das Quantil"},
    "document.u_pq_periods": {"en": "Periods of the PQ measurement", "de": "Abschnitte der PQ-Messung"},
    "document.u_pq_periods_value": {
        "en": "I_[RES](Z_[E]) scaled with one period alone {low} … {high} A; quantile of the {order}th order "
              "{qlow} … {qhigh} V",
        "de": "I_[RES](Z_[E]) skaliert mit einem Abschnitt allein {low} … {high} A; Quantil der {order}. Ordnung "
              "{qlow} … {qhigh} V"},
    "document.u_pq_periods_treat": {"en": "quantile over all periods", "de": "Quantil über alle Abschnitte"},
    "document.u_methods": {"en": "Method of the scaling", "de": "Verfahren der Skalierung"},
    "document.u_methods_value": {"en": "I_[RES] {min} … {max} A", "de": "I_[RES] {min} … {max} A"},
    "document.u_methods_treat": {"en": "per order (result)", "de": "je Ordnung (Ergebnis)"},
    "document.u_extrapolation": {"en": "Range of the equivalent sources", "de": "Gültigkeitsbereich der Ersatzquellen"},
    "document.u_extrapolation_value": {"en": "Z_[E] = {Z} Ω {inside} {zmin} … {zmax} Ω", "de": "Z_[E] = {Z} Ω {inside} {zmin} … {zmax} Ω"},
    # ---- summary
    "document.sec_summary": {"en": "Summary", "de": "Zusammenfassung"},
    "document.sum_ice": {
        "en": "In its normal switching state the network has a capacitive earth-fault current of I_[CE] ≈ {ice} A "
              "({source}).",
        "de": "Das Netzgebiet zeigt im Normalschaltzustand einen kapazitiven Erdschlussstrom von I_[CE] ≈ {ice} A "
              "({source})."},
    "document.sum_ice_both": {
        "en": "the coil controller displays {display} A, the measured coil currents add up to {coils} A",
        "de": "Anzeige des Spulenreglers {display} A, Summe der gemessenen Spulenströme {coils} A"},
    "document.sum_ires": {
        "en": "The residual earth-fault current of the tests reaches {ires} A in {name} – {share} % of I_[CE] –, with "
              "margin {a} A; the damping of the network is d = {d} %.",
        "de": "Der Erdschlussreststrom der Versuche erreicht {ires} A in {name} – {share} % von I_[CE] –, mit Zuschlag "
              "{a} A; die Dämpfung des Netzes beträgt d = {d} %."},
    "document.sum_composition": {
        "en": "As measured, {measured} predominates (I_[G] = {ig} A, I_[H] = {ih} A, mainly the {order}th order){scaled}.",
        "de": "Gemessen überwiegt {measured} (I_[G] = {ig} A, I_[H] = {ih} A, vor allem die {order}. Ordnung){scaled}."},
    "document.sum_composition_scaled": {
        "en": "; scaled to the power-quality measurement, {which} predominates (I_[G,ass] = {ig} A, I_[H,ass] = {ih} A)",
        "de": "; nach der Skalierung auf die PQ-Messung überwiegt {which} (I_[G,Bew] = {ig} A, I_[H,Bew] = {ih} A)"},
    "document.sum_composition_same": {
        "en": "; likewise after the scaling to the power-quality measurement (I_[G,ass] = {ig} A, I_[H,ass] = {ih} A)",
        "de": "; ebenso nach der Skalierung auf die PQ-Messung (I_[G,Bew] = {ig} A, I_[H,Bew] = {ih} A)"},
    "document.the_fundamental": {"en": "the fundamental", "de": "die Grundschwingung"},
    "document.the_harmonics": {"en": "the harmonic content", "de": "der Oberschwingungsanteil"},
    "document.sum_scaled": {
        "en": "Scaled order by order to the power-quality measurement (factor {k}), the largest test gives "
              "I_[RES,ass] = {ires} A; limit {limit} A: {status}.",
        "de": "Je Ordnung auf die PQ-Messung skaliert (Faktor {k}) ergibt der größte Versuch I_[RES,Bew] = {ires} A; "
              "Grenzwert {limit} A: {status}."},
    "document.sum_es": {
        "en": "With the equivalent sources the residual current at the earthing impedance Z_[E] = {Z} Ω is "
              "I_[RES,ass] = {ires} A{scaled}; limit {limit} A: {status}.",
        "de": "Mit den Ersatzquellen beträgt der Reststrom an der Erdungsimpedanz Z_[E] = {Z} Ω I_[RES,Bew] = {ires} A"
              "{scaled}; Grenzwert {limit} A: {status}."},
    "document.sum_es_value": {
        "en": "With the equivalent sources the residual current at the earthing impedance Z_[E] = {Z} Ω is "
              "I_[RES,ass] = {ires} A{scaled}.",
        "de": "Mit den Ersatzquellen beträgt der Reststrom an der Erdungsimpedanz Z_[E] = {Z} Ω I_[RES,Bew] = {ires} A"
              "{scaled}."},
    "document.sum_es_scaled": {"en": " (scaled)", "de": " (skaliert)"},
    "document.sum_decisive": {
        "en": "Decisive for the assessment is {which}.", "de": "Maßgeblich für die Bewertung ist {which}."},
    "document.decisive_ze": {
        "en": "the value at the earthing impedance Z_[E] (faults at installations in the network)",
        "de": "der Wert an der Erdungsimpedanz Z_[E] (Fehler an Anlagen im Netz)"},
    "document.decisive_busbar": {
        "en": "the value of the largest test (fault at the busbar)", "de": "der Wert des größten Versuchs (Fehler an der Sammelschiene)"},
    "document.sum_station_ok": {
        "en": "{station}: every approach keeps the permissible touch voltage U_[Tp] = {utp} V – at most {ut} V "
              "({case}, {share} % of U_[Tp]){measured}.",
        "de": "{station}: Alle Ansätze halten die zulässige Berührungsspannung U_[Tp] = {utp} V ein – höchstens "
              "{ut} V ({case}, {share} % von U_[Tp]){measured}."},
    "document.sum_station_exceeded": {
        "en": "{station}: {cases} exceed the permissible touch voltage U_[Tp] = {utp} V – up to {ut} V ({case}, "
              "{share} % of U_[Tp]){measured}.",
        "de": "{station}: {cases} überschreiten die zulässige Berührungsspannung U_[Tp] = {utp} V – bis {ut} V "
              "({case}, {share} % von U_[Tp]){measured}."},
    "document.sum_station_measured": {
        "en": "; with the values measured at the station ({cases}) at most {ut} V",
        "de": "; mit den an der Station gemessenen Werten ({cases}) höchstens {ut} V"},
    "document.sum_station_simple": {
        "en": "{station}: the measured earth potential rise is {ue} V at I_[F] = {i} A.",
        "de": "{station}: Die gemessene Erdungsspannung beträgt {ue} V bei I_[F] = {i} A."},
    "document.sum_forecast": {
        "en": "With the present damping the active current alone reaches {limit} A at I_[CE] ≈ {ice} A{ig}.",
        "de": "Unter Annahme der aktuellen Dämpfung erreicht der Wirkstrom allein {limit} A bei I_[CE] ≈ {ice} A{ig}."},
    "document.sum_forecast_ig": {
        "en": ", with the detuning of the tests the fundamental at I_[CE] ≈ {ice} A",
        "de": ", mit der Verstimmung der Versuche die Grundschwingung bei I_[CE] ≈ {ice} A"},
    "document.sum_ze": {
        "en": "At I_[RES,ass] = {i} A, earthing impedances up to {z} Ω keep the permissible touch voltage (r = 1: "
              "U_[E] ≤ F·U_[Tp] = {u} V).",
        "de": "Bei I_[RES,Bew] = {i} A halten Erdungsimpedanzen bis {z} Ω die zulässige Berührungsspannung ein "
              "(r = 1: U_[E] ≤ F·U_[Tp] = {u} V)."},
    "document.sum_hints": {"en": "Notes of the evaluation:", "de": "Hinweise der Auswertung:"},
    "document.recommendation": {"en": "Recommendation", "de": "Empfehlung"},
    "document.recommendation_placeholder": {
        "en": "Recommendation for the network area (engineering judgement of the author)",
        "de": "Empfehlung für das Netzgebiet (fachliche Einschätzung des Bearbeiters)"},
    # ------------------------------------------------------------------ appendices
    "document.app_data": {"en": "Data of the recordings", "de": "Messdaten je Messung"},
    "document.app_data_intro": {
        "en": "Recording, intervals and settings of every evaluated measurement.",
        "de": "Aufzeichnung, Bereiche und Einstellungen jeder ausgewerteten Messung."},
    "document.tab_data": {"en": "Data of the recordings", "de": "Messdaten der Aufzeichnungen"},
    "document.d_file": {"en": "Recording", "de": "Aufzeichnung"},
    "document.d_site": {"en": "Measuring point", "de": "Messort"},
    "document.d_start": {"en": "Start of the recording ({tz})", "de": "Beginn der Aufzeichnung ({tz})"},
    "document.d_fs": {"en": "Sampling rate in kHz", "de": "Abtastrate in kHz"},
    "document.d_duration": {"en": "Duration in s", "de": "Dauer in s"},
    "document.d_f1": {"en": "Fundamental frequency in Hz", "de": "Grundfrequenz in Hz"},
    "document.d_steady": {"en": "Steady state in s", "de": "Stationärer Bereich in s"},
    "document.d_pre": {"en": "Before the fault in s", "de": "Vor dem Fehler in s"},
    "document.d_post": {"en": "After the fault in s", "de": "Nach dem Fehler in s"},
    "document.d_window": {"en": "Time of the assessed window in s", "de": "Zeitpunkt des bewerteten Fensters in s"},
    "document.d_polarity": {"en": "Phasor of the fault current", "de": "Zeiger des Fehlerstroms"},
    "document.d_polarity_rotated": {"en": "turned by 180°", "de": "um 180° gedreht"},
    "document.d_twin": {"en": "Window length in ms", "de": "Fensterlänge in ms"},
    "document.d_phase": {"en": "Faulted phase", "de": "Fehlerleiter"},
    "document.d_z": {"en": "Additional impedance in Ω", "de": "Zusatzimpedanz in Ω"},
    "document.d_main": {"en": "Fault current (channel)", "de": "Fehlerstrom (Kanal)"},
    "document.d_corrections": {"en": "Corrections", "de": "Korrekturen"},
    "document.d_reproduced": {"en": "Evaluated again from the raw data", "de": "Aus den Rohdaten neu ausgewertet"},
    "document.d_reproduced_value": {"en": "largest deviation {dev} %", "de": "größte Abweichung {dev} %"},
    "document.app_plausibility": {"en": "Plausibility checks", "de": "Plausibilitätsprüfung"},
    "document.app_plausibility_intro": {
        "en": "Checks of the single evaluations (only those with a finding), between the tests and of this report.",
        "de": "Prüfungen der Einzelauswertungen (nur die mit Befund), zwischen den Versuchen und dieses Berichts."},
    "document.tab_plaus_single": {"en": "Findings of the single evaluations", "de": "Befunde der Einzelauswertungen"},
    "document.tab_plaus_campaign": {"en": "Consistency between the tests", "de": "Konsistenz zwischen den Versuchen"},
    "document.plaus_campaign_intro": {
        "en": "The consistency between the tests covers the tests with an additional impedance: {names}.",
        "de": "Die Konsistenz zwischen den Versuchen umfasst die Versuche mit Zusatzimpedanz: {names}."},
    "document.tab_plaus_report": {"en": "Checks of the report", "de": "Prüfungen des Berichts"},
    "document.col_check": {"en": "Check", "de": "Prüfung"},
    "document.col_status": {"en": "Status", "de": "Status"},
    "document.col_value": {"en": "Value", "de": "Wert"},
    "document.col_comment": {"en": "Comment", "de": "Kommentar"},
    "document.col_measurement": {"en": "Measurement", "de": "Messung"},
    "document.plaus_ok_count": {
        "en": "{n} further checks of the single evaluations without a finding.",
        "de": "{n} weitere Prüfungen der Einzelauswertungen ohne Befund."},
    "document.chk_split": {"en": "I_[w]² + I_[v]² = I_[G]²", "de": "I_[w]² + I_[v]² = I_[G]²"},
    "document.chk_rms": {"en": "I_[RES]² = I_[G]² + I_[H]² + I_[lf]²", "de": "I_[RES]² = I_[G]² + I_[H]² + I_[nf]²"},
    "document.chk_deviation": {"en": "largest deviation {dev} %", "de": "größte Abweichung {dev} %"},
    "document.chk_settings": {"en": "Same PQ settings in single evaluations and campaign", "de": "Gleiche PQ-Einstellungen in Einzel- und Kampagnenauswertung"},
    "document.chk_settings_ok": {"en": "factor {k} everywhere", "de": "Faktor {k} überall"},
    "document.chk_settings_unscaled": {
        "en": "single evaluations not scaled – the report scales with the campaign (factor {k})",
        "de": "Einzelauswertungen ohne Skalierung – der Bericht skaliert mit der Kampagne (Faktor {k})"},
    "document.chk_settings_diff": {
        "en": "single evaluations factor {stored}, campaign {used} – the report uses the campaign",
        "de": "Einzelauswertungen Faktor {stored}, Kampagne {used} – der Bericht verwendet die Kampagne"},
    "document.chk_margin": {"en": "Same margins m_[I], m_[U] in single evaluations and report",
                            "de": "Gleiche Zuschläge m_[I], m_[U] in Einzelauswertungen und Bericht"},
    "document.chk_margin_value": {"en": "m_[I] {single} ↔ {report}; m_[U] {single_u} ↔ {report_u}",
                                  "de": "m_[I] {single} ↔ {report}; m_[U] {single_u} ↔ {report_u}"},
    "document.chk_utp": {"en": "Same U_[Tp] and F in single evaluations and report",
                         "de": "Gleiche U_[Tp] und F in Einzelauswertungen und Bericht"},
    "document.chk_utp_value": {"en": "U_[Tp] {single} ↔ {report} V; F {single_f} ↔ {report_f}",
                               "de": "U_[Tp] {single} ↔ {report} V; F {single_f} ↔ {report_f}"},
    "document.below": {"en": "< {x}", "de": "< {x}"},
    "document.status_OK": {"en": "OK", "de": "OK"},
    "document.status_WARNING": {"en": "warning", "de": "Warnung"},
    "document.status_NOTE": {"en": "note", "de": "Hinweis"},
    "document.status_CORRECTED": {"en": "corrected", "de": "korrigiert"},
    "document.status_INFO": {"en": "info", "de": "Info"},
    "document.status_ERROR": {"en": "error", "de": "Fehler"},
    "document.chk_reproduced": {"en": "Evaluated again from the raw data", "de": "Neuberechnung aus den Rohdaten"},
    "document.chk_model": {"en": "Equivalent sources: I_[RES] at the impedances of the tests (level and I_[G] of each test)",
                           "de": "Ersatzquellen: I_[RES] an den Impedanzen der Versuche (Pegel und I_[G] je Versuch)"},
    "document.chk_residual": {"en": "Equivalent sources: residual of the currents per order",
                              "de": "Ersatzquellen: Residuum der Ströme je Ordnung"},
    "document.chk_residual_value": {"en": "largest {dev} % ({order}th order)",
                                    "de": "größtes {dev} % ({order}. Ordnung)"},
    "document.chk_loo": {"en": "Equivalent sources: leave-one-out per order",
                         "de": "Ersatzquellen: Leave-one-out je Ordnung"},
    "document.chk_loo_value": {"en": "{bad} of {n} outside {lim}; up to {dev} % ({order}th order)",
                               "de": "{bad} von {n} außerhalb {lim}; bis {dev} % ({order}. Ordnung)"},
    "document.chk_loo_value_ok": {"en": "all within {lim}; up to {dev} % ({order}th order)",
                                  "de": "alle innerhalb {lim}; bis {dev} % ({order}. Ordnung)"},
    "document.chk_probe": {"en": "Touch voltage (probe) ≤ earth potential rise", "de": "Berührungsspannung (Spieß) ≤ Erdungsspannung"},
    "document.chk_probe_value": {"en": "{station}: {ut} V ↔ {ue} V", "de": "{station}: {ut} V ↔ {ue} V"},
    "document.chk_pq": {"en": "Power-quality measurement valid", "de": "PQ-Messung gültig"},
    # ---- review round: processing, context, coil controller, limits, compact station
    "document.proc_dc_highpass": {
        "en": "removal of the DC component with a high-pass (Butterworth, {fc} Hz, zero-phase) on every interval, then "
              "the mean per window",
        "de": "Entfernen des Gleichanteils mit einem Hochpass (Butterworth, {fc} Hz, phasenfrei) je Auswertebereich, "
              "danach je Fenster der Mittelwert"},
    "document.proc_dc_highpass_median": {
        "en": "removal of the DC component with a high-pass (Butterworth, {fc} Hz, zero-phase) on every interval, then "
              "the median per window",
        "de": "Entfernen des Gleichanteils mit einem Hochpass (Butterworth, {fc} Hz, phasenfrei) je Auswertebereich, "
              "danach je Fenster der Median"},
    "document.ctrl_intro": {"en": "Display of the coil controller{when}: {parts}.",
                            "de": "Anzeige des Spulenreglers{when}: {parts}."},
    "document.ctrl_when": {"en": " ({when})", "de": " ({when})"},
    "document.ctrl_coils": {"en": "coils {coils} (together {sum} A)", "de": "Spulen {coils} (zusammen {sum} A)"},
    "document.ctrl_ice": {"en": "I_[CE] = {ice} A", "de": "I_[CE] = {ice} A"},
    "document.ctrl_detuning": {"en": "detuning v = {v} %", "de": "Verstimmung v = {v} %"},
    "document.ctrl_active": {"en": "active residual current {a} A", "de": "Wirkreststrom {a} A"},
    "document.ctrl_derived": {
        "en": "Hence I_[CE] = I_[L]/(1 + v) = {sum} A/{den} ≈ {ice} A (v > 0: overcompensation).",
        "de": "Daraus I_[CE] = I_[L]/(1 + v) = {sum} A/{den} ≈ {ice} A (v > 0: Überkompensation)."},
    "document.ctrl_compare": {
        "en": "The measured values (median of the tests with an additional impedance) differ from the display: "
              "{items} – parameters of the coils and errors of the transducers and clamps.",
        "de": "Die Messwerte (Median der Versuche mit Zusatzimpedanz) weichen von der Anzeige ab: {items} – "
              "Parametrierung der Spulen und Messfehler der Wandler und Zangen."},
    "document.ctrl_cmp_coils": {"en": "sum of the coil currents {x} A ({dev} %)",
                                "de": "Summe der Spulenströme {x} A ({dev} %)"},
    "document.ctrl_cmp_active": {"en": "active current {x} A ({dev} %)", "de": "Wirkstrom {x} A ({dev} %)"},
    "document.ctrl_cmp_detuning": {"en": "detuning {x} % (magnitude)", "de": "Verstimmung {x} % (Betrag)"},
    "document.controller_unknown": {
        "en": "Unknown key(s) of coil_controller: {keys} – known: {known}.",
        "de": "Unbekannte Angabe(n) in coil_controller: {keys} – bekannt: {known}."},
    "document.detuning_change": {
        "en": "The reactive current changes between the tests ({ivs} A) at an almost equal active current ({lo} … {hi} "
              "A){within} – {cause}.",
        "de": "Der Blindstrom ändert sich zwischen den Versuchen ({ivs} A) bei nahezu gleichem Wirkstrom ({lo} … {hi} "
              "A){within} – {cause}."},
    "document.detuning_cause_unknown": {
        "en": "the tuning changed between the tests; the cause is not known",
        "de": "die Abstimmung hat sich zwischen den Versuchen geändert; die Ursache ist nicht geklärt"},
    "document.detuning_cause_coils": {
        "en": "during the fault the coil currents change by the same amount as the reactive current at a constant zero-"
              "sequence voltage (coil current balance ΔI_[v]/ΔΣI_[L]: {slopes}): the coils were adjusted, e.g. by "
              "the controller, the capacitive earth-fault current of the network stayed the same",
        "de": "im Fehler ändern sich die Spulenströme bei gleicher Nullspannung um denselben Betrag wie der Blindstrom "
              "(Spulenstrombilanz ΔI_[v]/ΔΣI_[L]: {slopes}): Die Spulen wurden verstellt, z. B. durch den Regler, "
              "der kapazitive Erdschlussstrom des Netzes blieb gleich"},
    "document.detuning_within": {"en": "; in {names} also during the fault (appendix B)",
                                 "de": "; in {names} auch während des Fehlers (Anhang B)"},
    "document.station_vs_busbar": {
        "en": "{station} shows a larger current than {test} at the {orders} order ({currents}), although the line to "
              "the station and its earthing add to the fault impedance: the voltage before the fault was larger "
              "({voltages}). Per order the network acts from the fault location as an active two-terminal network "
              "(equivalent source, I_[m] = U_[q,m]/|Z_[q,m] + Z|) – the current follows the source voltage. Referred "
              "to the voltage the current at the station is smaller ({ratios} A/V).",
        "de": "{station} zeigt bei der {orders} Ordnung einen größeren Strom als {test} ({currents}), obwohl die "
              "Leitung zur Station und deren Erdung die Fehlerimpedanz vergrößern: Die Spannung vor dem Fehler war "
              "größer ({voltages}). Je Ordnung wirkt das Netz von der Fehlerstelle aus als aktiver Zweipol "
              "(Ersatzquelle, I_[m] = U_[q,m]/|Z_[q,m] + Z|) – der Strom folgt der Quellenspannung. Bezogen auf die "
              "Spannung ist der Strom an der Station kleiner ({ratios} A/V)."},
    "document.versus": {"en": "{a} against {b}", "de": "{a} gegenüber {b}"},
    "document.even_orders_note": {
        "en": "In {name} the even orders are clearly larger than in the other tests (together {x} A against at most "
              "{y} A){decay}; even orders are not scaled.",
        "de": "In {name} sind die geraden Ordnungen deutlich größer als in den übrigen Versuchen (zusammen {x} A "
              "gegenüber höchstens {y} A){decay}; gerade Ordnungen werden nicht skaliert."},
    "document.even_orders_decay": {
        "en": "; they decay in the steady state from {a} A to {b} A ({t0} … {t1} s) – a transient after the fault "
              "inception, the assessed window ({t} s) lies at its beginning",
        "de": "; im stationären Bereich klingen sie von {a} A auf {b} A ab ({t0} … {t1} s) – ein Ausgleichsvorgang "
              "nach dem Fehlereintritt, das bewertete Fenster ({t} s) liegt an seinem Anfang"},
    "document.scaling_context": {
        "en": "Classification of the factors: in the tests the {order}th order lay at the {p_all} percentile of the "
              "power-quality measurement{window}. Its median is {u_day} V on the weekday of the tests ({day}) and "
              "{u_lo} … {u_hi} V over the days of the week; the quantile ({q} V) stands for the times of a high "
              "harmonic voltage. Hence the factors k_[m] ({ks}).",
        "de": "Einordnung der Faktoren: In den Versuchen lag die {order}. Ordnung im {p_all} Perzentil der "
              "PQ-Messung{window}. Ihr Median beträgt am Wochentag der Versuche ({day}) {u_day} V, über die Wochentage "
              "{u_lo} … {u_hi} V; das Quantil ({q} V) steht für die Zeiten hoher Oberschwingungsspannung. Daraus "
              "folgen die Faktoren k_[m] ({ks})."},
    "document.scaling_low": {
        "en": "In the tests the {orders} order lay at the {p} percentile of the power-quality measurement – below almost "
              "all of its values –, and its factor k_[m] = {k} raises it to a level the network did not have on the day "
              "of the tests. Check that the power-quality measurement and the tests have the same network and "
              "switching state (feeding HV network, coupling of the busbars) before relying on it.",
        "de": "Die {orders}. Ordnung lag in den Versuchen im {p} Perzentil der PQ-Messung – unter fast allen ihren "
              "Werten –, und ihr Faktor k_[m] = {k} hebt sie auf einen Pegel, den das Netz am Versuchstag nicht hatte. "
              "Vor der Verwendung prüfen, ob PQ-Messung und Versuche denselben Netz- und Schaltzustand haben "
              "(vorgelagertes HS-Netz, Kupplung der Sammelschienen)."},
    "document.scaling_context_window": {
        "en": ", within the same time window of the week at the {p} percentile (table {t})",
        "de": ", im selben Zeitfenster der Woche im {p} Perzentil (Tabelle {t})"},
    "document.es_basis": {"en": "[^{basis}^]", "de": "[^{basis}^]"},
    "document.es_decomposition": {
        "en": "Compared with {ires} A as measured, only the harmonic content grows ({ih} → {ihs} A); the fundamental "
              "({ig} A) stays{ref}.",
        "de": "Gegenüber {ires} A wie gemessen wächst allein der Oberschwingungsanteil ({ih} → {ihs} A); die "
              "Grundschwingung ({ig} A) bleibt{ref}."},
    "document.es_decomposition_ref": {"en": " – the factors are classified in section {sec}",
                                      "de": " – Einordnung der Faktoren in Abschnitt {sec}"},
    "document.lc_intro": {
        "en": "The limits of the network operator hold under conditions on the earthing systems. Every row assesses "
              "I_[RES,ass] at the largest earthing impedance Z_[E] of its case – the residual current falls with "
              "Z_[E], the earth potential rise U_[E] = r·I_[RES,ass]·Z_[E] rises – and compares U_[E] with "
              "F·U_[Tp] = {u} V.",
        "de": "Die Grenzwerte des Netzbetreibers gelten unter Bedingungen an die Erdungsanlagen. Jede Zeile bewertet "
              "I_[RES,Bew] an der größten Erdungsimpedanz Z_[E] ihres Falls – der Reststrom sinkt mit Z_[E], die "
              "Erdungsspannung U_[E] = r·I_[RES,Bew]·Z_[E] steigt – und vergleicht U_[E] mit F·U_[Tp] = {u} V."},
    "document.lc_intro_no_u": {
        "en": "The limits of the network operator hold under conditions on the earthing systems. Every row assesses "
              "I_[RES,ass] at the largest earthing impedance Z_[E] of its case (the residual current falls with "
              "Z_[E]); U_[Tp] or F is not set.",
        "de": "Die Grenzwerte des Netzbetreibers gelten unter Bedingungen an die Erdungsanlagen. Jede Zeile bewertet "
              "I_[RES,Bew] an der größten Erdungsimpedanz Z_[E] ihres Falls (der Reststrom sinkt mit Z_[E]); U_[Tp] "
              "oder F ist nicht gesetzt."},
    "document.lc_busbar": {
        "en": "Without equivalent sources the assessment value of the largest test (busbar) is used – an upper bound "
              "for every Z_[E].",
        "de": "Ohne Ersatzquellen wird der Bewertungswert des größten Versuchs (Sammelschiene) verwendet – eine obere "
              "Schranke für jedes Z_[E]."},
    "document.tab_limit_cases": {"en": "Limits and their conditions", "de": "Grenzwerte und ihre Bedingungen"},
    "document.fig_ires_ze_ranges": {"en": "; mark: the largest Z_[E] of the condition of a limit",
                                    "de": "; Strich: größte Z_[E] der Bedingung eines Grenzwerts"},
    "document.col_limit": {"en": "Limit in A", "de": "Grenzwert in A"},
    "document.col_condition": {"en": "Condition", "de": "Bedingung"},
    "document.col_vs_limit": {"en": "Current", "de": "Strom"},
    "document.col_vs_utp": {"en": "U_[E] ≤ F·U_[Tp]", "de": "U_[E] ≤ F·U_[Tp]"},
    "document.lc_standard": {"en": "standard", "de": "Standard"},
    "document.lc_applies": {"en": " (applies)", "de": " (gilt)"},
    "document.lc_result_ok": {
        "en": "Applies: the limit {limit} A – kept with I_[RES,ass] = {i} A{u_text}.",
        "de": "Es gilt der Grenzwert {limit} A: Mit I_[RES,Bew] = {i} A eingehalten{u_text}."},
    "document.lc_result_exceeded": {
        "en": "Applies: the limit {limit} A – I_[RES,ass] = {i} A exceeds it by {pct} %{u_text}.",
        "de": "Es gilt der Grenzwert {limit} A: I_[RES,Bew] = {i} A überschreitet ihn um {pct} %{u_text}."},
    "document.lc_u_ok": {
        "en": "; the condition behind it, U_[E] ≤ F·U_[Tp], is kept with {ue} V ({share} % of {u} V)",
        "de": "; die Bedingung dahinter, U_[E] ≤ F·U_[Tp], ist mit {ue} V eingehalten ({share} % von {u} V)"},
    "document.lc_u_exceeded": {
        "en": "; U_[E] = {ue} V exceeds F·U_[Tp] = {u} V",
        "de": "; U_[E] = {ue} V überschreitet F·U_[Tp] = {u} V"},
    "document.lc_others": {
        "en": "The other limits hold only if their condition is met: {items}.",
        "de": "Die übrigen Grenzwerte gelten nur, wenn ihre Bedingung erfüllt ist: {items}."},
    "document.lc_item": {"en": "{limit} A ({condition}): {i} A {status}",
                         "de": "{limit} A ({condition}): {i} A {status}"},
    # ---- network expansion
    "document.sec_expansion": {"en": "Network expansion", "de": "Netzausbau"},
    "document.expansion_intro": {
        "en": "Network planning asks whether further cables can be connected at the substation: up to which capacitive "
              "earth-fault current I_[CE] does the residual current keep the limits? Further cables raise the earth "
              "and the operating capacitance of the network, and the coils are retuned to the same detuning. The "
              "equivalent sources of the tests (section {sec}) describe the network at every order m as an active "
              "two-terminal network with the source voltage U_[q] and the source impedance Z_[q] = (Z_[0] + "
              "2·Z_[1])/3 of the fault loop. The expansion changes its parts; for the expanded network they are "
              "converted as follows (Δ: change by the expansion, d: damping I_[w]/I_[CE] of the new cables, I_[L]: "
              "current of the coils).",
        "de": "Für die Netzplanung stellt sich die Frage, ob am Umspannwerk weitere Kabel angeschlossen werden können: "
              "Bis zu welchem kapazitiven Erdschlussstrom I_[CE] hält der Erdschlussreststrom die Grenzwerte ein? "
              "Weitere Kabel erhöhen die Erd- und die Betriebskapazität des Netzes, die Spulen werden auf dieselbe "
              "Verstimmung nachgestellt. Die Ersatzquellen der Versuche (Abschnitt {sec}) beschreiben das Netz bei "
              "jeder Ordnung m als aktiven Zweipol mit der Quellenspannung U_[q] und der Quellenimpedanz Z_[q] = "
              "(Z_[0] + 2·Z_[1])/3 der Fehlerschleife. Der Ausbau ändert seine Teile; für das ausgebaute Netz werden "
              "sie so umgerechnet (Δ: Änderung durch den Ausbau, d: Dämpfung I_[w]/I_[CE] der neuen Kabel, I_[L]: "
              "Strom der Spulen)."},
    "document.expansion_z0": {
        "en": "In the zero sequence the added earth capacitance with its losses and the retuned coils lie in parallel "
              "with Z_[0]:",
        "de": "Im Nullsystem liegen die zusätzliche Erdkapazität mit ihren Verlusten und die nachgestellten Spulen "
              "parallel zu Z_[0]:"},
    "document.eq_expansion_z0": {
        "en": "1/Z_[0,new] = 1/Z_[0] + (j·m·ΔI_[CE] + d·ΔI_[CE] − j·ΔI_[L]/m) / (3·U_[LE])",
        "de": "1/Z_[0,neu] = 1/Z_[0] + (j·m·ΔI_[CE] + d·ΔI_[CE] − j·ΔI_[L]/m) / (3·U_[LE])"},
    "document.expansion_z1": {
        "en": "in the positive and negative sequence the operating capacitance of the new cables (c = C_[1]/C_[E] = "
              "{c}) lies in parallel with the busbar – it shifts the parallel resonance of the busbar to lower orders:",
        "de": "im Mit- und Gegensystem liegt die Betriebskapazität der neuen Kabel (c = C_[1]/C_[E] = {c}) parallel "
              "zur Sammelschiene – sie verschiebt deren Parallelresonanz zu niedrigeren Ordnungen:"},
    "document.eq_expansion_z1": {
        "en": "1/Z_[1,new] = 1/Z_[1] + (j·m·c·ΔI_[CE] + d·ΔI_[CE]) / (3·U_[LE])",
        "de": "1/Z_[1,neu] = 1/Z_[1] + (j·m·c·ΔI_[CE] + d·ΔI_[CE]) / (3·U_[LE])"},
    "document.expansion_zq": {"en": "and with both the source impedance changes:",
                              "de": "und mit beiden ändert sich die Quellenimpedanz:"},
    "document.eq_expansion_zq": {
        "en": "Z_[q,new] = Z_[q] + (Z_[0,new] − Z_[0] + 2·(Z_[1,new] − Z_[1])) / 3",
        "de": "Z_[q,neu] = Z_[q] + (Z_[0,neu] − Z_[0] + 2·(Z_[1,neu] − Z_[1])) / 3"},
    "document.expansion_uq": {
        "en": "The harmonic voltage of the busbar comes from sources outside the new cables (HV network, loads) and "
              "changes with the impedance it acts on,",
        "de": "Die Oberschwingungsspannung der Sammelschiene stammt von Quellen außerhalb der neuen Kabel (HS-Netz, "
              "Lasten) und ändert sich mit der Impedanz, an der sie liegt,"},
    "document.eq_expansion_uq": {"en": "U_[q,new] = U_[q]·Z_[1,new] / Z_[1]", "de": "U_[q,neu] = U_[q]·Z_[1,neu] / Z_[1]"},
    "document.expansion_uq0_coil": {
        "en": "at the orders of the zero sequence (3, 9, 15 …), however, the source arises in the fault (nonlinear "
              "current of the coils) and grows with Z_[0] and with the current of the coils:",
        "de": "bei den Ordnungen des Nullsystems (3, 9, 15 …) entsteht die Quelle dagegen im Erdschluss (nichtlinearer "
              "Strom der Spulen) und wächst mit Z_[0] und mit dem Strom der Spulen:"},
    "document.expansion_uq0_network": {
        "en": "at the orders of the zero sequence (3, 9, 15 …), however, the source arises in the fault (nonlinear "
              "current of the coils) and grows with Z_[0] – the current of the coils taken as unchanged:",
        "de": "bei den Ordnungen des Nullsystems (3, 9, 15 …) entsteht die Quelle dagegen im Erdschluss (nichtlinearer "
              "Strom der Spulen) und wächst mit Z_[0] – der Strom der Spulen als unverändert angenommen:"},
    "document.eq_expansion_uq0_coil": {
        "en": "U_[q,new] = U_[q]·(Z_[0,new] / Z_[0])·(I_[L,new] / I_[L])",
        "de": "U_[q,neu] = U_[q]·(Z_[0,neu] / Z_[0])·(I_[L,neu] / I_[L])"},
    "document.eq_expansion_uq0_network": {"en": "U_[q,new] = U_[q]·Z_[0,new] / Z_[0]",
                                          "de": "U_[q,neu] = U_[q]·Z_[0,neu] / Z_[0]"},
    "document.expansion_ig": {
        "en": "The fundamental grows with the active current of the new cables and with the detuning current, which is "
              "proportional to I_[CE]:",
        "de": "Die Grundschwingung wächst mit dem Wirkstrom der neuen Kabel und mit dem Verstimmungsstrom, der "
              "proportional zu I_[CE] ist:"},
    "document.eq_expansion_ig": {
        "en": "I_[G,new] = √((I_[w] + d·ΔI_[CE])² + (I_[v]·I_[CE,new] / I_[CE])²)",
        "de": "I_[G,neu] = √((I_[w] + d·ΔI_[CE])² + (I_[v]·I_[CE,neu] / I_[CE])²)"},
    "document.expansion_split": {
        "en": "The split of Z_[q] into Z_[0] and Z_[1] is not part of the equivalent sources. It is taken in two ways; "
              "the larger residual current of the two is the result:",
        "de": "Die Aufteilung von Z_[q] auf Z_[0] und Z_[1] ist nicht Teil der Ersatzquellen. Sie wird auf zwei Wegen "
              "bestimmt; maßgebend ist der größere Reststrom der beiden:"},
    "document.expansion_split_one": {
        "en": "The split of Z_[q] into Z_[0] and Z_[1] is not part of the equivalent sources; it is taken from the "
              "network data (the busbar voltages of the tests needed for a second way are not available):",
        "de": "Die Aufteilung von Z_[q] auf Z_[0] und Z_[1] ist nicht Teil der Ersatzquellen; sie wird aus den "
              "Netzdaten bestimmt (die Spannungen der Sammelschiene in den Versuchen für einen zweiten Weg fehlen):"},
    "document.expansion_split_sources": {
        "en": "**network data** – Z_[0] = 1/(j·(m/X_[C0] − 1/(m·X_[L0]))) + d·X_[C0] with X_[C0] = 3·U_[LE]/I_[CE] and "
              "X_[L0] = 3·U_[LE]/I_[L] (I_[CE] = {ice} A, I_[L] = {il} A, d = {d} %, U_[LE] = {ule} V), and Z_[1] = "
              "(3·Z_[q] − Z_[0])/2;",
        "de": "**Netzdaten** – Z_[0] = 1/(j·(m/X_[C0] − 1/(m·X_[L0]))) + d·X_[C0] mit X_[C0] = 3·U_[LE]/I_[CE] und "
              "X_[L0] = 3·U_[LE]/I_[L] (I_[CE] = {ice} A, I_[L] = {il} A, d = {d} %, U_[LE] = {ule} V) und Z_[1] = "
              "(3·Z_[q] − Z_[0])/2;"},
    "document.expansion_split_busbar": {
        "en": "**busbar voltages** – Z_[0] and Z_[1] from the change of the busbar voltages in the tests, Z_[s] = "
              "3·ΔU_[s]/I_[F] (related to the faulted phase; median of the tests; Z_[1] from the sequence that is not "
              "the natural one of the order; at the orders of the zero sequence Z_[0] = 3·Z_[q] − 2·Z_[1], because a "
              "source arises there in the fault).",
        "de": "**Sammelschienenspannungen** – Z_[0] und Z_[1] aus der Änderung der Spannungen der Sammelschiene in "
              "den Versuchen, Z_[s] = 3·ΔU_[s]/I_[F] (bezogen auf den Fehlerleiter; Median der Versuche; Z_[1] aus dem "
              "Symmetriesystem, das nicht das natürliche der Ordnung ist; bei den Ordnungen des Nullsystems Z_[0] = "
              "3·Z_[q] − 2·Z_[1], da dort im Erdschluss eine Quelle entsteht)."},
    "document.expansion_assumptions": {"en": "Assumptions", "de": "Annahmen"},
    "document.expansion_a_cables": {
        "en": "The new cables are connected at the substation; their operating capacitance is C_[1] = c·C_[E] with "
              "c = {c}{screen}, their damping d = {d} %{same}.",
        "de": "Die neuen Kabel werden am Umspannwerk angeschlossen; ihre Betriebskapazität ist C_[1] = c·C_[E] mit "
              "c = {c}{screen}, ihre Dämpfung d = {d} %{same}."},
    "document.expansion_a_screen": {"en": " (cables with a screen per core)", "de": " (Kabel mit Schirm je Ader)"},
    "document.expansion_a_same": {"en": " as that of the network today", "de": " wie die des heutigen Netzes"},
    "document.expansion_a_coils": {
        "en": "The coils are retuned to the detuning of today: I_[L,new] = I_[L]·I_[CE,new]/I_[CE] (today I_[L] = "
              "{il} A){range}.",
        "de": "Die Spulen werden auf die heutige Verstimmung nachgestellt: I_[L,neu] = I_[L]·I_[CE,neu]/I_[CE] (heute "
              "I_[L] = {il} A){range}."},
    "document.expansion_a_range": {
        "en": "; their range reaches I_[L] = {ilmax} A, that is I_[CE] ≈ {icemax} A",
        "de": "; ihr Stellbereich reicht bis I_[L] = {ilmax} A, d. h. bis I_[CE] ≈ {icemax} A"},
    "document.expansion_a_range_unknown": {
        "en": "; whether their range suffices is to be checked",
        "de": "; ob ihr Stellbereich reicht, ist zu prüfen"},
    "document.expansion_a_sources": {
        "en": "No new harmonic sources come with the cables (e.g. converters); the harmonic voltage of the HV network "
              "and of the loads stays{scaling}.",
        "de": "Mit den Kabeln kommen keine neuen Oberschwingungsquellen hinzu (z. B. Umrichter); die "
              "Oberschwingungsspannung des HS-Netzes und der Lasten bleibt{scaling}."},
    "document.expansion_a_scaling": {
        "en": ", and the scaling to the PQ measurement (factor {k}) and the margins hold as in section {sec}",
        "de": ", die Skalierung auf die PQ-Messung (Faktor {k}) und die Zuschläge gelten wie in Abschnitt {sec}"},
    "document.expansion_a_rest": {
        "en": "Orders outside the model (remainder, even orders) stay as measured.",
        "de": "Ordnungen außerhalb des Modells (Rest, gerade Ordnungen) bleiben wie gemessen."},
    "document.tab_expansion": {
        "en": "Residual current at Z_[E] = {Z} Ω after the network expansion",
        "de": "Erdschlussreststrom an Z_[E] = {Z} Ω nach dem Netzausbau"},
    "document.col_today": {"en": "Today", "de": "Heute"},
    "document.col_planned": {"en": "I_[CE] = {ice} A", "de": "I_[CE] = {ice} A"},
    "document.xrow_ice": {"en": "Capacitive earth-fault current I_[CE] in A",
                          "de": "Kapazitiver Erdschlussstrom I_[CE] in A"},
    "document.xrow_delta": {"en": "Change ΔI_[CE]", "de": "Änderung ΔI_[CE]"},
    "document.xrow_il": {"en": "Current of the coils I_[L] in A", "de": "Strom der Spulen I_[L] in A"},
    "document.xrow_ig": {"en": "Fundamental I_[G] in A", "de": "Grundschwingung I_[G] in A"},
    "document.xrow_ih": {"en": "{ih} in A – {split}", "de": "{ih} in A – {split}"},
    "document.xrow_ires": {"en": "I_[RES,ass] in A – {split}", "de": "I_[RES,Bew] in A – {split}"},
    "document.xrow_prop": {"en": "I_[RES,ass] in A – proportional to I_[CE]", "de": "I_[RES,Bew] in A – proportional zu I_[CE]"},
    "document.xrow_status": {"en": "Limit {limit} A", "de": "Grenzwert {limit} A"},
    "document.split_sources": {"en": "network data", "de": "Netzdaten"},
    "document.split_busbar": {"en": "busbar voltages", "de": "Sammelschiene"},
    "document.split_busbar_short": {"en": "busbar voltages", "de": "Sammel\u00adschiene"},
    "document.tab_expansion_orders": {
        "en": "Harmonic currents at Z_[E] = {Z} Ω and factor of the source voltage at I_[CE] = {ice} A",
        "de": "Oberschwingungsströme an Z_[E] = {Z} Ω und Faktor der Quellenspannung bei I_[CE] = {ice} A"},
    "document.col_i_today": {"en": "Today in A", "de": "Heute in A"},
    "document.col_i_split": {"en": "{split} in A", "de": "{split} in A"},
    "document.col_u_split": {"en": "U_[q,new]/U_[q] {split}", "de": "U_[q,neu]/U_[q] {split}"},
    "document.fig_expansion": {
        "en": "Residual current I_[RES,ass] over the capacitive earth-fault current I_[CE] after a network expansion, "
              "at the earthing impedances of the limits (line: the larger value of the two splits; band: between "
              "them)",
        "de": "Erdschlussreststrom I_[RES,Bew] über dem kapazitiven Erdschlussstrom I_[CE] nach einem Netzausbau, an "
              "den Erdungsimpedanzen der Grenzwerte (Linie: der größere Wert der beiden Aufteilungen; Band: zwischen "
              "ihnen)"},
    "document.tab_expansion_limits": {
        "en": "Capacitive earth-fault current at which the residual current reaches a limit",
        "de": "Kapazitiver Erdschlussstrom, bei dem der Erdschlussreststrom einen Grenzwert erreicht"},
    "document.col_allowed": {"en": "Permissible in A", "de": "Zulässig in A"},
    "document.col_ires_today": {"en": "I_[RES,ass] today in A", "de": "I_[RES,Bew] heute in A"},
    "document.col_reached": {"en": "I_[CE] in A – {split}", "de": "I_[CE] in A – {split}"},
    "document.col_margin": {"en": "Margin in A", "de": "Reserve in A"},
    "document.expansion_now_exceeded": {"en": "today exceeded", "de": "heute überschritten"},
    "document.expansion_beyond": {"en": "> {ice}", "de": "> {ice}"},
    "document.expansion_allowed_u": {
        "en": "Permissible is the smaller of the limit and F·U_[Tp]/(r·Z_[E]) = {u} V/(r·Z_[E]) – the earth potential "
              "rise U_[E] = r·I_[RES,ass]·Z_[E] keeps F·U_[Tp] up to it.",
        "de": "Zulässig ist der kleinere Wert aus Grenzwert und F·U_[Tp]/(r·Z_[E]) = {u} V/(r·Z_[E]) – bis dahin hält "
              "die Erdungsspannung U_[E] = r·I_[RES,Bew]·Z_[E] den Wert F·U_[Tp] ein."},
    "document.expansion_result": {
        "en": "At I_[CE] = {ice} A ({delta}) the residual current at Z_[E] = {Z} Ω is expected at I_[RES,ass] = "
              "{range} A (today {now} A{limit}). The fundamental grows to I_[G] = {ig} A (today "
              "{ig_now} A), the harmonics change to {ih} = {ih_range} A (today {ih_now} A){order}. In proportion to "
              "I_[CE] – every part growing like I_[CE] – it would be {prop} A.",
        "de": "Bei I_[CE] = {ice} A ({delta}) ist an Z_[E] = {Z} Ω ein Erdschlussreststrom von I_[RES,Bew] = {range} A "
              "zu erwarten (heute {now} A{limit}). Die Grundschwingung wächst auf I_[G] = "
              "{ig} A (heute {ig_now} A), die Oberschwingungen ändern sich auf {ih} = {ih_range} A (heute "
              "{ih_now} A){order}. Proportional zu I_[CE] hochgerechnet – jeder Anteil wächst wie I_[CE] – wären es "
              "{prop} A."},
    "document.expansion_vs_limit": {"en": "; limit {limit} A: {status}", "de": "; Grenzwert {limit} A: {status}"},
    "document.expansion_order": {
        "en": "; the voltage of the {m}th order at the busbar changes by the factor {k} (table {t})",
        "de": "; die Spannung der {m}. Ordnung an der Sammelschiene ändert sich um den Faktor {k} (Tabelle {t})"},
    "document.expansion_limit_reached": {
        "en": "Limit {limit} A ({condition}, Z_[E] = {Z} Ω{u}): reached at I_[CE] ≈ {ice} A – a margin of about {margin} A "
              "over today{split}.",
        "de": "Grenzwert {limit} A ({condition}, Z_[E] = {Z} Ω{u}): erreicht bei I_[CE] ≈ {ice} A – eine Reserve von etwa "
              "{margin} A gegenüber heute{split}."},
    "document.expansion_limit_u": {"en": "; permissible {allowed} A because of U_[E] ≤ F·U_[Tp]",
                                   "de": "; zulässig {allowed} A wegen U_[E] ≤ F·U_[Tp]"},
    "document.expansion_limit_split": {
        "en": " (network data {a} A, busbar voltages {b} A)",
        "de": " (Netzdaten {a} A, Sammelschiene {b} A)"},
    "document.expansion_limit_exceeded": {
        "en": "Limit {limit} A ({condition}, Z_[E] = {Z} Ω{u}): exceeded already today (I_[RES,ass] = {now} A).",
        "de": "Grenzwert {limit} A ({condition}, Z_[E] = {Z} Ω{u}): bereits heute überschritten (I_[RES,Bew] = {now} A)."},
    "document.expansion_limit_not_reached": {
        "en": "Limit {limit} A ({condition}, Z_[E] = {Z} Ω{u}): not reached up to I_[CE] = {ice} A.",
        "de": "Grenzwert {limit} A ({condition}, Z_[E] = {Z} Ω{u}): bis I_[CE] = {ice} A nicht erreicht."},
    "document.expansion_coils_short": {
        "en": "The range of the coils ends at I_[CE] ≈ {ice} A; beyond it a further coil is needed.",
        "de": "Der Stellbereich der Spulen endet bei I_[CE] ≈ {ice} A; darüber ist eine weitere Spule nötig."},
    "document.expansion_negative": {
        "en": "At the orders {orders} one split gives an impedance with a negative real part – their data do not fully "
              "match a passive network (phase errors at the harmonic frequency, a drifting source), and their change "
              "is uncertain.",
        "de": "Bei den Ordnungen {orders} ergibt eine Aufteilung eine Impedanz mit negativem Realteil – ihre Daten "
              "passen nicht ganz zu einem passiven Netz (Winkelfehler bei der Oberschwingungsfrequenz, eine driftende "
              "Quelle), ihre Änderung ist unsicher."},
    "document.expansion_negative_one": {
        "en": "At the {orders}th order one split gives an impedance with a negative real part – its data do not fully "
              "match a passive network (phase errors at the harmonic frequency, a drifting source), and its change is "
              "uncertain.",
        "de": "Bei der {orders}. Ordnung ergibt eine Aufteilung eine Impedanz mit negativem Realteil – ihre Daten "
              "passen nicht ganz zu einem passiven Netz (Winkelfehler bei der Oberschwingungsfrequenz, eine driftende "
              "Quelle), ihre Änderung ist unsicher."},
    "document.tz_local": {"en": "local time", "de": "Ortszeit"},
    "document.expansion_rejected_still": {
        "en": " With the network data, too, the real part stays negative at the orders {orders} – their change is "
              "uncertain.",
        "de": " Auch mit den Netzdaten bleibt der Realteil bei den Ordnungen {orders} negativ – ihre Änderung ist "
              "unsicher."},
    "document.expansion_rejected_still_one": {
        "en": " With the network data, too, its real part stays negative – its change is uncertain.",
        "de": " Auch mit den Netzdaten bleibt ihr Realteil negativ – ihre Änderung ist unsicher."},
    "document.expansion_rejected": {
        "en": "At the orders {orders} the split from the busbar voltages gives an impedance with a negative real part – "
              "no passive network, and a parallel resonance computed with it would be an artefact. For these orders "
              "both splits use the network data.",
        "de": "Bei den Ordnungen {orders} ergibt die Aufteilung über die Sammelschienenspannungen eine Impedanz mit "
              "negativem Realteil – das ist kein passives Netz, und eine damit berechnete Parallelresonanz wäre ein "
              "Rechenartefakt. Für diese Ordnungen verwenden beide Aufteilungen die Netzdaten."},
    "document.expansion_rejected_one": {
        "en": "At the {orders}th order the split from the busbar voltages gives an impedance with a negative real "
              "part – no passive network, and a parallel resonance computed with it would be an artefact. For this "
              "order both splits use the network data.",
        "de": "Bei der {orders}. Ordnung ergibt die Aufteilung über die Sammelschienenspannungen eine Impedanz mit "
              "negativem Realteil – das ist kein passives Netz, und eine damit berechnete Parallelresonanz wäre ein "
              "Rechenartefakt. Für diese Ordnung verwenden beide Aufteilungen die Netzdaten."},
    "document.expansion_validity": {
        "en": "The conversion extrapolates the network model of the tests: it holds for a network built as today and "
              "for the assumptions above; the difference of the two splits shows how well the split of Z_[q] is "
              "known. After the expansion a new measurement checks the value.",
        "de": "Die Umrechnung extrapoliert das Netzmodell der Versuche: Sie gilt für ein Netz, das wie heute aufgebaut "
              "ist, und für die Annahmen oben; der Unterschied der beiden Aufteilungen zeigt, wie gut die Aufteilung "
              "von Z_[q] bekannt ist. Nach dem Ausbau prüft eine neue Messung den Wert."},
    "document.expansion_failed": {
        "en": "Network expansion not computed: {error}", "de": "Netzausbau nicht berechnet: {error}"},
    "document.expansion_invalid": {
        "en": "expansion is not valid: {error} – expected True, False or a dict with I_CE (planned capacitive "
              "earth-fault current(s) in A), damping_added (≥ 0), c1_ratio (> 0), I_L_max (A), I_CE_max (A) and "
              "zero_sequence (\"coil\" or \"network\").",
        "de": "expansion ist ungültig: {error} – erwartet: True, False oder ein dict mit I_CE (geplanter kapazitiver "
              "Erdschlussstrom in A, auch mehrere), damping_added (≥ 0), c1_ratio (> 0), I_L_max (A), I_CE_max (A) "
              "und zero_sequence (\"coil\" oder \"network\")."},
    "document.kr_expansion": {
        "en": "Network expansion to I_[CE] = {ice} A – I_[RES,ass] at Z_[E] = {Z} Ω",
        "de": "Netzausbau auf I_[CE] = {ice} A – I_[RES,Bew] an Z_[E] = {Z} Ω"},
    "document.kr_expansion_limit": {
        "en": "I_[CE] at which I_[RES,ass] reaches {limit} A (Z_[E] = {Z} Ω)",
        "de": "I_[CE], bei dem I_[RES,Bew] {limit} A erreicht (Z_[E] = {Z} Ω)"},
    "document.kr_today": {"en": "today {ice}\u00a0A", "de": "heute {ice}\u00a0A"},
    "document.sum_expansion": {
        "en": "Network expansion: at I_[CE] = {ice} A the residual current at Z_[E] = {Z} Ω is expected at "
              "I_[RES,ass] ≈ {value} A{limit}.",
        "de": "Netzausbau: Bei I_[CE] = {ice} A ist an Z_[E] = {Z} Ω ein Erdschlussreststrom von I_[RES,Bew] ≈ "
              "{value} A zu erwarten{limit}."},
    "document.sum_expansion_limit": {"en": " (limit {limit} A: {status}{already})",
                                     "de": " (Grenzwert {limit} A: {status}{already})"},
    "document.sum_expansion_already": {"en": ", already today", "de": ", bereits heute"},
    "document.sum_expansion_reached": {"en": "It reaches {items}.", "de": "Er erreicht {items}."},
    "document.sum_expansion_reached_item": {
        "en": "the limit {limit} A (Z_[E] = {Z} Ω) at I_[CE] ≈ {ice} A",
        "de": "den Grenzwert {limit} A (Z_[E] = {Z} Ω) bei I_[CE] ≈ {ice} A"},
    "document.sum_expansion_not_reached": {
        "en": "It does not reach {items} up to I_[CE] = {ice} A.",
        "de": "Bis I_[CE] = {ice} A erreicht er {items} nicht."},
    "document.sum_expansion_exceeded": {"en": "It exceeds {items} already today.",
                                        "de": "Bereits heute überschreitet er {items}."},
    "document.sum_expansion_item": {"en": "the limit {limit} A (Z_[E] = {Z} Ω)",
                                    "de": "den Grenzwert {limit} A (Z_[E] = {Z} Ω)"},
    "document.u_expansion": {"en": "Network expansion: split of Z_[q]", "de": "Netzausbau: Aufteilung von Z_[q]"},
    "document.u_expansion_value": {
        "en": "I_[RES,ass] at I_[CE] = {ice} A: {a} A (network data), {b} A (busbar voltages)",
        "de": "I_[RES,Bew] bei I_[CE] = {ice} A: {a} A (Netzdaten), {b} A (Sammelschiene)"},
    "document.u_expansion_treat": {"en": "the larger value", "de": "der größere Wert"},
    "document.forecast_3_expansion": {
        "en": "Growth of the network: the active part of the fundamental grows with I_[CE] (I_[w] = d·I_[CE], d = {d} %); "
              "the harmonics do not grow in proportion (resonance) – the next section converts the equivalent "
              "sources to the expanded network.",
        "de": "Netzausbau: Der Wirkanteil der Grundschwingung wächst mit I_[CE] (I_[w] = d·I_[CE], d = {d} %); die "
              "Oberschwingungen wachsen nicht proportional (Resonanz) – der folgende Abschnitt rechnet die "
              "Ersatzquellen auf das ausgebaute Netz um."},
    "document.limit_case_invalid": {
        "en": "Limit case {k} is not valid: {error} – expected a dict with limit_A > 0, Z_E > 0 (optional, default: "
              "Z_E of the report), r in (0, 1] (optional, default 1) and condition (text).",
        "de": "Grenzwertfall {k} ist ungültig: {error} – erwartet: ein dict mit limit_A > 0, Z_E > 0 (optional, "
              "Standard: Z_E des Berichts), r in (0, 1] (optional, Standard 1) und condition (Text)."},
    "document.station_detail_invalid": {
        "en": "station_detail must be \"full\" or \"compact\" – got {value!r}.",
        "de": "station_detail muss \"full\" oder \"compact\" sein – erhalten: {value!r}."},
    "document.kr_ue_case": {"en": "U_[E] = r·I_[RES,ass]·Z_[E] at Z_[E] = {Z} Ω, r = {r}",
                            "de": "U_[E] = r·I_[RES,Bew]·Z_[E] an Z_[E] = {Z} Ω, r = {r}"},
    "document.kr_limit_case": {"en": "I_[RES,ass] at Z_[E] = {Z} Ω if {condition}",
                               "de": "I_[RES,Bew] an Z_[E] = {Z} Ω, wenn {condition}"},
    "document.ze_permissible_flat": {"en": ", {z} Ω with the flat r = {r}", "de": ", {z} Ω mit dem pauschalen r = {r}"},
    "document.station_probes_compact": {
        "en": "Measured against the earth probes {probes}: at most {umax} V{limit}{diff}.",
        "de": "Gemessen wurde gegen die Erdungsspieße {probes}: höchstens {umax} V{limit}{diff}."},
    "document.station_probe_compact": {
        "en": "Measured against the earth probe {probes}{limit}.",
        "de": "Gemessen wurde gegen den Erdungsspieß {probes}{limit}."},
    "document.probes_limit": {"en": ", {share} % of F·U_[Tp] = {u} V", "de": ", {share} % von F·U_[Tp] = {u} V"},
    "document.probes_diff": {
        "en": "; the differences between the probes (e.g. contact resistance or position of the probes) are without "
              "significance at this level",
        "de": "; die Unterschiede zwischen den Spießen (z. B. Übergangswiderstände oder Lage der Spieße) sind bei "
              "dieser Größe ohne Bedeutung"},
    "document.station_earthing_compact": {
        "en": "The earthing measurement of the station[^Current injection{details}; groundmeas measurement {ids}.^] "
              "gives Z_[E] = {z} Ω at 50 Hz.",
        "de": "Die Erdungsmessung der Station[^Stromeinspeisung{details}; groundmeas-Messung {ids}.^] ergibt "
              "Z_[E] = {z} Ω bei 50 Hz."},
    "document.k_item": {"en": "{m}th order {k}", "de": "{m}. Ordnung {k}"},
    "document.order_range": {"en": "{m}th order {lo} … {hi} A", "de": "{m}. Ordnung {lo} … {hi} A"},
    "document.weekdays": {"en": "Monday,Tuesday,Wednesday,Thursday,Friday,Saturday,Sunday",
                          "de": "Montag,Dienstag,Mittwoch,Donnerstag,Freitag,Samstag,Sonntag"},
    "document.u_stationarity_orders": {"en": "; orders: {items}", "de": "; Ordnungen: {items}"},
    "document.app_symbols": {"en": "Symbols", "de": "Formelzeichen"},
    "document.tab_symbols": {"en": "Symbols", "de": "Formelzeichen"},
    "document.col_symbol": {"en": "Symbol", "de": "Zeichen"},
    "document.col_meaning": {"en": "Meaning", "de": "Bedeutung"},
})

#: symbols of the appendix (symbol in the mark-up of the report texts, message key of the meaning)
_SYMBOLS = (
    ("I_[RES]", {"en": "residual earth-fault current (RMS value)", "de": "Erdschlussreststrom (Effektivwert)"}),
    ("I_[G]", {"en": "fundamental of the residual current", "de": "Grundschwingung des Reststroms"}),
    ("I_[w], I_[v]", {"en": "active and reactive (detuning) current", "de": "Wirk- und Blindstrom (Verstimmung)"}),
    ("I_[H]", {"en": "harmonic content: everything above the fundamental (orders ≥ 2 and the content between the "
                     "orders), without the low-frequency part",
               "de": "Oberschwingungsanteil: alles oberhalb der Grundschwingung (Ordnungen ≥ 2 und Anteile zwischen "
                     "den Ordnungen), ohne den niederfrequenten Anteil"}),
    ({"en": "I_[lf]", "de": "I_[nf]"}, {"en": "low-frequency part of the residual current (below 50 Hz)",
                                        "de": "niederfrequenter Anteil des Reststroms (unter 50 Hz)"}),
    ("THD_[I]", {"en": "harmonic distortion of the current, I_[H]/I_[G]", "de": "Oberschwingungsgehalt des Stroms, "
                                                                            "I_[H]/I_[G]"}),
    ({"en": "I_[H,scaled], I_[H,ass]", "de": "I_[H,skaliert], I_[H,Bew]"},
     {"en": "harmonic content scaled to the PQ measurement; the same with the margins (assessment value)",
      "de": "auf die PQ-Messung skalierter Oberschwingungsanteil; derselbe mit den Zuschlägen (Bewertungswert)"}),
    ({"en": "I_[zero], I_[even], I_[other]", "de": "I_[Null], I_[gerade], I_[sonst]"},
     {"en": "harmonic content not scaled: zero sequence, even orders, content between the orders",
      "de": "nicht skalierter Oberschwingungsanteil: Nullsystem, gerade Ordnungen, Anteile zwischen den Ordnungen"}),
    ("I_[RES,max]", {"en": "residual current with the scaled harmonic content, without the margins",
                     "de": "Erdschlussreststrom mit dem skalierten Oberschwingungsanteil, ohne Zuschläge"}),
    ({"en": "I_[RES,ass], I_[G,ass]", "de": "I_[RES,Bew], I_[G,Bew]"},
     {"en": "assessment value of the residual current (with the margins m_[I], m_[U]) – of a test or at Z_[E]; "
            "fundamental with the margin",
      "de": "Bewertungswert des Erdschlussreststroms (mit den Zuschlägen m_[I], m_[U]) – eines Versuchs bzw. an "
            "Z_[E]; Grundschwingung mit Zuschlag"}),
    ("I_[m]", {"en": "current of the harmonic order m", "de": "Strom der Ordnung m"}),
    ("I_[CE]", {"en": "capacitive earth-fault current", "de": "kapazitiver Erdschlussstrom"}),
    ("d, v", {"en": "damping and detuning (share of I_[CE])", "de": "Dämpfung und Verstimmung (Anteil an I_[CE])"}),
    ("U_[0]", {"en": "zero-sequence (displacement) voltage", "de": "Nullspannung (Verlagerungsspannung)"}),
    ("U_[m], U_[95,m]", {"en": "voltage of order m in the test (positive- or negative-sequence component, mean of "
                               "before and after the fault) and the quantile of the largest phase of the PQ measurement",
                         "de": "Spannung der Ordnung m im Versuch (Mit- bzw. Gegensystemkomponente, Mittel aus vor und "
                               "nach dem Fehler) und Quantil des größten Leiters der PQ-Messung"}),
    ("k_[m], k_[s]", {"en": "scaling factor of order m, factor on the quantile",
                      "de": "Skalierungsfaktor der Ordnung m, Faktor auf das Quantil"}),
    ("m_[I], m_[U]", {"en": "measurement-uncertainty margins of currents and voltages",
                      "de": "Messunsicherheitszuschläge der Ströme und Spannungen"}),
    ("U_[q,m], Z_[q,m] = R_[q] + jX_[q]", {"en": "source voltage and impedance of the equivalent source of order m",
                                         "de": "Quellenspannung und -impedanz der Ersatzquelle der Ordnung m"}),
    ("Z_[E]", {"en": "earthing impedance", "de": "Erdungsimpedanz"}),
    ({"en": "Z_[E,perm]", "de": "Z_[E,zul]"}, {"en": "permissible earthing impedance",
                                             "de": "zulässige Erdungsimpedanz"}),
    ("r, r_[eff]", {"en": "reduction factor (per order; effective over the orders)",
                    "de": "Reduktionsfaktor (je Ordnung; wirksam über die Ordnungen)"}),
    ("U_[E], U_[T], U_[Tp]", {"en": "earth potential rise, touch voltage, permissible touch voltage",
                              "de": "Erdungsspannung, Berührungsspannung, zulässige Berührungsspannung"}),
    ("F", {"en": "ratio of earth potential rise to touch voltage", "de": "Verhältnis Erdungs- zu Berührungsspannung"}),
    ("I_[F], I_[S]", {"en": "fault current and shield current at the station",
                      "de": "Fehlerstrom und Schirmstrom an der Station"}),
    ("I_[ON], Z_[E,ON]", {"en": "fault-related current into the LV network (PEN) and the earthing impedance of the LV "
                                "network (U_[E]/I_[ON])",
                          "de": "fehlerbedingter Strom ins Ortsnetz (PEN) und Erdungsimpedanz des Ortsnetzes "
                                "(U_[E]/I_[ON])"}),
    ({"en": None, "de": "HES"}, {"en": "", "de": "Haupterdungsschiene der Station (Legende einer Abbildung)"}),
)


# =================================================================================================== information
@dataclass
class ReportInfo:
    """What the report needs besides the results: title page, texts and the parameters of the assessment.

    Everything is optional; leave out what does not apply. Names (network, places, persons, organisation) appear only
    if you give them. The texts may contain the mark-up of the report: ``**bold**``, ``*italic*``, ``I_[RES]``
    (subscript), ``^[2]`` (superscript) and a footnote ``[^text of the note^]``.

    Attributes
    ----------
    title : str
        Title (default: "Evaluation of staged earth-fault tests" / "Auswertung Erdschlussmessung").
    network : str
        Network area, e.g. "20-kV network of substation X" (line under the title).
    subtitle : str
        A further line on the title page.
    organisation : str
        Organisation (title page; header, footer and logo come from your template).
    authors : sequence of str
        Authors (title page and document properties).
    date : str or datetime.date, optional
        Date of the report (default: today).
    version : str
        Version or status, e.g. "draft".
    further_info : str
        Text of the field "further information" (e.g. order number).
    places : mapping of str to str
        Display name per measuring point (site folder of the recordings, e.g. ``{"SUB": "Substation North"}``);
        default: the location of the sites file, else the folder name.
    timezone : str, optional
        Time zone of the times in the report (e.g. "Europe/Berlin"); the recorder clock is taken as UTC.
    purpose : str
        Text on the purpose of the tests (replaces the standard sentence).
    normal_state, coil_unchanged, auto_tuning : bool
        Standard sentences on the network state (normal switching state; coils not changed; automatic tuning).
    instrument : str
        Recording instrument, e.g. "a transient recorder of type …".
    transducers : mapping of str to str
        Transducer per channel role (``"u_l1e"``, ``"i_res"``, ``"i_coil"``, ``"u_e"`` …) for the table of measured
        quantities (without: the table has no column for them).
    quantity_labels : mapping of str to str
        Own name of a measured quantity, per key of the table of measured quantities (``"u_le"``, ``"i_res"``,
        ``"i_coil"``, ``"i_shield"``, ``"i_pen"``, ``"u_e"``, ``"u_t"`` …), e.g.
        ``{"i_shield": "shield current of a parallel feeder"}``.
    procedure : sequence of str, optional
        Steps of the procedure (default: tuning, fault on, fault off, impedance changed).
    uncertainty_components : sequence of (str, float)
        Accuracy of every device of the measurement chain in %, e.g. ``[("transient recorder", 2.0), …]``.
    coil_controller : mapping, optional
        Display of the coil controller (read before the tests): ``ice_A`` (capacitive earth-fault current I_CE in A),
        ``coils_A`` (current per coil in A, ``{"coil 1": 155.0, …}``), ``detuning_pct`` (detuning v in %, v > 0:
        overcompensation), ``active_A`` (active residual current in A) and ``when`` (text, e.g. "before test 1"); all
        optional. I_CE is ``ice_A`` or – with coil currents and detuning – I_L/(1 + v); it replaces the sum of the
        measured coil currents as I_CE of damping and detuning.
    limit_A : float, optional
        Limit of the residual current of this network (default: the limit of the first of ``limit_cases``, else the
        smallest limit of the call).
    limit_cases : sequence of mapping
        Limits with their conditions on the earthing systems, e.g. ``[{"limit_A": 60, "Z_E": 2.5}, {"limit_A": 90,
        "Z_E": 1.7, "condition": "all Z_E < 1.7 Ω"}, {"limit_A": 120, "Z_E": 2.5, "r": 0.5, "condition": "fully
        cabled, r ≥ 0.5 at every fault location"}]``: for every case the report assesses I_RES,ass at its Z_E (largest
        earthing impedance of the case; default: ``Z_E`` of the call) from the equivalent sources and the earth
        potential rise U_E = r·I_RES,ass·Z_E against F·U_Tp (``r`` default 1).
    Z_E_basis : str
        Where the earthing impedance Z_E of the assessment comes from, e.g. "design value of the earthing systems of
        the network" (a footnote to the result at Z_E).
    assessment : {"auto", "Z_E", "busbar"}
        Decisive value: at the earthing impedance (equivalent sources) or of the largest test (busbar); ``"auto"``:
        Z_E if the equivalent sources and Z_E are given.
    expansion : bool or mapping
        Section *Network expansion*: the residual current at the earthing impedances of the limits when further
        cables are connected at the substation and the coils are retuned to the same detuning, and the capacitive
        earth-fault current at which every limit is reached
        ([`expansion_curve`][groundfault.equivalent_source.expansion_curve], both splits of the source impedance).
        True (default): with the defaults (planned values 1.25, 1.5 and 2 times I_CE of today); False: left out; a
        mapping: ``I_CE`` (planned capacitive earth-fault current in A, one value or several, e.g. ``800``),
        ``damping_added`` (damping I_w/I_C of the new cables; default: that of the network), ``c1_ratio`` (C_1/C_E of
        the new cables; default 1, cables with a screen per core), ``I_L_max`` (largest current of the coils in A),
        ``I_CE_max`` (end of the curve in A; default four times I_CE of today) and ``zero_sequence`` (``"coil"``:
        the source of the zero-sequence orders grows with the current of the coils – default; ``"network"``: not).
        Needs the equivalent sources and I_CE.
    reference_frequency_Hz : float, optional
        Frequency of Z_E of the earthing measurement in the cases 1 and 2 of the touch voltage (default: 50 Hz;
        interpolated if the earthing measurement has no value there, see
        [`touch_voltage_cases`][groundfault.document.touch_voltage_cases]).
    r_flat : float, optional
        Flat reduction factor of case 2 (None: case 2 is left out).
    station_detail : {"full", "compact"}
        Sections of a station: ``"full"`` (default) with the comparison against the earthing measurement, the
        reduction factor, the LV network and the notes; ``"compact"`` only the measured profile (waveforms, earth
        potential rise per order, probes) and the approaches to the touch voltage – for a station that is shown for
        information.
    recommendation : str
        Recommendation of the author (without: a marked placeholder).
    notes : sequence of str
        Further notes on the procedure (special events).
    key_results : bool
        Summary of results in front of the contents.
    equations : bool
        The displayed formulas as Word equations (Office Math, editable in the equation editor of Word; default);
        False: as text.
    appendix : bool
        Appendices (data of the recordings, plausibility checks, symbols).
    waveforms : bool
        Waveforms from the raw recordings (if found).
    threshold_A : float
        Smallest current of an order in the table of the residual current.
    """

    title: str = ""
    network: str = ""
    subtitle: str = ""
    organisation: str = ""
    authors: Sequence[str] = ()
    date: Optional[Union[str, _dt.date]] = None
    version: str = ""
    further_info: str = ""
    places: Mapping[str, str] = field(default_factory=dict)
    timezone: Optional[str] = None
    purpose: str = ""
    normal_state: bool = True
    coil_unchanged: bool = True
    auto_tuning: bool = True
    instrument: str = ""
    transducers: Mapping[str, str] = field(default_factory=dict)
    quantity_labels: Mapping[str, str] = field(default_factory=dict)
    procedure: Optional[Sequence[str]] = None
    uncertainty_components: Sequence[Tuple[str, float]] = ()
    coil_controller: Optional[Mapping[str, Any]] = None
    limit_A: Optional[float] = None
    limit_cases: Sequence[Mapping[str, Any]] = ()
    Z_E_basis: str = ""
    assessment: str = "auto"
    expansion: Union[bool, Mapping[str, Any]] = True
    reference_frequency_Hz: Optional[float] = None
    r_flat: Optional[float] = None
    station_detail: str = "full"
    recommendation: str = ""
    notes: Sequence[str] = ()
    key_results: bool = True
    equations: bool = True
    appendix: bool = True
    waveforms: bool = True
    threshold_A: float = 0.5

    @classmethod
    def from_value(cls, value: Union[None, "ReportInfo", Mapping[str, Any]]) -> "ReportInfo":
        """A ReportInfo from None, a ReportInfo or a dict (unknown keys raise a ValueError)."""
        if value is None:
            return cls()
        if isinstance(value, cls):
            return value
        known = [f.name for f in fields(cls)]
        unknown = [k for k in value if k not in known]
        if unknown:
            raise ValueError(tr("document.info_unknown", keys=", ".join(map(str, unknown)), known=", ".join(known)))
        return cls(**dict(value))


def docx_available() -> bool:
    """True if python-docx can be imported (needed for the Word report)."""
    try:
        import docx  # noqa: F401
        return True
    except Exception:
        return False


# =================================================================================================== formatting
def _n(x, digits: int = 3) -> str:
    """Number with ``digits`` significant digits in the output language (``"–"`` for None/NaN)."""
    return report.number(x, digits)


def _hz(f) -> str:
    """Frequency of an order as an integer in Hz (nominal frequency × order)."""
    return "%d" % int(round(float(f))) if _finite(f) else "–"


def _span(values) -> str:
    """Range of numbers as the text shows them (47,5–48,1); one number when both ends look the same (63,3)."""
    lo, hi = _n(min(values)), _n(max(values))
    return lo if lo == hi else "%s–%s" % (lo, hi)


def _es_scaling_text(rs: Mapping[str, Any]) -> Tuple[str, str]:
    """How the scaled result of the equivalent sources at Z_E scales the orders of the positive or negative sequence –
    its own quantile, the THD factor or both, from the rules of its table (orders of the model and remainder) – and
    what stays as measured: the zero sequence, the even orders and the content between the orders, or the rest
    outside the orders of the model when it is not resolved per order."""
    own_rule, thd_rule, none_rule = fml.SCALING_RULES[0], fml.SCALING_RULES[1], fml.SCALING_RULES[3]
    tab = rs.get("table")
    if tab is not None and len(tab):
        pn = tab[tab["category"] == "pos_neg"]
        rules = set(pn["rule"])
        thd = pn.loc[pn["rule"] == thd_rule, "factor"]
        between = bool((tab["part"] == "remainder").any())
    else:                                                    # no table: the rules of the orders of the model
        rules = {r for m, r in (rs.get("rules") or {}).items() if fml.harmonic_category(int(m)) == "pos_neg"}
        f = [k for m, k in (rs.get("factors") or {}).items() if (rs.get("rules") or {}).get(m) == thd_rule]
        thd = pd.Series(f, dtype=float)
        between = False
    f_thd = _n(float(thd.iloc[0])) if len(thd) else "–"
    if own_rule in rules:
        how = (tr("document.es_how_mixed", f=f_thd) if thd_rule in rules else
               tr("document.es_how_own_unscaled") if none_rule in rules else tr("document.es_how_own"))
    else:
        how = tr("document.es_how_thd", f=f_thd) if thd_rule in rules else tr("document.es_how_none")
    return how, tr("document.es_kept_between" if between else "document.es_kept_rest")


def _signed(x, digits: int = 2) -> str:
    """Number with its sign (+1,7 / −2,1)."""
    return ("+" if _finite(x) and float(x) > 0 else "") + _n(x, digits)


def _fixed(x, decimals: int = 1) -> str:
    """Number with a fixed number of decimals."""
    return report.number(x, decimals, "fixed")


#: the home of the package – the reference in the report
_HOMEPAGE = "https://github.com/Ce1ectric/groundfault"


def _sentence(text: str) -> str:
    """A text as a sentence of its own (a footnote): first letter upper case, a full stop at the end."""
    t = str(text).strip()
    if not t:
        return t
    t = t[0].upper() + t[1:]
    return t if t[-1] in ".!?" else t + "."


def _g(x) -> str:
    """Number with up to three significant digits and without trailing zeros (0,95 – not 0,950)."""
    s = _n(x, 3)
    sep = "," if i18n.get_language() == "de" else "."
    return s.rstrip("0").rstrip(sep) if sep in s and "e" not in s.lower() else s


_SYMBOL_TOKEN = re.compile(r"(?<![A-Za-z0-9_])([A-Za-z])_([A-Za-z0-9]+)(?![A-Za-z0-9_\[])")


_ORDINAL = re.compile(r"(?<![\d.,])(\d+)th\b")


def _ordinals(text: str) -> str:
    """English ordinals of the texts: "3th" → "3rd", "21th" → "21st" (the messages write "{order}th")."""
    def fix(m):
        n = int(m.group(1))
        suffix = "th" if 11 <= n % 100 <= 13 else {1: "st", 2: "nd", 3: "rd"}.get(n % 10, "th")
        return "%d%s" % (n, suffix)
    return _ORDINAL.sub(fix, text)


def _markup_symbols(text) -> str:
    """Plain symbols of the evaluation texts (I_G, Z_F, I_RES) in the mark-up of the report (I_[G] …); channel
    names (ISp22_27, UE_5m) stay as they are."""
    return _SYMBOL_TOKEN.sub(r"\1_[\2]", str(text))


def _small(x, digits: int = 2, below: float = 0.001) -> str:
    """A deviation in %: ``< 0,001`` below ``below`` (a rounding difference), else ``digits`` significant digits."""
    if _finite(x) and 0 <= abs(float(x)) < below:
        return tr("document.below", x=report.number(below, 1))
    return _n(x, digits)


def _share(x) -> str:
    """A share in %: two significant digits below 10 %, else an integer."""
    if not _finite(x):
        return "–"
    return _n(x, 2) if abs(float(x)) < 10 else _fixed(x, 0)


_STATUS_ALIASES = {"WARN": "WARNING", "FAIL": "ERROR", "FAILED": "ERROR"}


#: comments of the plausibility checks that result files of earlier versions store in an older wording
_OLD_COMMENTS = {
    "erdschlussfreier Betrieb: Strom nahezu 0 A (DC-Offset/Zangen prüfen)": "evaluation.pre_fault_comment",
    "operation without earth fault: current nearly 0 A (check DC offset/clamps)": "evaluation.pre_fault_comment",
}


def _updated_comment(text) -> str:
    """The comment of a stored plausibility check in the current wording (and the language of the report)."""
    key = _OLD_COMMENTS.get(str(text or "").strip())
    if key:
        from . import evaluation  # noqa: F401  (registers the message)
        return tr(key)
    return str(text or "")


def _status(x) -> str:
    """Status of a check in the language of the report (OK, warning, note, corrected …)."""
    key = _STATUS_ALIASES.get(str(x or "").strip().upper(), str(x or "").strip().upper())
    return tr("document.status_" + key) if i18n.has_message("document.status_" + key) else str(x or "")


def _finite(x) -> bool:
    try:
        return x is not None and math.isfinite(float(x))
    except (TypeError, ValueError):
        return False


_MONTHS = {"de": ["Januar", "Februar", "März", "April", "Mai", "Juni", "Juli", "August", "September", "Oktober",
                  "November", "Dezember"],
           "en": ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
                  "November", "December"]}


def _date_text(d, lang: str) -> str:
    """Date in words: "4. Oktober 2026" / "4 October 2026"."""
    if d is None:
        return "–"
    if isinstance(d, str):
        try:
            d = pd.Timestamp(d).date()
        except (ValueError, TypeError):
            return d
    if isinstance(d, (pd.Timestamp, _dt.datetime)):
        d = d.date()
    m = _MONTHS["de" if lang == "de" else "en"][d.month - 1]
    return ("%d. %s %d" if lang == "de" else "%d %s %d") % (d.day, m, d.year)


def _date_short(d, lang: str) -> str:
    """Date as 04.10.2026 / 2026-10-04."""
    if d is None:
        return "–"
    if isinstance(d, str):
        try:
            d = pd.Timestamp(d)
        except (ValueError, TypeError):
            return d
    return d.strftime("%d.%m.%Y" if lang == "de" else "%Y-%m-%d")


def _local(ts: Optional[pd.Timestamp], tz: Optional[str]) -> Optional[pd.Timestamp]:
    """A recorder time (UTC, naive) in the time zone of the report."""
    if ts is None or pd.isna(ts):
        return None
    ts = pd.Timestamp(ts)
    if ts.tzinfo is None:
        ts = ts.tz_localize("UTC")
    return ts.tz_convert(tz) if tz else ts


def _tz_label(tz: Optional[str], ts: Optional[pd.Timestamp] = None) -> str:
    """Short label of the time zone (abbreviation of the time if known, e.g. CEST/MESZ)."""
    if not tz:
        return "UTC"
    if ts is not None and ts.tzinfo is not None:
        name = ts.tzname() or tz
        return {"CEST": "MESZ", "CET": "MEZ"}.get(name, name) if i18n.get_language() == "de" else name
    return tz


# ---------------------------------------------------------------------------------------------- mark-up
#: inline mark-up of the report texts: **bold**, *italic*, X_[sub], ^[sup], [[placeholder]] and a footnote
#: [^text of the note^] (no braces: the texts are str.format templates)
_TOKEN = re.compile(r"\*\*(?P<b>.+?)\*\*|\*(?P<i>[^*\s][^*]*?)\*|\[\[(?P<ph>.+?)\]\]|"
                    r"(?P<base>[A-Za-zΔφθ]+)_\[(?P<sub>[^\]]*)\]|\^\[(?P<sup>[^\]]*)\]")
_FOOTNOTE = re.compile(r"\[\^(?P<fn>.+?)\^\]", re.S)


def _runs(text: str) -> List[Tuple[str, dict]]:
    """Split a text with the report mark-up into runs ``(text, format)``."""
    out: List[Tuple[str, dict]] = []
    pos = 0
    for m in _TOKEN.finditer(text):
        if m.start() > pos:
            out.append((text[pos:m.start()], {}))
        if m.group("b") is not None:
            out.append((m.group("b"), {"bold": True}))
        elif m.group("i") is not None:
            out.append((m.group("i"), {"italic": True}))
        elif m.group("base") is not None:                       # a symbol: one letter italic, a word upright
            base = m.group("base")
            out.append((base, {"italic": True} if len(base) == 1 else {}))
            out.append((m.group("sub"), {"subscript": True}))
        elif m.group("sup") is not None:
            out.append((m.group("sup"), {"superscript": True}))
        else:
            out.append((m.group("ph"), {"highlight": True}))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], {}))
    return out


def plain(text: str) -> str:
    """Text without the report mark-up and without its footnotes (e.g. for the document properties)."""
    return "".join(t for t, _ in _runs(_FOOTNOTE.sub("", str(text))))


# ----------------------------------------------------------------------------------------- order of the WordprocessingML
# Word expects the children of a property element in the order of the schema (ECMA-376); an element added by hand goes
# before the first sibling that follows it in the sequence
_PPR_ORDER = ("pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
              "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct",
              "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
              "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment",
              "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange")
_TCPR_ORDER = ("cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar", "textDirection",
               "tcFitText", "vAlign", "hideMark", "headers", "cellIns", "cellDel", "cellMerge", "tcPrChange")
_RPR_ORDER = ("rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline", "shadow",
              "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing", "w", "kern",
              "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign", "rtl", "cs",
              "em", "lang", "eastAsianLayout", "specVanish", "oMath")
# w:settings: the whole sequence (prefix w: unless given)
_SETTINGS_ORDER = (
    "writeProtection", "view", "zoom", "removePersonalInformation", "removeDateAndTime", "doNotDisplayPageBoundaries",
    "displayBackgroundShape", "printPostScriptOverText", "printFractionalCharacterWidth", "printFormsData",
    "embedTrueTypeFonts", "embedSystemFonts", "saveSubsetFonts", "saveFormsData", "mirrorMargins",
    "alignBordersAndEdges", "bordersDoNotSurroundHeader", "bordersDoNotSurroundFooter", "gutterAtTop",
    "hideSpellingErrors", "hideGrammaticalErrors", "activeWritingStyle", "proofState", "formsDesign",
    "attachedTemplate", "linkStyles", "stylePaneFormatFilter", "stylePaneSortMethod", "documentType", "mailMerge",
    "revisionView", "trackRevisions", "doNotTrackMoves", "doNotTrackFormatting", "documentProtection",
    "autoFormatOverride", "styleLockTheme", "styleLockQFSet", "defaultTabStop", "autoHyphenation",
    "consecutiveHyphenLimit", "hyphenationZone", "doNotHyphenateCaps", "showEnvelope", "summaryLength",
    "clickAndTypeStyle", "defaultTableStyle", "evenAndOddHeaders", "bookFoldRevPrinting", "bookFoldPrinting",
    "bookFoldPrintingSheets", "drawingGridHorizontalSpacing", "drawingGridVerticalSpacing",
    "displayHorizontalDrawingGridEvery", "displayVerticalDrawingGridEvery", "doNotUseMarginsForDrawingGridOrigin",
    "drawingGridHorizontalOrigin", "drawingGridVerticalOrigin", "doNotShadeFormData", "noPunctuationKerning",
    "characterSpacingControl", "printTwoOnOne", "strictFirstAndLastChars", "noLineBreaksAfter", "noLineBreaksBefore",
    "savePreviewPicture", "doNotValidateAgainstSchema", "saveInvalidXml", "ignoreMixedContent",
    "alwaysShowPlaceholderText", "doNotDemarcateInvalidXml", "saveXmlDataOnly", "useXSLTWhenSaving", "saveThroughXslt",
    "showXMLTags", "alwaysMergeEmptyNamespace", "updateFields", "hdrShapeDefaults", "footnotePr", "endnotePr", "compat",
    "docVars", "rsids", "m:mathPr", "attachedSchema", "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats",
    "doNotAutoCompressPictures", "forceUpgrade", "captions", "readModeInkLockDown", "smartTagType", "sl:schemaLibrary",
    "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator")
_SETTINGS_TAGS = tuple(n if ":" in n else "w:" + n for n in _SETTINGS_ORDER)
# w:settings: the elements after w:updateFields
_SETTINGS_AFTER_UPDATE = _SETTINGS_TAGS[_SETTINGS_TAGS.index("w:updateFields") + 1:]


def _insert_ordered(parent, child, order: Sequence[str] = _PPR_ORDER):
    """Insert ``child`` into ``parent`` at its place in the schema sequence ``order`` (local names in the w namespace);
    an existing element of the same name is replaced."""
    from docx.oxml.ns import qn
    name = child.tag.split("}", 1)[1]
    old = parent.find(child.tag)
    if old is not None:
        parent.remove(old)
    later = {qn("w:" + n) for n in order[order.index(name) + 1:]}
    for i, x in enumerate(parent):
        if x.tag in later:
            parent.insert(i, child)
            return child
    parent.append(child)
    return child


def _settings_insert(settings, child):
    """Insert ``child`` into ``w:settings`` at its place in the schema sequence (an element of the same name is
    replaced); elements of extensions (w14, w15 …) stay behind it."""
    from docx.oxml.ns import qn
    tags = [qn(n) for n in _SETTINGS_TAGS]
    old = settings.find(child.tag)
    if old is not None:
        settings.remove(old)
    later = set(tags[tags.index(child.tag) + 1:])
    for i, x in enumerate(settings):
        if isinstance(x.tag, str) and (x.tag in later or x.tag not in tags):
            settings.insert(i, child)
            return child
    settings.append(child)
    return child


def _append_num(numbering, el) -> None:
    """Append a ``w:abstractNum`` or ``w:num`` to ``w:numbering`` in front of ``w:numIdMacAtCleanup`` (the last
    element of the schema sequence)."""
    from docx.oxml.ns import qn
    last = numbering.find(qn("w:numIdMacAtCleanup"))
    if last is not None:
        last.addprevious(el)
    else:
        numbering.append(el)


def _num_pr(num_id, ilvl: int = 0):
    """``<w:numPr>`` with the level and the numbering (0 = numbering off)."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    numpr = OxmlElement("w:numPr")
    for tag, val in (("w:ilvl", str(ilvl)), ("w:numId", str(num_id))):
        x = OxmlElement(tag)
        x.set(qn("w:val"), val)
        numpr.append(x)
    return numpr


# ----------------------------------------------------------------------------------------- Word equations (OMML)
_OMML_SYMBOL = re.compile(r"([A-Za-zΔφθ]+)_\[([^\]]*)\]")
_OMML_NUMBER = re.compile(r"\d+(?:[.,]\d+)*")
_OMML_WORD = re.compile(r"[A-Za-zÄÖÜäöüßΩΔφθ]+")
_OMML_FUNCTIONS = {"cos", "sin", "tan", "exp", "ln", "log", "arg"}
_OMML_UNITS = {"A", "V", "Ω", "Hz", "kHz", "kV", "kA", "mA", "mV", "ms", "s", "%"}
_M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"


def _m(tag: str, **attrs):
    """An element of the math namespace (``m:``)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    el = OxmlElement("m:" + tag)
    for k, v in attrs.items():
        el.set(qn("m:" + k), str(v))
    return el


def _m_run(text: str, style: Optional[str] = None, normal: bool = False):
    """A math run: ``style`` "p" upright; ``normal`` = normal text (upright, spaces kept, no math spacing)."""
    from docx.oxml.ns import qn
    r = _m("r")
    if style or normal:
        pr = _m("rPr")
        pr.append(_m("nor") if normal else _m("sty", val=style))
        r.append(pr)
    t = _m("t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    return r


def _m_wrap(tag: str, items):
    el = _m(tag)
    for x in items:
        el.append(x)
    return el


def _m_sub(base: str, sub: str):
    """A symbol with its index (as normal text: abbreviations and commas stay as written)."""
    return _m_wrap("sSub", [_m_wrap("e", [_m_run(base) if len(base) == 1 else _m_run(base, style="p")]),
                            _m_wrap("sub", [_m_run(sub, normal=True)])])


def _m_delimiter(items, open_: str = "(", close: str = ")"):
    d = _m("d")
    if (open_, close) != ("(", ")"):
        pr = _m("dPr")
        pr.append(_m("begChr", val=open_))
        pr.append(_m("endChr", val=close))
        d.append(pr)
    d.append(_m_wrap("e", items))
    return d


def _m_radical(items):
    rad = _m("rad")
    pr = _m("radPr")
    pr.append(_m("degHide", val=1))
    rad.append(pr)
    rad.append(_m("deg"))
    rad.append(_m_wrap("e", items))
    return rad


def _m_sum(item):
    """∑ without limits over one element."""
    nary = _m("nary")
    pr = _m("naryPr")
    pr.append(_m("chr", val="∑"))
    pr.append(_m("subHide", val=1))
    pr.append(_m("supHide", val=1))
    nary.append(pr)
    nary.append(_m("sub"))
    nary.append(_m("sup"))
    nary.append(_m_wrap("e", [item]))
    return nary


def _closing(text: str, i: int, open_: str = "(", close: str = ")") -> int:
    """Index of the delimiter that closes the one at ``i`` (ValueError if unbalanced)."""
    depth = 0
    for j in range(i, len(text)):
        if text[j] == open_ and (open_ != close or depth == 0):
            depth += 1
        elif text[j] == close:
            depth -= 1
        if depth == 0:
            return j
    raise ValueError("unbalanced " + open_)


def _text_only(inner: str) -> bool:
    """A remark in words, e.g. "(without the margins)"."""
    words = inner.split()
    return bool(words) and all(_OMML_WORD.fullmatch(w) and len(w) > 1 for w in words)


def _omml_items(text: str, start: int = 0, stop: Optional[int] = None) -> list:
    """The math elements of a formula part in the report mark-up: symbols ``I_[G]``, numbers, ``√(…)``, ``(…)``,
    ``|…|``, exponents ``²``, functions (cos, sin) with their argument, ``Σ`` over the next element, units."""
    end = len(text) if stop is None else stop
    out: list = []
    kinds: List[str] = []                            # sym, num, var, group, op, text, sum
    i = start
    while i < end:
        ch = text[i]
        if text.startswith("√(", i):
            j = _closing(text, i + 1)
            out.append(_m_radical(_omml_items(text, i + 2, j)))
            kinds.append("group")
            i = j + 1
        elif ch == "(":
            j = _closing(text, i)
            if j >= end:
                raise ValueError("parenthesis outside the part")
            out.append(_m_delimiter(_omml_items(text, i + 1, j)))
            kinds.append("group")
            i = j + 1
        elif ch == "|":
            j = text.find("|", i + 1, end)
            if j < 0:
                raise ValueError("unbalanced |")
            out.append(_m_delimiter(_omml_items(text, i + 1, j), "|", "|"))
            kinds.append("group")
            i = j + 1
        elif ch in "²³":
            sup = _m_run("2" if ch == "²" else "3")
            if out and kinds[-1] == "sym":                # I_[G]² → base, index and exponent
                el = _m("sSubSup")
                el.append(out[-1].find(_M + "e"))
                el.append(out[-1].find(_M + "sub"))
                el.append(_m_wrap("sup", [sup]))
                out[-1], kinds[-1] = el, "group"
            elif out and kinds[-1] in ("num", "var", "group"):
                out[-1], kinds[-1] = _m_wrap("sSup", [_m_wrap("e", [out[-1]]), _m_wrap("sup", [sup])]), "group"
            else:
                out.append(_m_run(ch))
                kinds.append("op")
            i += 1
        elif ch == "Σ":
            out.append(None)                              # ∑ over the next element (see below)
            kinds.append("sum")
            i += 1
        elif _OMML_SYMBOL.match(text, i) and _OMML_SYMBOL.match(text, i).end() <= end:
            m = _OMML_SYMBOL.match(text, i)
            out.append(_m_sub(m.group(1), m.group(2)))
            kinds.append("sym")
            i = m.end()
        elif _OMML_NUMBER.match(text, i):
            m = _OMML_NUMBER.match(text, i)
            out.append(_m_run(m.group(), normal=True))      # decimal comma without the spacing of a separator
            kinds.append("num")
            i = m.end()
        elif _OMML_WORD.match(text, i):
            m = _OMML_WORD.match(text, i)
            w, j = m.group(), min(m.end(), end)
            if w in _OMML_FUNCTIONS:                      # cos φ(…): the function with its argument
                k = j
                while k < end and text[k] == " ":
                    k += 1
                arg = _OMML_SYMBOL.match(text, k) or _OMML_WORD.match(text, k)
                arg_end = arg.end() if arg else k
                if arg_end < end and text[arg_end] == "(":
                    arg_end = _closing(text, arg_end) + 1
                func = _m("func")
                func.append(_m_wrap("fName", [_m_run(w, style="p")]))
                func.append(_m_wrap("e", _omml_items(text, k, arg_end)))
                out.append(func)
                kinds.append("group")
                i = arg_end
                continue
            if w in _OMML_UNITS and i > start and text[i - 1] == " " and out:
                out.append(_m_run("\u00a0" + w, normal=True))
                kinds.append("text")
            elif len(w) == 1:
                out.append(_m_run(w))
                kinds.append("var")
            else:
                out.append(_m_run(w, normal=True))
                kinds.append("text")
            i = j
        elif ch == " ":
            i += 1                                        # Word sets the spaces of a formula itself
        elif ch == "%":
            out.append(_m_run("\u00a0%" if i > start and text[i - 1] == " " else "%", normal=True))
            kinds.append("text")
            i += 1
        else:
            out.append(_m_run(ch))
            kinds.append("op")
            i += 1
    for k in range(len(out) - 1, -1, -1):                  # ∑ over the element after it
        if kinds[k] == "sum":
            if k + 1 >= len(out):
                raise ValueError("Σ without a term")
            out[k:k + 2] = [_m_sum(out[k + 1])]
            kinds[k:k + 2] = ["group"]
    return out


def _omml_parts(text: str) -> List[Tuple[str, str]]:
    """A formula split into math and text at the top level: words of two or more letters that are no function or
    unit ("with"), a remark in parentheses, a gap of several spaces between two equations."""
    parts: List[Tuple[str, str]] = []
    buf, i, depth = "", 0, 0

    def flush():
        nonlocal buf
        if buf.strip():
            parts.append(("math", buf))
        elif buf and parts:
            parts.append(("text", " "))
        buf = ""

    while i < len(text):
        ch = text[i]
        if depth == 0 and ch == "(" and not text.startswith("√(", i - 1 if i else 0):
            j = _closing(text, i)
            if _text_only(text[i + 1:j]):
                flush()
                parts.append(("text", " " + text[i:j + 1]))
                i = j + 1
                continue
        if depth == 0 and ch == " " and text.startswith("  ", i):
            flush()
            j = i
            while j < len(text) and text[j] == " ":
                j += 1
            parts.append(("text", "\u2003\u2003"))
            i = j
            continue
        m = _OMML_WORD.match(text, i) if depth == 0 else None
        if m and len(m.group()) > 1 and m.group() not in _OMML_FUNCTIONS and m.group() not in _OMML_UNITS and \
                (i == 0 or not (text[i - 1].isalpha() or text[i - 1] == "_")) and \
                not text.startswith("_[", m.end()):
            flush()
            parts.append(("text", " " + m.group() + " "))
            i = m.end()
            continue
        if ch in "([":
            depth += 1
        elif ch in ")]":
            depth -= 1
        buf += ch
        i += 1
    flush()
    return parts


def omml(text: str) -> Optional[list]:
    """A displayed formula of the report mark-up as Word equations: a list of ``m:oMath`` elements (the math parts)
    and texts (the words between them), or None if the text cannot be converted (unbalanced delimiters) – it is then
    written as text."""
    try:
        out = []
        for kind, part in _omml_parts(str(text)):
            if kind == "text":
                out.append(part)
            else:
                items = _omml_items(part)
                if items:
                    out.append(_m_wrap("oMath", items))
        return out if any(not isinstance(x, str) for x in out) else None
    except (ValueError, IndexError, AttributeError):
        return None


# =================================================================================================== Word document
class _Doc:
    """Thin layer over python-docx: styles, numbered headings, captions with fields, tables, figures, placeholders,
    footnotes – in the layout of a scientific report: body text justified (hyphenation on), lists indented, captions
    centred in the size of the body text, a blank line between a displayed formula, a figure or a table and the text."""

    #: lists: number or bullet at 0.635 cm, the text at 1.27 cm (twips)
    LIST_INDENT, LIST_HANGING = 720, 360

    def __init__(self, template: Optional[str] = None, styles: Optional[Mapping[str, str]] = None,
                 font: Optional[str] = None, size_pt: float = 10.0, language: Optional[str] = None):
        import docx
        from docx.shared import Cm, Pt
        self.docx = docx
        self.Cm, self.Pt = Cm, Pt
        if template:
            if not os.path.isfile(template):
                raise FileNotFoundError(tr("document.template_missing", path=template))
            self.doc = docx.Document(template)
            _clear_body(self.doc)
        else:
            self.doc = docx.Document()
            sec = self.doc.sections[0]
            sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
            sec.left_margin = sec.right_margin = Cm(2.5)
            sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
            st = self.doc.styles["Normal"]
            st.font.name = font or _docx_font()
            st.font.size = Pt(size_pt)
            st.paragraph_format.space_after = Pt(6)
            for name, size in (("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11)):
                try:
                    h = self.doc.styles[name]
                    h.font.name = font or _docx_font()
                    h.font.size = Pt(size)
                except KeyError:
                    pass
        self.templated = bool(template)
        self.equations = True                                     # displayed formulas as Word equations
        self.styles = {"title": "Title", "subtitle": "Subtitle", "toc_heading": "TOC Heading", "caption": "Caption",
                       "table": "Table Grid", "list_bullet": "List Bullet", "list_number": "List Number",
                       "heading1": "Heading 1", "heading2": "Heading 2", "heading3": "Heading 3",
                       "body": "Body Text", "list_paragraph": "List Paragraph"}
        self.styles.update(styles or {})
        self.numbers = [0, 0, 0]
        self.headings: List[Tuple[int, str, str]] = []           # (level, number, text) for the contents
        self.counters = {"figure": 0, "table": 0}
        self.toc_paragraph = None
        self._numbered = self._heading_numbering()
        sec = self.doc.sections[-1]
        self.text_width_cm = (sec.page_width - sec.left_margin - sec.right_margin) / 360000.0
        self._footnotes = None                                    # (part, w:footnotes, next id) once used
        self._list_abstract: Dict[bool, int] = {}                 # own list numbering: numbered → abstractNumId
        self._bullet_num: Optional[int] = None                    # … and the instance of the bulleted lists
        self._layout(language)

    # ------------------------------------------------------------------ styles
    def style(self, key: str) -> Optional[str]:
        name = self.styles.get(key, key)
        try:
            self.doc.styles[name]
            return name
        except KeyError:
            return None

    @staticmethod
    def _resolved(style, get):
        """A property of a style or else of its base styles (None: set in none of them)."""
        for _ in range(30):
            if style is None:
                return None
            v = get(style)
            if v is not None:
                return v
            style = style.base_style
        return None

    def _defaults(self) -> dict:
        """The document defaults of the styles part: font size (pt), spacing before/after (pt), line spacing."""
        from docx.oxml.ns import qn
        out = dict(size=None, before=None, after=None, line=None, rule=None)
        dd = self.doc.styles.element.find(qn("w:docDefaults"))
        if dd is None:
            return out
        sz = dd.find("/".join(qn(t) for t in ("w:rPrDefault", "w:rPr", "w:sz")))
        if sz is not None and sz.get(qn("w:val")):
            out["size"] = int(sz.get(qn("w:val"))) / 2.0
        sp = dd.find("/".join(qn(t) for t in ("w:pPrDefault", "w:pPr", "w:spacing")))
        if sp is not None:
            for key, attr in (("before", "w:before"), ("after", "w:after"), ("line", "w:line")):
                if sp.get(qn(attr)):
                    out[key] = int(sp.get(qn(attr))) / (20.0 if key != "line" else 1.0)
            out["rule"] = sp.get(qn("w:lineRule")) or ("auto" if out["line"] is not None else None)
        return out

    def _layout(self, language: Optional[str]):
        """The layout of a scientific report on top of the template: a justified body-text style, centred captions
        in the size of the body text, the language of the report (spelling, hyphenation), automatic hyphenation;
        the measures of the body text for the spacing of formulas, figures and tables."""
        from docx.enum.style import WD_STYLE_TYPE
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        from docx.shared import RGBColor
        styles, Pt = self.doc.styles, self.Pt
        default = styles.default(WD_STYLE_TYPE.PARAGRAPH)
        dd = self._defaults()
        # body text: justified, its own style ("Body Text" – created on the default style if the template has none)
        name = self.styles["body"]
        try:
            body = styles[name]
        except KeyError:
            body = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH, builtin=name == "Body Text")
            if default is not None:
                body.base_style = default
            body.next_paragraph_style = body
            body.quick_style = True
            if self._resolved(body, lambda s: s.paragraph_format.space_after) is None and not dd["after"]:
                body.paragraph_format.space_after = Pt(6)          # paragraphs apart (the template has no spacing)
        body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        body.paragraph_format.widow_control = True
        self.body_style = body.name
        size = self._resolved(body, lambda s: s.font.size)
        self.body_pt = size.pt if size is not None else (dd["size"] or 10.0)
        after = self._resolved(body, lambda s: s.paragraph_format.space_after)
        before = self._resolved(body, lambda s: s.paragraph_format.space_before)
        self.body_after_pt = after.pt if after is not None else (dd["after"] or 0.0)
        self.body_before_pt = before.pt if before is not None else (dd["before"] or 0.0)
        line = self._resolved(body, lambda s: s.paragraph_format.line_spacing)
        rule = self._resolved(body, lambda s: s.paragraph_format.line_spacing_rule)
        natural = 1.2 * self.body_pt                              # a single line of the usual fonts
        if line is None and dd["line"] is not None:               # from the defaults: twips or 240ths of a line
            line, rule = (dd["line"] / 240.0, None) if dd["rule"] == "auto" else (Pt(dd["line"] / 20.0), dd["rule"])
        if line is None:
            self.line_pt = natural
        elif isinstance(line, float):
            self.line_pt = natural * line
        else:
            exact = rule is not None and "EXACTLY" in str(rule).upper() or rule == "exact"
            self.line_pt = line.pt if exact else max(line.pt, natural)
        # captions: centred, in the size of the body text (own layout: black, not bold)
        cap = self.style("caption")
        if cap:
            cs = styles[cap]
            cs.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cs.font.size = Pt(self.body_pt)
            if not self.templated:
                cs.font.bold = False
                cs.font.color.rgb = RGBColor(0, 0, 0)
        if language:
            self._language({"de": "de-DE", "en": "en-GB"}[i18n.normalize_language(language)], default)
        # automatic hyphenation for the justified text (unless the template switches it off)
        settings = self.doc.settings.element
        if settings.find(qn("w:autoHyphenation")) is None:
            hy = OxmlElement("w:autoHyphenation")
            hy.set(qn("w:val"), "true")
            _settings_insert(settings, hy)
        self._style_ids = None

    def _language(self, tag: str, default) -> None:
        """The language of the text (spelling, hyphenation) in the document defaults and in the default paragraph
        style if that sets one; a template keeps its own if it is the same language (de-DE, de-AT …)."""
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        targets = []
        dd = self.doc.styles.element.find(qn("w:docDefaults"))
        if dd is not None:
            rpd = dd.find(qn("w:rPrDefault"))
            if rpd is None:
                rpd = OxmlElement("w:rPrDefault")
                dd.insert(0, rpd)
            rpr = rpd.find(qn("w:rPr"))
            if rpr is None:
                rpr = OxmlElement("w:rPr")
                rpd.append(rpr)
            targets.append((rpr, True))
        if default is not None and default.element.rPr is not None:
            targets.append((default.element.rPr, False))
        for rpr, create in targets:
            lang = rpr.find(qn("w:lang"))
            if lang is None and not create:
                continue                                          # the default style follows the defaults
            if lang is None:
                lang = _insert_ordered(rpr, OxmlElement("w:lang"), _RPR_ORDER)
            current = lang.get(qn("w:val")) or ""
            if not self.templated or current.split("-")[0].lower() != tag[:2]:
                lang.set(qn("w:val"), tag)

    def _style_id(self, key: str) -> Optional[str]:
        """The style id of a style key (or name) – None if the document has no such style."""
        name = self.style(key)
        return self.doc.styles[name].style_id if name else None

    def _style_of(self, style_id: Optional[str]):
        """The style with this id (None: the default paragraph style)."""
        from docx.enum.style import WD_STYLE_TYPE
        if self._style_ids is None or (style_id and style_id not in self._style_ids):
            self._style_ids = {s.style_id: s for s in self.doc.styles if s.type == WD_STYLE_TYPE.PARAGRAPH}
        if not style_id:
            return self.doc.styles.default(WD_STYLE_TYPE.PARAGRAPH)
        return self._style_ids.get(style_id)

    # ------------------------------------------------------------------ spacing (a blank line around a formula …)
    def _last_block(self):
        """The last paragraph or table of the body (None at the start)."""
        from docx.oxml.ns import qn
        for el in reversed(self.doc.element.body):
            if el.tag != qn("w:sectPr"):
                return el
        return None

    def _space_below(self, el) -> float:
        """The space a block leaves below itself (pt): an empty paragraph, a table (followed by a spacer) or a
        heading counts as a blank line, otherwise the spacing after the paragraph."""
        from docx.oxml.ns import qn
        if el is None or el.tag != qn("w:p"):
            return self.line_pt
        has_text = any((t.text or "").strip() for t in el.iter(qn("w:t")))
        if not has_text and el.find(".//" + qn("w:drawing")) is None and el.find(".//" + _M + "oMath") is None:
            return self.line_pt
        ppr = el.find(qn("w:pPr"))
        sid = None
        if ppr is not None:
            ps = ppr.find(qn("w:pStyle"))
            sid = ps.get(qn("w:val")) if ps is not None else None
            if ppr.find(qn("w:outlineLvl")) is not None:
                return self.line_pt
            sp = ppr.find(qn("w:spacing"))
            if sp is not None and sp.get(qn("w:after")) is not None:
                return int(sp.get(qn("w:after"))) / 20.0
        if sid and sid in {self._style_id("heading%d" % k) for k in (1, 2, 3)}:
            return self.line_pt
        st = self._style_of(sid)
        v = self._resolved(st, lambda s: s.paragraph_format.space_after) if st is not None else None
        return v.pt if v is not None else (self._defaults()["after"] or 0.0)

    def _blank_line_before(self, p, prev) -> None:
        """Space before ``p`` so that a blank line separates it from the block ``prev`` above it."""
        p.paragraph_format.space_before = self.Pt(round(max(0.0, self.line_pt - self._space_below(prev)), 1))

    def _heading_numbering(self) -> bool:
        """Numbering of the headings: that of the template, else a multilevel list 1, 1.1, 1.1.1 is added."""
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        names = [self.style("heading%d" % k) for k in (1, 2, 3)]
        if None in names:
            return False
        first = self.doc.styles[names[0]].element
        ppr = first.find(qn("w:pPr"))
        if ppr is not None and ppr.find(qn("w:numPr")) is not None:
            return True                                            # the template numbers its headings
        numbering = self.doc.part.numbering_part.element
        ids = [int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))]
        nids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
        aid, nid = max(ids + [0]) + 1, max(nids + [0]) + 1
        absn = OxmlElement("w:abstractNum")
        absn.set(qn("w:abstractNumId"), str(aid))
        mlt = OxmlElement("w:multiLevelType")
        mlt.set(qn("w:val"), "multilevel")
        absn.append(mlt)
        for lvl, name in enumerate(names):
            el = OxmlElement("w:lvl")
            el.set(qn("w:ilvl"), str(lvl))
            for tag, val in (("w:start", "1"), ("w:numFmt", "decimal"),
                             ("w:pStyle", self.doc.styles[name].style_id),
                             ("w:lvlText", ".".join("%%%d" % (k + 1) for k in range(lvl + 1))),
                             ("w:lvlJc", "left")):
                x = OxmlElement(tag)
                x.set(qn("w:val"), val)
                el.append(x)
            ppr_l = OxmlElement("w:pPr")
            ind = OxmlElement("w:ind")
            ind.set(qn("w:left"), str(567 + 142 * lvl))
            ind.set(qn("w:hanging"), str(567 + 142 * lvl))
            ppr_l.append(ind)
            el.append(ppr_l)
            absn.append(el)
        first_num = numbering.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(absn)
        else:
            _append_num(numbering, absn)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(nid))
        an = OxmlElement("w:abstractNumId")
        an.set(qn("w:val"), str(aid))
        num.append(an)
        _append_num(numbering, num)
        for lvl, name in enumerate(names):
            _insert_ordered(self.doc.styles[name].element.get_or_add_pPr(), _num_pr(nid, lvl))
        return True

    # ------------------------------------------------------------------ text
    def _add_runs(self, p, text: str, size: Optional[float] = None, bold: Optional[bool] = None,
                  color: Optional[str] = None, footnotes: bool = True):
        """The runs of a text with the report mark-up at the end of ``p``; ``[^…^]`` becomes a footnote."""
        text = _ordinals(str(text)) if i18n.get_language() == "en" else str(text)
        pos = 0
        for m in (_FOOTNOTE.finditer(text) if footnotes else ()):
            self._add_marked_runs(p, text[pos:m.start()].rstrip(" "), size, bold, color)
            self.footnote(p, m.group("fn"))
            pos = m.end()
        self._add_marked_runs(p, text[pos:], size, bold, color)
        return p

    def _add_marked_runs(self, p, text: str, size, bold, color):
        from docx.enum.text import WD_COLOR_INDEX
        from docx.shared import RGBColor
        for t, f in _runs(text):
            r = p.add_run(t)
            if f.get("bold") or bold:
                r.bold = True
            if f.get("italic"):
                r.italic = True
            if f.get("subscript"):
                r.font.subscript = True
            if f.get("superscript"):
                r.font.superscript = True
            if f.get("highlight"):
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            if size:
                r.font.size = self.Pt(size)
            if color:
                r.font.color.rgb = RGBColor.from_string(color)

    def para(self, text: str = "", style: Optional[str] = None, align: Optional[str] = None,
             size: Optional[float] = None, bold: Optional[bool] = None, keep_next: bool = False,
             space_after: Optional[float] = None, color: Optional[str] = None, body: Optional[bool] = None):
        """A paragraph; a text without ``style`` and ``align`` is body text (justified, ``body=False``: the
        default paragraph style)."""
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        if body is None:
            body = style is None and align is None and bool(str(text).strip())
        name = self.body_style if body and style is None else (self.style(style) if style else None)
        p = self.doc.add_paragraph(style=name)
        self._add_runs(p, text, size, bold, color)
        if align:
            p.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
                           "justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
        if keep_next:
            p.paragraph_format.keep_with_next = True
        if space_after is not None:
            p.paragraph_format.space_after = self.Pt(space_after)
        return p

    def formula(self, text: str):
        """A displayed formula, centred, with a blank line to the text above and below: Word equations (Office Math)
        for the math and normal text for the words between them; as text if it cannot be converted."""
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        prev = self._last_block()
        parts = omml(_ordinals(str(text)) if i18n.get_language() == "en" else str(text)) if self.equations else None
        if parts is None:
            p = self.para(text, align="center")
        else:
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for x in parts:
                if isinstance(x, str):
                    p.add_run(x)
                else:
                    p._p.append(x)
        self._blank_line_before(p, prev)                       # formula below formula: one blank line between
        p.paragraph_format.space_after = self.Pt(round(max(0.0, self.line_pt - self.body_before_pt), 1))
        self._keep_lead(prev)
        return p

    def _keep_lead(self, el) -> None:
        """Keep the body text that leads to a formula or a list with it (not alone at the end of a page)."""
        from docx.oxml.ns import qn
        from docx.text.paragraph import Paragraph
        if el is None or el.tag != qn("w:p") or not any((t.text or "").strip() for t in el.iter(qn("w:t"))):
            return
        ps = el.find(qn("w:pPr") + "/" + qn("w:pStyle"))
        if (ps.get(qn("w:val")) if ps is not None else None) == self.doc.styles[self.body_style].style_id:
            Paragraph(el, self.doc).paragraph_format.keep_with_next = True

    def placeholder_text(self, what: str):
        """A text still to write, highlighted (yellow) – in the style of the body text."""
        from docx.enum.text import WD_COLOR_INDEX
        p = self.doc.add_paragraph(style=self.body_style)
        r = p.add_run(tr("document.placeholder_text", what=plain(what)))
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        return p

    def bullets(self, items: Iterable[str], numbered: bool = False):
        """Bulleted or numbered list (a numbered list starts at 1), indented: number or bullet at 0.635 cm, the
        text at 1.27 cm, justified. With the list styles of the document (``list_bullet``, ``list_number``) or else
        a numbering of its own in the style ``list_paragraph``."""
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        items = list(items)
        if not items:
            return
        key = "list_number" if numbered else "list_bullet"
        st = self.style(key)
        if st and self._style_numbering(st) is not None:
            num_id = self._restart_numbering(st) if numbered else None
        else:
            st, num_id = self.style("list_paragraph"), self._own_list(numbered)
            key = "list_paragraph" if st else None
        self._keep_lead(self._last_block())
        for i, it in enumerate(items):
            p = self.para(it, style=key, body=key is None)
            if num_id is not None:
                _insert_ordered(p._p.get_or_add_pPr(), _num_pr(num_id))
            pf = p.paragraph_format
            pf.left_indent = self.Pt(self.LIST_INDENT / 20.0)
            pf.first_line_indent = self.Pt(-self.LIST_HANGING / 20.0)
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if i < len(items) - 1:
                pf.keep_with_next = len(items) <= 6               # a short list stays on one page
            else:
                pf.space_after = self.Pt(self.body_after_pt)      # the text after the list as after a paragraph

    def _style_numbering(self, style_name: str) -> Optional[str]:
        """The numId of a list style (None: the style is not numbered)."""
        from docx.oxml.ns import qn
        ppr = self.doc.styles[style_name].element.find(qn("w:pPr"))
        numpr = ppr.find(qn("w:numPr")) if ppr is not None else None
        nid = numpr.find(qn("w:numId")) if numpr is not None else None
        return nid.get(qn("w:val")) if nid is not None else None

    def _own_list(self, numbered: bool) -> int:
        """A numbering of the report itself (the document has no list styles): bullets "•" or numbers "1." at
        level 0, indented like the lists of the report; every numbered list gets an instance of its own (it starts
        at 1), the bulleted lists share one."""
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        numbering = self.doc.part.numbering_part.element
        aid = self._list_abstract.get(numbered)
        if aid is None:
            aid = max([int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))] + [0]) + 1
            absn = OxmlElement("w:abstractNum")
            absn.set(qn("w:abstractNumId"), str(aid))
            mlt = OxmlElement("w:multiLevelType")
            mlt.set(qn("w:val"), "singleLevel")
            absn.append(mlt)
            lvl = OxmlElement("w:lvl")
            lvl.set(qn("w:ilvl"), "0")
            for tag, val in (("w:start", "1"), ("w:numFmt", "decimal" if numbered else "bullet"),
                             ("w:lvlText", "%1." if numbered else "\u2022"), ("w:lvlJc", "left")):
                x = OxmlElement(tag)
                x.set(qn("w:val"), val)
                lvl.append(x)
            ppr = OxmlElement("w:pPr")
            ind = OxmlElement("w:ind")
            ind.set(qn("w:left"), str(self.LIST_INDENT))
            ind.set(qn("w:hanging"), str(self.LIST_HANGING))
            ppr.append(ind)
            lvl.append(ppr)
            absn.append(lvl)
            first_num = numbering.find(qn("w:num"))
            if first_num is not None:
                first_num.addprevious(absn)
            else:
                _append_num(numbering, absn)
            self._list_abstract[numbered] = aid
        if not numbered and self._bullet_num is not None:
            return self._bullet_num
        nid = max([int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))] + [0]) + 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(nid))
        a = OxmlElement("w:abstractNumId")
        a.set(qn("w:val"), str(aid))
        num.append(a)
        if numbered:                                           # without the override Word continues the count
            ov = OxmlElement("w:lvlOverride")
            ov.set(qn("w:ilvl"), "0")
            so = OxmlElement("w:startOverride")
            so.set(qn("w:val"), "1")
            ov.append(so)
            num.append(ov)
        _append_num(numbering, num)
        if not numbered:
            self._bullet_num = nid
        return nid

    def _restart_numbering(self, style_name: str) -> Optional[int]:
        """A new numbering instance of the list style (so that the list starts at 1)."""
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        el = self.doc.styles[style_name].element
        ppr = el.find(qn("w:pPr"))
        numpr = ppr.find(qn("w:numPr")) if ppr is not None else None
        nid = numpr.find(qn("w:numId")) if numpr is not None else None
        if nid is None:
            return None
        numbering = self.doc.part.numbering_part.element
        base = next((n for n in numbering.findall(qn("w:num")) if n.get(qn("w:numId")) == nid.get(qn("w:val"))), None)
        if base is None:
            return None
        aid = base.find(qn("w:abstractNumId")).get(qn("w:val"))
        new_id = max(int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))) + 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(new_id))
        a = OxmlElement("w:abstractNumId")
        a.set(qn("w:val"), aid)
        num.append(a)
        ov = OxmlElement("w:lvlOverride")
        ov.set(qn("w:ilvl"), "0")
        so = OxmlElement("w:startOverride")
        so.set(qn("w:val"), "1")
        ov.append(so)
        num.append(ov)
        _append_num(numbering, num)
        return new_id

    def heading(self, text: str, level: int = 1, numbered: bool = True) -> str:
        """Heading of level 1–3; returns its number ("2.3") or ""."""
        key = self.style("heading%d" % level) or "Heading %d" % level
        p = self.doc.add_paragraph(style=key)
        self._add_runs(p, text)
        number = ""
        if numbered:
            self.numbers[level - 1] += 1
            for k in range(level, 3):
                self.numbers[k] = 0
            number = ".".join(str(x) for x in self.numbers[:level])
        elif self._numbered:                                       # no number: numbering off for this paragraph
            _insert_ordered(p._p.get_or_add_pPr(), _num_pr(0))
        self.headings.append((level, number, plain(text)))
        return number

    def page_break(self):
        from docx.enum.text import WD_BREAK
        p = self.doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)

    # ------------------------------------------------------------------ fields
    @staticmethod
    def _field(paragraph, instr: str, cached: str):
        """Complex field ``instr`` with the cached result ``cached`` at the end of a paragraph."""
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement

        def fld(kind):
            r = OxmlElement("w:r")
            f = OxmlElement("w:fldChar")
            f.set(qn("w:fldCharType"), kind)
            r.append(f)
            return r

        paragraph._p.append(fld("begin"))
        r = OxmlElement("w:r")
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = " %s " % instr
        r.append(it)
        paragraph._p.append(r)
        paragraph._p.append(fld("separate"))
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = cached
        r.append(t)
        paragraph._p.append(r)
        paragraph._p.append(fld("end"))

    def caption(self, kind: str, text: str, keep_next: bool = False) -> int:
        """Caption "Figure 3: …" / "Table 2: …" with a SEQ field, centred in the size of the body text; returns the
        number. The title of a table (above it) has a blank line to the text above, the caption of a figure (below
        it) one to the text below."""
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        self.counters[kind] += 1
        n = self.counters[kind]
        label = tr("document.figure" if kind == "figure" else "document.table")
        prev = self._last_block()
        style = self.style("caption")
        p = self.doc.add_paragraph(style=style)
        size = None if style else self.body_pt
        self._add_runs(p, label + " ", size=size)
        self._field(p, "SEQ %s \\* ARABIC" % label, str(n))
        self._add_runs(p, ": " + text, size=size)
        pf = p.paragraph_format
        if not style:
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if kind == "table":
            self._blank_line_before(p, prev)
            pf.space_after = self.Pt(3)
        else:
            pf.space_before = self.Pt(3)
            pf.space_after = self.Pt(round(max(0.0, self.line_pt - self.body_before_pt), 1))
        if keep_next:
            pf.keep_with_next = True
        return n

    def toc(self, title: str):
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        st = self.style("toc_heading")
        if st:
            hp = self.para(title, style="toc_heading")
            ppr = hp._p.get_or_add_pPr()                       # not numbered, not in the contents itself
            _insert_ordered(ppr, _num_pr(0))
            ol = OxmlElement("w:outlineLvl")
            ol.set(qn("w:val"), "9")
            _insert_ordered(ppr, ol)
        else:
            self.para(title, bold=True, size=14, body=False)
        p = self.doc.add_paragraph()
        self._field(p, 'TOC \\o "1-3" \\h \\z \\u', tr("document.toc_update"))
        self.toc_paragraph = p

    def fill_toc(self):
        """Cached contents of the TOC field: the headings (page numbers come when Word updates the field)."""
        if self.toc_paragraph is None:
            return
        from docx.oxml.ns import qn
        p = self.toc_paragraph._p
        runs = p.findall(qn("w:r"))
        sep = [i for i, r in enumerate(runs) if r.find(qn("w:fldChar")) is not None
               and r.find(qn("w:fldChar")).get(qn("w:fldCharType")) == "separate"]
        if not sep:
            return
        cached = runs[sep[0] + 1]
        t = cached.find(qn("w:t"))
        lines = [("%s  %s" % (num, text)).strip() if num else text for level, num, text in self.headings
                 if level <= 3]
        t.text = lines[0] if lines else ""
        from docx.oxml import OxmlElement
        for line in lines[1:]:
            br = OxmlElement("w:br")
            cached.append(br)
            tt = OxmlElement("w:t")
            tt.set(qn("xml:space"), "preserve")
            tt.text = line
            cached.append(tt)

    def update_fields_on_open(self):
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        settings = self.doc.settings.element
        if settings.find(qn("w:updateFields")) is None:
            u = OxmlElement("w:updateFields")
            u.set(qn("w:val"), "true")
            _settings_insert(settings, u)
        zoom = settings.find(qn("w:zoom"))                    # w:percent is required (missing in some templates)
        if zoom is not None and zoom.get(qn("w:percent")) is None:
            zoom.set(qn("w:percent"), "100")

    # ------------------------------------------------------------------ footnotes
    def _footnote_root(self):
        """The ``w:footnotes`` element of the document – created with the separators if the document has none."""
        if self._footnotes is not None:
            return self._footnotes[1]
        from docx.opc.constants import CONTENT_TYPE as CT, RELATIONSHIP_TYPE as RT
        from docx.opc.packuri import PackURI
        from docx.opc.part import XmlPart
        from docx.oxml import OxmlElement, parse_xml
        from docx.oxml.ns import nsdecls, qn
        part, root = _footnotes_part(self.doc)
        if root is None:
            sep = ('<w:footnote w:type="%s" w:id="%d"><w:p><w:pPr><w:spacing w:after="0" w:line="240" '
                   'w:lineRule="auto"/></w:pPr><w:r><w:%s/></w:r></w:p></w:footnote>')
            root = parse_xml("<w:footnotes %s>%s%s</w:footnotes>" % (
                nsdecls("w"), sep % ("separator", -1, "separator"),
                sep % ("continuationSeparator", 0, "continuationSeparator")))
            names = {str(p.partname) for p in self.doc.part.package.iter_parts()}
            name = "/word/footnotes.xml" if "/word/footnotes.xml" not in names else \
                str(self.doc.part.package.next_partname("/word/footnotes%d.xml"))
            part = XmlPart(PackURI(name), CT.WML_FOOTNOTES, root, self.doc.part.package)
            self.doc.part.relate_to(part, RT.FOOTNOTES)
            settings = self.doc.settings.element
            if settings.find(qn("w:footnotePr")) is None:           # the separators of the notes (as Word writes it)
                pr = OxmlElement("w:footnotePr")
                for i in (-1, 0):
                    x = OxmlElement("w:footnote")
                    x.set(qn("w:id"), str(i))
                    pr.append(x)
                _settings_insert(settings, pr)
        ids = [int(f.get(qn("w:id"))) for f in root.findall(qn("w:footnote"))]
        self._footnotes = [part, root, max(ids + [0]) + 1]
        return root

    def _footnote_styles(self) -> Tuple[str, str]:
        """Style ids of the footnote text and the footnote mark – those of the template or created: the text
        justified, 2 pt smaller than the body text, the mark superscript."""
        from docx.enum.style import WD_STYLE_TYPE
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        styles = self.doc.styles
        try:
            text = styles["footnote text"]
        except KeyError:
            text = styles.add_style("footnote text", WD_STYLE_TYPE.PARAGRAPH, builtin=True)
            text.base_style = styles.default(WD_STYLE_TYPE.PARAGRAPH)
            text.font.size = self.Pt(max(7.0, self.body_pt - 2.0))
            pf = text.paragraph_format
            pf.space_before, pf.space_after, pf.line_spacing = self.Pt(0), self.Pt(0), 1.0
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.widow_control = True
        try:
            mark = styles["footnote reference"]
        except KeyError:
            mark = styles.add_style("footnote reference", WD_STYLE_TYPE.CHARACTER, builtin=True)
            base = styles.default(WD_STYLE_TYPE.CHARACTER)
            if base is not None:
                mark.base_style = base
            mark.font.superscript = True
        return text.style_id, mark.style_id

    def footnote(self, p, text: str) -> int:
        """A footnote with ``text`` (mark-up allowed): the mark at the end of the paragraph ``p``, the note at the
        foot of the page (Word numbers them); returns the id of the note."""
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.text.paragraph import Paragraph
        root = self._footnote_root()
        text_style, mark_style = self._footnote_styles()
        fid = self._footnotes[2]
        self._footnotes[2] += 1

        def mark(child):
            r = OxmlElement("w:r")
            rpr = OxmlElement("w:rPr")
            rs = OxmlElement("w:rStyle")
            rs.set(qn("w:val"), mark_style)
            rpr.append(rs)
            r.append(rpr)
            r.append(child)
            return r

        ref = OxmlElement("w:footnoteReference")
        ref.set(qn("w:id"), str(fid))
        p._p.append(mark(ref))
        note = OxmlElement("w:footnote")
        note.set(qn("w:id"), str(fid))
        fp = OxmlElement("w:p")
        ppr = OxmlElement("w:pPr")
        ps = OxmlElement("w:pStyle")
        ps.set(qn("w:val"), text_style)
        ppr.append(ps)
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), "both")
        ppr.append(jc)
        fp.append(ppr)
        fp.append(mark(OxmlElement("w:footnoteRef")))
        note.append(fp)
        root.append(note)
        q = Paragraph(fp, self.doc)
        q.add_run(" ")
        self._add_runs(q, str(text).strip(), footnotes=False)
        return fid

    # ------------------------------------------------------------------ tables
    def table(self, header: Sequence[str], rows: Sequence[Sequence[Any]], caption: Optional[str] = None,
              widths_cm: Optional[Sequence[float]] = None, align: Optional[Sequence[str]] = None,
              font_size: Optional[float] = None, first_bold: bool = False, keep_together: Optional[bool] = None):
        """Table with a bold, shaded header row (repeated on every page); cells are text (mark-up allowed). A short
        table (``keep_together``, default: up to 20 rows) is kept on one page."""
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        if caption:
            self.caption("table", caption, keep_next=True)
        n = len(header)
        t = self.doc.add_table(rows=1, cols=n)
        st = self.style("table")
        if st:
            t.style = st
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        size = font_size or 9.0
        align = list(align) if align else ["left"] + ["right"] * (n - 1)

        def fill(cell, value, a, bold=False, shade=None):
            cell.text = ""
            p = cell.paragraphs[0]
            self._add_runs(p, "" if value is None else str(value), size=size, bold=bold)
            p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT,
                           "center": WD_ALIGN_PARAGRAPH.CENTER}[a]
            p.paragraph_format.space_after = self.Pt(0)
            if shade:
                tcpr = cell._tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), shade)
                _insert_ordered(tcpr, shd, _TCPR_ORDER)

        for j, h in enumerate(header):
            fill(t.rows[0].cells[j], h, "left" if j == 0 else "center", bold=True, shade="E7E6E6")
        trpr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        trpr.append(th)
        for row in rows:
            cells = t.add_row().cells
            for j in range(n):
                v = row[j] if j < len(row) else ""
                fill(cells[j], v, align[j] if j < len(align) else "right", bold=first_bold and j == 0)
        for r in t.rows:
            trpr = r._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit")
            cs.set(qn("w:val"), "true")
            trpr.append(cs)
        if keep_together if keep_together is not None else len(rows) <= 20:
            for r in t.rows[:-1]:                              # every row with the next one: the table on one page
                for cell in r.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.keep_with_next = True
        if widths_cm:
            widths = [float(w) for w in list(widths_cm)[:n]]
            total = sum(widths)
            if total > self.text_width_cm:                     # fit into the text width
                widths = [w * self.text_width_cm / total for w in widths]
            t.autofit = False
            for j, w in enumerate(widths):
                t.columns[j].width = self.Cm(w)                # grid (LibreOffice) …
            for r in t.rows:
                for j, w in enumerate(widths):
                    r.cells[j].width = self.Cm(w)              # … and cells (Word)
        self.doc.add_paragraph().paragraph_format.space_after = self.Pt(2)
        return t

    def frame(self, header: Sequence[str], df: pd.DataFrame, caption: Optional[str] = None,
              digits: Union[int, Mapping[str, int]] = 3, widths_cm=None, align=None, font_size=None):
        """Table of a DataFrame with numbers formatted for the report."""
        rows = []
        for _, r in df.iterrows():
            line = []
            for c in df.columns:
                v = r[c]
                d = digits.get(c, 3) if isinstance(digits, Mapping) else digits
                if isinstance(v, (float, np.floating)) or (isinstance(v, (int, np.integer)) and not isinstance(v, bool)):
                    line.append(_n(v, d) if isinstance(v, (float, np.floating)) else str(int(v)))
                else:
                    line.append("" if v is None or (isinstance(v, float) and np.isnan(v)) else str(v))
            rows.append(line)
        return self.table(header, rows, caption, widths_cm, align, font_size)

    # ------------------------------------------------------------------ figures and photos
    def figure(self, path: str, caption: str, width_cm: Optional[float] = None) -> int:
        """The figure at ``path`` with its caption; returns the number of the figure."""
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        w = min(width_cm or report.STYLE.full_width_cm, self.text_width_cm)
        prev = self._last_block()
        self.doc.add_picture(path, width=self.Cm(w))
        pic = self.doc.paragraphs[-1]
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.paragraph_format.keep_with_next = True
        self._blank_line_before(pic, prev)
        return self.caption("figure", caption)

    def set_cell(self, cell, text: str, size: Optional[float] = None, align: str = "left"):
        """Replace the text of a table cell (mark-up allowed)."""
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        p = cell.paragraphs[0]
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        self._add_runs(p, text, size=size)
        p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT,
                       "center": WD_ALIGN_PARAGRAPH.CENTER}[align]

    def photo(self, path: Optional[str], caption: str, height_cm: float = 7.0) -> bool:
        """The photo at ``path`` (scaled to the text width, at most ``height_cm`` high) or a framed placeholder;
        returns True if the photo was inserted."""
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        if path and os.path.isfile(path):
            try:
                from PIL import Image
                with Image.open(path) as im:
                    w_px, h_px = im.size
                aspect = h_px / float(w_px) if w_px else 0.75
            except Exception:
                aspect = 0.75
            w = min(self.text_width_cm, height_cm / aspect if aspect > 0 else self.text_width_cm)
            prev = self._last_block()
            self.doc.add_picture(path, width=self.Cm(w))
            pic = self.doc.paragraphs[-1]
            pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pic.paragraph_format.keep_with_next = True
            self._blank_line_before(pic, prev)
            self.caption("figure", caption)
            return True
        if self._space_below(self._last_block()) < 0.5 * self.line_pt:
            self.doc.add_paragraph().paragraph_format.space_after = self.Pt(0)    # a blank line above the frame
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = t.rows[0].cells[0]
        cell.width = self.Cm(min(12.0, self.text_width_cm))
        t.rows[0].height = self.Cm(height_cm)
        t.rows[0].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._add_runs(p, tr("document.placeholder_photo", caption=plain(caption)), size=9, color="7F7F7F")
        tcpr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        for side in ("top", "left", "bottom", "right"):
            b = OxmlElement("w:%s" % side)
            b.set(qn("w:val"), "dashed")
            b.set(qn("w:sz"), "6")
            b.set(qn("w:color"), "7F7F7F")
            borders.append(b)
        _insert_ordered(tcpr, borders, _TCPR_ORDER)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "F2F2F2")
        _insert_ordered(tcpr, shd, _TCPR_ORDER)
        cap = self.doc.add_paragraph()
        cap.paragraph_format.space_after = self.Pt(0)
        self.caption("figure", caption)
        return False

    # ------------------------------------------------------------------ header and footer (own layout only)
    def header_footer(self, header_text: str):
        if self.templated:
            return
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        sec = self.doc.sections[0]
        sec.different_first_page_header_footer = True
        hp = sec.header.paragraphs[0]
        hp.text = ""
        self._add_runs(hp, header_text, size=8, color="595959")
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp = sec.footer.paragraphs[0]
        fp.text = ""
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._field(fp, "PAGE", "1")
        fp.add_run(" / ")
        self._field(fp, "NUMPAGES", "1")

    def properties(self, title: str, authors: Sequence[str], subject: str = ""):
        cp = self.doc.core_properties
        cp.title = plain(title)
        cp.subject = plain(subject)
        cp.author = ", ".join(authors) if authors else ""
        cp.last_modified_by = cp.author
        cp.comments = "groundfault %s" % __version__
        cp.keywords = ""
        cp.revision = 1                 # Word and LibreOffice count up on every save – see _keep_edited
        cp.created = cp.modified = _dt.datetime.now().replace(microsecond=0)

    def save(self, path: str):
        self.fill_toc()
        self.update_fields_on_open()
        if self._footnotes is not None:
            _store_footnotes(self._footnotes[0], self._footnotes[1])
        folder = os.path.dirname(os.path.abspath(path))
        os.makedirs(folder, exist_ok=True)
        self.doc.save(path)


def _keep_edited(path: str) -> Optional[str]:
    """Keep an existing report that was edited after it was written: saved in Word or LibreOffice (revision > 1) or
    not a report of this module – it is renamed to ``<name>_<date>_<time>.docx`` (its modification time). An
    unchanged report is overwritten. Returns the new file name or None."""
    if not os.path.isfile(path):
        return None
    try:
        import docx
        edited = int(docx.Document(path).core_properties.revision or 0) != 1
    except Exception:                   # not readable as a Word file → keep it
        edited = True
    if not edited:
        return None
    stamp = _dt.datetime.fromtimestamp(os.path.getmtime(path)).strftime("%Y-%m-%d_%H%M%S")
    stem, ext = os.path.splitext(path)
    target, k = "%s_%s%s" % (stem, stamp, ext), 2
    while os.path.exists(target):
        target, k = "%s_%s_%d%s" % (stem, stamp, k, ext), k + 1
    os.replace(path, target)
    return target


def _file_stem(text: str) -> str:
    """Text as part of a file name: letters (also umlauts), digits, "-" and "_" – blanks become "_"."""
    t = re.sub(r"\s+", "_", plain(str(text)).strip())
    return re.sub(r"[^\w\-]+", "", t, flags=re.UNICODE).strip("_") or "report"


def _docx_font() -> str:
    """Font of the text: the font of the figures (a metric-compatible system font counts as Arial)."""
    f = report.used_font()
    return "Arial" if f in ("Liberation Sans", "Helvetica", "DejaVu Sans", "Nimbus Sans") else f


def _clear_body(doc) -> None:
    """Remove the content of a template (keeps the section settings, headers, footers and styles) and the
    relationships of pictures, charts and objects only the removed content used."""
    from docx.oxml.ns import qn
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    xml = body.xml
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    removable = {getattr(RT, k) for k in ("IMAGE", "CHART", "PACKAGE", "OLE_OBJECT", "HYPERLINK", "DIAGRAM_DATA",
                                          "DIAGRAM_LAYOUT", "DIAGRAM_QUICK_STYLE", "DIAGRAM_COLORS") if hasattr(RT, k)}
    removable.add("http://schemas.microsoft.com/office/2007/relationships/diagramDrawing")
    rels = doc.part.rels
    for rid in [k for k, r in rels.items() if r.reltype in removable]:
        if ('"%s"' % rid) not in xml:
            rels.pop(rid)
    part, root = _footnotes_part(doc)
    if root is not None:                                   # the notes of the removed text (the separators stay)
        for fn in list(root.findall(qn("w:footnote"))):
            if fn.get(qn("w:type")) in (None, "normal"):
                root.remove(fn)
        _store_footnotes(part, root)


def _footnotes_part(doc):
    """The footnotes part of a document and its root element ``w:footnotes`` – (None, None) without one. python-docx
    reads the part as bytes; [`_store_footnotes`] writes a changed root back."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import parse_xml
    for rel in doc.part.rels.values():
        if rel.reltype == RT.FOOTNOTES and not rel.is_external:
            part = rel.target_part
            root = getattr(part, "element", None)
            return part, (root if root is not None else parse_xml(part.blob))
    return None, None


def _store_footnotes(part, root) -> None:
    """Write the root of a footnotes part read as bytes back into the part."""
    if getattr(part, "element", None) is None:
        from lxml import etree
        part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


# =================================================================================================== campaign data
def _harmonics(e) -> Dict[int, float]:
    """{order: magnitude} of the main channel of the fault current (assessed window)."""
    return report._main_channel_harmonics(e)


def _assessed(e, quantity: str) -> float:
    """Assessment value of a quantity of the table ``assessment_quantities`` (NaN if missing)."""
    t = e.tables.get("assessment_quantities")
    if t is None or t.empty or "quantity" not in t:
        return float("nan")
    row = t[t["quantity"] == quantity]
    return float(row["assessment"].iloc[0]) if len(row) else float("nan")


def _dc_method(meas: Mapping[str, Any]) -> str:
    """The DC removal of a result file: ``"highpass"`` (high-pass, then the mean per window) or
    ``"highpass_median"`` – a result of groundfault 3.4 with ``"highpass"`` (no ``dc_window`` stored) used the median
    per window – or the method per window."""
    dc = str(meas.get("dc_method", "") or "")
    if dc == "highpass" and meas.get("dc_window", "median") == "median":
        return "highpass_median"
    return dc


def _stored_settings(e) -> dict:
    meas = getattr(e, "measurement", None)
    if isinstance(meas, dict):
        return dict(meas.get("settings") or {})
    s = getattr(e, "settings", None)
    if s is not None and hasattr(s, "__dataclass_fields__"):
        return {k: getattr(s, k) for k in s.__dataclass_fields__}
    return {}


def _measurement(e) -> dict:
    meas = getattr(e, "measurement", None)
    if isinstance(meas, dict):
        return meas
    data = getattr(e, "data", None)
    return dict((data or {}).get("measurement") or {})


def _meta(e) -> dict:
    data = getattr(e, "data", None)
    if isinstance(data, dict):
        return dict(data.get("meta") or {})
    return dict(getattr(e, "meta", {}) or {})


def _is_station(e) -> bool:
    for t in ("station_frequency", "station_voltages"):
        df = e.tables.get(t)
        if df is not None and not df.empty:
            return True
    return False


_CONTROLLER_KEYS = ("ice_A", "coils_A", "detuning_pct", "active_A", "when")
_LIMIT_CASE_KEYS = ("limit_A", "Z_E", "r", "condition")
_EXPANSION_KEYS = ("I_CE", "damping_added", "c1_ratio", "I_L_max", "I_CE_max", "zero_sequence")


def _controller(value: Optional[Mapping[str, Any]]) -> dict:
    """The display of the coil controller (``ReportInfo.coil_controller``) with its I_CE: given, or I_L/(1 + v)."""
    if not value:
        return {}
    unknown = [k for k in value if k not in _CONTROLLER_KEYS]
    if unknown:
        raise ValueError(tr("document.controller_unknown", keys=", ".join(map(str, unknown)),
                            known=", ".join(_CONTROLLER_KEYS)))
    coils = {str(k): float(v) for k, v in dict(value.get("coils_A") or {}).items() if _finite(v)}
    total = float(sum(coils.values())) if coils else float("nan")
    v = float(value["detuning_pct"]) if _finite(value.get("detuning_pct")) else float("nan")
    ice = float(value["ice_A"]) if _finite(value.get("ice_A")) else float("nan")
    how = "display" if _finite(ice) else ""
    if not _finite(ice) and coils and _finite(v) and 1.0 + v / 100.0 > 0:
        ice, how = total / (1.0 + v / 100.0), "derived"
    active = float(value["active_A"]) if _finite(value.get("active_A")) else float("nan")
    return dict(coils=coils, sum=total, v=v, active=active, when=str(value.get("when") or ""), ice=ice, how=how)


def _expansion_settings(value) -> Optional[dict]:
    """The settings of the network expansion (``ReportInfo.expansion``), checked; None: no section."""
    if value is None or value is False:
        return None
    value = {} if value is True else value
    try:
        if not isinstance(value, Mapping):
            raise TypeError(type(value).__name__)
        unknown = [k for k in value if k not in _EXPANSION_KEYS]
        if unknown:
            raise KeyError(", ".join(map(str, unknown)))
        t = value.get("I_CE")
        targets = [] if t is None else [float(t)] if np.ndim(t) == 0 else [float(x) for x in t]
        da = None if value.get("damping_added") is None else float(value["damping_added"])
        c1 = float(value.get("c1_ratio", 1.0))
        il_max = None if value.get("I_L_max") is None else float(value["I_L_max"])
        top = None if value.get("I_CE_max") is None else float(value["I_CE_max"])
        zs = str(value.get("zero_sequence", "coil"))
        if any(not (math.isfinite(x) and x > 0) for x in targets + [c1] + [x for x in (il_max, top) if x is not None]) \
                or (da is not None and not (math.isfinite(da) and da >= 0)) or zs not in ("coil", "network"):
            raise ValueError(", ".join("%s=%r" % (k, value.get(k)) for k in _EXPANSION_KEYS if k in value))
    except (KeyError, TypeError, ValueError) as ex:
        raise ValueError(tr("document.expansion_invalid", error=ex)) from None
    return dict(targets=sorted(set(targets)), damping_added=da, c1_ratio=c1, I_L_max=il_max, I_CE_max=top,
                zero_sequence=zs)


def _limit_cases(cases) -> List[dict]:
    """The limit cases of ``ReportInfo.limit_cases``, checked: ``limit_A``, ``Z_E`` (None: Z_E of the report), ``r``,
    ``condition``."""
    out = []
    for k, case in enumerate(cases or (), start=1):
        try:
            if not isinstance(case, Mapping):
                raise TypeError(type(case).__name__)
            unknown = [x for x in case if x not in _LIMIT_CASE_KEYS]
            if unknown:
                raise KeyError(", ".join(map(str, unknown)))
            lim = float(case["limit_A"])
            z = None if case.get("Z_E") is None else float(case["Z_E"])
            r = float(case.get("r", 1.0))
            if not (math.isfinite(lim) and lim > 0) or (z is not None and not (math.isfinite(z) and z > 0)) \
                    or not (0.0 < r <= 1.0):
                raise ValueError("limit_A=%r, Z_E=%r, r=%r" % (case.get("limit_A"), case.get("Z_E"), case.get("r", 1.0)))
        except (KeyError, TypeError, ValueError) as ex:
            raise ValueError(tr("document.limit_case_invalid", k=k, error=ex)) from None
        out.append(dict(limit_A=lim, Z_E=z, r=r, condition=_markup_symbols(str(case.get("condition") or ""))))
    return out


class _Campaign:
    """The data of the report, prepared from the results and the objects of the campaign evaluation."""

    def __init__(self, results, *, measurements=None, info: ReportInfo, series=None, comparison=None,
                 monte_carlo=None, loo=None, Z_E=None, IG="max", limits=(), u_tp=None, f_factor=None,
                 uncertainty_I=None, uncertainty_U=None, scaling=None, pq=None, pq_levels=None, pq_comparison=None,
                 pq_scaling=None, earthing=None, reduction=None, campaign_folder=None, raw=True, bootstrap=200):
        from . import equivalent_source as es
        from . import evaluation
        if not results:
            raise ValueError(tr("document.no_results"))
        self.info = info
        self.lang = i18n.get_language()
        self.results = dict(results)
        self.names = list(self.results)
        self.meas = {m["name"]: dict(m) for m in (measurements or []) if "name" in m}
        self.folder = str(campaign_folder) if campaign_folder else None
        self.hints: List[str] = []
        self.Z = {}
        for n, e in self.results.items():
            if n in self.meas:
                self.Z[n] = self.meas[n].get("Z")
            else:
                self.Z[n] = _measurement(e).get("z_fault_ohm")
        self.stations = [n for n in self.names if _is_station(self.results[n])]
        self.subs = [n for n in self.names if n not in self.stations and self.Z.get(n) is not None]
        if not self.subs:
            self.subs = [n for n in self.names if n not in self.stations]
        first = self.results[self.names[0]]
        stored = _stored_settings(first)
        self.mI = float(uncertainty_I if uncertainty_I is not None else stored.get("uncertainty_I", 0.0) or 0.0)
        self.mU = float(uncertainty_U if uncertainty_U is not None else stored.get("uncertainty_U", 0.0) or 0.0)
        lim = tuple(float(x) for x in (limits or stored.get("ires_limits") or ()))
        self.limits = lim
        if info.station_detail not in ("full", "compact"):
            raise ValueError(tr("document.station_detail_invalid", value=info.station_detail))
        cases_in = _limit_cases(info.limit_cases)
        expansion = _expansion_settings(info.expansion)
        if info.limit_A is not None:
            self.limit = float(info.limit_A)
        elif cases_in:
            self.limit = float(cases_in[0]["limit_A"])
        else:
            self.limit = min(lim) if lim else None
        self.u_tp = u_tp if u_tp is not None else stored.get("u_tp")
        self.f_factor = f_factor if f_factor is not None else stored.get("f_factor")
        self.series, self.comparison = series, comparison
        self.Z_E = Z_E
        self.IG = IG
        self.scaling = dict(scaling) if scaling else None
        self.pq = pq
        self.pq_levels = pq_levels or {}
        self.pq_comparison = pq_comparison
        self.pq_scaling = pq_scaling
        self.earthing = earthing or {}
        self.reduction = reduction or {}
        self.sites = None
        if self.folder:
            try:
                from . import sites as _sites
                self.sites = _sites.load(self.folder)
            except Exception:
                self.sites = None
        # ---- raw recordings (waveforms, trends, fault times, reproduction)
        self.full: Dict[str, Any] = {}
        self.recordings: Dict[str, Any] = {}
        self.fault_times: Dict[str, Tuple[Optional[float], Optional[float]]] = {}
        self.reproduction: Dict[str, float] = {}
        if raw and info.waveforms and self.folder:
            self._load_raw()
        # ---- per test
        self.tests = {n: self._test(n) for n in self.names}
        # ---- scaling of every test with the PQ measurement of the campaign (consistent within the report)
        self.pq_valid = bool(pq is not None and getattr(pq, "valid", False))
        self.scaled: Dict[str, dict] = {}
        self.thd95_used: Dict[str, Optional[float]] = {}               # THD quantile of the scaling per test
        self.stored_factors = sorted({float(_stored_settings(self.results[n]).get("pq_safety_factor") or 1.0)
                                      for n in self.names if _finite(self.results[n].values.get("IH_scaled"))})
        if self.pq_valid:
            sc = pq.scaling()
            for n in self.names:
                if n in self.stations and not _finite(self.results[n].values.get("IH_scaled")):
                    continue                                    # a station test is scaled only if its evaluation was
                try:
                    thd95 = pq.thd95() if _finite(pq.thd95()) else None
                    self.scaled[n] = evaluation.rescale_harmonics(
                        self.results[n], u95=sc["u95"], u95_reference=sc["u95_reference"],
                        thd95_percent=thd95, safety_factor=pq.factor,
                        uncertainty_I=self.mI, uncertainty_U=self.mU)
                    self.thd95_used[n] = thd95
                except Exception as ex:                         # a test without voltages: not scaled
                    self.hints.append("%s: %s" % (n, ex))
        else:
            for n in self.names:
                e = self.results[n]
                t = e.tables.get("harmonic_scaling")
                if _finite(e.values.get("IH_scaled")) and t is not None and not t.empty:
                    st = _stored_settings(e)
                    self.scaled[n] = evaluation.rescale_harmonics(
                        e, u95=st.get("pq_u95"), u95_reference=st.get("pq_u95_reference") or "seq",
                        thd95_percent=st.get("u_thd95_percent"), safety_factor=float(st.get("pq_safety_factor") or 1),
                        uncertainty_I=self.mI, uncertainty_U=self.mU)
                    self.thd95_used[n] = st.get("u_thd95_percent")
        for n, sres in self.scaled.items():
            big = sres["table"][sres["table"]["factor"] > 20.0]
            if len(big):
                self.hints.append(tr("document.factor_large", name=n, orders=", ".join(str(int(o)) for o in big["order"]),
                                     factor=_n(float(big["factor"].max()))))
        self.factor = float(pq.factor) if self.pq_valid else (self.stored_factors[0] if self.stored_factors else None)
        # ---- basis test: the largest assessment value
        def worst(n):
            s = self.scaled.get(n)
            if s is not None:
                return s["IRES_scaled_assessed"]
            return self.tests[n]["IRES_a"]
        cand = [n for n in self.subs if _finite(worst(n))] or [n for n in self.names if _finite(worst(n))]
        self.basis = max(cand, key=worst) if cand else self.names[0]
        self.largest = max((n for n in self.subs if _finite(self.tests[n]["IRES"])),
                           key=lambda n: self.tests[n]["IRES"], default=self.basis)
        # ---- capacitive earth-fault current, damping, detuning
        sums = [self.tests[n]["coil_sum"] for n in self.subs if _finite(self.tests[n]["coil_sum"])]
        self.coil_sum = float(np.median(sums)) if sums else float("nan")
        self.controller = _controller(info.coil_controller)
        self.coil_display = float(self.controller.get("ice", float("nan")))
        self.ice = self.coil_display if _finite(self.coil_display) else self.coil_sum
        iw = [self.tests[n]["Iw"] for n in self.subs if _finite(self.tests[n]["Iw"])]
        iv = [abs(self.tests[n]["Iv"]) for n in self.subs if _finite(self.tests[n]["Iv"])]
        self.iw = float(np.median(iw)) if iw else float("nan")
        self.iv = float(np.median(iv)) if iv else float("nan")
        self.d = self.iw / self.ice if _finite(self.ice) and self.ice > 0 else float("nan")
        self.v = self.iv / self.ice if _finite(self.ice) and self.ice > 0 else float("nan")
        self.ice_limit = (self.limit / self.d) if (self.limit and _finite(self.d) and self.d > 0) else float("nan")
        # ---- equivalent sources
        self.es = self.es_scaled = self.es_other = None
        self.mc = monte_carlo
        self.loo = loo
        self.boot = None
        self.pq_periods: Optional[dict] = None
        if series is not None and Z_E is not None:
            margins = dict(uncertainty_I=self.mI, uncertainty_U=self.mU)
            self.es = es.ires_at_impedance(series, Z_E, IG=IG, **margins)
            if self.scaling:
                self.es_scaled = es.ires_at_impedance(series, Z_E, IG=IG, **margins, **self.scaling)
            if comparison is not None:
                self.es_other = es.ires_at_impedance(comparison, Z_E, IG=IG, **margins, **(self.scaling or {}))
            if self.mc is None:
                try:
                    self.mc = es.monte_carlo(series, Z_E, n=1000, IG=IG)
                except Exception:
                    self.mc = None
            if self.loo is None:
                try:
                    self.loo = series.loo_table()
                except Exception:
                    self.loo = None
            if self.pq_valid and self.scaling and bootstrap:
                self.boot = self._bootstrap(int(bootstrap))
            if self.pq_valid and self.scaling:
                self.pq_periods = self._periods()
            if self.pq_valid and self.scaling and self.pq_scaling is None:
                try:                                    # the methods with the margins of this report
                    self.pq_scaling = pq.scaling_methods(series, Z_E, IG=IG, uncertainty_I=self.mI,
                                                         uncertainty_U=self.mU, scaling=self.scaling)
                except Exception as ex:
                    self.hints.append(str(ex))
        elif series is not None and self.loo is None:
            try:
                self.loo = series.loo_table()
            except Exception:
                self.loo = None
        # the band of the Monte Carlo simulation with the fundamental of this report: I_RES,p = √(I_G² + I_H,p²) per
        # quantile p (monotonic in I_H) – a simulation of the notebook may use another I_G
        self.mc_band: Dict[int, float] = {}
        self.mc_band_ih: Dict[int, float] = {}
        if self.mc and self.es is not None:
            ih = self.mc.get("IH") or {}
            self.mc_band_ih = {p: float(ih[p]) for p in (5, 50, 95) if _finite(ih.get(p))}
            self.mc_band = {p: math.hypot(float(self.es["IG"]), v) for p, v in self.mc_band_ih.items()}
        # ---- limits with their conditions: I_RES,ass at the Z_E of every case, U_E = r·I_RES,ass·Z_E ↔ F·U_Tp
        self.limit_cases: List[dict] = []
        fu = float(self.f_factor) * float(self.u_tp) if _finite(self.u_tp) and _finite(self.f_factor) else float("nan")
        for case in cases_in:
            z = case["Z_E"] if case["Z_E"] is not None else Z_E
            i, basis = float("nan"), ""
            if series is not None and z is not None:
                try:
                    rr = es.ires_at_impedance(series, z, IG=IG, uncertainty_I=self.mI, uncertainty_U=self.mU,
                                              **(self.scaling or {}))
                    i, basis = float(rr["IRES_scaled_assessed"] if self.scaling else rr["IRES_assessed"]), "Z_E"
                except Exception as ex:
                    self.hints.append(str(ex))
            if not _finite(i):                          # without equivalent sources: the busbar value (upper bound)
                sb = self.scaled.get(self.basis)
                i = float(sb["IRES_scaled_assessed"]) if sb is not None else float(self.tests[self.basis]["IRES_a"])
                basis = "busbar"
            za = abs(complex(z)) if z is not None else float("nan")
            ue = case["r"] * i * za if _finite(i) and _finite(za) else float("nan")
            self.limit_cases.append(dict(case, Z_E=za, Z_E_value=z, I_A=i, basis=basis, U_E_V=ue, FU_V=fu,
                                         ok_I=bool(i <= case["limit_A"]) if _finite(i) else None,
                                         ok_U=bool(ue <= fu) if _finite(ue) and _finite(fu) else None,
                                         applies=self.limit is not None and abs(case["limit_A"] - self.limit) < 1e-9))
        # ---- network expansion: I_RES,ass over I_CE at the Z_E of the limits, where every limit is reached
        self.expansion: Optional[dict] = None
        if expansion is not None and series is not None and series.fits and _finite(self.ice) and self.ice > 0 \
                and (Z_E is not None or any(_finite(lc["Z_E"]) for lc in self.limit_cases)):
            try:
                self.expansion = self._expansion_study(expansion)
            except Exception as ex:
                self.hints.append(tr("document.expansion_failed", error=ex))
        # ---- touch voltage per station
        self.cases: Dict[str, pd.DataFrame] = {}
        for n in self.stations:
            ratio = None
            if n in self.earthing:
                ratio = self.earthing[n][1]
            elif n in self.reduction and self.reduction[n].get("reduction_factor") is not None:
                ratio = self.reduction[n]["reduction_factor"]
            try:
                self.cases[n] = touch_voltage_cases(
                    self.results[n], currents=self.assessed_currents(self.basis), limit_A=self.limit,
                    U_Tp=self.u_tp, F=self.f_factor, r_flat=info.r_flat,
                    reference_frequency_Hz=info.reference_frequency_Hz, earthing=ratio)
            except Exception as ex:
                self.hints.append("%s: %s" % (n, ex))
        # ---- the decisive value
        self.assessment = info.assessment if info.assessment in ("Z_E", "busbar") else (
            "Z_E" if self.es is not None else "busbar")
        if self.assessment == "Z_E" and self.es is None:
            self.assessment = "busbar"

    # ------------------------------------------------------------------ helpers
    def place(self, n: str) -> str:
        """Display name of the measuring point of a test."""
        site = self.tests[n]["site"] if n in getattr(self, "tests", {}) else _meta(self.results[n]).get("site", "")
        if site in self.info.places:
            return str(self.info.places[site])
        loc = self.meas.get(n, {}).get("location")
        if loc:
            return str(loc)
        if self.sites is not None:
            try:
                ent = self.sites.for_site(site)
                if ent is not None and ent.location:
                    return ent.location
            except Exception:
                pass
        return str(site).replace("_", " ") if site else n

    def _test(self, n: str) -> dict:
        e = self.results[n]
        v = e.values
        meas = _measurement(e)
        meta = _meta(e)
        coils = []
        cc = e.tables.get("coil_currents")
        if cc is not None and not cc.empty:
            coils = [(str(r["channel"]), float(r["H1_A"])) for _, r in cc.iterrows() if _finite(r.get("H1_A"))]
        ires = float(v.get("IRES", np.nan))
        ires_a = ires * (1 + self.mI) if _finite(ires) else float("nan")       # the margin of this report
        start = None
        if meta.get("start_time"):
            try:
                start = pd.Timestamp(meta["start_time"])
            except (ValueError, TypeError):
                start = None
        return dict(IRES=ires, IRES_a=ires_a, IG=float(v.get("IG", np.nan)), Iw=float(v.get("Iw", np.nan)),
                    Iv=float(v.get("Iv", np.nan)), IH=float(v.get("IH_total", np.nan)),
                    Ilf=float(v.get("ires_low_frequency_A", np.nan)), harmonics=_harmonics(e), coils=coils,
                    coil_sum=float(sum(c[1] for c in coils)) if coils else float("nan"),
                    phase=v.get("faulted_phase") or "", t_max=float(v.get("t_max_window_s", np.nan)),
                    start=start, site=meta.get("site", ""), cfg=meas.get("cfg_file", ""),
                    fs=float(meas.get("fs_hz", np.nan)), duration=float(meas.get("duration_s", np.nan)),
                    f1=float(meas.get("f1_hz", np.nan)), steady=meas.get("steady_state_s"),
                    pre=meas.get("pre_fault_s"), post=meas.get("post_fault_s"),
                    T_win=float(meas.get("T_win_s", np.nan)), window=meas.get("window", ""),
                    dc=_dc_method(meas), main=(meas.get("main_channel") or {}).get("i_res", ""),
                    corrections=list(meta.get("corrections") or []), Z=self.Z.get(n),
                    polarity=meas.get("polarity_i_res"))

    def assessed_currents(self, n: str) -> Dict[int, float]:
        """Assessment current per order of a test: I_G·(1 + m_I) and the scaled harmonics with the margins."""
        t = self.tests[n]
        out = {1: t["IG"] * (1 + self.mI)} if _finite(t["IG"]) else {}
        s = self.scaled.get(n)
        if s is not None:
            for _, r in s["table"].iterrows():
                ka = 1.0 / (1.0 - self.mU) if r["rule"] in ("own voltage", "THD") else 1.0
                out[int(r["order"])] = float(r["I_scaled_A"]) * ka * (1 + self.mI)
        else:
            for m, i in t["harmonics"].items():
                if m >= 2:
                    out[int(m)] = float(i) * (1 + self.mI)
        return out

    def _load_raw(self):
        """Evaluate every recording again from its raw data (waveforms, trends, fault times, reproduction)."""
        import groundfault as gf
        from . import corrections, evaluation, layout, station as _station
        try:
            lay = layout.detect_layout(self.folder)
            corr = corrections.load(self.folder)
        except Exception as ex:
            self.hints.append(str(ex))
            return
        known = set(evaluation.Settings.__dataclass_fields__)
        for n, e in self.results.items():
            meas, meta = _measurement(e), _meta(e)
            cfg = meas.get("cfg_file") or ""
            if not cfg:
                continue
            path = None
            for raw in dict.fromkeys((lay.raw, layout.ENGLISH.raw, layout.GERMAN.raw)):
                cand = os.path.join(self.folder, raw, meta.get("site") or "", cfg)
                if os.path.isfile(cand):
                    path = cand
                    break
                base = os.path.join(self.folder, raw)
                if os.path.isdir(base):
                    for dp, _, files in os.walk(base):
                        if cfg in files:
                            path = os.path.join(dp, cfg)
                            break
                if path:
                    break
            if path is None:
                self.hints.append(tr("document.raw_missing", name=n, file=cfg))
                continue
            try:
                rec = gf.open_recording(path)
                corr.apply(rec)
                roles = gf.detect_roles(rec.channel_names, rec.units, corr.role_overrides(rec), corr.main_channel(rec),
                                        corr.distances(rec))
                st = dict(meas.get("settings") or {})
                if st.get("pq_u95"):
                    st["pq_u95"] = {int(k): x for k, x in st["pq_u95"].items()}
                for k in ("harmonics", "export_orders", "voltage_harmonics", "ires_limits"):
                    if isinstance(st.get(k), list):
                        st[k] = tuple(st[k])
                settings = evaluation.Settings(**{k: x for k, x in st.items() if k in known})
                iv = evaluation.Intervals(tuple(meas["steady_state_s"]),
                                          tuple(meas["pre_fault_s"]) if meas.get("pre_fault_s") else None,
                                          tuple(meas["post_fault_s"]) if meas.get("post_fault_s") else None)
                with warnings.catch_warnings():
                    warnings.simplefilter("ignore")
                    res = gf.evaluate(rec, iv, settings, roles=roles, z_fault_ohm=meas.get("z_fault_ohm"))
                self.full[n] = res
                self.recordings[n] = rec
                dev = []
                for k in ("IRES", "IG", "IH_total"):
                    a, b = e.values.get(k), res.values.get(k)
                    if _finite(a) and _finite(b) and abs(float(a)) > 0:
                        dev.append(abs(float(b) - float(a)) / abs(float(a)))
                self.reproduction[n] = 100.0 * max(dev) if dev else float("nan")
                k = res.roles.main("i_res")
                if k is not None:
                    ss = meas.get("steady_state_s") or [None, None]
                    mid = 0.5 * (ss[0] + ss[1]) if ss[0] is not None else None
                    self.fault_times[n] = _station.fault_times(rec, k.index, t_inside=mid)
            except Exception as ex:
                self.hints.append(tr("document.raw_failed", name=n, error=ex))

    def _scaled_with(self, row) -> float:
        """I_RES(Z_E) scaled with the quantiles of ``row`` (per order and ``THD_LL_pct``) instead of those of the
        measurement (the factor and the rules of the scaling as in this report); NaN if it cannot be formed."""
        from . import equivalent_source as es
        sc = dict(self.scaling)
        u = dict(sc.get("u95") or {})
        for m in list(u):
            if m in row.index and _finite(row[m]):
                u[m] = float(row[m])
        sc["u95"] = u
        if "thd95_percent" in sc and "THD_LL_pct" in row.index and _finite(row["THD_LL_pct"]):
            sc["thd95_percent"] = float(row["THD_LL_pct"])
        try:
            return float(es.ires_at_impedance(self.series, self.Z_E, IG=self.IG, uncertainty_I=self.mI,
                                              uncertainty_U=self.mU, **sc)["IRES_scaled"])
        except Exception:
            return float("nan")

    def _periods(self) -> Optional[dict]:
        """Measuring periods of the PQ measurement (two or more): the quantile of the dominant order per period and
        I_RES(Z_E) scaled with every period alone – how much the period of the measurement matters."""
        try:
            p = self.pq.periods()
        except Exception:
            return None
        if len(p) < 2:
            return None
        dom = self.pq.dominant_order
        ires = [self._scaled_with(row) for _, row in p.iterrows()]
        return dict(table=p, order=dom, ires=ires, covered=float(p["days"].sum()))

    def _bootstrap(self, n: int) -> Optional[dict]:
        """I_RES(Z_E) scaled with the quantiles of every bootstrap sample of the PQ measurement."""
        try:
            b = self.pq.quantile_bootstrap(n=n)
        except Exception:
            return None
        vals = [v for v in (self._scaled_with(row) for _, row in b.iterrows()) if _finite(v)]
        if not vals:
            return None
        dom = self.pq.dominant_order
        out = dict(IRES_p5=float(np.percentile(vals, 5)), IRES_p95=float(np.percentile(vals, 95)), n=len(vals))
        if dom is not None and dom in b.columns:
            q = self.pq.u95().get(dom)
            if _finite(q) and q > 0:
                out.update(order=dom, low=100 * (float(b[dom].quantile(0.05)) / q - 1),
                           high=100 * (float(b[dom].quantile(0.95)) / q - 1))
        return out

    def _expansion_study(self, xp: dict) -> dict:
        """The network expansion (``ReportInfo.expansion``): for both splits of the source impedance the assessment
        value over I_CE at the earthing impedance of every limit, the values at the planned I_CE and where every limit
        is reached (the larger value of the splits decides)."""
        from . import equivalent_source as es
        ice = float(self.ice)
        ctrl = self.controller
        if _finite(ctrl.get("sum")) and ctrl["sum"] > 0:
            il = float(ctrl["sum"])                            # the coils of the controller display
        elif _finite(ctrl.get("v")):
            il = ice * (1.0 + float(ctrl["v"]) / 100.0)
        elif _finite(self.coil_sum) and self.coil_sum > 0:
            il = float(self.coil_sum)
        else:
            il = ice
        d = float(self.d) if _finite(self.d) else 0.0
        da = d if xp["damping_added"] is None else float(xp["damping_added"])
        targets = list(xp["targets"]) or [10.0 * round(f * ice / 10.0) for f in (1.25, 1.5, 2.0)]
        kw = dict(I_L=il, damping=d, damping_added=da, c1_ratio=xp["c1_ratio"], zero_sequence=xp["zero_sequence"])
        base = dict(IG=self.IG, uncertainty_I=self.mI, uncertainty_U=self.mU, **(self.scaling or {}))
        col = "IRES_scaled_assessed_A" if self.scaling else "IRES_assessed_A"
        ih_col = "IH_scaled_A" if self.scaling else "IH_A"
        busbar = not es.busbar_impedances(self.series).empty
        splits = ("sources", "busbar") if busbar else ("sources",)
        top = float(xp["I_CE_max"]) if xp["I_CE_max"] else max(4.0 * ice, 1.1 * max(targets))
        grid = np.linspace(ice, max(top, ice * 1.01), 241)
        # the earthing impedances (as given, complex possible): that of the report first, then those of the limits
        key = lambda z: (round(complex(z).real, 9), round(complex(z).imag, 9))
        with_z = [lc for lc in self.limit_cases if _finite(lc["Z_E"])]
        z_rep = self.Z_E if self.Z_E is not None else with_z[0]["Z_E_value"]
        zs = {key(z_rep): z_rep}
        for lc in sorted(with_z, key=lambda lc: -lc["Z_E"]):
            zs.setdefault(key(lc["Z_E_value"]), lc["Z_E_value"])
        curves: Dict[Tuple[float, float], dict] = {}
        for k, z in zs.items():
            per = {sp: es.expansion_curve(self.series, z, ice, grid, **base, split=sp, **kw) for sp in splits}
            dec = pd.DataFrame({"I_C_A": per["sources"]["I_C_A"],
                                "decisive": np.nanmax(np.vstack([per[sp][col].to_numpy(float) for sp in splits]),
                                                      axis=0)})
            curves[k] = dict(Z_E=abs(complex(z)), frames=per, decisive=dec)
        # the planned values at the earthing impedance of the report
        rows = {sp: es.expansion_curve(self.series, z_rep, ice, [ice] + targets, **base, split=sp, **kw)
                for sp in splits}
        first = {sp: es.network_expansion(self.series, ice, targets[0], split=sp, **kw) for sp in splits}
        negative = sorted({int(r["order"]) for sp in splits for _, r in first[sp].table.iterrows()
                           if complex(r["Z0"]).real < 0 or complex(r["Z1"]).real < 0})
        rejected = sorted({int(r["order"]) for sp in splits for _, r in first[sp].table.iterrows()
                           if bool(r.get("busbar_rejected", False))})
        now = float(curves[key(z_rep)]["decisive"]["decisive"].iloc[0])
        # the limits: those of the cases (at their Z_E), else the limits of the call at the Z_E of the report
        fu = float(self.f_factor) * float(self.u_tp) if _finite(self.u_tp) and _finite(self.f_factor) else float("nan")
        cases = [dict(limit_A=lc["limit_A"], Z_E=float(lc["Z_E"]), key=key(lc["Z_E_value"]), r=lc["r"],
                      condition=lc["condition"], applies=lc["applies"]) for lc in with_z]
        if not cases:
            cases = [dict(limit_A=float(x), Z_E=abs(complex(z_rep)), key=key(z_rep), r=1.0, condition="",
                          applies=self.limit is not None and abs(float(x) - self.limit) < 1e-9)
                     for x in sorted(self.limits)]
        for lc in cases:
            allowed, governs = float(lc["limit_A"]), "I"
            if _finite(fu) and lc["Z_E"] > 0:
                iu = fu / (lc["r"] * lc["Z_E"])
                if iu < allowed:
                    allowed, governs = iu, "U"
            cv = curves[lc["key"]]
            lc.update(allowed=allowed, governs=governs, now=float(cv["decisive"]["decisive"].iloc[0]),
                      reach=es.expansion_limit(cv["decisive"], allowed, "decisive"),
                      per={sp: es.expansion_limit(cv["frames"][sp], allowed, col) for sp in splits})
        reached = [lc["reach"]["I_C_A"] for lc in cases if lc["reach"]["status"] == "reached"]
        end = float(xp["I_CE_max"]) if xp["I_CE_max"] else min(grid[-1], max([1.5 * ice, 1.1 * max(targets)] +
                                                                              [1.1 * x for x in reached]))
        coil_end = float(xp["I_L_max"]) * ice / il if xp["I_L_max"] else float("nan")
        return dict(ice=ice, il=il, d=d, damping_added=da, c1_ratio=float(xp["c1_ratio"]),
                    zero_sequence=xp["zero_sequence"], targets=targets, planned=bool(xp["targets"]), splits=splits,
                    Z_E=abs(complex(z_rep)), curves=curves, rows=rows, first=first, negative=negative,
                    rejected=rejected, now=now,
                    cases=cases, column=col, ih_column=ih_col, end=end, top=float(grid[-1]), I_L_max=xp["I_L_max"],
                    coil_end=coil_end, U_LE=float(first["sources"].U_LE), fu=fu)

    def ok(self, value: float) -> str:
        """Status of a current against the limit of the network."""
        if self.limit is None or not _finite(value):
            return "–"
        return tr("document.met") if value <= self.limit else tr("document.exceeded")


# =================================================================================================== touch voltage
REFERENCE_TOLERANCE_HZ = 5.0
"""Largest distance in Hz of a measured frequency of the earthing measurement from the reference frequency of the
cases 1 and 2 of the touch voltage; farther, Z_E is interpolated linearly between the two neighbouring frequencies
(outside the measured range: the nearest)."""


def touch_voltage_cases(station, currents: Optional[Mapping[int, float]] = None, limit_A: Optional[float] = None,
                        U_Tp: Optional[float] = None, F: Optional[float] = None, r_flat: Optional[float] = None,
                        reference_frequency_Hz: Optional[float] = None,
                        earthing: Optional[pd.DataFrame] = None) -> pd.DataFrame:
    """Touch voltage at a station by four approaches, U_T = r·I·Z_E/F = U_E/F.

    1. the limit current with r = 1 and Z_E of the earthing measurement at ``reference_frequency_Hz``,
    2. the same with the flat reduction factor ``r_flat``,
    3. per frequency: r·Z_E as measured in the earth-fault test (fault-related, the earth-potential probe – role
       ``u_e`` – with the largest value) and the assessment current per order (``currents``) –
       U_E = √(Σ (r·Z_E)_m²·I_m²); orders without their own value take the largest measured r·Z_E,
    4. the measured earth potential rise (the same probe) at the fault current of the test, and – 4b – extrapolated
       to the limit current; with a touch-voltage probe also its measured value.

    Parameters
    ----------
    station : Result or ResultFile
        Evaluated test at the station (tables ``station_voltages``, ``station_frequency``, ``station_metrics``).
    currents : mapping of int to float, optional
        Assessment current per order in A (e.g. the fundamental and the scaled harmonics of the largest test, with
        the margin); without: case 3 uses the fault current of the station test.
    limit_A : float, optional
        Limit of the residual current (cases 1, 2 and 4b).
    U_Tp : float, optional
        Permissible touch voltage in V (column ``share_pct``).
    F : float, optional
        Ratio of earth potential rise to touch voltage (without: U_T is not formed).
    r_flat : float, optional
        Flat reduction factor (case 2; None: left out).
    reference_frequency_Hz : float, optional
        Frequency of Z_E in the cases 1 and 2 (default: 50 Hz) – measured (±``REFERENCE_TOLERANCE_HZ``), else
        interpolated linearly between the neighbouring frequencies of the earthing measurement, outside its range the
        nearest one.
    earthing : pandas.DataFrame, optional
        Earthing measurement per order: the ratio table of
        [`groundfault.earthing.compare`][groundfault.earthing.compare] or the table of
        [`groundfault.earthing.reduction_factor`][groundfault.earthing.reduction_factor] (column
        ``Z_E_injection_Ohm``); without: cases 1 and 2 are left out.

    Returns
    -------
    pandas.DataFrame
        Columns ``case``, ``label``, ``description``, ``I_A``, ``r``, ``Z_E_Ohm``, ``U_E_V``, ``F``, ``U_T_V``,
        ``U_Tp_V``, ``share_pct``; ``attrs``: ``orders_measured`` (orders with r·Z_E of the test), ``covered_pct``
        (share of Σ I² of the assessment currents at measured orders), ``probe``, ``reference_frequency_Hz`` (the
        frequency of Z_E in the cases 1 and 2), ``reference_frequency_requested_Hz``, ``reference_method``
        (``"measured"``, ``"interpolated"`` or ``"nearest"``) and ``reference_bracket_Hz`` (the neighbouring
        frequencies of an interpolation).
    """
    rows = []
    F_ = float(F) if _finite(F) and float(F) > 0 else None
    utp = float(U_Tp) if _finite(U_Tp) else float("nan")

    def add(case, i, r, z, ue):
        ut = ue / F_ if F_ is not None and _finite(ue) else float("nan")
        rows.append(dict(case=case, label=tr("document.case_short_" + case), description="", I_A=i, r=r, Z_E_Ohm=z,
                         U_E_V=ue, F=F_ if F_ is not None else float("nan"), U_T_V=ut, U_Tp_V=utp,
                         share_pct=100.0 * ut / utp if _finite(ut) and _finite(utp) and utp > 0 else float("nan")))

    f_req = float(reference_frequency_Hz) if _finite(reference_frequency_Hz) else 50.0
    f_ref, z_ref, how, bracket = f_req, float("nan"), "", None
    if earthing is not None and len(earthing) and "Z_E_injection_Ohm" in earthing:
        df = earthing
        for col in ("channel", "probe"):
            if col in df and (df[col] == "max(UE)").any():
                df = df[df[col] == "max(UE)"]
                break
        zs = pd.DataFrame(dict(f=pd.to_numeric(df["f_Hz"], errors="coerce"),
                               z=pd.to_numeric(df["Z_E_injection_Ohm"], errors="coerce"))).dropna()
        if len(zs):
            zs = zs.groupby(zs["f"].round(1))["z"].mean()             # several earthing measurements: their mean
            f_, z_ = zs.index.to_numpy(float), zs.to_numpy(float)
            j = int(np.argmin(np.abs(f_ - f_req)))
            if abs(f_[j] - f_req) <= REFERENCE_TOLERANCE_HZ:           # measured at the reference frequency
                f_ref, z_ref, how = float(f_[j]), float(z_[j]), "measured"
            elif f_[0] < f_req < f_[-1]:                                # between two measured frequencies
                z_ref, how = float(np.interp(f_req, f_, z_)), "interpolated"
                bracket = (float(f_[f_ < f_req].max()), float(f_[f_ > f_req].min()))
            else:                                                       # outside: the nearest measured frequency
                f_ref, z_ref, how = float(f_[j]), float(z_[j]), "nearest"
    lim = float(limit_A) if _finite(limit_A) else float("nan")
    if _finite(lim) and _finite(z_ref):
        add("1", lim, 1.0, z_ref, lim * z_ref)
        if _finite(r_flat):
            add("2", lim, float(r_flat), z_ref, float(r_flat) * lim * z_ref)
    # ---- case 3: r·Z_E per order of the earth-fault test (largest probe, fault-related)
    sv = station.tables.get("station_voltages")
    sf = station.tables.get("station_frequency")
    probe = None
    if sv is not None and not sv.empty:
        ue = sv[sv["role"] == "u_e"] if (sv["role"] == "u_e").any() else sv
        probe = str(ue.loc[pd.to_numeric(ue["RMS_V"], errors="coerce").idxmax(), "channel"])
    rze: Dict[int, float] = {}
    if sf is not None and not sf.empty and probe is not None:
        g = sf[sf["channel"] == probe]
        for _, r in g.iterrows():
            z = r.get("Z_fault_Ohm")
            if not _finite(z):
                z = r.get("Z_direct_Ohm")
            if _finite(z):
                rze[int(r["order"])] = float(z)
    if not rze and sv is not None and probe is not None:
        row = sv[sv["channel"] == probe].iloc[0]
        for c in sv.columns:
            mm = re.fullmatch(r"rZE_H(\d+)_Ohm", c)
            if mm and _finite(row[c]):
                rze[int(mm.group(1))] = float(row[c])
    metrics = station.tables.get("station_metrics")
    i_station = float(metrics["IF_A"].iloc[0]) if metrics is not None and not metrics.empty else float("nan")
    cur = dict(currents or {})
    if not cur:
        cur = {int(m): float(i) for m, i in report._main_channel_harmonics(station).items()}
    covered = float("nan")
    if rze and cur:
        zmax = max(rze.values())
        s2 = sum((rze.get(m, zmax) * i) ** 2 for m, i in cur.items())
        tot = sum(i ** 2 for i in cur.values())
        covered = 100.0 * sum(i ** 2 for m, i in cur.items() if m in rze) / tot if tot > 0 else float("nan")
        add("3", math.sqrt(tot), float("nan"), float("nan"), math.sqrt(s2))
    if metrics is not None and not metrics.empty and _finite(metrics["UE_max_V"].iloc[0]):
        ue_m = float(metrics["UE_max_V"].iloc[0])
        add("4", i_station, float("nan"), float("nan"), ue_m)
        if _finite(lim) and _finite(i_station) and i_station > 0:
            add("4b", lim, float("nan"), float("nan"), ue_m * lim / i_station)
        ut_m = metrics["UT_V"].iloc[0] if "UT_V" in metrics else float("nan")
        if _finite(ut_m):
            rows.append(dict(case="ut", label=tr("document.case_short_ut"), description="", I_A=i_station,
                             r=float("nan"), Z_E_Ohm=float("nan"), U_E_V=float("nan"), F=float("nan"),
                             U_T_V=float(ut_m), U_Tp_V=utp,
                             share_pct=100.0 * float(ut_m) / utp if _finite(utp) and utp > 0 else float("nan")))
    desc = {"1": tr("document.case_1", f=_n(f_ref, 3)), "2": tr("document.case_2", f=_n(f_ref, 3)),
            "3": tr("document.case_3"), "4": tr("document.case_4"), "4b": tr("document.case_4b"),
            "ut": tr("document.case_ut_measured")}
    for r in rows:
        r["description"] = desc[r["case"]]
    out = pd.DataFrame(rows, columns=["case", "label", "description", "I_A", "r", "Z_E_Ohm", "U_E_V", "F", "U_T_V",
                                      "U_Tp_V", "share_pct"])
    out.attrs["orders_measured"] = sorted(rze)
    out.attrs["covered_pct"] = covered
    out.attrs["probe"] = probe
    out.attrs["reference_frequency_Hz"] = f_ref
    out.attrs["reference_frequency_requested_Hz"] = f_req
    out.attrs["reference_method"] = how
    out.attrs["reference_bracket_Hz"] = bracket
    return out


# =================================================================================================== sections
#: groups of channel roles for the table of measured quantities (key, roles, English and German label)
_ROLE_GROUPS = (
    ("u_le", ("u_l1e", "u_l2e", "u_l3e"), "line-to-earth voltages (MV)", "Leiter-Erde-Spannungen (MS)"),
    ("u_hv", ("u_hv_l1e", "u_hv_l2e", "u_hv_l3e"), "line-to-earth voltages (HV)", "Leiter-Erde-Spannungen (HS)"),
    ("u_ll", ("u_l12", "u_l23", "u_l31"), "line-to-line voltages", "Leiter-Leiter-Spannungen"),
    ("u_ne", ("u_ne",), "neutral displacement voltage", "Verlagerungsspannung"),
    ("i_res", ("i_res",), "fault current (residual earth-fault current)", "Fehlerstrom (Erdschlussreststrom)"),
    ("i_coil", ("i_coil",), "coil currents", "Spulenströme"),
    ("i_l", ("i_l1", "i_l2", "i_l3"), "line currents of the transformer", "Leiterströme des Transformators"),
    ("i_shield", ("i_shield",), "shield current", "Schirmstrom"),
    ("i_pen", ("i_pen",), "PEN currents", "PEN-Ströme"),
    ("u_e", ("u_e",), "earth potential rise (probe)", "Erdungsspannung (Sonde)"),
    ("u_t", ("u_t",), "touch voltage (probe)", "Berührungsspannung (Sonde)"),
    ("other", ("u_gen", "i_gen"), "further channels", "weitere Kanäle"),
)


class _Writer:
    """Writes the sections of the report."""

    def __init__(self, doc: _Doc, ctx: _Campaign, figures_folder: str, photos: Mapping[str, Any],
                 media_folder: Optional[str]):
        self.d, self.c = doc, ctx
        self.fig_dir = figures_folder
        self.photos = dict(photos or {})
        self.media = media_folder
        self.figure_files: List[str] = []
        self.placeholders: List[str] = []
        self.order_threshold = float(ctx.info.threshold_A)
        self.sections: Dict[str, str] = {}                 # key → number of its heading (key results)
        self.tables: Dict[str, int] = {}                   # key → number of a table referred to in the text
        self.kr_cells: List[Tuple[Any, str]] = []          # cells of the key results still to get their section

    # ------------------------------------------------------------------ helpers
    def fig(self, make, name: str, caption: str, width_cm: Optional[float] = None) -> int:
        """Draw a figure with ``make()``, save it as PNG and insert it (left out with a hint if drawing fails);
        returns its number (0: left out)."""
        import matplotlib.pyplot as plt
        try:
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                f = make()
        except Exception as ex:
            self.c.hints.append(tr("document.figure_failed", name=name, error=ex))
            plt.close("all")
            return 0
        os.makedirs(self.fig_dir, exist_ok=True)
        paths = report.save(f, os.path.join(self.fig_dir, name), ("png",), close=True)
        png = [p for p in paths if p.lower().endswith(".png")]
        self.figure_files += paths
        if png:
            return self.d.figure(png[0], caption, width_cm)
        return 0

    def heading(self, text: str, level: int, key: Optional[str] = None) -> str:
        """A numbered heading; ``key`` records its number for the references of the key results."""
        number = self.d.heading(text, level)
        if key:
            self.sections[key] = number
        return number

    def table_number(self) -> int:
        """Number of the table written last."""
        return int(self.d.counters["table"])

    def photo(self, slot: str, caption: str, key: Optional[str] = None):
        value = self.photos.get(key or slot, self.photos.get(slot))
        if value is False:
            return
        path, cap = None, caption
        if isinstance(value, Mapping):
            path, cap = value.get("path"), value.get("caption") or caption
        elif value:
            path = str(value)
        if path and not os.path.isabs(path) and self.media:
            cand = os.path.join(self.media, path)
            path = cand if os.path.isfile(cand) or not os.path.isfile(path) else path
        if path and not os.path.isfile(path):
            self.c.hints.append(tr("document.photo_missing", path=path))
            path = None
        if not self.d.photo(path, cap):
            self.placeholders.append(plain(cap))

    def extra_photos(self, section: str):
        for item in self.photos.get("extra") or []:
            if isinstance(item, Mapping) and str(item.get("section", "1.1")) == section:
                path = item.get("path")
                if path and not os.path.isabs(path) and self.media:
                    path = os.path.join(self.media, path)
                if path and not os.path.isfile(path):
                    self.c.hints.append(tr("document.photo_missing", path=path))
                    path = None
                if not self.d.photo(path, item.get("caption", "")):
                    self.placeholders.append(plain(item.get("caption", "")))

    def f_text(self, m: int) -> str:
        return _n(m * 50.0 if not self.c.results else m * self._f_nominal(), 4)

    def _f_nominal(self) -> float:
        e = self.c.results[self.c.names[0]]
        s = getattr(e, "settings", None)
        return float(getattr(s, "f_nominal", 50.0) or 50.0)

    def time_text(self, n: str) -> str:
        t = self.c.tests[n]
        if t["start"] is None:
            return "–"
        on = self.c.fault_times.get(n, (None, None))[0]
        ts = t["start"] + pd.Timedelta(seconds=float(on)) if on is not None else t["start"]
        loc = _local(ts, self.c.info.timezone)
        if loc is None:
            return "–"
        return loc.strftime("%H:%M:%S") + ((" " + self.tz_labels()[n]) if self.tz_mixed() else "")

    def tz_labels(self) -> Dict[str, str]:
        """Label of the time zone per test with a start time (e.g. MESZ/MEZ – they differ across a clock change)."""
        return {n: _tz_label(self.c.info.timezone, _local(self.c.tests[n]["start"], self.c.info.timezone))
                for n in self.c.names if self.c.tests[n]["start"] is not None}

    def tz_mixed(self) -> bool:
        """True if the tests lie on both sides of a clock change (summer and standard time)."""
        return len(set(self.tz_labels().values())) > 1

    def tz_text(self) -> str:
        """Label of the time zone in a column heading – "local time" if the tests lie on both sides of a clock
        change, each time then carries its own label."""
        if self.tz_mixed():
            return tr("document.tz_local")
        n = next((x for x in self.c.names if self.c.tests[x]["start"] is not None), None)
        ts = _local(self.c.tests[n]["start"], self.c.info.timezone) if n else None
        return _tz_label(self.c.info.timezone, ts)

    def duration(self, n: str) -> float:
        on, off = self.c.fault_times.get(n, (None, None))
        return float(off - on) if on is not None and off is not None else float("nan")

    def dates(self) -> List[Tuple[str, List[str]]]:
        """(date, places) of the tests, by day."""
        out: Dict[str, List[str]] = {}
        for n in self.c.names:
            t = self.c.tests[n]
            loc = _local(t["start"], self.c.info.timezone)
            key = _date_text(loc.date() if loc is not None else None, self.c.lang)
            out.setdefault(key, [])
            p = self.c.place(n)
            if p not in out[key]:
                out[key].append(p)
        return list(out.items())

    @staticmethod
    def join(items: Sequence[str]) -> str:
        items = list(items)
        if len(items) <= 1:
            return "".join(items)
        word = " und " if i18n.get_language() == "de" else " and "
        return ", ".join(items[:-1]) + word + items[-1]

    def orders_shown(self, names: Sequence[str]) -> List[int]:
        """Orders ≥ 2 with at least ``threshold_A`` in one of the tests."""
        orders = set()
        for n in names:
            for m, i in self.c.tests[n]["harmonics"].items():
                if m >= 2 and _finite(i) and i >= self.order_threshold:
                    orders.add(int(m))
        return sorted(orders)

    def ordinal(self, m: int) -> str:
        """An order as an ordinal: "5." / "5th" (corrected to "3rd" … when written)."""
        return "%d." % int(m) if self.c.lang == "de" else "%dth" % int(m)

    @staticmethod
    def status(ok) -> str:
        """Kept / exceeded / – for a check."""
        return "–" if ok is None else tr("document.met") if ok else tr("document.exceeded")

    def percentile_range(self, values) -> str:
        """A percentile or a range of percentiles as ordinals ("15.–23." / "15th–23rd"); "" without values."""
        v = [float(x) for x in values if _finite(x)]
        if not v:
            return ""
        lo, hi = int(round(min(v))), int(round(max(v)))
        return self.ordinal(lo) if lo == hi else "%s–%s" % (self.ordinal(lo), self.ordinal(hi))

    def u_level(self, n: str, m: int) -> float:
        """Voltage of order ``m`` before the fault: the component of the natural sequence (positive for 1, 7, 13 …,
        negative for 5, 11 …), mean of before and after the fault (NaN without)."""
        v = self.c.results[n].values.get("voltage") or {}
        ords = [int(o) for o in v.get("orders", [])]
        if int(m) not in ords:
            return float("nan")
        seq = "U_pos" if int(m) % 3 == 1 else "U_neg"
        j = ords.index(int(m))
        vals = [float(v[b][seq][j]) for b in ("pre_fault", "post_fault")
                if isinstance(v.get(b), dict) and seq in v[b] and _finite(v[b][seq][j])]
        return float(np.mean(vals)) if vals else float("nan")

    def earthing_info(self, n: str):
        """The earthing measurement of a station: (rows of the probe maximum, ids, details text, probe)."""
        c = self.c
        long_df, ratio = c.earthing[n]
        rr = ratio[ratio["channel"] == "max(UE)"] if "channel" in ratio and (ratio["channel"] == "max(UE)").any() \
            else ratio
        ids = sorted({int(x) for x in pd.to_numeric(rr.get("injection_id", pd.Series(dtype=float)),
                                                     errors="coerce").dropna()}) if len(rr) else []
        details = ""
        if len(ids) == 1:
            try:
                from . import earthing as _earthing
                inj = long_df[(long_df["source"] == _earthing.SOURCE_INJECTION)
                              & (pd.to_numeric(long_df["measurement_id"], errors="coerce") == ids[0])]
                date = str(inj["date"].iloc[0])[:10] if len(inj) and "date" in inj else ""
                if date:
                    details += tr("document.earthing_details", date=_date_short(pd.Timestamp(date), c.lang))
            except Exception:
                pass
            dist = pd.to_numeric(rr.get("distance_injection_m", pd.Series(dtype=float)), errors="coerce").dropna()
            if len(dist):
                details += tr("document.earthing_distance", d=_n(float(dist.iloc[0])))
        probe = str(rr["probe"].iloc[0]) if "probe" in rr and len(rr) else ""
        return rr, ids, details, probe

    def context_station_busbar(self) -> List[str]:
        """Why a station shows a larger harmonic current than the busbar test: the larger voltage before the fault
        (active two-terminal network) – one sentence per station, for the orders of the natural sequence with a
        larger current and a larger voltage, but a smaller current per volt."""
        c = self.c
        subs = [n for n in c.subs if _finite(c.Z.get(n))]
        if not subs or not c.stations:
            return []
        n0 = min(subs, key=lambda n: float(c.Z[n]))
        out = []
        for ns in c.stations:
            items = []
            for m in sorted(c.tests[ns]["harmonics"]):
                if m < 2 or fml.harmonic_category(int(m)) != "pos_neg":
                    continue
                i_s, i_0 = c.tests[ns]["harmonics"].get(m), c.tests[n0]["harmonics"].get(m)
                if not (_finite(i_s) and _finite(i_0)) or i_s < self.order_threshold or i_s <= i_0:
                    continue
                u_s, u_0 = self.u_level(ns, m), self.u_level(n0, m)
                if not (_finite(u_s) and _finite(u_0)) or u_0 <= 0 or u_s <= u_0 or i_s / u_s >= i_0 / u_0:
                    continue
                items.append((int(m), float(i_s), float(i_0), float(u_s), float(u_0)))
            if not items:
                continue
            items = sorted(sorted(items, key=lambda x: -x[1])[:2])
            vs = lambda a, b: tr("document.versus", a=a, b=b)
            out.append(tr("document.station_vs_busbar", station=ns, test=n0,
                          orders=self.join([self.ordinal(m) for m, *_ in items]),
                          currents="; ".join(vs("%s A" % _n(a), "%s A" % _n(b)) for _, a, b, _, _ in items),
                          voltages="; ".join(vs("%s V" % _n(us), "%s V" % _n(u0)) for _, _, _, us, u0 in items),
                          ratios="; ".join(vs(_n(a / us, 2), _n(b / u0, 2)) for _, a, b, us, u0 in items)))
        return out

    @staticmethod
    def even_trend(res):
        """Even content of the fault current in the first and the last reference window of the steady state:
        (first, last, t_first, t_last) or None."""
        try:
            st = res.steady_state
            k = res.roles.main("i_res").index
            idx = [j for j, m in enumerate(st.orders) if int(m) >= 2 and int(m) % 2 == 0]
            ok = np.asarray(res.reference_windows, bool) & ~np.asarray(st.nan_window, bool)
            if not idx or ok.sum() < 3:
                return None
            e = np.sqrt((np.abs(st.X[:, k, :][:, idx]) ** 2).sum(axis=1))
            w = np.nonzero(ok)[0]
            return float(e[w[0]]), float(e[w[-1]]), float(st.t_mid[w[0]]), float(st.t_mid[w[-1]])
        except Exception:
            return None

    def context_even_orders(self) -> List[str]:
        """A test at the substation with much larger even orders than the others (and whether they decay)."""
        c = self.c
        even = {n: math.sqrt(sum(float(i) ** 2 for m, i in c.tests[n]["harmonics"].items()
                                 if int(m) >= 2 and int(m) % 2 == 0 and _finite(i))) for n in c.subs}
        out = []
        for n, x in even.items():
            others = [v for k, v in even.items() if k != n]
            if not others or x < max(1.0, self.order_threshold) or x < 3.0 * max(others):
                continue
            decay = ""
            trend = self.even_trend(c.full[n]) if n in c.full else None
            if trend is not None and trend[1] < 0.5 * trend[0]:
                decay = tr("document.even_orders_decay", a=_n(trend[0]), b=_n(trend[1]), t0=_n(trend[2]),
                           t1=_n(trend[3]), t=_n(c.tests[n]["t_max"]))
            out.append(tr("document.even_orders_note", name=n, x=_n(x), y=_n(max(others)), decay=decay))
        return out

    def scaling_context(self, n: str, s: dict) -> str:
        """Why the factors k_m are as large as they are: percentile of the tests in the PQ measurement, median of
        the dominant order on the weekday of the tests and over the days of the week."""
        c = self.c
        pq = c.pq
        if pq is None or not c.pq_valid or pq.dominant_order is None:
            return ""
        dom = int(pq.dominant_order)
        pc = c.pq_comparison
        if pc is None or not len(pc):
            return ""
        g = pc[(pd.to_numeric(pc["order"], errors="coerce") == dom) & pc["test"].astype(str).isin(c.subs)]
        p_all = self.percentile_range(pd.to_numeric(g.get("percentile_all", pd.Series(dtype=float)), errors="coerce"))
        p_win = self.percentile_range(pd.to_numeric(g.get("percentile_window", pd.Series(dtype=float)),
                                                    errors="coerce"))
        if not p_all:
            return ""
        ser = pq.largest(dom)
        idx = ser.index
        if not isinstance(idx, pd.DatetimeIndex) or len(ser) < 2:
            return ""
        tz = c.info.timezone
        if tz and idx.tz is not None:
            idx = idx.tz_convert(tz)
        med = ser.groupby(idx.dayofweek).median()
        loc = _local(c.tests[n]["start"], tz)
        if loc is None or len(med) < 2 or int(loc.dayofweek) not in med.index:
            return ""
        wd = int(loc.dayofweek)
        days = tr("document.weekdays").split(",")
        tab = s["table"]
        own = tab[tab["rule"] == "own voltage"].sort_values("I_scaled_A", ascending=False).head(3)
        ks = "; ".join(tr("document.k_item", m=int(r["order"]), k=_n(r["factor"])) for _, r in
                       own.sort_values("order").iterrows())
        window = tr("document.scaling_context_window", p=p_win, t=self.tables["pq_tests"]) \
            if p_win and "pq_tests" in self.tables else ""
        text = tr("document.scaling_context", order=dom, p_all=p_all, window=window, u_day=_n(med[wd]),
                  day=days[wd], u_lo=_n(med.min()), u_hi=_n(med.max()), q=_n(pq.u95().get(dom)), ks=ks or "–")
        # an order whose test level lay below (almost) every value of the PQ measurement: a large factor that
        # rests on an untypical test – the reader should check the network state
        for _, r in tab[tab["rule"] == "own voltage"].sort_values("order").iterrows():
            m, k = int(r["order"]), float(r["factor"])
            gm = pc[(pd.to_numeric(pc["order"], errors="coerce") == m) & pc["test"].astype(str).isin(c.subs)]
            pa = pd.to_numeric(gm.get("percentile_all", pd.Series(dtype=float)), errors="coerce").dropna()
            if _finite(k) and k >= 2.0 and len(pa) and float(pa.max()) <= 5.0:
                text += " " + tr("document.scaling_low", orders=m, p=self.percentile_range(pa), k=_n(k))
        return text

    def order_spread(self, n: str) -> str:
        """Orders of the fault current that change by more than 20 % over the windows of the steady state."""
        c = self.c
        stats = c.results[n].tables.get("statistics")
        if stats is None or stats.empty or not {"role", "quantity", "min", "max"} <= set(stats.columns):
            return ""
        s = stats[(stats["role"] == "i_res") & stats["quantity"].astype(str).str.fullmatch(r"H\d+")]
        main = c.tests[n]["main"]
        short = (getattr(c.results[n], "data", {}) or {}).get("signals", {}).get(main, {}).get("short") if main else None
        if short and (s["channel"] == short).any():
            s = s[s["channel"] == short]
        elif len(s):
            s = s[s["channel"] == s["channel"].iloc[0]]
        items = []
        for _, r in s.iterrows():
            m = int(str(r["quantity"])[1:])
            lo, hi = float(r["min"]), float(r["max"])
            if m >= 2 and _finite(lo) and _finite(hi) and hi >= max(1.0, self.order_threshold) and hi - lo > 0.2 * hi:
                items.append(tr("document.order_range", m=m, lo=_n(lo), hi=_n(hi)))
        return tr("document.u_stationarity_orders", items="; ".join(items)) if items else ""

    def reactive_drift(self, n: str) -> bool:
        """True if the stationarity check of a test warned (the fundamental changed during the fault)."""
        pr = self.c.results[n].tables.get("plausibility")
        if pr is None or pr.empty or not {"check", "status"} <= set(pr.columns):
            return False
        names = {tr("evaluation.check_stationarity", lang=x) for x in ("en", "de")}
        rows = pr[pr["check"].astype(str).isin(names)]
        return bool(len(rows) and (rows["status"].astype(str).str.upper() == "WARNING").any())

    def periods_sentence(self, periods: pd.DataFrame, order: int) -> str:
        """Quantile of the dominant order per measuring period, their spread and the assessment voltage against the
        highest period."""
        pq = self.c.pq
        q = pd.to_numeric(periods[order], errors="coerce")
        whole = pq.u95().get(order, float("nan"))
        extra = ""
        if _finite(pq.factor) and _finite(whole) and q.max() > 0:
            ua = whole * float(pq.factor)
            pct = 100 * (ua / q.max() - 1)
            extra = tr("document.pq_periods_above" if pct >= 0 else "document.pq_periods_below", ua=_n(ua),
                       pct=_n(abs(pct), 2))
        return tr("document.pq_periods", order=order, values=" / ".join(_n(v) for v in q), whole=_n(whole),
                  spread=_n(100 * (q.max() / q.min() - 1), 2) if q.min() > 0 else "–", assessment=extra)

    def controller_text(self) -> str:
        """The display of the coil controller and how the measured values differ from it."""
        c = self.c
        ctl = c.controller
        if not ctl:
            return ""
        parts = []
        if ctl["coils"]:
            parts.append(tr("document.ctrl_coils", coils=", ".join("%s %s A" % (k, _n(v)) for k, v in
                                                                    ctl["coils"].items()), sum=_n(ctl["sum"])))
        if ctl["how"] == "display":
            parts.append(tr("document.ctrl_ice", ice=_n(ctl["ice"])))
        if _finite(ctl["v"]):
            parts.append(tr("document.ctrl_detuning", v=_g(ctl["v"])))
        if _finite(ctl["active"]):
            parts.append(tr("document.ctrl_active", a=_n(ctl["active"])))
        if not parts:
            return ""
        text = tr("document.ctrl_intro", when=tr("document.ctrl_when", when=ctl["when"]) if ctl["when"] else "",
                  parts=", ".join(parts))
        if ctl["how"] == "derived":
            text += " " + tr("document.ctrl_derived", sum=_n(ctl["sum"]), den=_n(1.0 + ctl["v"] / 100.0, 4),
                             ice=_n(ctl["ice"]))
        cmp = []
        pct = lambda x: ("+" if x >= 0.05 else "") + _fixed(x, 1) if abs(x) >= 0.05 else _fixed(0.0, 1)
        if _finite(c.coil_sum) and _finite(ctl["sum"]) and ctl["sum"] > 0:
            cmp.append(tr("document.ctrl_cmp_coils", x=_n(c.coil_sum), dev=pct(100 * (c.coil_sum / ctl["sum"] - 1))))
        if _finite(c.iw) and _finite(ctl["active"]) and ctl["active"] > 0:
            cmp.append(tr("document.ctrl_cmp_active", x=_n(c.iw), dev=pct(100 * (c.iw / ctl["active"] - 1))))
        if _finite(c.v) and _finite(ctl["v"]):
            cmp.append(tr("document.ctrl_cmp_detuning", x=_fixed(100 * c.v, 1)))
        if cmp:
            text += " " + tr("document.ctrl_compare", items=", ".join(cmp))
        return text

    def detuning_text(self) -> str:
        """The reactive current changes between the tests at an almost equal active current (or "")."""
        c = self.c
        rows = [(abs(float(c.tests[n]["Iv"])), float(c.tests[n]["Iw"])) for n in c.subs
                if _finite(c.tests[n]["Iv"]) and _finite(c.tests[n]["Iw"])]
        if len(rows) < 2:
            return ""
        iv, iw = [r[0] for r in rows], [r[1] for r in rows]
        d_iv, d_iw = max(iv) - min(iv), max(iw) - min(iw)
        if d_iv < max(2.0, 0.1 * float(np.median(iv))) or d_iw > 0.5 * d_iv:
            return ""
        within = [n for n in c.subs if self.reactive_drift(n)]
        # coil current balance of the single evaluations: slope ΔIv/ΔΣI_L ≈ 1 (0.8 … 1.25, the range of the check)
        # → the inductance of the coils changed during the fault, not the capacitance of the network
        slopes = {n: float(c.results[n].values.get("coil_balance_slope")) for n in c.subs
                  if _finite(c.results[n].values.get("coil_balance_slope"))}
        cause = tr("document.detuning_cause_unknown")
        if slopes and all(0.8 <= k <= 1.25 for k in slopes.values()) and (not within or set(within) <= set(slopes)):
            text = (" / ".join(_fixed(slopes[n], 2) for n in c.subs) if len(slopes) == len(c.subs) else
                    "; ".join("%s %s" % (n, _fixed(k, 2)) for n, k in slopes.items()))
            cause = tr("document.detuning_cause_coils", slopes=text)
        return tr("document.detuning_change", ivs=" / ".join(_n(x) for x in iv), lo=_n(min(iw)), hi=_n(max(iw)),
                  within=tr("document.detuning_within", names=self.join(within)) if within else "", cause=cause)

    def limit_case_sentences(self) -> List[str]:
        """The result of the applying limit and the other limits with their conditions."""
        c = self.c
        out, others = [], []
        for lc in c.limit_cases:
            if not _finite(lc["I_A"]):
                continue
            if lc["applies"]:
                u_text = ""
                if lc["ok_U"] is not None:
                    u_text = (tr("document.lc_u_ok", ue=_n(lc["U_E_V"]), share=_share(100 * lc["U_E_V"] / lc["FU_V"]),
                                 u=_n(lc["FU_V"])) if lc["ok_U"] else
                              tr("document.lc_u_exceeded", ue=_n(lc["U_E_V"]), u=_n(lc["FU_V"])))
                if lc["ok_I"]:
                    out.append(tr("document.lc_result_ok", limit=_g(lc["limit_A"]), i=_n(lc["I_A"]), u_text=u_text))
                else:
                    out.append(tr("document.lc_result_exceeded", limit=_g(lc["limit_A"]), i=_n(lc["I_A"]),
                                  pct=_n(100 * (lc["I_A"] / lc["limit_A"] - 1), 2), u_text=u_text))
            else:
                others.append(tr("document.lc_item", limit=_g(lc["limit_A"]),
                                 condition=lc["condition"] or tr("document.lc_standard"), i=_n(lc["I_A"]),
                                 status=self.status(lc["ok_I"] and lc["ok_U"] is not False)))
        if others:
            out.append(tr("document.lc_others", items="; ".join(others)))
        return out

    # ------------------------------------------------------------------ title page and key results
    def title_page(self):
        c, d, info = self.c, self.d, self.c.info
        d.para(info.title or tr("document.default_title"), style="title")
        if info.network:
            d.para(info.network, style="subtitle")
        if info.subtitle:
            d.para(info.subtitle, body=False)
        d.para("")
        for date, places in self.dates():
            d.para(tr("document.carried_out", date=date, places=self.join(places)), style="carried_out")
        if info.further_info:
            d.para("**%s:** %s" % (tr("document.further_info"), info.further_info), body=False)
        for _ in range(6):
            d.para("")
        date = info.date or _dt.date.today()
        d.para("%s %s" % (tr("document.date_label"), _date_short(date, c.lang)), body=False)
        if info.authors:
            d.para("%s %s" % (tr("document.authors_label"), ", ".join(info.authors)), body=False)
        if info.organisation:
            d.para(info.organisation, body=False)
        if info.version:
            d.para("%s %s" % (tr("document.version_label"), info.version), body=False)
        d.para(tr("document.software_line", version=__version__), size=8, color="7F7F7F", body=False)
        d.page_break()

    def key_results(self):
        c, d = self.c, self.d
        d.heading(tr("document.key_results"), 1, numbered=False)
        d.para(tr("document.key_results_intro"))
        rows, keys = [], []

        def add(row, key):
            rows.append(row)
            keys.append(key)

        lim = _n(c.limit, 3) + " A" if c.limit else "–"
        st = ""
        if c.stations:
            st = (tr("document.kr_tests_station") if len(c.stations) == 1 else
                  tr("document.kr_tests_stations", n=len(c.stations)))
        add((tr("document.kr_tests"), tr("document.kr_tests_value", n_sub=len(c.subs), stations=st), "–", "–"),
            "object")
        if _finite(c.ice):
            approx = "" if c.controller.get("how") == "display" else "≈ "
            add((tr("document.kr_ice"), approx + _n(c.ice, 3) + " A", "–", "–"), "frequency")
            add((tr("document.kr_damping"), _fixed(100 * c.d, 1) + " %", "–", "–"), "frequency")
            add((tr("document.kr_detuning"), _fixed(100 * c.v, 1) + " %", "–", "–"), "frequency")
        t = c.tests[c.largest]
        add((tr("document.kr_ires") + " – " + c.largest, "%s A" % _n(t["IRES"]), lim, c.ok(t["IRES"])), "frequency")
        if c.mI > 0:
            add((tr("document.kr_ires_margin"), "%s A" % _n(t["IRES_a"]), lim, c.ok(t["IRES_a"])), "ires")
        s = c.scaled.get(c.basis)
        both = s is not None and c.es is not None
        if s is not None:
            add((tr("document.kr_ires_scaled") + " – " + c.basis
                 + (tr("document.kr_decisive") if both and c.assessment == "busbar" else ""),
                 "%s A" % _n(s["IRES_scaled_assessed"]), lim, c.ok(s["IRES_scaled_assessed"])), "ires")
        if c.es is not None:
            v = c.es_scaled["IRES_scaled_assessed"] if c.es_scaled is not None else c.es["IRES_assessed"]
            add((tr("document.kr_ires_ze", Z=_n(c.Z_E))
                 + (tr("document.kr_decisive") if both and c.assessment == "Z_E" else ""),
                 "%s A" % _n(v), lim, c.ok(v)), "es")
        za = abs(complex(c.Z_E)) if c.Z_E is not None else float("nan")
        for lc in sorted(c.limit_cases, key=lambda x: not x["applies"]):
            if not _finite(lc["I_A"]):
                continue
            same = lc["applies"] and lc["basis"] == "Z_E" and c.es is not None and abs(lc["Z_E"] - za) < 1e-9
            if not same:
                add((tr("document.kr_limit_case", Z=_n(lc["Z_E"]), condition=lc["condition"] or tr("document.lc_standard")),
                     "%s A" % _n(lc["I_A"]), "%s A" % _n(lc["limit_A"]), self.status(lc["ok_I"])), "es")
            if lc["applies"] and lc["ok_U"] is not None:
                add((tr("document.kr_ue_case", Z=_n(lc["Z_E"]), r=_g(lc["r"])), "%s V" % _n(lc["U_E_V"]),
                     "%s V" % _n(lc["FU_V"]), self.status(lc["ok_U"])), "es")
        x = c.expansion
        if x is not None:
            col = x["column"]
            for i, t in enumerate(x["targets"] if x["planned"] else []):
                v = max(float(x["rows"][sp][col].iloc[i + 1]) for sp in x["splits"])
                add((tr("document.kr_expansion", ice=_g(t), Z=_n(x["Z_E"])), "%s A" % _n(v), lim, c.ok(v)), "expansion")
            for lc in x["cases"]:
                r = lc["reach"]
                value = ("–" if r["status"] == "exceeded" else
                         self.reach_text(r, x["top"]) + " A" if r["status"] == "not_reached" else
                         "≈ %s A" % _fixed(round(r["I_C_A"], -1), 0))
                result = (tr("document.exceeded") if r["status"] == "exceeded" else
                          "+%s A" % _fixed(round(r["margin_A"], -1), 0) if r["status"] == "reached" else
                          "> +%s A" % _fixed(round(x["top"] - x["ice"], -1), 0))
                add((tr("document.kr_expansion_limit", limit=_g(lc["allowed"]), Z=_n(lc["Z_E"])), value,
                     tr("document.kr_today", ice=_n(x["ice"])), result), "expansion")
        for n, cases in c.cases.items():
            ut_all = pd.to_numeric(cases["U_T_V"], errors="coerce")
            if ut_all.notna().any():
                top = cases.loc[ut_all.idxmax()]
                ut = float(ut_all.max())
                utp = float(c.u_tp) if _finite(c.u_tp) else float("nan")
                case = tr("document.probe") if top["case"] == "ut" else str(top["label"])
                add((tr("document.kr_ut", station=c.place(n), case=case), "%s V" % _n(ut),
                     ("%s V" % _n(utp)) if _finite(utp) else "–",
                     (tr("document.met") if ut <= utp else tr("document.exceeded")) if _finite(utp) else "–"),
                    "station:" + n)
        if c.pq is not None:
            cov = c.pq.coverage
            n_per = len(cov.get("periods") or [])
            value = (tr("document.kr_pq_value_periods", covered=_fixed(cov.get("covered_days"), 1), n=n_per,
                        days=_fixed(cov.get("gapless_days"), 1), q=_g(100 * c.pq.quantile), k=_n(c.pq.factor))
                     if n_per >= 2 else
                     tr("document.kr_pq_value", days=_fixed(cov.get("gapless_days"), 1), q=_g(100 * c.pq.quantile),
                        k=_n(c.pq.factor)))
            add((tr("document.kr_pq"), value, "–", tr("document.valid") if c.pq.valid else tr("document.not_valid")),
                "pq")
        tab = d.table([tr("document.kr_quantity"), tr("document.kr_value"), tr("document.kr_reference"),
                       tr("document.kr_result"), tr("document.kr_section")], [r + ("",) for r in rows],
                      widths_cm=[6.2, 3.4, 2.2, 2.4, 2.0], align=["left", "right", "right", "center", "center"])
        self.kr_cells = [(tab.rows[i + 1].cells[4], k) for i, k in enumerate(keys)]
        d.page_break()

    def fill_sections(self):
        """The section of every row of the key results (known once the chapters are written)."""
        for cell, key in self.kr_cells:
            self.d.set_cell(cell, self.sections.get(key, "–"), size=9.0, align="center")

    # ------------------------------------------------------------------ chapter 1
    def chapter_setup(self):
        c, d, info = self.c, self.d, self.c.info
        d.heading(tr("document.ch_setup"), 1)
        self.heading(tr("document.sec_object"), 2, key="object")
        places = []
        for n in c.names:
            p = c.place(n)
            if p not in places:
                places.append(p)
        text = [tr("document.object_intro", places=self.join(places))]
        if info.normal_state:
            text.append(tr("document.normal_state"))
        if info.coil_unchanged:
            text.append(tr("document.coil_unchanged"))
        if info.auto_tuning:
            text.append(tr("document.auto_tuning"))
        text.append(tr("document.n_tests", n=len(c.names)))
        d.para(" ".join(text))
        rows = []
        for n in c.names:
            z = c.Z.get(n)
            rows.append((n, c.place(n), _n(z) if _finite(z) else "–", self.time_text(n), _fixed(self.duration(n), 1)))
        d.table([tr("document.col_test"), tr("document.col_place"), tr("document.col_z"),
                 tr("document.col_start", tz=self.tz_text()), tr("document.col_duration")], rows,
                caption=tr("document.tab_tests"), widths_cm=[4.4, 3.8, 3.2, 2.6, 2.0],
                align=["left", "left", "right", "right", "right"])
        d.para(info.purpose or tr("document.purpose_default",
                                  station=tr("document.purpose_station") if c.stations else ""))
        if any(_finite(c.Z.get(n)) and float(c.Z.get(n)) > 0 for n in c.names):
            self.photo("setup_impedance", tr("document.photo_setup_impedance"))
        fs = [c.tests[n]["fs"] for n in c.names if _finite(c.tests[n]["fs"])]
        d.para(tr("document.recorder_text", instrument=info.instrument or tr("document.instrument_default"),
                  fs=_n(max(fs) / 1000.0, 3) if fs else "–"))
        self.quantities_table()
        self.photo("recorder", tr("document.photo_recorder"))
        self.extra_photos("1.1")
        self.corrections_table()
        d.heading(tr("document.sec_procedure"), 2)
        d.para(tr("document.procedure_intro"))
        if info.procedure:
            steps = list(info.procedure)
        else:
            durs = [self.duration(n) for n in c.names if _finite(self.duration(n))]
            dur = ""
            if durs and max(durs) - min(durs) > 0.05 * max(durs):
                dur = tr("document.step_on_durations", lo=_fixed(min(durs), 1), hi=_fixed(max(durs), 1))
            elif durs:
                dur = tr("document.step_on_duration", d=_fixed(float(np.median(durs)), 1))
            steps = [tr("document.step_tuning"), tr("document.step_on", duration=dur), tr("document.step_off"),
                     tr("document.step_impedance")]
        d.bullets(steps, numbered=True)
        for note in info.notes or ():
            d.para(note)

    def quantities_table(self):
        c, d = self.c, self.d
        lang = c.lang
        rows = []
        seen_sites = []
        for n in c.names:
            site = c.tests[n]["site"]
            if site in seen_sites:
                continue
            seen_sites.append(site)
            e = c.results[n]
            meas = _measurement(e)
            roles = meas.get("roles") or {}
            signals = (getattr(e, "data", None) or {}).get("signals", {}) if isinstance(getattr(e, "data", None),
                                                                                       dict) else {}
            first = True
            for key, group, en, de in _ROLE_GROUPS:
                chans = [ch for ch, r in roles.items() if r in group]
                if not chans:
                    continue
                shorts = []
                for ch in chans:
                    sig = signals.get(ch, {})
                    s = sig.get("short") or ch
                    dist = sig.get("distance_m")
                    shorts.append("%s (%s m)" % (s, _n(dist)) if _finite(dist) else s)
                trans = c.info.transducers.get(key) or next((c.info.transducers.get(r) for r in group
                                                             if c.info.transducers.get(r)), "")
                label = c.info.quantity_labels.get(key) or (de if lang == "de" else en)
                rows.append((label, ", ".join(shorts), c.place(n) if first else "", trans))
                first = False
        if rows:
            header = [tr("document.col_quantity"), tr("document.col_channels"), tr("document.col_location")]
            if any(r[3] for r in rows):
                d.table(header + [tr("document.col_transducer")], rows, caption=tr("document.tab_quantities"),
                        widths_cm=[5.2, 4.2, 3.6, 3.2], align=["left", "left", "left", "left"], font_size=8.5)
            else:
                d.table(header, [r[:3] for r in rows], caption=tr("document.tab_quantities"),
                        widths_cm=[6.4, 5.0, 4.8], align=["left", "left", "left"], font_size=8.5)

    def corrections_table(self):
        c, d = self.c, self.d
        items: Dict[Tuple[str, str, str], List[str]] = {}
        for n in c.names:
            for corr in c.tests[n]["corrections"]:
                own = corr.get("factor", corr.get("factor_total"))
                total = corr.get("factor_total", own)
                factor = _g(own)
                if _finite(own) and _finite(total) and abs(float(total) - float(own)) > 1e-9:
                    factor = tr("document.factor_total", f=_g(own), t=_g(total))   # with a factor of another row
                key = (str(corr.get("short") or corr.get("channel")), factor, str(corr.get("remark") or ""))
                items.setdefault(key, []).append(n)
        rotated = [n for n in c.names if c.tests[n].get("polarity") == -1]
        if not items and not rotated:
            return
        if items:
            d.para(tr("document.corrections_intro"))
            rows = [(", ".join(names), k[0], k[1], _markup_symbols(k[2])) for k, names in items.items()]
            d.table([tr("document.col_files"), tr("document.col_channel"), tr("document.col_factor"),
                     tr("document.col_remark")], rows, caption=tr("document.tab_corrections"),
                    widths_cm=[3.8, 2.0, 2.2, 8.2], align=["left", "left", "right", "left"], font_size=8)
        if rotated:
            d.para(tr("document.polarity_rotated", tests=self.join(rotated)))

    # ------------------------------------------------------------------ chapter 2
    def chapter_results(self):
        c, d = self.c, self.d
        d.heading(tr("document.ch_results"), 1)
        d.heading(tr("document.sec_general"), 2)
        d.heading(tr("document.sec_uncertainty"), 3)
        comps = [(str(a), float(b)) for a, b in (c.info.uncertainty_components or ())]
        if comps:
            d.para(tr("document.uncertainty_components"))
            d.table([tr("document.col_component"), tr("document.col_accuracy")],
                    [(a, _n(b)) for a, b in comps], caption=tr("document.tab_uncertainty_components"),
                    widths_cm=[9.0, 3.0])
            lin = sum(b for _, b in comps)
            rss = math.sqrt(sum(b * b for _, b in comps))
            d.para(tr("document.uncertainty_sum", lin=_n(lin), rss=_n(rss)))
        if c.mI > 0:
            mu = tr("document.uncertainty_margin_u", mu=_n(100 * c.mU)) if c.mU > 0 else ""
            d.para(tr("document.uncertainty_margin", mi=_n(100 * c.mI), mu=mu))
        else:
            d.para(tr("document.uncertainty_none"))
        fs = [c.tests[n]["fs"] for n in c.names if _finite(c.tests[n]["fs"])]
        st = _stored_settings(c.results[c.basis])
        top = int(st.get("thd_max_order") or 40)
        t = c.tests[c.basis]
        T = t["T_win"] * 1000 if _finite(t["T_win"]) else float("nan")
        if fs:
            df = tr("document.sampling_df", df=_g(1000.0 / T)) if _finite(T) and T > 0 else ""
            d.para(tr("document.sampling_text", fs=_n(min(fs) / 1000.0, 3), ny=_n(min(fs) / 2000.0, 3), order=top,
                      f=_n(top * self._f_nominal(), 4), df=df))
        d.heading(tr("document.sec_processing"), 3)
        d.para(tr("document.processing_intro", version=__version__, url=_HOMEPAGE))
        has_corr = any(c.tests[n]["corrections"] for n in c.names)
        f1s = [c.tests[n]["f1"] for n in c.names if _finite(c.tests[n]["f1"])]
        f1_text = (tr("document.proc_f1_range", lo=_n(min(f1s), 5), hi=_n(max(f1s), 5))
                   if len(f1s) > 1 and max(f1s) - min(f1s) > 1e-3 else tr("document.proc_f1", f1=_n(t["f1"], 5)))
        res = c.full.get(c.basis)
        fig_ref = tr("document.proc_intervals_fig", name=c.basis) if res is not None else ""
        steps = [tr("document.proc_read", corr=tr("document.proc_read_corr") if has_corr else ""),
                 (tr("document.proc_dc_" + t["dc"], fc=_g(st.get("dc_cutoff_hz") or 5.0))
                  if t["dc"] in ("highpass", "highpass_median")
                  else tr("document.proc_dc", method=self.word("dc", t["dc"]))),
                 f1_text,
                 tr("document.proc_dft", T=_n(T), n=_n(T / 1000.0 * self._f_nominal(), 2) if _finite(T) else "–",
                    window=self.word("window", t["window"]), order=top),
                 tr("document.proc_intervals", fig=fig_ref), tr("document.proc_phasors")]
        d.bullets(steps, numbered=True)
        if res is not None:
            rec = c.recordings.get(c.basis)
            chans = [r for r in ("i_res", "i_coil") if res.roles.has(r)]
            self.fig(lambda: report.rms_trend(res, rec, channels=chans), "rms_trend_" + report.slug(c.basis),
                     tr("document.fig_rms_trend", name=c.basis))
        # ---- fault current
        d.heading(tr("document.sec_fault_current"), 2)
        d.heading(tr("document.sec_waveforms"), 3)
        if c.full:
            shown, dc = [], []
            for n in c.names:
                res, rec = c.full.get(n), c.recordings.get(n)
                k = res.roles.main("i_res") if res is not None and rec is not None else None
                if k is None:
                    continue
                t0 = max(0.0, float(res.t_max_window) - 0.05)
                shown.append((n, res, rec, k, t0))
                try:                                   # DC component in the excerpt (the evaluation removes it)
                    _, x = rec.read_time(t0, t0 + 0.1)
                    y = np.asarray(x[:, k.index], float)
                    med = float(np.median(y))
                    if abs(med) >= max(1.0, 0.05 * float(np.max(np.abs(y - med)))):
                        dc.append("%s %s A" % (n, _signed(med, 2)))
                except Exception:
                    pass
            d.para(tr("document.waveforms_intro") + " " + tr(
                "document.waveforms_dc_free", items=tr("document.waveforms_dc_items", items=", ".join(dc)) if dc else ""))
            for n, res, rec, k, t0 in shown:
                self.fig(lambda res=res, rec=rec, k=k, t0=t0: report.instantaneous(rec, res.roles, t0, 0.1,
                                                                                   channels=[k.index], remove_dc=True),
                         "waveform_" + report.slug(n), tr("document.fig_waveform", name=n))
        else:
            d.para(tr("document.waveforms_missing"))
        self.heading(tr("document.sec_frequency"), 3, key="frequency")
        names = c.names
        orders = self.orders_shown(names)
        top_order = max(orders) if orders else 1
        d.para(tr("document.frequency_intro", threshold=_n(self.order_threshold), order=top_order,
                  f=_n(top_order * self._f_nominal(), 4)))
        fn = self._f_nominal()
        rows = [[tr("document.row_rms"), ""] + [_n(c.tests[n]["IRES"]) for n in names],
                [tr("document.row_order", f=_hz(fn), m=1), tr("document.category_fundamental")]
                + [_n(c.tests[n]["harmonics"].get(1)) for n in names]]
        for m in orders:
            rows.append([tr("document.row_order", f=_hz(m * fn), m=m), tr("document.category_" + fml.harmonic_category(m))]
                        + [_n(c.tests[n]["harmonics"].get(m)) for n in names])
        rows.append([tr("document.row_harmonic_content"), ""] + [_n(c.tests[n]["IH"]) for n in names])
        if any(_finite(c.tests[n]["Ilf"]) and c.tests[n]["Ilf"] > 0.05 for n in names):
            rows.append([tr("document.row_low_frequency"), ""] + [_n(c.tests[n]["Ilf"]) for n in names])
        rows.append([tr("document.row_thd"), ""] + [
            _n(100 * c.tests[n]["IH"] / c.tests[n]["harmonics"].get(1, np.nan)) if c.tests[n]["harmonics"].get(1)
            else "–" for n in names])
        w = max(1.6, (d.text_width_cm - 5.4) / max(len(names), 1))
        d.table(["", tr("document.col_category")] + list(names), rows, caption=tr("document.tab_ires"),
                widths_cm=[3.4, 2.0] + [w] * len(names), align=["left", "left"] + ["right"] * len(names),
                font_size=8.5)
        subs = [c.results[n] for n in names]
        fig_orders = [1] + [m for m in orders if m <= 25][:14]
        self.fig(lambda: report.spectrum_comparison(subs, names, orders=fig_orders), "spectrum_comparison",
                 tr("document.fig_spectrum"))
        for text in self.context_station_busbar() + self.context_even_orders():
            d.para(text)
        d.para(tr("document.split_intro"))
        d.formula(tr("document.eq_iw"))
        rows = [[tr("document.row_ig")] + [_n(c.tests[n]["IG"]) for n in names],
                [tr("document.row_iw")] + [_n(c.tests[n]["Iw"]) for n in names],
                [tr("document.row_iv")] + [_n(c.tests[n]["Iv"]) for n in names]]
        channels = []
        for n in names:
            for ch, _ in c.tests[n]["coils"]:
                if ch not in channels:
                    channels.append(ch)
        for ch in channels:
            rows.append([tr("document.row_coil", channel=ch)] + [
                _n(dict(c.tests[n]["coils"]).get(ch)) if dict(c.tests[n]["coils"]).get(ch) is not None else "–"
                for n in names])
        if len(channels) > 1:
            rows.append([tr("document.row_coil_sum")] + [_n(c.tests[n]["coil_sum"]) for n in names])
        if _finite(c.ice):
            rows.append([tr("document.row_damping")] + [_fixed(100 * c.tests[n]["Iw"] / c.ice, 1)
                                                        if _finite(c.tests[n]["Iw"]) else "–" for n in names])
            rows.append([tr("document.row_detuning")] + [_fixed(100 * abs(c.tests[n]["Iv"]) / c.ice, 1)
                                                         if _finite(c.tests[n]["Iv"]) else "–" for n in names])
        w = max(1.8, (d.text_width_cm - 5.2) / max(len(names), 1))
        d.table([""] + list(names), rows, caption=tr("document.tab_split"), widths_cm=[5.2] + [w] * len(names),
                font_size=8.5)
        for text in (self.controller_text(), self.detuning_text()):
            if text:
                d.para(text)
        if _finite(c.ice):
            without = [n for n in names if not c.tests[n]["coils"]]
            others = tr("document.ice_others", names=self.join(without)) if without else ""
            src = (tr("document.ice_from_controller") if c.controller.get("how") == "derived" else
                   tr("document.ice_from_display") if _finite(c.coil_display)
                   else tr("document.ice_from_coils_long", v=_fixed(100 * c.v, 1)))
            d.para(tr("document.ice_text", source=src, others=others))
        for n in c.stations:
            self.station_results(n)
        self.voltages()

    @staticmethod
    def word(kind: str, value) -> str:
        """Name of a setting in the output language (DC method, window), e.g. "median" → "Median"."""
        v = str(value or "").strip()
        if not v:
            return "–"
        key = "document.%s_%s" % (kind, v.lower())
        return tr(key) if i18n.has_message(key) else v[:1].upper() + v[1:]

    def station_results(self, n: str):
        c, d = self.c, self.d
        e = c.results[n]
        place = c.place(n)
        d.heading(tr("document.sec_station", station=place), 2)
        d.heading(tr("document.sec_station_waveforms"), 3)
        self.photo("station", tr("document.photo_station", station=place), key="station:" + n)
        sv = e.tables.get("station_voltages")
        sv_e = sv[sv["role"] == "u_e"] if sv is not None and not sv.empty and (sv["role"] == "u_e").any() else sv
        probe = (str(sv_e.loc[pd.to_numeric(sv_e["RMS_V"], errors="coerce").idxmax(), "channel"])
                 if sv_e is not None and not sv_e.empty else None)
        res, rec = c.full.get(n), c.recordings.get(n)
        compact = c.info.station_detail == "compact"
        if res is not None and rec is not None:
            idx, what = [], []
            ue = res.roles.with_role("u_e") if hasattr(res.roles, "with_role") else []
            if ue:
                idx += [k.index for k in ue if k.short == probe] or [ue[0].index]
                what.append(c.info.quantity_labels.get("u_e") or tr("document.wf_u_e"))
            for role in ("i_res", "i_shield"):
                k = res.roles.main(role)
                if k is not None:
                    idx.append(k.index)
                    what.append(c.info.quantity_labels.get(role) or tr("document.wf_" + role))
            if what:
                d.para(tr("document.station_waveforms_intro", items=self.join(what)))
            if idx:
                t0 = max(0.0, float(res.t_max_window) - 0.05)
                panels = ""
                try:                                   # the panels in the order of the figure
                    groups = list(report._groups(res.roles, idx).values())
                    panels = ": " + ", ".join("(%s) %s" % (chr(97 + j), ", ".join(k.short for k in g))
                                              for j, g in enumerate(groups))
                except Exception:
                    panels = ""
                self.fig(lambda: report.instantaneous(rec, res.roles, t0, 0.1, channels=idx, remove_dc=True),
                         "station_waveform_" + report.slug(n),
                         tr("document.fig_station_waveform", station=place, panels=panels))
        d.heading(tr("document.sec_station_frequency"), 3)
        if sv_e is not None and not sv_e.empty:
            if compact:
                d.para(self.probes_compact(n))
            else:
                probes = ["%s (%s m)" % (r["channel"], _n(r["distance_m"])) if _finite(r.get("distance_m")) else
                          str(r["channel"]) for _, r in sv_e.iterrows()]
                d.para(tr("document.station_probes" if len(probes) == 1 else "document.station_probes_many",
                          probes=self.join(probes)))
            r = sv_e.loc[pd.to_numeric(sv_e["RMS_V"], errors="coerce").idxmax()]
            ords = sorted(int(cn[1:-2]) for cn in sv.columns if re.fullmatch(r"H\d+_V", cn))
            hz = _harmonics(e)
            sf = e.tables.get("station_frequency")
            fault, unsure = {}, set()
            if sf is not None and not sf.empty:
                for _, x in sf[sf["channel"] == r["channel"]].iterrows():
                    if _finite(x.get("Z_fault_Ohm")):
                        fault[int(x["order"])] = float(x["Z_fault_Ohm"])
                    if str(x.get("uncertain")).lower() in ("true", "1", "yes"):
                        unsure.add(int(x["order"]))
            star = lambda m: " *" if m in unsure else ""
            rows = [(tr("document.row_rms_plain"), _n(r["RMS_V"]), _n(r.get("V_per_100A_RMS")), "–",
                     _n(c.tests[n]["IRES"]))]
            for m in ords:
                rows.append((tr("document.row_order", f=_hz(m * self._f_nominal()), m=m), _n(r.get("H%d_V" % m)),
                             _n(100 * r.get("rZE_H%d_Ohm" % m, np.nan)),
                             (_n(100 * fault[m]) + star(m)) if m in fault else "–", _n(hz.get(m))))
            d.table(["", tr("document.col_ue_max"), tr("document.col_per_100"), tr("document.col_per_100_fault"),
                     tr("document.col_if")], rows, caption=tr("document.tab_epr", probe=r["channel"]),
                    widths_cm=[3.4, 3.2, 3.4, 3.6, 2.4])
            d.para(tr("document.epr_note", uncertain=tr("document.epr_uncertain")
                      if any(m in unsure for m in ords if m in fault) else ""))
            inv = self.probe_inversion(n) if not compact else ""
            if inv:
                d.para(inv)
        if compact:                                    # for information: the earthing impedance at 50 Hz only
            z50 = self.earthing_impedance_at(n, 50.0)
            if n in c.earthing and _finite(z50):
                _, ids, details, _ = self.earthing_info(n)
                d.para(tr("document.station_earthing_compact", ids=", ".join(map(str, ids)) or "–", details=details,
                          z=_n(z50)))
            return
        sf = e.tables.get("station_frequency")
        if sf is not None and not sf.empty:
            self.fig(lambda: report.station_frequency(e), "station_frequency_" + report.slug(n),
                     tr("document.fig_station_frequency", station=place))
        if n in c.earthing:
            long_df, ratio = c.earthing[n]
            rr, ids, details, pr = self.earthing_info(n)
            d.para(tr("document.station_comparison_intro", ids=", ".join(map(str, ids)) or "–", details=details,
                      probe=tr("document.probe_maximum_here", probe=pr) if pr else ""))
            rows = [(_n(x["f_Hz"], 4), _n(x["Z_E_injection_Ohm"]), _n(x["Z_fault_Ohm"]), _n(x["ratio_fault"]))
                    for _, x in rr.iterrows()]
            if rows:
                d.table([tr("document.col_f"), tr("document.col_ze_injection"), tr("document.col_rze_fault"),
                         tr("document.col_r")], rows, caption=tr("document.tab_earthing"),
                        widths_cm=[2.5, 4.0, 4.8, 2.0])
            self.fig(lambda: report.earthing_comparison(long_df, ratio), "earthing_comparison_" + report.slug(n),
                     tr("document.fig_earthing", station=place))
        else:
            d.para(tr("document.no_earthing"))
        red = c.reduction.get(n) or {}
        eff = red.get("reduction_factor_eff")
        if eff is not None and len(eff):
            row = eff[eff["channel"] == "max(UE)"].iloc[0] if (eff["channel"] == "max(UE)").any() else eff.iloc[0]
            cls = str(row.get("classification_r_50Hz") or "")
            key = "document.r_class_" + cls
            text = tr(key) if cls and i18n.has_message(key) else cls
            if text and text != "–" and cls in ("below", "typical", "above") and _finite(row.get("r_typical_min")) \
                    and _finite(row.get("r_typical_max")):
                text += tr("document.reduction_range", low=_g(row["r_typical_min"]), high=_g(row["r_typical_max"]))
            d.para(tr("document.reduction_text", r=_n(row.get("r_eff"), 3), r50=_n(row.get("r_50Hz"), 3),
                      classification=tr("document.reduction_class", text=text) if text and text != "–" else "",
                      shield=_n(100 * float(row.get("share_shield_eff", np.nan)), 2)))
            rf = red.get("reduction_factor")
            if rf is not None and len(rf):
                self.fig(lambda: report.reduction_factor(rf, eff), "reduction_factor_" + report.slug(n),
                         tr("document.fig_reduction", station=place))
        lv = e.tables.get("station_lv_network_eff")
        if lv is not None and not lv.empty:
            row = lv[lv["channel"] == "max(UE)"].iloc[0] if (lv["channel"] == "max(UE)").any() else lv.iloc[0]
            d.para(tr("document.lv_text", ion=_n(row.get("I_LV_eff_A")), share=_n(100 * row.get("share_I_LV_eff", np.nan)),
                      zon=_n(row.get("Z_E_LV_eff_Ohm"))))
            lvt = e.tables.get("station_lv_network")
            chk = red.get("lv_network_check")
            self.fig(lambda: report.lv_network(lvt, chk if chk is not None and len(chk) else None),
                     "lv_network_" + report.slug(n), tr("document.fig_lv", station=place))
        hints = []
        try:                                           # the notes of the comparison, as in the campaign notebook
            from . import earthing as _earthing
            hints = _earthing.reduction_hints(eff if eff is not None and len(eff) else None,
                                              red.get("lv_network_check"))
        except Exception:
            hints = []
        if hints:
            d.para("**%s**" % tr("document.earthing_hints"), keep_next=True)
            drop = (" Als Befund festhalten.", " Record it as a finding.")      # instructions of the notebook
            d.bullets([_markup_symbols(functools.reduce(lambda x, y: x.replace(y, ""), drop, h)) for h in hints])

    def probes_compact(self, n: str) -> str:
        """The probes of a station with their RMS values in one sentence, against F·U_Tp (compact station)."""
        c = self.c
        sv = c.results[n].tables.get("station_voltages")
        rows = sv[sv["role"].isin(["u_e", "u_t"])] if sv is not None and not sv.empty else sv
        if rows is None or rows.empty:
            return ""
        v = pd.to_numeric(rows["RMS_V"], errors="coerce")
        items = []
        for j, r in rows.iterrows():
            dist = " (%s m)" % _n(r["distance_m"]) if _finite(r.get("distance_m")) else ""
            items.append("%s%s %s V" % (r["channel"], dist, _n(v[j])))
        umax = float(v.max())
        fu = float(c.f_factor) * float(c.u_tp) if _finite(c.u_tp) and _finite(c.f_factor) else float("nan")
        limit = tr("document.probes_limit", share=_share(100 * umax / fu), u=_n(fu)) if _finite(fu) and fu > 0 else ""
        if len(items) == 1:
            return tr("document.station_probe_compact", probes=items[0], limit=limit)
        diff = tr("document.probes_diff") if _finite(fu) and fu > 0 and umax < 0.05 * fu else ""
        return tr("document.station_probes_compact", probes=self.join(items), umax=_n(umax), limit=limit, diff=diff)

    def probe_inversion(self, n: str) -> str:
        """The sentence on a touch-voltage probe above the earth potential rise (U_T > U_E), or ""."""
        sv = self.c.results[n].tables.get("station_voltages")
        if sv is None or sv.empty or not (sv["role"] == "u_e").any() or not (sv["role"] == "u_t").any():
            return ""
        v = pd.to_numeric(sv["RMS_V"], errors="coerce")
        ie, it = v[sv["role"] == "u_e"].idxmax(), v[sv["role"] == "u_t"].idxmax()
        if not (_finite(v[ie]) and _finite(v[it])) or v[it] <= v[ie]:
            return ""
        return tr("document.probe_inverted", ut_probe=str(sv.loc[it, "channel"]), ut=_n(v[it]),
                  ue_probe=str(sv.loc[ie, "channel"]), ue=_n(v[ie]))

    def given_quantiles(self) -> Optional[Tuple[str, str]]:
        """The quantiles of a scaling without a file of the PQ measurement – those of the campaign (``scaling``),
        else those stored in the evaluation of the basis test – as text, and the factor k_s; None without any."""
        c = self.c
        sc = dict(c.scaling or {})
        u95, ref, thd, k = sc.get("u95"), sc.get("u95_reference"), sc.get("thd95_percent"), sc.get("safety_factor")
        if not (u95 or _finite(thd)):
            if c.basis not in c.scaled:
                return None
            st = _stored_settings(c.results[c.basis])
            u95, ref, thd, k = (st.get("pq_u95"), st.get("pq_u95_reference"), st.get("u_thd95_percent"),
                                st.get("pq_safety_factor"))
            if not (u95 or _finite(thd)):
                return None
        items = []
        if u95:
            side = tr("document.given_ll") if str(ref or "seq") == "LL" else tr("document.given_seq")
            items.append(", ".join(tr("document.given_u95", m=int(m), u=_n(float(u)))
                                   for m, u in sorted(dict(u95).items(), key=lambda x: int(x[0]))) + " (%s)" % side)
        if _finite(thd):
            items.append(tr("document.given_thd", thd=_n(float(thd))) + " (%s)" % tr("document.given_ll"))
        return "; ".join(items), _n(float(k) if _finite(k) else 1.0)

    def voltages(self):
        c, d = self.c, self.d
        d.heading(tr("document.sec_voltages"), 2)
        self.heading(tr("document.sec_pq"), 3, key="pq")
        pq = c.pq
        if pq is None:
            given = self.given_quantiles()
            d.para(tr("document.no_pq_given", values=given[0], k=given[1]) if given else tr("document.no_pq"))
        else:
            cov = pq.coverage
            fmt = (lambda ts: _date_short(pd.Timestamp(ts), c.lang)) if cov.get("start") is not None else (lambda ts: "–")
            device = getattr(pq.data, "instrument", "") or ""
            device = tr("document.pq_device", device=device) if device else ""
            periods = pq.periods() if len(cov.get("periods") or []) >= 2 else None
            if periods is not None:
                items = [tr("document.pq_period_item", start=fmt(r.start), end=fmt(r.end), days=_fixed(r.days, 1))
                         for r in periods.itertuples()]
                d.para(tr("document.pq_intro_periods", k=len(periods), device=device, items=self.join(items),
                          covered=_fixed(cov.get("covered_days"), 1), n=str(int(cov.get("n") or 0)),
                          gapless=_fixed(cov.get("gapless_days"), 1)))
            else:
                d.para(tr("document.pq_intro", device=device, start=fmt(cov.get("start")), end=fmt(cov.get("end")),
                          span=_fixed(cov.get("span_days"), 1), n=str(int(cov.get("n") or 0)),
                          complete=_fixed(100 * float(cov.get("completeness", np.nan)), 1),
                          gapless=_fixed(cov.get("gapless_days"), 1)))
            nx = cov.get("n_excluded")
            excluded = (tr("document.pq_excluded", n=int(nx)) if _finite(nx) and int(nx) > 0 else
                        tr("document.pq_excluded_none") if _finite(nx) else "")
            d.para(tr("document.pq_method", q=_g(100 * pq.quantile), excluded=excluded))
            rule = getattr(pq, "factor_rule", None) or {}
            if rule:
                items = "; ".join(tr("document.pq_rule_item", days=_g(dd), k=_g(kk))
                                  for dd, kk in sorted(rule.items(), key=lambda x: -float(x[0])))
                hint = tr("document.pq_factor_rule", days=_fixed(cov.get("gapless_days"), 1), rule=items)
            else:
                hint = plain(pq.factor_hint).rstrip(".")
            d.para(tr("document.pq_factor", k=_n(pq.factor), hint=hint))
            if not pq.valid:
                d.para(tr("document.pq_invalid"))
            dom = pq.dominant_order
            t = pq.table()
            if dom is not None:
                row = t.set_index("order").loc[dom]
                d.para(tr("document.pq_dominant", order=dom, f=_hz(dom * pq.f1), share=_n(row["share_of_UH_pct"], 2),
                          u=_n(row["U_quantile_V"]), umax=_n(row["U_max_V"])))
                if periods is not None and dom in periods.columns:
                    d.para(self.periods_sentence(periods, dom))
            rows = []
            for _, r in t[t["order"] >= 2].iterrows():
                rows.append((str(int(r["order"])), _hz(r["f_Hz"]), tr("document.category_" + fml.harmonic_category(int(r["order"]))),
                             _n(r["U_quantile_V"]), _n(r["U_quantile_V"] * pq.factor), _n(r["U_median_V"]),
                             _n(r["U_max_V"])))
            d.table([tr("document.col_order"), tr("document.col_f"), tr("document.col_category"),
                     tr("document.col_u_quantile"), tr("document.col_u_factored"), tr("document.col_u_median"),
                     tr("document.col_u_max")], rows, caption=tr("document.tab_pq"),
                    widths_cm=[1.9, 1.6, 2.3, 2.6, 2.6, 2.6, 2.4], font_size=8.5)
            if _finite(pq.thd95()):
                d.para(tr("document.pq_thd", q=_n(pq.thd95()), k=_n(pq.thd95() * pq.factor)))
            self.fig(lambda: report.pq_series(pq), "pq_series", tr("document.fig_pq_series"))
            levels = {n: lv[dom] for n, lv in (c.pq_levels or {}).items() if dom in lv} if dom is not None else {}
            self.fig(lambda: report.pq_duration(pq, levels=levels), "pq_duration",
                     tr("document.fig_pq_duration", order=dom))
            self.fig(lambda: report.pq_profile(pq), "pq_profile", tr("document.fig_pq_profile", order=dom))
            pc = c.pq_comparison
            if pc is not None and len(pc):
                d.para(tr("document.pq_comparability"))
                rows = [(str(r["test"]), str(int(r["order"])), _n(r["U_test_V"]), _fixed(r["percentile_all"], 0),
                         str(r.get("window", "")), _fixed(r.get("percentile_window"), 0), _n(r.get("factor_to_test")))
                        for _, r in pc.iterrows()]
                d.table([tr("document.col_test"), tr("document.col_order"), tr("document.col_u_test"),
                         tr("document.col_percentile"), tr("document.col_window"), tr("document.col_percentile_window"),
                         tr("document.col_factor_test")], rows, caption=tr("document.tab_pq_tests"),
                        widths_cm=[3.6, 1.8, 2.0, 2.0, 2.8, 2.0, 2.0], align=["left"] + ["right"] * 3 + ["left"]
                        + ["right"] * 2, font_size=8)
                self.tables["pq_tests"] = self.table_number()
                listed = set(str(x) for x in pc["test"])
                missing = [n for n in c.stations if n not in listed]       # a station: not the busbar voltage
                d.para(tr("document.pq_tests_note", others=tr("document.pq_tests_others", names=self.join(missing))
                          if missing else ""))
        # ---- voltages of the tests
        d.heading(tr("document.sec_test_voltages"), 3)
        d.para(tr("document.test_voltages_intro"))
        header = [""]
        cols = []
        for n in c.names:
            v = c.results[n].values.get("voltage") or {}
            for b, key in (("pre_fault", "document.col_before"), ("post_fault", "document.col_after")):
                if isinstance(v.get(b), dict):
                    cols.append((n, b))
                    header.append("%s %s" % (tr(key), n))
        if cols:
            orders = sorted({int(o) for n in c.names for o in ((c.results[n].values.get("voltage") or {}).get("orders") or [])
                             if int(o) % 2 == 1 and int(o) % 3 != 0})

            def level(n, b, m):
                v = c.results[n].values.get("voltage") or {}
                ords = [int(o) for o in v.get("orders", [])]
                seq = "pos" if m % 3 == 1 else "neg"
                return v[b]["U_" + seq][ords.index(m)] if m in ords and "U_" + seq in v[b] else np.nan

            orders = [m for m in orders if m == 1 or any(_finite(level(n, b, m)) and level(n, b, m) >= 1.0
                                                         for n, b in cols)]
            rows = []
            fn = self._f_nominal()
            for m in orders:
                seq = "pos" if m % 3 == 1 else "neg"
                line = ["%s Hz (%s)" % (_hz(m * fn), tr("document.positive" if seq == "pos" else "document.negative"))]
                for n, b in cols:
                    v = c.results[n].values.get("voltage") or {}
                    ords = [int(o) for o in v.get("orders", [])]
                    val = v[b]["U_" + seq][ords.index(m)] if m in ords and "U_" + seq in v[b] else np.nan
                    line.append(_n(val, 4 if m == 1 else 3))
                rows.append(line)
            thd = [tr("document.row_thd_ll")]
            for n, b in cols:
                k = c.results[n].tables.get("voltage_metrics")
                kk = k[k["interval"] == b] if k is not None and not k.empty else None
                thd.append(_n(float(kk[["UTHD12_pct", "UTHD23_pct", "UTHD31_pct"]].max(axis=1).iloc[0]))
                           if kk is not None and len(kk) else "–")
            rows.append(thd)
            w = max(1.2, (d.text_width_cm - 3.0) / max(len(cols), 1))
            d.table(header, rows, caption=tr("document.tab_voltages"), widths_cm=[3.0] + [w] * len(cols),
                    font_size=7.5 if len(cols) > 6 else 8.5)
            self.tables["voltages"] = self.table_number()

    # ------------------------------------------------------------------ chapter 3
    def chapter_assessment(self):
        c, d = self.c, self.d
        d.heading(tr("document.ch_assessment"), 1)
        self.processing_ires()
        if c.subs:
            self.equivalent_sources()
        if c.expansion is not None:
            self.expansion()
        for n in c.stations:
            self.station_assessment(n)
        self.uncertainties()
        self.summary()

    def processing_ires(self):
        c, d = self.c, self.d
        n = c.basis
        t = c.tests[n]
        self.heading(tr("document.sec_ires"), 2, key="ires")
        d.para(tr("document.ires_intro", name=n))
        d.para(tr("document.step1", mi=_n(100 * c.mI)))
        orders = self.orders_shown([n])
        rows = [(tr("document.row_rms"), _n(t["IRES"]), _n(t["IRES_a"]))]
        fn = self._f_nominal()
        for m in [1] + orders:
            i = t["harmonics"].get(m)
            rows.append((tr("document.row_order", f=_hz(m * fn), m=m), _n(i), _n(i * (1 + c.mI) if _finite(i) else None)))
        d.table(["", tr("document.col_measured"), tr("document.col_with_margin")], rows,
                caption=tr("document.tab_step1"), widths_cm=[5.0, 4.0, 4.0])
        s = c.scaled.get(n)
        k = c.factor
        if s is None:
            d.para(tr("document.step2_none"))
            d.para(tr("document.step3_unscaled", ires=_n(t["IRES_a"])))
        else:
            u_ref = tr("document.ref_table", n=self.tables["voltages"]) if "voltages" in self.tables else ""
            safe = tr("document.step2_safe", n=self.tables["pq_tests"]) if "pq_tests" in self.tables else ""
            tab = s["table"]
            thd_rows = bool((tab["rule"] == "THD").any())
            if thd_rows and not (tab["rule"] == "own voltage").any():      # a PQ measurement with the THD only
                d.para(tr("document.step2_thd_only", k=_n(k)))
            else:
                d.para(tr("document.step2", k=_n(k), u_ref=u_ref, safe=safe))
            f_thd, u_meas, thd95 = s.get("f_thd"), s.get("UTHD_meas"), c.thd95_used.get(n)
            if thd_rows and _finite(f_thd) and _finite(u_meas) and _finite(thd95):
                ref = tr("document.ref_table_paren", n=self.tables["voltages"]) if "voltages" in self.tables else ""
                d.para(tr("document.step2_thd_value", k=_n(k), thd95=_n(thd95), thd=_n(u_meas), f=_n(f_thd),
                          u_ref=ref))
            d.para(tr("document.step2_condition"))
            shown = tab[tab["I_A"] >= self.order_threshold]
            rest = tab[(tab["I_A"] < self.order_threshold)]
            rows = []
            for _, r in shown.iterrows():
                rows.append((str(int(r["order"])), _hz(int(r["order"]) * fn), tr("document.category_" + r["category"]),
                             _n(r["I_A"]), _n(r["U_level_V"]), _n(r["U95_V"]), _n(r["factor"]),
                             tr("document.rule_" + str(r["rule"])), _n(r["I_scaled_A"])))
            d.table([tr("document.col_order"), tr("document.col_f"), tr("document.col_category"),
                     tr("document.col_i_measured"), tr("document.col_level"), tr("document.col_u95"),
                     tr("document.col_k"), tr("document.col_rule"), tr("document.col_i_scaled")], rows,
                    caption=tr("document.tab_step2", name=n), widths_cm=[1.6, 1.3, 1.8, 1.4, 1.6, 1.6, 1.1, 2.9, 2.1],
                    align=["right", "right", "left", "right", "right", "right", "right", "left", "right"], font_size=8)
            context = self.scaling_context(n, s)
            if context:
                d.para(context)
            rest_scaled = math.sqrt(float((rest["I_scaled_A"] ** 2).sum())) if len(rest) else 0.0
            thd_orders = [int(o) for o in rest.loc[rest["rule"] == "THD", "order"]]
            thd = tr("document.step2_thd_orders", orders=", ".join(map(str, thd_orders))) if thd_orders else ""
            d.para(tr("document.step2_rest", threshold=_n(self.order_threshold), rest=_n(rest_scaled), thd=thd,
                      zero=_n(s["IH_zero"]), even=_n(s["IH_even"]), other=_n(s["IH_other"])))
            d.formula(tr("document.eq_ih_scaled", ih=_n(s["IH_scaled"])))
            d.formula(tr("document.eq_ires_scaled", ig=_n(t["IG"]), ih=_n(s["IH_scaled"]), ires=_n(s["IRES_scaled"])))
            d.para(tr("document.step3_mu", mu=_n(100 * c.mU)) if c.mU > 0 else tr("document.step3"))
            d.formula(tr("document.eq_ih_assessed_mu" if c.mU > 0 else "document.eq_ih_assessed",
                         ih=_n(s["IH_scaled_assessed"])))
            ig_a = t["IG"] * (1 + c.mI)
            d.formula(tr("document.eq_ires_max", ig=_n(ig_a), ih=_n(s["IH_scaled_assessed"]),
                         ires=_n(s["IRES_scaled_assessed"])))
        # all tests
        rows = []
        for m in c.names:
            tm = c.tests[m]
            sm = c.scaled.get(m)
            rows.append((m, _n(tm["IRES"]), _n(tm["IRES_a"]), _n(sm["IH_scaled"]) if sm else "–",
                         _n(sm["IRES_scaled"]) if sm else "–", _n(sm["IRES_scaled_assessed"]) if sm else "–"))
        d.table([tr("document.col_test"), tr("document.col_ires_measured"), tr("document.col_ires_margin"),
                 tr("document.col_ih_scaled"), tr("document.col_ires_scaled"), tr("document.col_ires_scaled_a")],
                rows, caption=tr("document.tab_all_tests"), widths_cm=[5.0, 2.2, 2.2, 2.2, 2.2, 2.4], font_size=8.5)
        st_scaled = [m for m in c.stations if m in c.scaled]
        if st_scaled:
            d.para(tr("document.station_scaled_note", names=self.join(st_scaled)))
        if c.pq_valid and c.stored_factors and any(abs(x - c.pq.factor) > 1e-9 for x in c.stored_factors):
            d.para(tr("document.single_factor_note", stored=" / ".join(_n(x) for x in c.stored_factors),
                      used=_n(c.pq.factor)))

    def equivalent_sources(self):
        c, d = self.c, self.d
        self.heading(tr("document.sec_es"), 2, key="es")
        fn = self._f_nominal()
        if c.series is None:
            with_z = any(_finite(c.Z.get(n)) for n in c.subs)
            d.para(tr("document.no_es_given" if with_z else "document.no_es"))
            return
        d.para(tr("document.es_intro"))
        d.formula(tr("document.eq_es"))
        zs = sorted({float(c.Z[n]) for n in c.subs if _finite(c.Z.get(n))})
        zs_text = " / ".join(_n(z) for z in zs) if zs else "–"
        measured = (c.series.settings or {}).get("source") == "measured"
        how = tr("document.zq_smallest" if (c.series.settings or {}).get("impedance") == "smallest"
                 else "document.zq_median")
        d.para(tr("document.es_source_measured" if measured else "document.es_source_fitted", zs=zs_text, how=how))
        st = c.series.source_table()
        loo_max: Dict[int, float] = {}
        if c.loo is not None and len(c.loo) and "deviation_pct" in c.loo:
            for o, g in c.loo.groupby("order"):
                vals = pd.to_numeric(g["deviation_pct"], errors="coerce").abs().dropna()
                if len(vals):
                    loo_max[int(o)] = float(vals.max())
        parts = (c.es or {}).get("parts", {}) if c.es else {}
        scaled_tab = c.es_scaled["table"] if c.es_scaled is not None else None
        rows = []
        for _, r in st.iterrows():
            m = int(r["order"])
            i_s = "–"
            if scaled_tab is not None:
                row = scaled_tab[(scaled_tab["order"] == m) & (scaled_tab["part"] == "model")]
                i_s = _n(row["I_scaled_A"].iloc[0]) if len(row) else "–"
            seq = {"pos": "document.positive", "neg": "document.negative", "zero": "document.zero"}.get(
                str(r.get("sequence")), "")
            rows.append(("%d (%s)" % (m, tr(seq)) if seq else str(m), _hz(m * fn),
                         tr("document.source_measured" if r.get("method") == "measured" else "document.source_fitted"),
                         _n(r["Uq_V"]), _n(r["Rq_Ohm"]), _n(r["Xq_Ohm"]),
                         _n(loo_max.get(m), 2) if m in loo_max else "–", _n(parts.get(m)), i_s))
        if c.Z_E is not None:
            header = [tr("document.col_order"), tr("document.col_f"), tr("document.col_source"), tr("document.col_uq"),
                      tr("document.col_rq"), tr("document.col_xq"), tr("document.col_loo"), tr("document.col_i_ze"),
                      tr("document.col_i_ze_scaled")]
            d.table(header, rows, caption=tr("document.tab_es", Z=_n(c.Z_E)),
                    widths_cm=[2.0, 1.3, 1.9, 1.7, 1.6, 1.6, 2.2, 1.8, 1.9], font_size=8,
                    align=["left", "right", "left"] + ["right"] * 6)
        self.fig(lambda: report.source_fit(c.series, Z_E=c.Z_E, comparison=c.comparison), "source_fit",
                 tr("document.fig_source_fit"))
        if c.es is not None:
            r0 = c.es
            zmin, zmax = c.series.Z_range
            inside = tr("document.inside") if zmin <= abs(c.Z_E) <= zmax else tr("document.outside")
            d.para(tr("document.es_result", Z=_n(c.Z_E), ires=_n(r0["IRES"]), ig=_n(r0["IG"]), ih=_n(r0["IH"]),
                      assessed=tr("document.es_result_assessed", a=_n(r0["IRES_assessed"])) if c.mI > 0 else "",
                      inside=inside, zmin=_n(zmin), zmax=_n(zmax),
                      basis=tr("document.es_basis", basis=_sentence(c.info.Z_E_basis)) if c.info.Z_E_basis else ""))
            if c.es_scaled is not None:
                rs = c.es_scaled
                how, kept = _es_scaling_text(rs)
                d.para(tr("document.es_scaled", how=how, kept=kept, ih=_n(rs["IH_scaled"]),
                          ires=_n(rs["IRES_scaled"]), a=_n(rs["IRES_scaled_assessed"])))
                ref = tr("document.es_decomposition_ref", sec=self.sections["ires"]) if "ires" in self.sections else ""
                d.para(tr("document.es_decomposition", ires=_n(r0["IRES"]), ih=_n(r0["IH"]), ihs=_n(rs["IH_scaled"]),
                          ig=_n(r0["IG"]), ref=ref))
            if c.mc_band and 5 in c.mc_band and 95 in c.mc_band:
                cfg = dict(sigma_magnitude=0.01, sigma_phase_deg=0.5, sigma_Z=0.02, sigma_level=0.0, model_error=True)
                cfg.update((c.mc or {}).get("settings") or {})          # a result of an older version: defaults
                d.para(tr("document.es_mc", p5=_n(c.mc_band[5]), p95=_n(c.mc_band[95]), h5=_n(c.mc_band_ih[5]),
                          h95=_n(c.mc_band_ih[95]), sm=_g(100 * cfg["sigma_magnitude"]),
                          sp=_g(cfg["sigma_phase_deg"]), sz=_g(100 * cfg["sigma_Z"]),
                          level=tr("document.es_mc_level", sl=_g(100 * cfg["sigma_level"]))
                          if cfg["sigma_level"] else "",
                          model=tr("document.es_mc_model") if cfg["model_error"] else "",
                          n=int((c.mc or {}).get("n") or 0) or "–"))
        if loo_max:
            o = max(loo_max, key=loo_max.get)
            lim = ""
            if c.loo is not None and "within_limit" in c.loo:
                w = c.loo["within_limit"]
                bad = int(w.eq(False).sum())
                if bad:
                    lim = tr("document.es_loo_limit", bad=bad, n=int(w.notna().sum()), lim=self.loo_limit_text())
            d.para(tr("document.es_loo", dev=_n(loo_max[o], 2), order=o, limit=lim))
        fig_ze = 0
        if c.Z_E is not None or c.limits:
            zmax_plot = max(10.0, 2 * (c.series.Z_range[1] or 0), 1.5 * abs(c.Z_E or 0))
            lims = " / ".join(_g(x) for x in c.limits)
            ranges = {lc["limit_A"]: lc["Z_E"] for lc in c.limit_cases if _finite(lc["Z_E"])}
            fig_ze = self.fig(lambda: report.ires_over_ze(c.series, zmax_plot, c.limits or tuple(ranges), c.IG,
                                                          c.scaling, comparison=c.comparison, limit_ranges=ranges,
                                                          Z_E=c.Z_E, margins=dict(uncertainty_I=c.mI,
                                                                                  uncertainty_U=c.mU)),
                              "ires_over_ze", tr("document.fig_ires_ze") + (" (%s A)" % lims if lims else "")
                              + (tr("document.fig_ires_ze_ranges") if ranges else ""))
        self.limit_cases_table()
        ps = c.pq_scaling
        if ps is not None and len(ps) and c.Z_E is not None:
            rows = [(_markup_symbols(r["method"]), _n(r["IH_A"]), _n(r["IRES_A"]), _n(r["IRES_assessed_A"]))
                    for _, r in ps.iterrows()]
            d.table([tr("document.col_method"), tr("document.col_ih"), tr("document.col_ires"),
                     tr("document.col_ires_assessed")], rows, caption=tr("document.tab_methods", Z=_n(c.Z_E)),
                    widths_cm=[7.6, 2.4, 2.4, 3.0])
        d.para("**%s**" % tr("document.forecast_title"), keep_next=True)
        xq = pd.to_numeric(st.get("Xq_Ohm", pd.Series(dtype=float)), errors="coerce").dropna()
        reactive = tr("document.forecast_reactive") if len(xq) and (xq < 0).all() else ""
        items = [tr("document.forecast_1", fig=tr("document.ref_figure", n=fig_ze) if fig_ze else "",
                    reactive=reactive), tr("document.forecast_2" if c.scaling else "document.forecast_2_unscaled")]
        if c.expansion is not None and _finite(c.d):
            items.append(tr("document.forecast_3_expansion", d=_fixed(100 * c.d, 1)))
        elif _finite(c.ice_limit):
            ig = ""
            if _finite(c.v) and c.v > 0 and _finite(c.d) and c.d > 0:
                ig = tr("document.forecast_3_ig", v=_fixed(100 * c.v, 1), ice=_n(c.limit / math.hypot(c.d, c.v), 2))
            items.append(tr("document.forecast_3", d=_fixed(100 * c.d, 1), limit=_n(c.limit), ice=_n(c.ice_limit, 2),
                            ice_now=_n(c.ice, 3), ig=ig))
        d.bullets(items)

    def expansion(self):
        """Network expansion: the equivalent sources converted to a network with more cable at the substation."""
        c, d = self.c, self.d
        x = c.expansion
        self.heading(tr("document.sec_expansion"), 2, key="expansion")
        d.para(tr("document.expansion_intro", sec=self.sections.get("es", "")))
        d.para(tr("document.expansion_z0"), keep_next=True)
        d.formula(tr("document.eq_expansion_z0"))
        d.para(tr("document.expansion_z1", c=_g(x["c1_ratio"])), keep_next=True)
        d.formula(tr("document.eq_expansion_z1"))
        d.para(tr("document.expansion_zq"), keep_next=True)
        d.formula(tr("document.eq_expansion_zq"))
        d.para(tr("document.expansion_uq"), keep_next=True)
        d.formula(tr("document.eq_expansion_uq"))
        coil = x["zero_sequence"] == "coil"
        d.para(tr("document.expansion_uq0_coil" if coil else "document.expansion_uq0_network"), keep_next=True)
        d.formula(tr("document.eq_expansion_uq0_coil" if coil else "document.eq_expansion_uq0_network"))
        d.para(tr("document.expansion_ig"), keep_next=True)
        d.formula(tr("document.eq_expansion_ig"))
        two = len(x["splits"]) > 1
        d.para(tr("document.expansion_split" if two else "document.expansion_split_one"), keep_next=True)
        items = [tr("document.expansion_split_sources", ice=_n(x["ice"]), il=_n(x["il"]), d=_fixed(100 * x["d"], 1),
                    ule=_n(x["U_LE"], 4))]
        if two:
            items.append(tr("document.expansion_split_busbar"))
        d.bullets(items)
        # ---- assumptions
        d.para("**%s**" % tr("document.expansion_assumptions"), keep_next=True)
        rng = (tr("document.expansion_a_range", ilmax=_n(x["I_L_max"]), icemax=_n(x["coil_end"]))
               if x["I_L_max"] else tr("document.expansion_a_range_unknown"))
        scaling = ""
        if c.scaling:
            scaling = tr("document.expansion_a_scaling", k=_n(c.factor) if _finite(c.factor) else "–",
                         sec=self.sections.get("ires", ""))
        d.bullets([tr("document.expansion_a_cables", c=_g(x["c1_ratio"]),
                      screen=tr("document.expansion_a_screen") if abs(x["c1_ratio"] - 1.0) < 1e-9 else "",
                      d=_fixed(100 * x["damping_added"], 1),
                      same=tr("document.expansion_a_same") if abs(x["damping_added"] - x["d"]) < 1e-12 else ""),
                   tr("document.expansion_a_coils", il=_n(x["il"]), range=rng),
                   tr("document.expansion_a_sources", scaling=scaling),
                   tr("document.expansion_a_rest")])
        # ---- the planned values at the earthing impedance of the report
        col, ihc = x["column"], x["ih_column"]
        rows_ = x["rows"]
        name = {sp: tr("document.split_" + sp) for sp in x["splits"]}
        ih_sym = "I_[H,skaliert]" if (c.scaling and c.lang == "de") else "I_[H,scaled]" if c.scaling else "I_[H]"
        ref = rows_["sources"]
        n_t = len(x["targets"])
        dec = np.nanmax(np.vstack([rows_[sp][col].to_numpy(float) for sp in x["splits"]]), axis=0)
        prop = [x["now"] * float(v) / x["ice"] for v in ref["I_C_A"]]
        table = [(tr("document.xrow_ice"),) + tuple(_n(v) for v in ref["I_C_A"]),
                 (tr("document.xrow_delta"), "–") + tuple("%s A (%s %%)" % (_signed(t - x["ice"], 3),
                                                                           _signed(100 * (t / x["ice"] - 1), 2))
                                                        for t in x["targets"]),
                 (tr("document.xrow_il"),) + tuple(_n(v) for v in ref["I_L_A"]),
                 (tr("document.xrow_ig"),) + tuple(_n(v) for v in ref["IG_A"])]
        for sp in x["splits"]:
            table.append((tr("document.xrow_ih", ih=ih_sym, split=name[sp]),) + tuple(_n(v) for v in rows_[sp][ihc]))
        for sp in x["splits"]:
            table.append((tr("document.xrow_ires", split=name[sp]),) + tuple(_n(v) for v in rows_[sp][col]))
        table.append((tr("document.xrow_prop"),) + tuple(_n(v) for v in prop))
        if c.limit:
            table.append((tr("document.xrow_status", limit=_g(c.limit)),) + tuple(c.ok(v) for v in dec))
        head = ["", tr("document.col_today")] + [tr("document.col_planned", ice=_g(t)) for t in x["targets"]]
        w0 = 6.6 if n_t <= 3 else 5.4
        d.table(head, table, caption=tr("document.tab_expansion", Z=_n(x["Z_E"])),
                widths_cm=[w0] + [(16.2 - w0) / (n_t + 1)] * (n_t + 1), align=["left"] + ["right"] * (n_t + 1),
                font_size=8.5)
        # ---- the orders at the first planned value
        t1 = x["targets"][0]
        orders = sorted(c.series.fits)
        orows = []
        for m in orders:
            key = "I_%d_A" % m
            if key not in ref.columns:
                continue
            row = [str(m), _n(ref[key].iloc[0])]
            row += [_n(rows_[sp][key].iloc[1]) for sp in x["splits"]]
            for sp in x["splits"]:
                tb = x["first"][sp].table.set_index("order")
                row.append(_n(tb.loc[m, "source_factor"]) if m in tb.index else "–")
            orows.append(tuple(row))
        t_orders = 0
        if orows:
            head = [tr("document.col_order"), tr("document.col_i_today")]
            head += [tr("document.col_i_split", split=_sentence(name[sp]).rstrip(".")) for sp in x["splits"]]
            head += [tr("document.col_u_split", split=name[sp]) for sp in x["splits"]]
            k = len(x["splits"])
            d.table(head, orows, caption=tr("document.tab_expansion_orders", Z=_n(x["Z_E"]), ice=_g(t1)),
                    widths_cm=[1.7, 1.9] + [3.0] * k + [3.3] * k if k > 1 else [1.7, 2.4, 3.5, 3.9],
                    align=["right"] * len(head), font_size=8.5)
            t_orders = self.table_number()
        # ---- what it means
        t_rows = {sp: rows_[sp].iloc[1] for sp in x["splits"]}
        vals = [float(t_rows[sp][col]) for sp in x["splits"]]
        ihs = [float(t_rows[sp][ihc]) for sp in x["splits"]]
        span = _span
        order = ""
        parts = {m: float(ref["I_%d_A" % m].iloc[0]) for m in orders if "I_%d_A" % m in ref.columns}
        if parts and t_orders:
            dom = max(parts, key=parts.get)
            ks = []
            for sp in x["splits"]:
                tb = x["first"][sp].table.set_index("order")
                if dom in tb.index:
                    ks.append(float(tb.loc[dom, "source_factor"]))
            if ks:
                order = tr("document.expansion_order", m=dom, k=span(ks), t=t_orders)
        d.para(tr("document.expansion_result", ice=_g(t1), delta="%s A, %s %%" % (
            _signed(t1 - x["ice"], 3), _signed(100 * (t1 / x["ice"] - 1), 2)), Z=_n(x["Z_E"]), range=span(vals),
            now=_n(x["now"]),
            limit=tr("document.expansion_vs_limit", limit=_g(c.limit), status=c.ok(max(vals))) if c.limit else "",
            ig=_n(ref["IG_A"].iloc[1]), ig_now=_n(ref["IG_A"].iloc[0]), ih=ih_sym, ih_range=span(ihs),
            ih_now=_n(ref[ihc].iloc[0]), order=order, prop=_n(x["now"] * t1 / x["ice"])))
        # ---- I_RES,ass over I_CE
        curves = []
        for cv in x["curves"].values():
            frames = {sp: f[f["I_C_A"] <= x["end"] + 1e-9] for sp, f in cv["frames"].items()}
            curves.append(dict(Z_E=cv["Z_E"], sources=frames.get("sources"), busbar=frames.get("busbar")))
        lims = [dict(limit_A=lc["allowed"], Z_E=lc["Z_E"]) for lc in x["cases"]]
        self.fig(lambda: report.ires_over_ice(curves, lims, I_CE=x["ice"], targets=x["targets"] if x["planned"] else (),
                                              column=col, I_CE_max=x["end"]),
                 "ires_over_ice", tr("document.fig_expansion"))
        # ---- where every limit is reached
        if x["cases"]:
            u_col = any(lc["governs"] == "U" for lc in x["cases"])
            rows = []
            for lc in x["cases"]:
                cond = lc["condition"] or tr("document.lc_standard")
                row = [_g(lc["limit_A"]), cond, _n(lc["Z_E"]), _g(lc["r"])]
                if u_col:
                    row.append(_n(lc["allowed"]))
                row.append(_n(lc["now"]))
                for sp in x["splits"]:
                    row.append(self.reach_text(lc["per"][sp], x["top"]))
                r = lc["reach"]
                row.append("–" if r["status"] != "reached" else _fixed(round(r["margin_A"], -1), 0))
                rows.append(tuple(row))
            head = [tr("document.col_limit"), tr("document.col_condition"), tr("document.col_ze"), tr("document.col_r")]
            head += [tr("document.col_allowed")] if u_col else []
            short = {sp: tr("document.split_busbar_short") if sp == "busbar" else name[sp] for sp in x["splits"]}
            head += [tr("document.col_ires_today")] + [tr("document.col_reached", split=short[sp]) for sp in x["splits"]]
            head += [tr("document.col_margin")]
            k = len(x["splits"])
            wc = 16.2 - (1.9 + 1.3 + 1.0 + 1.9 + 2.5 * k + 1.7) - (1.6 if u_col else 0.0)
            d.table(head, rows, caption=tr("document.tab_expansion_limits"),
                    widths_cm=[1.9, wc, 1.3, 1.0] + ([1.6] if u_col else []) + [1.9] + [2.5] * k + [1.7],
                    align=["right", "left", "right", "right"] + ["right"] * (len(head) - 4), font_size=8)
            if u_col:
                d.para(tr("document.expansion_allowed_u", u=_n(x["fu"])))
        for text in self.expansion_limit_sentences():
            d.para(text)
        if _finite(x["coil_end"]) and x["coil_end"] < max(x["targets"] + [lc["reach"]["I_C_A"] for lc in x["cases"]
                                                                          if lc["reach"]["status"] == "reached"]):
            d.para(tr("document.expansion_coils_short", ice=_fixed(round(x["coil_end"], -1), 0)))
        rej = x.get("rejected") or []
        neg = [m for m in x["negative"] if m not in rej]       # orders with a negative real part in network data too
        still = [m for m in x["negative"] if m in rej]
        if neg:
            d.para(tr("document.expansion_negative_one", orders=neg[0]) if len(neg) == 1 else
                   tr("document.expansion_negative", orders=self.join([str(m) for m in neg])))
        if rej:
            text = (tr("document.expansion_rejected_one", orders=rej[0]) if len(rej) == 1 else
                    tr("document.expansion_rejected", orders=self.join([str(m) for m in rej])))
            if still:
                text += (tr("document.expansion_rejected_still_one") if len(rej) == 1 else
                         tr("document.expansion_rejected_still", orders=self.join([str(m) for m in still])))
            d.para(text)
        d.para(tr("document.expansion_validity"))

    @staticmethod
    def reach_text(r: dict, top: float) -> str:
        """The I_CE at which a curve reaches a limit (rounded to 10 A), "today exceeded" or "> end of the curve"."""
        if r["status"] == "exceeded":
            return tr("document.expansion_now_exceeded")
        if r["status"] == "not_reached":
            return tr("document.expansion_beyond", ice=_fixed(round(top, -1), 0))
        return _fixed(round(r["I_C_A"], -1), 0)

    def expansion_limit_sentences(self) -> List[str]:
        """Per limit: where the expanded network reaches it."""
        c = self.c
        x = c.expansion
        out = []
        for lc in x["cases"]:
            kw = dict(limit=_g(lc["limit_A"]), condition=lc["condition"] or tr("document.lc_standard"), Z=_n(lc["Z_E"]),
                      u=tr("document.expansion_limit_u", allowed=_n(lc["allowed"])) if lc["governs"] == "U" else "")
            r = lc["reach"]
            if r["status"] == "exceeded":
                out.append(tr("document.expansion_limit_exceeded", now=_n(lc["now"]), **kw))
            elif r["status"] == "not_reached":
                out.append(tr("document.expansion_limit_not_reached", ice=_fixed(round(x["top"], -1), 0), **kw))
            else:
                split = ""
                if len(x["splits"]) > 1 and all(lc["per"][sp]["status"] == "reached" for sp in x["splits"]):
                    a, b = (lc["per"][sp]["I_C_A"] for sp in x["splits"])
                    if abs(a - b) >= 10:
                        split = tr("document.expansion_limit_split", a=_fixed(round(a, -1), 0),
                                   b=_fixed(round(b, -1), 0))
                out.append(tr("document.expansion_limit_reached", ice=_fixed(round(r["I_C_A"], -1), 0),
                              margin=_fixed(round(r["margin_A"], -1), 0), split=split, **kw))
        return out

    def limit_cases_table(self):
        """The limits with their conditions: I_RES,ass at the Z_E of every case and U_E = r·I_RES,ass·Z_E."""
        c, d = self.c, self.d
        cases = [lc for lc in c.limit_cases if _finite(lc["I_A"])]
        if not cases:
            return
        fu = cases[0]["FU_V"]
        d.para(tr("document.lc_intro", u=_n(fu)) if _finite(fu) else tr("document.lc_intro_no_u"))
        if any(lc["basis"] == "busbar" for lc in cases):
            d.para(tr("document.lc_busbar"))
        rows = []
        for lc in cases:
            cond = (lc["condition"] or tr("document.lc_standard")) + (tr("document.lc_applies") if lc["applies"] else "")
            rows.append((_g(lc["limit_A"]), cond, _n(lc["Z_E"]), _g(lc["r"]), _n(lc["I_A"]), self.status(lc["ok_I"]),
                         _n(lc["U_E_V"]) if _finite(lc["U_E_V"]) else "–", self.status(lc["ok_U"])))
        d.table([tr("document.col_limit"), tr("document.col_condition"), tr("document.col_ze"), tr("document.col_r"),
                 tr("document.col_ires_assessed"), tr("document.col_vs_limit"), tr("document.col_ue"),
                 tr("document.col_vs_utp")], rows, caption=tr("document.tab_limit_cases"),
                widths_cm=[1.9, 4.0, 1.3, 1.0, 2.0, 2.2, 1.5, 2.1], font_size=8,
                align=["right", "left", "right", "right", "right", "center", "right", "center"])
        for text in self.limit_case_sentences():
            d.para(text)

    def loo_limit_text(self) -> str:
        """The permissible deviation of the leave-one-out check: ±15 % (or ±0,5 A)."""
        attrs = getattr(self.c.loo, "attrs", {}) or {}
        parts = []
        plain_ = lambda x: "%d" % int(round(float(x))) if float(x).is_integer() else _n(x)
        if _finite(attrs.get("limit_pct")):
            parts.append("±%s %%" % plain_(attrs["limit_pct"]))
        if _finite(attrs.get("limit_A")):
            parts.append("±%s A" % plain_(attrs["limit_A"]))
        return (" %s " % tr("document.or")).join(parts) if parts else tr("document.loo_limit_unknown")

    def station_assessment(self, n: str):
        c, d = self.c, self.d
        place = c.place(n)
        self.heading(tr("document.sec_station_assessment", station=place), 2, key="station:" + n)
        d.para(tr("document.ut_condition"))
        d.formula(tr("document.eq_ut"))
        if not (_finite(c.u_tp) and _finite(c.f_factor)):
            d.para(tr("document.ut_missing"))
            return
        d.para(tr("document.ut_parameters", utp=_n(c.u_tp), f=_n(c.f_factor),
                  r=tr("document.ut_r_flat", r=_g(c.info.r_flat)) if _finite(c.info.r_flat) else ""))
        cases = c.cases.get(n)
        if cases is None or not len(cases):
            return
        rows = [(str(r["description"]), _n(r["I_A"]), _g(r["r"]) if _finite(r["r"]) else "–",
                 _n(r["Z_E_Ohm"]) if _finite(r["Z_E_Ohm"]) else "–", _n(r["U_E_V"]) if _finite(r["U_E_V"]) else "–",
                 _n(r["U_T_V"]), _share(r["share_pct"])) for _, r in cases.iterrows()]
        d.table([tr("document.col_case"), tr("document.col_i"), tr("document.col_r"), tr("document.col_ze"),
                 tr("document.col_ue"), tr("document.col_ut"), tr("document.col_share")], rows,
                caption=tr("document.tab_ut", station=place), widths_cm=[6.6, 1.6, 1.2, 1.6, 1.6, 1.6, 1.8],
                font_size=8.5)
        f_ref = cases.attrs.get("reference_frequency_Hz")
        z_ref = float("nan")
        if (cases["case"] == "1").any():
            z_ref = float(cases.loc[cases["case"] == "1", "Z_E_Ohm"].iloc[0])
            z50 = self.earthing_impedance_at(n, 50.0)
            z50_text = tr("document.ut_cases_z50", z=_n(z50)) if _finite(z50) and _finite(f_ref) and \
                abs(float(f_ref) - 50.0) > 5.0 else ""
            method = cases.attrs.get("reference_method")
            br = cases.attrs.get("reference_bracket_Hz")
            req = cases.attrs.get("reference_frequency_requested_Hz", f_ref)
            how = (tr("document.ut_ref_interpolated", lo=_hz(br[0]), hi=_hz(br[1])) if method == "interpolated" and br
                   else tr("document.ut_ref_nearest", used=_hz(f_ref), req=_hz(req)) if method == "nearest" else "")
            d.para(tr("document.ut_cases_note", limit=_g(c.limit), f=_hz(req), z=_n(z_ref), how=how, z50=z50_text))
            if not (cases["case"] == "2").any():
                d.para(tr("document.ut_case2_missing"))
        if (cases["case"] == "3").any():
            om = cases.attrs.get("orders_measured") or []
            d.para(tr("document.ut_case3_note", name=c.basis, orders=", ".join(map(str, om)) or "–",
                      probe=cases.attrs.get("probe") or "–", covered=_fixed(cases.attrs.get("covered_pct"), 0)))
        inv = self.probe_inversion(n) if c.info.station_detail != "compact" else ""
        if inv:
            d.para(inv)
        self.fig(lambda: report.touch_voltage_cases(cases, c.u_tp), "touch_voltage_" + report.slug(n),
                 tr("document.fig_ut", station=place))
        i_max = self.decisive_current()
        if _finite(i_max) and i_max > 0:
            z1 = float(c.f_factor) * float(c.u_tp) / i_max
            zr = ""
            eff = (c.reduction.get(n) or {}).get("reduction_factor_eff") if c.info.station_detail != "compact" else None
            if _finite(c.info.r_flat) and float(c.info.r_flat) > 0:
                zr += tr("document.ze_permissible_flat", z=_n(z1 / float(c.info.r_flat)), r=_g(c.info.r_flat))
            if eff is not None and len(eff):
                row = eff[eff["channel"] == "max(UE)"].iloc[0] if (eff["channel"] == "max(UE)").any() else eff.iloc[0]
                r = row.get("r_eff")
                if _finite(r) and float(r) > 0:
                    zr += tr("document.ze_permissible_r", z=_n(z1 / float(r)), r=_n(r, 3))
            cmp = tr("document.ze_permissible_cmp", f=_hz(f_ref), z=_n(z_ref)) if _finite(z_ref) else ""
            d.para(tr("document.ze_permissible", i=_n(i_max), z1=_n(z1), zr=zr, cmp=cmp))

    def earthing_impedance_at(self, n: str, f: float) -> float:
        """Z_E of the earthing measurement of a station at the frequency closest to ``f`` (NaN without one)."""
        if n not in self.c.earthing:
            return float("nan")
        ratio = self.c.earthing[n][1]
        if ratio is None or not len(ratio) or "Z_E_injection_Ohm" not in ratio:
            return float("nan")
        rr = ratio[ratio["channel"] == "max(UE)"] if "channel" in ratio and (ratio["channel"] == "max(UE)").any() \
            else ratio
        fz = pd.to_numeric(rr["f_Hz"], errors="coerce")
        if fz.notna().sum() == 0:
            return float("nan")
        j = (fz - f).abs().idxmin()
        return float(rr.loc[j, "Z_E_injection_Ohm"]) if abs(float(fz[j]) - f) <= 5.0 else float("nan")

    def decisive_current(self) -> float:
        """The decisive assessment value of the residual current."""
        c = self.c
        if c.assessment == "Z_E" and c.es is not None:
            return float(c.es_scaled["IRES_scaled_assessed"] if c.es_scaled is not None else c.es["IRES_assessed"])
        s = c.scaled.get(c.basis)
        return float(s["IRES_scaled_assessed"]) if s is not None else float(c.tests[c.basis]["IRES_a"])

    def uncertainties(self):
        c, d = self.c, self.d
        d.heading(tr("document.sec_uncertainties"), 2)
        d.para(tr("document.uncertainties_intro"))
        rows = []
        comps = [float(b) for _, b in (c.info.uncertainty_components or ())]
        val = tr("document.u_chain_value", lin=_n(sum(comps)), rss=_n(math.sqrt(sum(b * b for b in comps)))) \
            if comps else "–"
        treat = (tr("document.u_margin_value_mu", mi=_n(100 * c.mI), mu=_n(100 * c.mU)) if c.mU > 0 else
                 tr("document.u_margin_value", mi=_n(100 * c.mI)) if c.mI > 0 else tr("document.informative"))
        rows.append((tr("document.u_chain"), val, treat))
        st = self.stationarity(c.basis)
        if st is not None:
            rows.append((tr("document.u_stationarity", name=c.basis),
                         tr("document.u_stationarity_value", n=str(st["n"]), min=_n(st["min"]), max=_n(st["max"]),
                            std=_n(st["std"], 2)) + self.order_spread(c.basis), tr("document.u_stationarity_treat")))
        others = [(m, self.stationarity(m)) for m in c.subs if m != c.basis]
        others = [(m, x) for m, x in others if x is not None]
        if others:
            rows.append((tr("document.u_stationarity_others"),
                         "; ".join(tr("document.u_stationarity_item", name=m, min=_n(x["min"]), max=_n(x["max"]),
                                      std=_n(x["std"], 2)) for m, x in others),
                         tr("document.u_stationarity_others_treat")))
        igs = [c.tests[n]["IG"] for n in c.subs if _finite(c.tests[n]["IG"])]
        if len(igs) > 1:
            span = 100 * (max(igs) - min(igs)) / float(np.median(igs))
            choice = ("%s A" % _n(c.IG)) if not isinstance(c.IG, str) else (
                tr("document.ig_" + c.IG) if c.IG in ("max", "median") else c.IG)
            rows.append((tr("document.u_tests"), tr("document.u_tests_value", min=_n(min(igs)), max=_n(max(igs)),
                                                    span=_n(span, 2)), tr("document.u_tests_treat", choice=choice)))
        if c.mc_band and 5 in c.mc_band and 95 in c.mc_band:
            loo = ""
            if c.loo is not None and "deviation_pct" in c.loo:
                vals = pd.to_numeric(c.loo["deviation_pct"], errors="coerce").abs().dropna()
                if len(vals):
                    loo = "; " + tr("document.u_loo_value", dev=_n(vals.max(), 2))
            rows.append((tr("document.u_model"), tr("document.u_model_value", p5=_n(c.mc_band[5]),
                                                    p95=_n(c.mc_band[95])) + loo, tr("document.informative")))
        if c.boot:
            v = tr("document.u_pq_value", p5=_n(c.boot["IRES_p5"]), p95=_n(c.boot["IRES_p95"]))
            if "order" in c.boot:
                v += "; " + tr("document.u_pq_value_order", order=c.boot["order"], low=_signed(c.boot["low"], 2),
                               high=_signed(c.boot["high"], 2))
            rows.append((tr("document.u_pq"), v, tr("document.u_pq_treat", k=_n(c.pq.factor))))
        per = c.pq_periods
        if per and per["order"] in per["table"].columns and any(_finite(x) for x in per["ires"]):
            ires = [x for x in per["ires"] if _finite(x)]
            qd = per["table"][per["order"]]
            rows.append((tr("document.u_pq_periods"),
                         tr("document.u_pq_periods_value", low=_n(min(ires)), high=_n(max(ires)), order=per["order"],
                            qlow=_n(qd.min()), qhigh=_n(qd.max())), tr("document.u_pq_periods_treat")))
        ps = c.pq_scaling
        if ps is not None and len(ps) > 1:
            vals = pd.to_numeric(ps["IRES_A"], errors="coerce").iloc[1:]
            rows.append((tr("document.u_methods"), tr("document.u_methods_value", min=_n(vals.min()), max=_n(vals.max())),
                         tr("document.u_methods_treat")))
        x = c.expansion
        if x is not None and len(x["splits"]) > 1:
            a, b = (float(x["rows"][sp][x["column"]].iloc[1]) for sp in x["splits"])
            rows.append((tr("document.u_expansion"), tr("document.u_expansion_value", ice=_g(x["targets"][0]),
                                                        a=_n(a), b=_n(b)), tr("document.u_expansion_treat")))
        if c.es is not None:
            zmin, zmax = c.series.Z_range
            inside = tr("document.inside") if zmin <= abs(c.Z_E) <= zmax else tr("document.outside")
            rows.append((tr("document.u_extrapolation"), tr("document.u_extrapolation_value", Z=_n(c.Z_E),
                                                            inside=inside, zmin=_n(zmin), zmax=_n(zmax)),
                         tr("document.informative")))
        d.table([tr("document.col_contribution"), tr("document.col_size"), tr("document.col_treatment")], rows,
                caption=tr("document.tab_uncertainties"), widths_cm=[5.0, 7.0, 4.2], align=["left", "left", "left"],
                font_size=8.5)

    def stationarity(self, n: str) -> Optional[dict]:
        """I_RES over the windows of the steady state of a test: ``n``, ``min``, ``max``, ``std`` (% of the mean)."""
        c = self.c
        stats = c.results[n].tables.get("statistics")
        if stats is None or stats.empty or not {"role", "quantity"} <= set(stats.columns):
            return None
        s = stats[(stats["role"] == "i_res") & (stats["quantity"] == "RMS")]
        main = c.tests[n]["main"]
        if len(s) and main:
            short = (getattr(c.results[n], "data", {}) or {}).get("signals", {}).get(main, {}).get("short")
            if short and (s["channel"] == short).any():
                s = s[s["channel"] == short]
        if not len(s):
            return None
        r = s.iloc[0]
        std = 100 * float(r["std"]) / float(r["mean"]) if _finite(r["mean"]) and float(r["mean"]) else np.nan
        return dict(n=int(r["n"]), min=float(r["min"]), max=float(r["max"]), std=std)

    def summary(self):
        c, d = self.c, self.d
        d.heading(tr("document.sec_summary"), 2)
        if _finite(c.ice):
            both = "document.sum_ice_controller" if c.controller.get("how") == "derived" else "document.sum_ice_both"
            src = tr(both, display=_n(c.coil_display), coils=_n(c.coil_sum)) \
                if _finite(c.coil_display) and _finite(c.coil_sum) else (
                tr("document.ice_from_display") if _finite(c.coil_display) else tr("document.ice_from_coils"))
            d.para(tr("document.sum_ice", ice=_n(c.ice), source=src))
        t = c.tests[c.largest]
        if _finite(t["IRES"]):
            share = 100 * t["IRES"] / c.ice if _finite(c.ice) and c.ice > 0 else float("nan")
            d.para(tr("document.sum_ires", ires=_n(t["IRES"]), name=c.largest, a=_n(t["IRES_a"]), share=_fixed(share, 1),
                      d=_fixed(100 * c.d, 1)))
            tb = c.tests[c.basis]
            if _finite(tb["IG"]) and _finite(tb["IH"]):
                hz = {m: i for m, i in tb["harmonics"].items() if m >= 2 and _finite(i)}
                dom = max(hz, key=hz.get) if hz else "–"
                which = lambda g, h: tr("document.the_fundamental" if g >= h else "document.the_harmonics")
                sb = c.scaled.get(c.basis)
                scaled = ""
                if sb is not None:
                    ig_a, ih_a = tb["IG"] * (1 + c.mI), sb["IH_scaled_assessed"]
                    same = (ig_a >= ih_a) == (tb["IG"] >= tb["IH"])
                    scaled = (tr("document.sum_composition_same", ig=_n(ig_a), ih=_n(ih_a)) if same else
                              tr("document.sum_composition_scaled", which=which(ig_a, ih_a), ig=_n(ig_a), ih=_n(ih_a)))
                d.para(tr("document.sum_composition", measured=which(tb["IG"], tb["IH"]), ig=_n(tb["IG"]),
                          ih=_n(tb["IH"]), order=dom, scaled=scaled))
        s = c.scaled.get(c.basis)
        lim = _n(c.limit) if c.limit else "–"
        if s is not None:
            d.para(tr("document.sum_scaled", k=_n(c.factor), ires=_n(s["IRES_scaled_assessed"]), limit=lim,
                      status=c.ok(s["IRES_scaled_assessed"])))
        if c.es is not None:
            v = c.es_scaled["IRES_scaled_assessed"] if c.es_scaled is not None else c.es["IRES_assessed"]
            # with limits and their conditions the verdict follows in the limit sentences – stated once
            judged = any(lc["applies"] and _finite(lc["I_A"]) for lc in c.limit_cases)
            d.para(tr("document.sum_es_value" if judged else "document.sum_es", Z=_n(c.Z_E), ires=_n(v), limit=lim,
                      status=c.ok(v), scaled=tr("document.sum_es_scaled") if c.es_scaled is not None else ""))
        if s is not None or c.es is not None:
            d.para(tr("document.sum_decisive", which=tr("document.decisive_ze" if c.assessment == "Z_E"
                                                          else "document.decisive_busbar")))
        for text in self.limit_case_sentences():
            d.para(text)
        for n in c.stations:
            cases = c.cases.get(n)
            if cases is not None and len(cases) and _finite(c.u_tp) and pd.to_numeric(
                    cases["U_T_V"], errors="coerce").notna().any():
                d.para(self.station_sentence(n, cases))
                continue
            m = c.results[n].tables.get("station_metrics")
            if m is not None and not m.empty:
                d.para(tr("document.sum_station_simple", station=c.place(n), ue=_n(m["UE_max_V"].iloc[0]),
                          i=_n(m["IF_A"].iloc[0])))
        if c.expansion is not None:
            d.para(" ".join(self.expansion_summary()))
        elif _finite(c.ice_limit):
            ig = ""
            if _finite(c.v) and c.v > 0 and _finite(c.d) and c.d > 0:
                ig = tr("document.sum_forecast_ig", ice=_n(c.limit / math.hypot(c.d, c.v), 2))
            d.para(tr("document.sum_forecast", limit=_n(c.limit), ice=_n(c.ice_limit, 2), ig=ig))
        i_max = self.decisive_current()
        if _finite(c.u_tp) and _finite(c.f_factor) and _finite(i_max) and i_max > 0:
            d.para(tr("document.sum_ze", z=_n(float(c.f_factor) * float(c.u_tp) / i_max),
                      u=_n(float(c.f_factor) * float(c.u_tp)), i=_n(i_max)))
        d.para("**%s**" % tr("document.recommendation"), keep_next=True)
        if c.info.recommendation:
            for part in str(c.info.recommendation).split("\n\n"):
                d.para(part.strip())
        else:
            d.placeholder_text(tr("document.recommendation_placeholder"))
            self.placeholders.append(plain(tr("document.recommendation_placeholder")))

    def expansion_summary(self) -> List[str]:
        """The network expansion in the summary: the first planned value against the limit of the network, then where
        the residual current reaches every limit."""
        c = self.c
        x = c.expansion
        t = x["targets"][0]
        v = max(float(x["rows"][sp][x["column"]].iloc[1]) for sp in x["splits"])
        already = c.limit is not None and x["now"] > c.limit
        lim = tr("document.sum_expansion_limit", limit=_g(c.limit), status=c.ok(v),
                 already=tr("document.sum_expansion_already") if already else "") if c.limit else ""
        out = [tr("document.sum_expansion", ice=_g(t), Z=_n(x["Z_E"]), value=_n(v), limit=lim)]
        same = lambda lc: lc["applies"] and abs(lc["Z_E"] - x["Z_E"]) < 1e-9
        reached, beyond, over = [], [], []
        for lc in x["cases"]:
            kw = dict(limit=_g(lc["allowed"]), Z=_n(lc["Z_E"]))
            st = lc["reach"]["status"]
            if st == "reached":
                reached.append(tr("document.sum_expansion_reached_item",
                                  ice=_fixed(round(lc["reach"]["I_C_A"], -1), 0), **kw))
            elif st == "not_reached":
                beyond.append(tr("document.sum_expansion_item", **kw))
            elif not (same(lc) and already):
                over.append(tr("document.sum_expansion_item", **kw))
        if reached:
            out.append(tr("document.sum_expansion_reached", items=self.join(reached)))
        if beyond:
            out.append(tr("document.sum_expansion_not_reached", items=self.join(beyond),
                          ice=_fixed(round(x["top"], -1), 0)))
        if over:
            out.append(tr("document.sum_expansion_exceeded", items=self.join(over)))
        return out

    def station_sentence(self, n: str, cases: pd.DataFrame) -> str:
        """The touch voltage of a station over all approaches (summary)."""
        c = self.c
        ut = pd.to_numeric(cases["U_T_V"], errors="coerce")
        utp = float(c.u_tp)
        labels = dict(zip(cases["case"], cases["label"]))
        label = lambda case: tr("document.probe") if case == "ut" else str(labels.get(case, case))
        top = cases.loc[ut.idxmax()]
        measured = cases[cases["case"].isin(["3", "4", "ut"]) & ut.notna()]
        m_text = ""
        if len(measured) and not (len(measured) == 1 and measured.index[0] == top.name):
            mm = pd.to_numeric(measured["U_T_V"], errors="coerce")
            m_text = tr("document.sum_station_measured", cases=", ".join(label(x) for x in measured["case"]),
                        ut=_n(mm.max()))
        over = cases[ut > utp]
        kw = dict(station=c.place(n), utp=_n(utp), ut=_n(ut.max()), case=label(top["case"]),
                  share=_share(100 * float(ut.max()) / utp), measured=m_text)
        if len(over):
            return tr("document.sum_station_exceeded", cases=", ".join(label(x) for x in over["case"]), **kw)
        return tr("document.sum_station_ok", **kw)

    def corrections_text(self, n: str) -> str:
        """The corrections of a test per channel with the total factor, e.g. "ISp22_27 × 30 (10 × 3)"."""
        per: Dict[str, List[float]] = {}
        for x in self.c.tests[n]["corrections"]:
            f = x.get("factor", x.get("factor_total"))
            if _finite(f):
                per.setdefault(str(x.get("short") or x.get("channel")), []).append(float(f))
        items = []
        for ch, fs in per.items():
            total = float(np.prod(fs))
            items.append("%s × %s" % (ch, _g(total)) + (" (%s)" % " × ".join(_g(f) for f in fs) if len(fs) > 1 else ""))
        return "; ".join(items) or "–"

    # ------------------------------------------------------------------ appendices
    def appendices(self):
        c, d = self.c, self.d
        d.page_break()
        d.heading(tr("document.appendix", letter="A", title=tr("document.app_data")), 1, numbered=False)
        d.para(tr("document.app_data_intro"))
        names = c.names
        tzl = self.tz_text()

        def iv(x):
            return "%s … %s" % (_n(x[0]), _n(x[1])) if x else "–"

        rows = []
        add = rows.append
        add([tr("document.d_file")] + [c.tests[n]["cfg"] for n in names])
        add([tr("document.d_site")] + [c.place(n) for n in names])
        labels = self.tz_labels() if self.tz_mixed() else {}
        add([tr("document.d_start", tz=tzl)] + [
            _local(c.tests[n]["start"], c.info.timezone).strftime("%d.%m.%Y %H:%M:%S" if c.lang == "de"
                                                                  else "%Y-%m-%d %H:%M:%S")
            + ((" " + labels[n]) if n in labels else "")
            if c.tests[n]["start"] is not None else "–" for n in names])
        add([tr("document.d_fs")] + [_n(c.tests[n]["fs"] / 1000.0) for n in names])
        add([tr("document.d_duration")] + [_n(c.tests[n]["duration"]) for n in names])
        add([tr("document.d_f1")] + [_n(c.tests[n]["f1"], 6) for n in names])
        add([tr("document.d_steady")] + [iv(c.tests[n]["steady"]) for n in names])
        add([tr("document.d_pre")] + [iv(c.tests[n]["pre"]) for n in names])
        add([tr("document.d_post")] + [iv(c.tests[n]["post"]) for n in names])
        add([tr("document.d_window")] + [_n(c.tests[n]["t_max"], 4) for n in names])
        add([tr("document.d_twin")] + [_n(1000 * c.tests[n]["T_win"]) for n in names])
        add([tr("document.d_z")] + [_n(c.Z.get(n)) if _finite(c.Z.get(n)) else "–" for n in names])
        add([tr("document.d_phase")] + [str(c.tests[n]["phase"] or "–") for n in names])
        add([tr("document.d_main")] + [
            str(((getattr(c.results[n], "data", {}) or {}).get("signals", {}).get(c.tests[n]["main"], {}) or {}).get(
                "short") or c.tests[n]["main"] or "–") for n in names])
        add([tr("document.d_corrections")] + [self.corrections_text(n) for n in names])
        if any(c.tests[n].get("polarity") == -1 for n in names):
            add([tr("document.d_polarity")] + [tr("document.d_polarity_rotated") if c.tests[n].get("polarity") == -1
                                               else "–" for n in names])
        if c.reproduction:
            add([tr("document.d_reproduced")] + [tr("document.d_reproduced_value", dev=_small(c.reproduction[n]))
                                                 if n in c.reproduction else "–" for n in names])
        w = max(2.0, (d.text_width_cm - 4.0) / max(len(names), 1))
        d.table([""] + list(names), rows, caption=tr("document.tab_data"), widths_cm=[4.0] + [w] * len(names),
                align=["left"] * (len(names) + 1), font_size=7.5)
        # ---- plausibility
        d.heading(tr("document.appendix", letter="B", title=tr("document.app_plausibility")), 1, numbered=False)
        d.para(tr("document.app_plausibility_intro"), keep_next=True)
        rows, n_ok = [], 0
        for n in names:
            for p in (getattr(c.results[n], "data", {}) or {}).get("plausibility", []) or []:
                if str(p.get("status", "")).upper() in ("OK", ""):
                    n_ok += 1
                    continue
                rows.append((n, _markup_symbols(p.get("check", "")), _status(p.get("status", "")),
                             _markup_symbols(p.get("value", "")), _markup_symbols(_updated_comment(p.get("comment", "")))))
        if rows:
            d.table([tr("document.col_measurement"), tr("document.col_check"), tr("document.col_status"),
                     tr("document.col_value"), tr("document.col_comment")], rows,
                    caption=tr("document.tab_plaus_single"), widths_cm=[3.0, 3.4, 2.0, 3.2, 4.6],
                    align=["left"] * 5, font_size=7.5)
        d.para(tr("document.plaus_ok_count", n=n_ok))
        if len(c.subs) > 1:
            from . import campaign as _campaign
            try:
                cons = _campaign.consistency({n: c.results[n] for n in c.subs})
                rows = [(_markup_symbols(r["check"]), _status(r["status"]), _markup_symbols(r["value"]),
                         _markup_symbols(r["comment"])) for _, r in cons.iterrows()]
                d.para(tr("document.plaus_campaign_intro", names=self.join(c.subs)), keep_next=True)
                d.table([tr("document.col_check"), tr("document.col_status"), tr("document.col_value"),
                         tr("document.col_comment")], rows, caption=tr("document.tab_plaus_campaign"),
                        widths_cm=[4.0, 1.6, 4.4, 6.2], align=["left"] * 4, font_size=7.5)
            except Exception as ex:
                c.hints.append(str(ex))
        rows = [(a, _status(b), v) for a, b, v in self.report_checks()]
        d.table([tr("document.col_check"), tr("document.col_status"), tr("document.col_value")], rows,
                caption=tr("document.tab_plaus_report"), widths_cm=[7.0, 1.8, 7.4], align=["left"] * 3, font_size=8)
        # ---- symbols
        d.heading(tr("document.appendix", letter="C", title=tr("document.app_symbols")), 1, numbered=False)
        lang = c.lang
        key = "de" if lang == "de" else "en"
        skip = ("I_[ON], Z_[E,ON]", "HES") if c.info.station_detail == "compact" or not c.stations else ()
        d.table([tr("document.col_symbol"), tr("document.col_meaning")],
                [(sym if isinstance(sym, str) else sym[key], txt[key]) for sym, txt in _SYMBOLS
                 if (isinstance(sym, str) or sym[key]) and (sym if isinstance(sym, str) else sym[key]) not in skip],
                caption=tr("document.tab_symbols"),
                widths_cm=[4.0, 12.0], align=["left", "left"], font_size=8.5)

    def report_checks(self) -> List[Tuple[str, str, str]]:
        c = self.c
        rows = []
        dev = []
        for n in c.names:
            t = c.tests[n]
            if _finite(t["Iw"]) and _finite(t["Iv"]) and _finite(t["IG"]) and t["IG"] > 0:
                dev.append(abs(math.hypot(t["Iw"], t["Iv"]) - t["IG"]) / t["IG"])
        if dev:
            rows.append((tr("document.chk_split"), "OK" if max(dev) < 0.01 else "WARN",
                         tr("document.chk_deviation", dev=_n(100 * max(dev), 2))))
        dev = []
        for n in c.names:
            t = c.tests[n]
            if _finite(t["IRES"]) and _finite(t["IG"]) and _finite(t["IH"]) and t["IRES"] > 0:
                lf = t["Ilf"] if _finite(t["Ilf"]) else 0.0
                dev.append(abs(math.sqrt(t["IG"] ** 2 + t["IH"] ** 2 + lf ** 2) - t["IRES"]) / t["IRES"])
        if dev:
            rows.append((tr("document.chk_rms"), "OK" if max(dev) < 0.03 else "WARN",
                         tr("document.chk_deviation", dev=_n(100 * max(dev), 2))))
        if c.pq_valid:
            stored = c.stored_factors
            if stored and any(abs(x - c.pq.factor) > 1e-9 for x in stored):
                rows.append((tr("document.chk_settings"), "WARN",
                             tr("document.chk_settings_diff", stored=" / ".join(_n(x) for x in stored),
                                used=_n(c.pq.factor))))
            elif stored:
                rows.append((tr("document.chk_settings"), "OK", tr("document.chk_settings_ok", k=_n(c.pq.factor))))
            else:
                rows.append((tr("document.chk_settings"), "OK",
                             tr("document.chk_settings_unscaled", k=_n(c.pq.factor))))
        mi = sorted({float(_stored_settings(c.results[n]).get("uncertainty_I") or 0.0) for n in c.names})
        mu = sorted({float(_stored_settings(c.results[n]).get("uncertainty_U") or 0.0) for n in c.names})
        same = all(abs(x - c.mI) < 1e-9 for x in mi) and all(abs(x - c.mU) < 1e-9 for x in mu)
        pct = lambda xs: " / ".join("%s %%" % _n(100 * x) for x in xs)
        rows.append((tr("document.chk_margin"), "OK" if same else "WARN",
                     tr("document.chk_margin_value", single=pct(mi), report=pct([c.mI]), single_u=pct(mu),
                        report_u=pct([c.mU]))))
        ut = sorted({float(x) for x in (_stored_settings(c.results[n]).get("u_tp") for n in c.names) if _finite(x)})
        ff = sorted({float(x) for x in (_stored_settings(c.results[n]).get("f_factor") for n in c.names)
                     if _finite(x)})
        if (ut or ff) and _finite(c.u_tp) and _finite(c.f_factor):
            same_t = all(abs(x - float(c.u_tp)) < 1e-9 for x in ut) and all(abs(x - float(c.f_factor)) < 1e-9
                                                                              for x in ff)
            rows.append((tr("document.chk_utp"), "OK" if same_t else "WARN",
                         tr("document.chk_utp_value", single=" / ".join(_g(x) for x in ut) or "–",
                            report=_g(c.u_tp), single_f=" / ".join(_g(x) for x in ff) or "–", report_f=_g(c.f_factor))))
        if c.reproduction:
            mx = max(c.reproduction.values())
            rows.append((tr("document.chk_reproduced"), "OK" if mx < 0.1 else "WARN",
                         tr("document.chk_deviation", dev=_small(mx))))
        if c.series is not None:
            try:
                ct = c.series.check_table()
                v = pd.to_numeric(ct["deviation_pct"], errors="coerce").abs().dropna()
                if len(v):
                    rows.append((tr("document.chk_model"), "OK" if v.max() < 5 else "WARN",
                                 tr("document.chk_deviation", dev=_n(v.max(), 2))))
            except Exception:
                pass
            attrs = getattr(c.loo, "attrs", {}) or {}
            lim_pct = float(attrs["limit_pct"]) if _finite(attrs.get("limit_pct")) else 15.0
            try:                                    # how well the model reproduces the currents per order
                stab = c.series.source_table()
                rp = pd.to_numeric(stab.get("residual_pct", pd.Series(dtype=float)), errors="coerce")
                if rp.notna().any():
                    j = rp.idxmax()
                    rows.append((tr("document.chk_residual"), "OK" if rp.max() <= lim_pct else "WARN",
                                 tr("document.chk_residual_value", dev=_n(rp.max(), 2), order=int(stab.loc[j, "order"]))))
            except Exception:
                pass
            if c.loo is not None and len(c.loo) and "deviation_pct" in c.loo:
                dv = pd.to_numeric(c.loo["deviation_pct"], errors="coerce").abs()
                if dv.notna().any():
                    j = dv.idxmax()
                    order = int(c.loo.loc[j, "order"]) if "order" in c.loo else "–"
                    w = c.loo["within_limit"] if "within_limit" in c.loo else None
                    bad = int(w.eq(False).sum()) if w is not None else 0
                    if w is not None and w.notna().any():
                        value = (tr("document.chk_loo_value", bad=bad, n=int(w.notna().sum()),
                                    lim=self.loo_limit_text(), dev=_n(dv.max(), 2), order=order) if bad else
                                 tr("document.chk_loo_value_ok", lim=self.loo_limit_text(), dev=_n(dv.max(), 2),
                                    order=order))
                        rows.append((tr("document.chk_loo"), "WARN" if bad else "OK", value))
        for n in c.stations:
            sv = c.results[n].tables.get("station_voltages")
            if sv is None or sv.empty or not (sv["role"] == "u_e").any() or not (sv["role"] == "u_t").any():
                continue
            v = pd.to_numeric(sv["RMS_V"], errors="coerce")
            ue, ut = v[sv["role"] == "u_e"].max(), v[sv["role"] == "u_t"].max()
            if _finite(ue) and _finite(ut):
                fu = float(c.f_factor) * float(c.u_tp) if _finite(c.u_tp) and _finite(c.f_factor) else float("nan")
                small = c.info.station_detail == "compact" and _finite(fu) and max(ue, ut) < 0.05 * fu
                rows.append((tr("document.chk_probe"), "OK" if ut <= ue else ("NOTE" if small else "WARN"),
                             tr("document.chk_probe_value", station=c.place(n), ut=_n(ut), ue=_n(ue))))
        if c.pq is not None:
            rows.append((tr("document.chk_pq"), "OK" if c.pq.valid else "WARN",
                         tr("document.valid") if c.pq.valid else tr("document.not_valid")))
        return rows


# =================================================================================================== public API
def campaign_report(results: Mapping[str, Any], path: Optional[str] = None, *,
                    info: Union[None, ReportInfo, Mapping[str, Any]] = None,
                    measurements: Optional[Sequence[Mapping[str, Any]]] = None,
                    campaign_folder: Optional[str] = None,
                    series=None, comparison=None, monte_carlo=None, loo: Optional[pd.DataFrame] = None,
                    Z_E: Optional[float] = None, IG: Union[str, float] = "max", limits: Sequence[float] = (),
                    u_tp: Optional[float] = None, f_factor: Optional[float] = None,
                    uncertainty_I: Optional[float] = None, uncertainty_U: Optional[float] = None,
                    scaling: Optional[Mapping[str, Any]] = None,
                    pq=None, pq_levels: Optional[Mapping[str, Mapping[int, float]]] = None,
                    pq_comparison: Optional[pd.DataFrame] = None, pq_scaling: Optional[pd.DataFrame] = None,
                    earthing: Optional[Mapping[str, Tuple[pd.DataFrame, pd.DataFrame]]] = None,
                    reduction: Optional[Mapping[str, Mapping[str, pd.DataFrame]]] = None,
                    photos: Optional[Mapping[str, Any]] = None, template: Optional[str] = None,
                    styles: Optional[Mapping[str, str]] = None, language: Optional[str] = None,
                    figures_folder: Optional[str] = None, raw: bool = True, pdf: bool = False,
                    bootstrap: int = 200) -> Dict[str, Any]:
    """Write the Word report of a campaign of staged earth-fault tests.

    The arguments are the objects of the campaign evaluation (as in the campaign notebook); everything except the
    results is optional – sections without data are left out.

    Parameters
    ----------
    results : mapping of str to Result or ResultFile
        Results in report order, e.g. ``{"Test 1 (substation 2.5 Ω)": result, …}``; the keys name the tests.
    path : str, optional
        The ``.docx`` file (default: ``<campaign>/<report>/<title>.docx``, without a campaign folder in the
        working directory).
    info : ReportInfo or dict, optional
        Title page, texts and parameters (see [`ReportInfo`][groundfault.document.ReportInfo]).
    measurements : sequence of dict, optional
        The measurement list of the campaign notebook (``name``, ``measurement``, ``Z``, optional ``location``):
        the additional impedance (None: not a test with an additional impedance) and the place of every test.
    campaign_folder : str, optional
        Folder of the campaign: raw recordings (waveforms, trends, fault times), media (photos), report and
        figures folder.
    series, comparison : SourceSeries, optional
        Equivalent sources of the result and – for comparison – the other source (measured/fitted).
    monte_carlo : dict, optional
        Result of [`groundfault.equivalent_source.monte_carlo`][groundfault.equivalent_source.monte_carlo] (default:
        computed with 1000 samples).
    loo : pandas.DataFrame, optional
        Leave-one-out table of the series (with your permissible deviations).
    Z_E : float, optional
        Earthing impedance in Ω at which I_RES is assessed.
    IG : {"max", "median"} or float
        Fundamental of the residual current at Z_E.
    limits : sequence of float
        Limits of the residual current in A (lines in the figures; the smallest is the limit of the network unless
        ``info.limit_A`` is set).
    u_tp, f_factor : float, optional
        Permissible touch voltage in V and the factor F (default: the values stored in the results).
    uncertainty_I, uncertainty_U : float, optional
        Margins m_I of the currents and m_U of the voltages (default: the settings stored in the results); m_U
        enters an order scaled with a measured voltage as k_m/(1 − m_U) – in the single tests and at Z_E alike.
    scaling : dict, optional
        Scaling of the equivalent sources to the PQ measurement (keyword arguments of
        [`groundfault.equivalent_source.ires_at_impedance`][groundfault.equivalent_source.ires_at_impedance]).
    pq : PQAssessment, optional
        The power-quality measurement of the busbar; a valid one also scales every single test (with its factor).
    pq_levels : mapping, optional
        Pre-fault voltage per order of every test (for the duration curve).
    pq_comparison, pq_scaling : pandas.DataFrame, optional
        Tests in the distribution of the PQ measurement; I_RES at Z_E by the methods of the scaling (default:
        [`PQAssessment.scaling_methods`][groundfault.pq.PQAssessment.scaling_methods] with the margins of the
        report).
    earthing, reduction : mapping, optional
        Per station: ``(long_df, ratio)`` of [`groundfault.earthing.compare`][groundfault.earthing.compare] and the
        tables of [`groundfault.earthing.reduction_factor`][groundfault.earthing.reduction_factor] /
        [`groundfault.earthing.check_lv_network`][groundfault.earthing.check_lv_network] (keys
        ``reduction_factor``, ``reduction_factor_eff``, ``lv_network_check``, ``lv_network_check_eff``).
    photos : mapping, optional
        Photos on site: ``{"setup_impedance": path, "recorder": path, "station": path}`` (``"station:<test>"`` per
        station; a dict ``{"path": …, "caption": …}`` sets the caption; False leaves the photo out; None: placeholder)
        and ``"extra"``: a list of ``{"path", "caption", "section"}`` (section ``"1.1"``). Relative paths are taken in
        the media folder of the campaign.
    template : str, optional
        Word file whose styles, headers, footers and logo the report uses (its content is replaced) – e.g. made with
        [`make_template`][groundfault.document.make_template].
    styles : mapping of str to str, optional
        Style names of the template for ``title``, ``subtitle``, ``toc_heading``, ``caption``, ``table``,
        ``list_bullet``, ``list_number``, ``list_paragraph``, ``body``, ``heading1`` … ``heading3``. The report sets
        the layout of a scientific report on them: the body text (default "Body Text", created if the template has
        none) justified, the caption centred in the size of the body text, lists indented, a blank line between a
        formula, figure or table and the text, hyphenation on.
    language : {"de", "en"}, optional
        Language of the report (default: the output language).
    figures_folder : str, optional
        Folder of the figures (PNG; default: ``<campaign>/<figures>/report`` or ``…/Bericht``).
    raw : bool
        Evaluate the raw recordings again for waveforms, trends and fault times (if found in the campaign folder).
    pdf : bool
        Also write a PDF with LibreOffice (if installed; see [`to_pdf`][groundfault.document.to_pdf]).
    bootstrap : int
        Samples of the bootstrap of the PQ measurement (0: none).

    Returns
    -------
    dict
        ``path`` (docx), ``pdf`` (or None), ``figures`` (files), ``placeholders`` (captions of the photo and text
        placeholders still to fill), ``hints``, ``basis`` (the test of the assessment) and ``kept``: an existing
        report at ``path`` that was edited after it was written (saved in Word, for instance) is not overwritten but
        renamed to ``<name>_<date>_<time>.docx`` – ``kept`` is that file (None: there was none, or it was unchanged
        and is overwritten).

    Raises
    ------
    ImportError
        If python-docx is not installed.
    ValueError
        Without results or with unknown entries of ``info`` or ``photos``.
    """
    if not docx_available():
        raise ImportError(tr("document.no_docx"))
    info_ = ReportInfo.from_value(info)
    for key in (photos or {}):
        if key not in PHOTO_SLOTS and key != "extra" and not str(key).startswith("station:"):
            raise ValueError(tr("document.photo_unknown", slot=key, known=", ".join(PHOTO_SLOTS)))
    lang = i18n.normalize_language(language or i18n.get_language())
    with i18n.language(lang):
        from . import layout as _layout
        lay = _layout.detect_layout(campaign_folder) if campaign_folder else _layout.get_layout(lang)
        title = info_.title or tr("document.default_title")
        if path is None:
            stem = _file_stem(title) + ("_" + _file_stem(os.path.basename(os.path.normpath(str(campaign_folder))))
                                        if campaign_folder else "")
            folder = os.path.join(str(campaign_folder), lay.report) if campaign_folder else os.getcwd()
            path = os.path.join(folder, stem + ".docx")
        if figures_folder is None:
            sub = "Bericht" if lay.language == "de" else "report"
            figures_folder = (os.path.join(str(campaign_folder), lay.figures, sub) if campaign_folder
                              else os.path.join(os.path.dirname(os.path.abspath(path)), "figures"))
        media = os.path.join(str(campaign_folder), lay.media) if campaign_folder else None
        ctx = _Campaign(results, measurements=measurements, info=info_, series=series, comparison=comparison,
                        monte_carlo=monte_carlo, loo=loo, Z_E=Z_E, IG=IG, limits=limits, u_tp=u_tp, f_factor=f_factor,
                        uncertainty_I=uncertainty_I, uncertainty_U=uncertainty_U, scaling=scaling, pq=pq,
                        pq_levels=pq_levels,
                        pq_comparison=pq_comparison, pq_scaling=pq_scaling, earthing=earthing, reduction=reduction,
                        campaign_folder=campaign_folder, raw=raw, bootstrap=bootstrap)
        doc = _Doc(template, styles, language=lang)
        doc.equations = bool(info_.equations)
        w = _Writer(doc, ctx, figures_folder, photos or {}, media)
        header = title + (" – " + plain(info_.network) if info_.network else "")
        doc.header_footer(header)
        doc.properties(title, list(info_.authors), info_.network)
        w.title_page()
        if info_.key_results:
            w.key_results()
        doc.toc(tr("document.toc"))
        doc.page_break()
        w.chapter_setup()
        w.chapter_results()
        w.chapter_assessment()
        if info_.appendix:
            w.appendices()
        w.fill_sections()
        kept = _keep_edited(path)
        doc.save(path)
        out = dict(path=path, pdf=None, figures=w.figure_files, placeholders=w.placeholders, hints=ctx.hints,
                   basis=ctx.basis, kept=kept)
        if kept:
            out["hints"].append(tr("document.kept_edited", name=os.path.basename(kept)))
        if pdf:
            out["pdf"] = to_pdf(path)
            if out["pdf"] is None:
                out["hints"].append(tr("document.pdf_none"))
    return out


def make_template(source: str, target: str) -> str:
    """Word template from an earlier report: the content is removed, the styles, page layout, headers, footers and
    the logo stay – the template of your organisation for [`campaign_report`][groundfault.document.campaign_report]
    (``template=``).

    Parameters
    ----------
    source : str
        An existing ``.docx`` report (it is not changed).
    target : str
        The template to write (``.docx``).

    Returns
    -------
    str
        ``target``.
    """
    if not docx_available():
        raise ImportError(tr("document.no_docx"))
    import docx
    if not os.path.isfile(source):
        raise FileNotFoundError(tr("document.template_missing", path=source))
    d = docx.Document(source)
    _clear_body(d)
    cp = d.core_properties
    cp.title = cp.subject = cp.keywords = cp.comments = ""
    os.makedirs(os.path.dirname(os.path.abspath(target)), exist_ok=True)
    d.save(target)
    return target


def _soffice() -> Optional[str]:
    for name in ("soffice", "libreoffice"):
        p = shutil.which(name)
        if p:
            return p
    for p in ("/Applications/LibreOffice.app/Contents/MacOS/soffice",
              r"C:\Program Files\LibreOffice\program\soffice.exe"):
        if os.path.isfile(p):
            return p
    return None


def to_pdf(path: str, folder: Optional[str] = None, timeout: float = 180.0) -> Optional[str]:
    """PDF of a Word file with LibreOffice (``soffice --headless``), next to it or in ``folder``.

    Returns the path of the PDF, or None if LibreOffice is not installed or the conversion failed (a warning names
    the reason). Word: *Save as → PDF* gives the same with the contents and field numbers updated.
    """
    exe = _soffice()
    if exe is None:
        warnings.warn(tr("document.pdf_none"), stacklevel=2)
        return None
    out = folder or os.path.dirname(os.path.abspath(path))
    profile = tempfile.mkdtemp(prefix="gf_lo_")
    try:
        subprocess.run([exe, "-env:UserInstallation=" + pathlib.Path(profile).as_uri(), "--headless",
                        "--convert-to", "pdf", "--outdir", out, path], check=True, capture_output=True,
                       timeout=timeout)
    except Exception as ex:
        warnings.warn(tr("document.pdf_failed", error=ex), stacklevel=2)
        return None
    finally:
        shutil.rmtree(profile, ignore_errors=True)
    pdf = os.path.join(out, os.path.splitext(os.path.basename(path))[0] + ".pdf")
    return pdf if os.path.isfile(pdf) else None
