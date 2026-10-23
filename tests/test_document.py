# -*- coding: utf-8 -*-
"""Word report of a campaign (groundfault.document): structure, numbers from the results, placeholders and photos,
languages, templates, sections without data – and the building blocks it uses (scaling of a single test with another
PQ measurement, bootstrap of the PQ quantiles, assessment values with the voltage margin, methods of the scaling,
touch-voltage cases)."""
import importlib.util
import math
import os
import re
import zipfile

import numpy as np
import pandas as pd
import pytest

from groundfault import demo, document, equivalent_source, evaluation, i18n, pq, report

docx = pytest.importorskip("docx")
HERE = os.path.dirname(os.path.abspath(__file__))

#: user parameters of the demo report – invented for the test, not defaults of the package
PARAMETERS = dict(LIMITS=(60.0, 90.0), U_TP=80.0, F_FACTOR=2.0, UNCERTAINTY_I=0.05, LOO_LIMIT_PCT=15.0)


def _generator():
    spec = importlib.util.spec_from_file_location("make_notebooks", os.path.join(os.path.dirname(HERE), "tools",
                                                                                 "make_notebooks.py"))
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _run(cells, namespace):
    for cell in cells:
        exec("\n".join(line for line in cell.splitlines() if not line.startswith("%")), namespace)


@pytest.fixture(scope="module")
def campaign(tmp_path_factory):
    """The German example campaign notebook on the demo campaign, with the parameters above (namespace); section 8
    writes the report."""
    pytest.importorskip("IPython")
    import matplotlib
    folder = tmp_path_factory.mktemp("report_de")
    cells = [c.source for c in _generator().campaign_notebook("de", True).cells if c.cell_type == "code"]
    cwd = os.getcwd()
    namespace = {"__name__": "notebook"}
    os.chdir(str(folder))
    try:
        with matplotlib.rc_context():
            _run(cells[:2], namespace)                                     # setup and configuration
            namespace.update(display=lambda *a, **k: None, **PARAMETERS)
            _run(cells[2:], namespace)
    finally:
        os.chdir(cwd)
        i18n.set_language("de")
    return namespace


def _kwargs(ns, **extra):
    """Arguments of document.campaign_report from the notebook objects (as section 8)."""
    kw = dict(info=ns["REPORT"], measurements=ns["MEASUREMENTS"], campaign_folder=str(ns["CAMPAIGN_FOLDER"]),
              series=ns["MAIN"], comparison=ns["OTHER"], monte_carlo=ns["MC"], loo=ns["LOO"], Z_E=ns["Z_E"],
              IG=ns["IG_CHOICE"], limits=ns["LIMITS"], u_tp=ns["U_TP"], f_factor=ns["F_FACTOR"],
              uncertainty_I=ns["UNCERTAINTY_I"], uncertainty_U=ns["UNCERTAINTY_U"], scaling=ns["SCALING"], pq=ns["PQ"],
              pq_levels=ns["PQ_LEVELS"], pq_comparison=ns["PQ_COMPARISON"], earthing=ns["COMPARISONS"],
              reduction=ns["REDUCTION"])
    kw.update(extra)
    return kw


def _headings(path):
    d = docx.Document(str(path))
    return [p.text.strip() for p in d.paragraphs if p.style.name.startswith("Heading") and p.text.strip()]


def _text(path):
    """All text of the document, its headers and footers."""
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    for s in d.sections:
        for hf in (s.header, s.footer, s.first_page_header, s.first_page_footer):
            parts.extend(p.text for p in hf.paragraphs)
    return "\n".join(parts)


def _order_violations(path):
    """Children of w:pPr, w:rPr and w:tcPr out of the schema order (document, styles, footnotes, numbering) and of
    w:settings – Word expects the order of the schema."""
    from docx.oxml.ns import qn
    from lxml import etree
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    bad = []
    with zipfile.ZipFile(str(path)) as z:
        parts = ["word/document.xml", "word/styles.xml", "word/numbering.xml"] + _note_parts(z)
        for part in parts:
            root = etree.fromstring(z.read(part))
            for tag, order in (("pPr", document._PPR_ORDER), ("tcPr", document._TCPR_ORDER),
                               ("rPr", document._RPR_ORDER)):
                for el in root.iter(w + tag):
                    idx = [order.index(c.tag[len(w):]) for c in el if c.tag.startswith(w) and c.tag[len(w):] in order]
                    if idx != sorted(idx):
                        bad.append((part, tag, [c.tag.split("}")[1] for c in el]))
        tags = [qn(x) for x in document._SETTINGS_TAGS]
        idx = [tags.index(c.tag) for c in etree.fromstring(z.read("word/settings.xml")) if c.tag in tags]
        if idx != sorted(idx):
            bad.append(("word/settings.xml", [document._SETTINGS_TAGS[i] for i in idx]))
    return bad


def _note_parts(z):
    return [n for n in z.namelist() if re.fullmatch(r"word/footnotes\d*\.xml", n)]


def _notes(path):
    """The footnotes of a document in the order of their ids: [(id, text)] (without the separators)."""
    from lxml import etree
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(str(path)) as z:
        parts = _note_parts(z)
        assert len(parts) <= 1
        if not parts:
            return []
        root = etree.fromstring(z.read(parts[0]))
    return [(int(f.get(w + "id")), "".join(t.text or "" for t in f.iter(w + "t")).strip())
            for f in root.iter(w + "footnote") if f.get(w + "type") in (None, "normal")]


def _note_refs(path):
    """The ids of the footnote marks in the order of the text."""
    with zipfile.ZipFile(str(path)) as z:
        body = z.read("word/document.xml").decode("utf-8")
    return [int(x) for x in re.findall(r'<w:footnoteReference w:id="(-?\d+)"', body)]


def _table_with(path, first_cell):
    d = docx.Document(str(path))
    for t in d.tables:
        if t.rows and t.rows[0].cells[0].text.strip() == first_cell:
            return [[c.text.strip() for c in r.cells] for r in t.rows]
    raise AssertionError(first_cell)


def _n(x, digits=3):
    return report.number(x, digits, lang="de")


# --------------------------------------------------------------------------------------------------------- report
def test_report_of_the_notebook_keeps_the_structure_of_the_reports(campaign):
    """Section 8 writes the report into the report folder: title page, key results, contents, chapters 1–3 with the
    sections of the reports so far plus equivalent sources and uncertainties, and the appendices – numbered headings,
    captions as fields, a table of contents."""
    path = campaign["REPORT_FILE"]
    assert path.exists() and path.parent.name == "4_Bericht" and path.name.endswith("_2026-01-15_Demo.docx")
    h = _headings(path)
    expected = ["Kurzfassung", "Aufbau und Durchführung", "Untersuchungsgegenstand und Messaufbau", "Durchführung",
                "Ergebnisse", "Allgemeines zur Messung", "Messunsicherheiten", "Messwert-Aufbereitung", "Fehlerstrom",
                "Netzspannungen", "Power-Quality-Messung", "Erdschlussmessung", "Bewertung",
                "Aufbereitung Erdschlussreststrom", "Ersatzquellenverfahren und Prognose", "Unsicherheiten",
                "Zusammenfassung"]
    positions = [h.index(x) for x in expected]
    assert positions == sorted(positions), h
    assert any(x.startswith("Erdungsspannung – ") for x in h) and any(x.startswith("Aufbereitung Erdung – ") for x in h)
    assert [x for x in h if x.startswith("Anhang")] == ["Anhang A – Messdaten je Messung",
                                                         "Anhang B – Plausibilitätsprüfung", "Anhang C – Formelzeichen"]
    with zipfile.ZipFile(str(path)) as z:
        body = z.read("word/document.xml").decode("utf-8")
        settings = z.read("word/settings.xml").decode("utf-8")
    assert body.count("SEQ Abbildung") >= 10 and body.count("SEQ Tabelle") >= 10 and "TOC \\o" in body
    assert "updateFields" in settings                                          # Word renews the fields on opening
    assert _order_violations(path) == []                                       # elements in the order of the schema
    out = campaign["REPORT_OUT"]
    assert out["basis"] in {m["name"] for m in campaign["MEASUREMENTS"]} and len(out["figures"]) >= 10
    assert all(os.path.isfile(f) for f in out["figures"])


def test_numbers_come_from_the_results(campaign):
    """The key results are those of the result files and of the equivalent sources with the margins of the
    configuration – recomputed here from the objects of the notebook."""
    ns = campaign
    path = ns["REPORT_FILE"]
    rows = {r[0]: r for r in _table_with(path, "Größe")}
    subs = [m["name"] for m in ns["MEASUREMENTS"] if m["Z"] is not None]
    largest = max(subs, key=lambda n: ns["RESULTS"][n].values["IRES"])
    ires = ns["RESULTS"][largest].values["IRES"]
    row = next(r for k, r in rows.items() if k.endswith(largest) and "gemessen" in k)
    assert row[1] == _n(ires) + " A"
    margin = next(r for k, r in rows.items() if k.startswith("… mit Messunsicherheitszuschlag"))
    assert margin[1] == _n(ires * 1.05) + " A"
    r = equivalent_source.ires_at_impedance(ns["MAIN"], ns["Z_E"], IG=ns["IG_CHOICE"], uncertainty_I=0.05,
                                            **ns["SCALING"])
    ze = next(r_ for k, r_ in rows.items() if k.startswith("IRES,Bew an ZE"))
    assert ze[1] == _n(r["IRES_scaled_assessed"]) + " A"
    assert ze[3] == ("eingehalten" if r["IRES_scaled_assessed"] <= 60.0 else "überschritten")
    assert "maßgeblich" in ze[0]                                     # with equivalent sources the value at Z_E
    text = _text(path)
    assert "I_RES" not in text and "Z_E" not in text                 # symbols as subscripts, not as plain text
    assert re.search(r"größte Abweichung (< 0,001|0,0\d*) %", text)  # raw data evaluated again: as the result files


def test_placeholders_and_photos(campaign, tmp_path):
    """Without photos the report has a framed placeholder per photo slot and a highlighted one for the
    recommendation; a photo named in PHOTOS (in the media folder) is inserted instead of its placeholder."""
    out = campaign["REPORT_OUT"]
    assert len(out["placeholders"]) == 4 and any("Empfehlung" in p for p in out["placeholders"])
    media = campaign["CAMPAIGN_FOLDER"] / "5_Medien"
    media.mkdir(exist_ok=True)
    from PIL import Image
    Image.new("RGB", (400, 300), (200, 120, 40)).save(str(media / "rekorder.png"))
    path = tmp_path / "with_photo.docx"
    res = document.campaign_report(campaign["RESULTS"], str(path), raw=False, bootstrap=0,
                                   **_kwargs(campaign, photos={"recorder": "rekorder.png"},
                                             info=dict(campaign["REPORT"], recommendation="Erste Zeile.\n\nZweite.")))
    assert len(res["placeholders"]) == 2 and not any("Empfehlung" in p for p in res["placeholders"])
    with zipfile.ZipFile(str(path)) as z:
        assert sum(n.startswith("word/media/") for n in z.namelist()) >= len(res["figures"]) + 1
    text = _text(path)
    assert "Erste Zeile." in text and "Zweite." in text
    with pytest.raises(ValueError, match="photo|Foto"):
        document.campaign_report(campaign["RESULTS"], str(tmp_path / "x.docx"), photos={"logo": "x.png"})


def test_english_report_with_a_template(campaign, tmp_path):
    """English texts; the template of an organisation keeps its styles, headers and footers, its body is replaced
    and its style names are mapped with ``styles``."""
    tpl_src = docx.Document()
    tpl_src.styles.add_style("Org Title", 1).font.size = docx.shared.Pt(24)
    tpl_src.sections[0].header.paragraphs[0].text = "Header of the organisation"
    tpl_src.add_paragraph("Old content that must not stay")
    src = tmp_path / "old_report.docx"
    tpl_src.save(str(src))
    template = document.make_template(str(src), str(tmp_path / "template.docx"))
    assert "Old content" not in _text(template)
    path = tmp_path / "report_en.docx"
    info = dict(network="20 kV network of a test substation", authors=["A. Author"])
    res = document.campaign_report(campaign["RESULTS"], str(path), language="en", template=template,
                                   styles={"title": "Org Title"}, raw=False, bootstrap=0,
                                   **_kwargs(campaign, info=info))
    d = docx.Document(str(path))
    titles = [p for p in d.paragraphs if p.style.name == "Org Title"]
    assert titles and titles[0].text == "Evaluation of staged earth-fault tests"
    h = _headings(path)
    assert "Assessment" in h and "Equivalent-source method and forecast" in h and "Summary" in h
    assert not {"Bewertung", "Zusammenfassung", "Ergebnisse"} & set(h)
    captions = [p.text for p in d.paragraphs if p.style.name == "Caption"]
    assert captions and all(c.startswith(("Figure ", "Table ")) for c in captions), captions[:3]
    text = _text(path)
    assert "Header of the organisation" in text and "Old content" not in text
    assert res["path"] == str(path)
    assert _order_violations(path) == []


def test_sections_without_data_are_left_out(campaign, tmp_path):
    """Only the substation tests, no PQ measurement, no equivalent sources: no station sections, the PQ section says
    that there was no measurement (the scaling is left out), no figure that needs them – the report is complete."""
    ns = campaign
    subs = {m["name"]: ns["RESULTS"][m["name"]] for m in ns["MEASUREMENTS"] if m["Z"] is not None}
    path = tmp_path / "minimal.docx"
    res = document.campaign_report(subs, str(path), measurements=ns["MEASUREMENTS"], raw=False, bootstrap=0,
                                   info=dict(key_results=False))
    h = _headings(path)
    assert "Kurzfassung" not in h and "Ersatzquellenverfahren und Prognose" in h
    assert not any(x.startswith(("Erdungsspannung – ", "Aufbereitung Erdung – ")) for x in h)
    text = _text(path)
    assert "Power-Quality-Messung der Sammelschiene lag für diese Auswertung nicht vor" in text
    assert "Schritt 2 entfällt" in text and "Für diese Auswertung wurden keine Ersatzquellen bestimmt" in text
    assert not any("pq_" in os.path.basename(f) or "source_fit" in f for f in res["figures"])
    assert "Zusammenfassung" in h and "Anhang C – Formelzeichen" in h
    assert len(res["placeholders"]) == 3                             # setup, recorder, recommendation
    with pytest.raises(ValueError):
        document.campaign_report({}, str(tmp_path / "none.docx"))
    with pytest.raises(ValueError, match="unknown|unbekannt"):
        document.ReportInfo.from_value(dict(titel="x"))


def test_a_report_edited_in_word_is_kept(campaign, tmp_path):
    """Writing the report again overwrites an unchanged one; one saved in Word since (revision > 1) is kept under a
    new name with its modification time – nothing filled in by hand is lost."""
    ns = campaign
    subs = {m["name"]: ns["RESULTS"][m["name"]] for m in ns["MEASUREMENTS"] if m["Z"] is not None}
    path = tmp_path / "report.docx"
    kw = dict(measurements=ns["MEASUREMENTS"], raw=False, bootstrap=0, info=dict(key_results=False))
    docx_files = lambda: sorted(f for f in os.listdir(str(tmp_path)) if f.endswith(".docx"))
    assert document.campaign_report(subs, str(path), **kw)["kept"] is None
    assert docx.Document(str(path)).core_properties.revision == 1
    assert document.campaign_report(subs, str(path), **kw)["kept"] is None     # unchanged → overwritten
    assert docx_files() == ["report.docx"]
    d = docx.Document(str(path))                                               # edited and saved in Word
    d.add_paragraph("Ergänzung von Hand")
    d.core_properties.revision = 2
    d.save(str(path))
    res = document.campaign_report(subs, str(path), **kw)
    kept = res["kept"]
    assert kept and re.fullmatch(r"report_\d{4}-\d\d-\d\d_\d{6}\.docx", os.path.basename(kept))
    assert docx_files() == sorted(["report.docx", os.path.basename(kept)])
    assert "Ergänzung von Hand" in _text(kept) and "Ergänzung von Hand" not in _text(path)
    assert any(os.path.basename(kept) in h for h in res["hints"])


def test_no_organisation_in_the_report_without_configuration(campaign, tmp_path):
    """Without names in the configuration the report names no organisation, address or person – everything of that
    kind comes from the user's REPORT and template."""
    path = tmp_path / "neutral.docx"
    document.campaign_report(campaign["RESULTS"], str(path), raw=False, bootstrap=0, **_kwargs(campaign, info=None))
    with zipfile.ZipFile(str(path)) as z:
        xml = "\n".join(z.read(n).decode("utf-8", "replace") for n in z.namelist() if n.endswith(".xml"))
    for word in ("GmbH", " AG ", "Ltd", "Inc.", "@", "Straße", "Street", "python-docx"):
        assert word not in xml, word
    assert "Demo-Bearbeitung" not in xml                             # the author only with REPORT


# ------------------------------------------------------------------------------------------------- building blocks
def test_rescale_harmonics_reproduces_the_stored_scaling(demo_results):
    """The scaling of a single test with the quantiles and the factor of its evaluation gives the stored values (the
    margins are taken from the result); another factor scales the orders with their own voltage, not the zero
    sequence."""
    from groundfault import Intervals, Settings, evaluate
    res0, truth, _, rec = demo_results[0.0]
    b = truth["intervals"]
    iv = Intervals(tuple(b["steady_state"]), tuple(b["pre_fault"]), tuple(b["post_fault"]))
    u95 = {5: 400.0, 7: 120.0, 11: 60.0, 13: 40.0}
    res = evaluate(rec, iv, Settings(pq_u95=u95, pq_u95_reference="seq", pq_safety_factor=1.2, uncertainty_I=0.05,
                                     uncertainty_U=0.05), z_fault_ohm=0.0)
    s = evaluation.rescale_harmonics(res, u95=u95, u95_reference="seq", safety_factor=1.2)
    assert s["uncertainty_I"] == 0.05 and s["uncertainty_U"] == 0.05
    assert s["IH_scaled"] == pytest.approx(res.values["IH_scaled"], rel=1e-9)
    assert s["IRES_scaled"] == pytest.approx(res.values["IRES_scaled"], rel=1e-9)
    stored = res.tables["assessment_quantities"].set_index("quantity")["assessment"]
    assert s["IH_scaled_assessed"] == pytest.approx(stored["IH_scaled"], rel=1e-9)
    assert s["IRES_scaled_assessed"] == pytest.approx(stored["IRMS_forecast"], rel=1e-9)
    s2 = evaluation.rescale_harmonics(res, u95=u95, safety_factor=2.4)
    t1, t2 = s["table"].set_index("order"), s2["table"].set_index("order")
    assert t2.loc[5, "I_scaled_A"] == pytest.approx(2 * t1.loc[5, "I_scaled_A"])
    assert t2.loc[3, "I_scaled_A"] == pytest.approx(t1.loc[3, "I_scaled_A"]) and t1.loc[3, "rule"] == "as measured"
    ka = 1 / (1 - 0.05)                                              # assessment: k_m/(1 − m_U), every current (1 + m_I)
    own = t1[t1["rule"].isin(["own voltage", "THD"])]
    rest = t1[~t1["rule"].isin(["own voltage", "THD"])]
    ih_a = 1.05 * math.sqrt(float(((own["I_scaled_A"] * ka) ** 2).sum() + (rest["I_scaled_A"] ** 2).sum())
                            + s["IH_other"] ** 2)
    assert s["IH_scaled_assessed"] == pytest.approx(ih_a, rel=1e-9)


def test_equivalent_sources_take_the_voltage_margin_as_the_single_evaluation(campaign):
    """At Z_E the assessment value has the same margins as in a single test: (1 + m_I) on every current and
    k_m/(1 − m_U) for the orders scaled with a measured voltage; m_U = 0 gives the former value."""
    ns = campaign
    sc = ns["SCALING"]
    r0 = equivalent_source.ires_at_impedance(ns["MAIN"], ns["Z_E"], IG="max", uncertainty_I=0.05, **sc)
    r1 = equivalent_source.ires_at_impedance(ns["MAIN"], ns["Z_E"], IG="max", uncertainty_I=0.05,
                                             uncertainty_U=0.05, **sc)
    assert r0["IRES_scaled_assessed"] == pytest.approx(1.05 * r0["IRES_scaled"])
    tab = r1["table"]
    scaled = tab["rule"].isin(["own voltage", "THD"])
    s2 = float(((tab.loc[scaled, "I_scaled_A"] / 0.95) ** 2).sum() + (tab.loc[~scaled, "I_scaled_A"] ** 2).sum())
    s2 += tab.attrs.get("other_A", 0.0) ** 2
    assert r1["IRES_scaled_assessed"] == pytest.approx(1.05 * math.sqrt(r1["IG"] ** 2 + s2), rel=1e-9)
    assert r1["IRES_scaled_assessed"] > r0["IRES_scaled_assessed"] and r1["IRES_scaled"] == r0["IRES_scaled"]
    with pytest.raises(ValueError):
        equivalent_source.ires_at_impedance(ns["MAIN"], ns["Z_E"], uncertainty_U=1.0)
    mc = equivalent_source.monte_carlo(ns["MAIN"], ns["Z_E"], n=50, IG="median")
    assert mc["IG"] == pytest.approx(float(np.nanmedian(ns["MAIN"].IG_measured)))
    loo = ns["MAIN"].loo_table(15.0, 0.5)
    assert loo.attrs == {"limit_pct": 15.0, "limit_A": 0.5}


def test_scaling_methods_of_the_pq_measurement(campaign):
    """Methods of the scaling at Z_E: without, per order (the result), the time series and the THD factor; the value
    per order is that of ires_at_impedance, the time series not above it."""
    ns = campaign
    t = ns["PQ"].scaling_methods(ns["MAIN"], ns["Z_E"], IG=ns["IG_CHOICE"], uncertainty_I=0.05, scaling=ns["SCALING"])
    assert t.attrs["keys"] == ["none", "order", "time", "thd"]
    r = equivalent_source.ires_at_impedance(ns["MAIN"], ns["Z_E"], IG=ns["IG_CHOICE"], uncertainty_I=0.05,
                                            **ns["SCALING"])
    assert t["IRES_A"].iloc[1] == pytest.approx(r["IRES_scaled"]) and t["IRES_assessed_A"].iloc[1] == pytest.approx(
        r["IRES_scaled_assessed"])
    assert t["IRES_A"].iloc[2] <= t["IRES_A"].iloc[1] * (1 + 1e-6) and t["IRES_A"].iloc[0] < t["IRES_A"].iloc[1]
    assert list(t["method"])[0] == "ohne Skalierung"
    assert ns["PQ_SCALING"]["IRES_A"].tolist() == pytest.approx(t["IRES_A"].tolist())     # the notebook: the same


def test_quantile_bootstrap(tmp_path):
    """Bootstrap over days: one row per sample, the quantiles spread around the quantile of the measurement,
    reproducible with the seed; invalid arguments raise."""
    from groundfault import demo
    p = demo.make_pq_export(str(tmp_path / "pq.csv"), days=21.0)
    a = pq.assess(pq.read(str(tmp_path / "pq.csv")), quantile=0.95)
    b = a.quantile_bootstrap(n=40, seed=1)
    assert len(b) == 40 and 5 in b.columns and "THD_LL_pct" in b.columns
    q5 = a.u95()[5]
    assert b[5].min() <= q5 * 1.02 and b[5].max() >= q5 * 0.98 and b[5].std() > 0
    pd.testing.assert_frame_equal(b, a.quantile_bootstrap(n=40, seed=1))
    with pytest.raises(ValueError):
        a.quantile_bootstrap(n=0)
    with pytest.raises(ValueError):
        a.quantile_bootstrap(block_days=0)
    assert p is not None


def _station_result(campaign):
    return next(e for e in campaign["RESULTS"].values() if not e.tables.get("station_frequency", pd.DataFrame()).empty)


def test_touch_voltage_cases(campaign):
    """Cases 1 and 2 with the limit current and Z_E of the earthing measurement, 3 per order, 4 measured and 4b
    extrapolated; U_T = U_E/F and the share of U_Tp."""
    st = _station_result(campaign)
    earthing = pd.DataFrame(dict(f_Hz=[50.0, 150.0, 250.0], Z_E_injection_Ohm=[0.5, 0.8, 1.1]))
    c = document.touch_voltage_cases(st, currents={1: 40.0, 3: 10.0, 5: 20.0}, limit_A=60.0, U_Tp=80.0, F=2.0,
                                     r_flat=0.5, reference_frequency_Hz=150.0, earthing=earthing).set_index("case")
    assert c.loc["1", "U_E_V"] == pytest.approx(60.0 * 0.8) and c.loc["1", "U_T_V"] == pytest.approx(24.0)
    assert c.loc["2", "U_E_V"] == pytest.approx(0.5 * 60.0 * 0.8)
    assert c.loc["1", "share_pct"] == pytest.approx(100 * 24.0 / 80.0)
    m = st.tables["station_metrics"]
    i_f, ue = float(m["IF_A"].iloc[0]), float(m["UE_max_V"].iloc[0])
    assert c.loc["4", "U_E_V"] == pytest.approx(ue) and c.loc["4b", "U_E_V"] == pytest.approx(ue * 60.0 / i_f)
    assert c.loc["3", "I_A"] == pytest.approx(math.sqrt(40 ** 2 + 10 ** 2 + 20 ** 2))
    assert c.attrs["reference_method"] == "measured" and c.attrs["reference_frequency_Hz"] == 150.0
    between = document.touch_voltage_cases(st, limit_A=60.0, U_Tp=80.0, F=2.0, reference_frequency_Hz=200.0,
                                           earthing=earthing)          # not measured at 200 Hz: interpolated
    assert between.attrs["reference_method"] == "interpolated" and between.attrs["reference_bracket_Hz"] == (150.0, 250.0)
    assert between.set_index("case").loc["1", "Z_E_Ohm"] == pytest.approx(0.95)
    outside = document.touch_voltage_cases(st, limit_A=60.0, U_Tp=80.0, F=2.0, reference_frequency_Hz=350.0,
                                           earthing=earthing)          # beyond the measured range: the nearest
    assert outside.attrs["reference_method"] == "nearest" and outside.attrs["reference_frequency_Hz"] == 250.0
    without = document.touch_voltage_cases(st, limit_A=60.0, U_Tp=80.0, F=2.0)
    assert not without["case"].isin(["1", "2"]).any()                # no earthing measurement: no cases 1 and 2
    fig = report.touch_voltage_cases(c.reset_index(), 80.0)
    ax = fig.axes[0]
    assert len(ax.patches) == int(np.isfinite(c["U_T_V"]).sum()) and ax.get_lines()[0].get_ydata()[0] == 80.0
    import matplotlib.pyplot as plt
    plt.close(fig)


def test_helpers_of_the_texts():
    """Small deviations, quantiles, states and plain symbols in the language of the report."""
    from groundfault.document import _g, _markup_symbols, _small, _status
    with i18n.language("de"):
        assert _small(5e-12) == "< 0,001" and _small(0.25) == "0,25" and _g(0.95) == "0,95"
        assert _status("WARN") == _status("WARNING") == "Warnung" and _status("CORRECTED") == "korrigiert"
    with i18n.language("en"):
        assert _small(5e-12) == "< 0.001" and _status("NOTE") == "note"
    assert _markup_symbols("I_H sinkt mit Z_F; ISp22_27, UE_5m") == "I_[H] sinkt mit Z_[F]; ISp22_27, UE_5m"
    from groundfault.document import _updated_comment
    with i18n.language("de"):                         # a comment of an older result file in the current wording
        new = _updated_comment("erdschlussfreier Betrieb: Strom nahezu 0 A (DC-Offset/Zangen prüfen)")
        assert new.startswith("im erdschlussfreien Betrieb erwartet ≈ 0 A") and _updated_comment("x") == "x"


def test_probe_above_the_earth_potential_rise_is_reported(campaign, tmp_path):
    """A touch-voltage probe that reads more than the earth-potential probe (U_T > U_E, physically inconsistent) is
    named in the text and flagged in the checks of the report; the key results give the section of every value."""
    ns = campaign
    station = next(m["name"] for m in ns["MEASUREMENTS"] if m["Z"] is None)
    res = ns["RESULTS"][station]
    sv = res.tables["station_voltages"]
    ue = sv[sv["role"] == "u_e"].iloc[0].copy()
    ut = ue.copy()
    ut["channel"], ut["role"], ut["RMS_V"] = "UT_test", "u_t", 1.3 * float(ue["RMS_V"])
    patched = pd.concat([sv[sv["role"] != "u_t"], ut.to_frame().T], ignore_index=True)
    old = res.tables["station_voltages"]
    try:
        res.tables["station_voltages"] = patched
        path = tmp_path / "probe.docx"
        document.campaign_report(ns["RESULTS"], str(path), raw=False, bootstrap=0, **_kwargs(ns))
    finally:
        res.tables["station_voltages"] = old
    text = _text(path)
    assert "Die Spannung gegen den Spieß UT_test" in text and "Sondenmaximum" in text
    rows_all = [[c.text.strip() for c in r.cells] for t in docx.Document(str(path)).tables for r in t.rows]
    check = next(r for r in rows_all if r[0].startswith("Berührungsspannung (Spieß) ≤ Erdungsspannung"))
    assert check[1] == "Warnung"
    rows = {r[0]: r for r in _table_with(path, "Größe")}
    assert rows["Erdschlussversuche"][4] == "1.1"
    assert next(r for k, r in rows.items() if k.startswith("IRES,Bew an ZE"))[4] == "3.2"


def test_limits_with_conditions_controller_and_compact_station(campaign, tmp_path):
    """Limits with their conditions (I_RES,ass at the Z_E of every case, U_E = r·I·Z_E against F·U_Tp), the display of
    the coil controller (I_CE = I_L/(1 + v)), a compact station, own names of quantities, no transducer column,
    where Z_E comes from – and the checks of the new keys."""
    ns = campaign
    cases = [dict(limit_A=60.0, Z_E=2.5), dict(limit_A=90.0, Z_E=1.7, condition="alle Z_E < 1,7 Ω"),
             dict(limit_A=120.0, Z_E=2.5, r=0.5, condition="volle Verkabelung")]
    info = dict(ns["REPORT"], limit_cases=cases, Z_E_basis="Dimensionierungswert", station_detail="compact",
                coil_controller=dict(coils_A={"C1": 300.0, "C2": 320.0}, detuning_pct=4.0, active_A=20.0,
                                     when="vor Messung 1"),
                quantity_labels={"i_shield": "Schirmstrom eines Abgangs"}, r_flat=0.5, reference_frequency_Hz=50.0)
    path = tmp_path / "limits.docx"
    kw = _kwargs(ns, info=info)
    res = document.campaign_report(ns["RESULTS"], str(path), raw=False, bootstrap=0, **kw)
    text = _text(path)
    rows = _table_with(path, "Grenzwert in A")
    assert [r[0] for r in rows[1:]] == ["60", "90", "120"] and rows[1][1] == "Standard (gilt)"
    sc = ns["SCALING"]
    for row, case in zip(rows[1:], cases):
        r = equivalent_source.ires_at_impedance(ns["MAIN"], case["Z_E"], IG=ns["IG_CHOICE"], uncertainty_I=kw["uncertainty_I"],
                                                uncertainty_U=kw["uncertainty_U"], **sc)
        i = r["IRES_scaled_assessed"]
        assert row[4] == _n(i) and row[6] == _n(case.get("r", 1.0) * i * case["Z_E"])
        assert row[5] == ("eingehalten" if i <= case["limit_A"] else "überschritten")
    assert "Es gilt der Grenzwert 60 A" in text and "Die übrigen Grenzwerte gelten nur, wenn ihre Bedingung" in text
    es_line = [p for p in text.split("\n") if p.startswith("Mit den Ersatzquellen beträgt der Reststrom")]
    assert es_line and "Grenzwert" not in es_line[0]                      # the verdict once, with its condition
    assert "alle ZE < 1,7 Ω" in text                                      # Z_E of the condition as a symbol
    key = {r[0]: r for r in _table_with(path, "Größe")}
    ue = next(r for k, r in key.items() if k.startswith("UE = r·IRES,Bew·ZE an ZE = 2,50 Ω"))
    assert ue[2] == "160 V"
    assert any(k.startswith("IRES,Bew an ZE = 1,70 Ω, wenn alle ZE < 1,7 Ω") for k in key)
    assert "Anzeige des Spulenreglers (vor Messung 1): Spulen C1 300 A, C2 320 A (zusammen 620 A)" in text
    assert "≈ %s A (v > 0: Überkompensation)" % _n(620 / 1.04) in text
    assert "Die Messwerte (Median der Versuche mit Zusatzimpedanz) weichen von der Anzeige ab" in text
    assert key["Kapazitiver Erdschlussstrom ICE"][1] == "≈ %s A" % _n(620 / 1.04)
    assert "ZE = 2,50 Ω (ohmsch)" in text and "ohmsch;" not in text
    assert "Dimensionierungswert." in [t for _, t in _notes(path)]            # where Z_E comes from: a footnote
    assert "Schirmstrom eines Abgangs" in text
    assert [len(r) for r in _table_with(path, "Netzgröße")][0] == 3          # no transducer column
    assert "Gemessen wurde gegen die Erdungsspieße" in text or "Gemessen wurde gegen den Erdungsspieß" in text
    assert "Hinweise zur Erdung" not in text and "Wirksamer Reduktionsfaktor" not in text
    assert not any("earthing_comparison" in f or "reduction_factor" in f for f in res["figures"])
    assert "mit dem pauschalen r = 0,5" in text
    for bad in (dict(limit_cases=[dict(limit_A=60.0, r=0.0)]), dict(limit_cases=[dict(limit=60.0)]),
                dict(station_detail="kurz"), dict(coil_controller=dict(coils={"C1": 1.0}))):
        with pytest.raises(ValueError):
            document.campaign_report(ns["RESULTS"], str(tmp_path / "bad.docx"), raw=False, bootstrap=0,
                                     **_kwargs(ns, info=dict(ns["REPORT"], **bad)))


def test_word_equations():
    """The displayed formulas are Word equations: math parts as m:oMath with indices, exponents, roots, sums,
    delimiters and functions; words and remarks between them as normal text; unbalanced text stays text."""
    from docx.oxml.ns import qn
    i18n.set_language("de")
    try:
        parts = document.omml(i18n.tr("document.eq_ires_max", ig="34,4", ih="63,9", ires="72,6"))
        kinds = ["text" if isinstance(x, str) else x.tag.split("}")[1] for x in parts]
        assert kinds == ["oMath", "text", "oMath"] and parts[1].strip() == "mit"
        xml = "".join(docx.oxml.xmlchemy.serialize_for_reading(x) for x in parts if not isinstance(x, str))
        assert xml.count("<m:rad>") == 2 and "<m:sSubSup>" in xml and ">RES,Bew<" in xml and ">72,6<" in xml
        parts = document.omml(i18n.tr("document.eq_ih_scaled", ih="58,0"))
        xml = docx.oxml.xmlchemy.serialize_for_reading(parts[0])
        assert '<m:chr m:val="∑"/>' in xml and "<m:sSup>" in xml and "<m:d>" in xml
        parts = document.omml(i18n.tr("document.eq_iw"))
        assert len([x for x in parts if not isinstance(x, str)]) == 2
        assert "<m:func>" in docx.oxml.xmlchemy.serialize_for_reading(parts[0])
        parts = document.omml(i18n.tr("document.eq_es"))
        assert 'm:begChr m:val="|"' in docx.oxml.xmlchemy.serialize_for_reading(parts[0])
        assert document.omml("I_[G] = √(1 + 2") is None and document.omml("") is None
        d = document._Doc()
        p = d.formula(i18n.tr("document.eq_ut"))
        assert p._p.findall(qn("m:oMath")) and p.alignment is not None
        d.equations = False
        p = d.formula(i18n.tr("document.eq_ut"))
        assert not p._p.findall(qn("m:oMath")) and "U" in p.text
    finally:
        i18n.set_language("de")


def test_layout_of_a_scientific_report(campaign):
    """The layout of a scientific report: body text in a justified style (hyphenation on, the language of the
    report), lists indented, captions centred in the size of the body text, a blank line between a displayed formula
    and the text, and the sources and side remarks as footnotes – every mark with its note, numbered in the order of
    the text."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt
    path = campaign["REPORT_FILE"]
    d = docx.Document(str(path))
    body = d.styles["Body Text"]
    assert body.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY and body.paragraph_format.widow_control
    assert sum(p.style.name == "Body Text" for p in d.paragraphs) > 40
    first = next(i for i, p in enumerate(d.paragraphs) if p.style.name.startswith("Heading"))
    left = [p.text[:40] for p in d.paragraphs[first:] if p.text.strip() and p.style.name == "Normal"
            and p.alignment != WD_ALIGN_PARAGRAPH.CENTER and 'TOC \\o' not in p._p.xml]
    assert left == []                                       # no text left-aligned in the default style
    cap = d.styles["Caption"]
    assert cap.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER and cap.font.size == Pt(10.0)
    assert d.styles["Normal"].font.size == Pt(10.0) and not cap.font.bold
    items = [p for p in d.paragraphs if p._p.pPr is not None and p._p.pPr.numPr is not None
             and p._p.pPr.numPr.numId.val != 0 and not p.style.name.startswith("Heading")]
    assert len(items) >= 10                                 # procedure and processing
    assert all(p.paragraph_format.left_indent == Pt(36) and p.paragraph_format.first_line_indent == Pt(-18)
               for p in items)
    ref = document._Doc(language="de")
    formulas = [p for p in d.paragraphs if p._p.find(qn("m:oMath")) is not None]
    assert len(formulas) >= 5
    for p in formulas:
        prev = p._p.getprevious()
        after_prev = 0.0 if prev.find(qn("m:oMath")) is not None else ref.body_after_pt
        if prev.find(qn("w:pPr") + "/" + qn("w:pStyle")) is not None and \
                prev.find(qn("w:pPr") + "/" + qn("w:pStyle")).get(qn("w:val")) == "BodyText":
            assert after_prev + p.paragraph_format.space_before.pt == pytest.approx(ref.line_pt, abs=0.1)
        assert p.paragraph_format.space_after.pt == pytest.approx(ref.line_pt - ref.body_before_pt, abs=0.1)
    notes, refs = _notes(path), _note_refs(path)
    assert [i for i, _ in notes] == refs == list(range(1, len(refs) + 1))
    texts = [t for _, t in notes]
    for part in ("Open-Source-Software (MIT-Lizenz): https://github.com/Ce1ectric/groundfault",
                 "nach IEEE C37.111 bzw. IEC 60255-24.", "quadratische Mittelwert gebildet (wie in IEC 61000-4-30).",
                 "Leave-one-out: Jeder Versuch wird aus den übrigen Versuchen vorhergesagt.",
                 "Zufallsfehler (1σ) der Ströme 1 % und 0,5°", "groundmeas-Messung"):
        assert any(part in t for t in texts), part
    text = _text(path)
    assert "(Leave-one-out" not in text and "Monte Carlo des Oberschwingungsanteils:" in text
    with zipfile.ZipFile(str(path)) as z:
        settings = z.read("word/settings.xml").decode("utf-8")
        styles = z.read("word/styles.xml").decode("utf-8")
        types = z.read("[Content_Types].xml").decode("utf-8")
    assert '<w:autoHyphenation w:val="true"/>' in settings and '<w:footnote w:id="-1"/>' in settings
    assert 'w:lang w:val="de-DE"' in styles and "footnotes+xml" in types
    assert _order_violations(path) == []


def test_paragraphs_lists_and_footnotes_of_the_layout(tmp_path):
    """The building blocks of the layout: a blank line between formula and text (one between two formulas), the
    sentence that leads to a formula or a list kept with it, a numbered list starting at 1 with a numbering of its
    own if the document has no list styles, the caption of a table and of a figure, footnotes from the mark-up
    [^…^] – no text of the paragraph, left out of the plain text."""
    from docx.oxml.ns import qn
    from docx.shared import Pt
    with i18n.language("de"):
        d = document._Doc(language="de")
        line = d.line_pt
        assert line == pytest.approx(1.2 * 10.0 * 1.15) and d.body_after_pt == 6.0 and d.body_before_pt == 0.0
        lead = d.para("Die Erdungsspannung folgt aus")
        f1 = d.formula("U_[E] = r·I_[RES]·Z_[E]")
        f2 = d.formula("U_[T] = U_[E]/F")
        after = d.para("mit dem Faktor F[^Ein Faktor der Vorgaben des Netzbetreibers.^].")
        assert lead.style.name == "Body Text" and lead.paragraph_format.keep_with_next
        assert d.body_after_pt + f1.paragraph_format.space_before.pt == pytest.approx(line, abs=0.1)
        assert f1.paragraph_format.space_after.pt + f2.paragraph_format.space_before.pt == pytest.approx(line, abs=0.1)
        assert f2.paragraph_format.space_after.pt + d.body_before_pt == pytest.approx(line, abs=0.1)
        assert after.text == "mit dem Faktor F." and after._p.find(".//" + qn("w:footnoteReference")) is not None
        assert document.plain("Faktor F[^Eine Anmerkung.^] gilt") == "Faktor F gilt"
        n = d.caption("table", "Werte")
        cap = d.doc.paragraphs[-1]
        assert n == 1 and cap.paragraph_format.space_after == Pt(3)
        assert cap.paragraph_format.space_before.pt == pytest.approx(line - d.body_after_pt, abs=0.1)
        d.para("Schritte:")
        d.bullets(["eins", "zwei"], numbered=True)
        d.bullets(["drei", "vier"], numbered=True)
        items = d.doc.paragraphs[-4:]
        assert [p.style.name for p in items] == ["List Number"] * 4
        ids = [p._p.pPr.numPr.numId.val for p in items]
        assert ids[0] == ids[1] != ids[2] == ids[3]          # every numbered list starts at 1
        assert items[-1].paragraph_format.space_after == Pt(6)
        d.save(str(tmp_path / "layout.docx"))
        assert _notes(tmp_path / "layout.docx") == [(1, "Ein Faktor der Vorgaben des Netzbetreibers.")]
        assert _order_violations(tmp_path / "layout.docx") == []
        own = document._Doc(styles={"list_number": "none", "list_bullet": "none"}, language="en")
        own.bullets(["one", "two"], numbered=True)
        own.bullets(["three"], numbered=True)
        own.bullets(["a", "b"])
        ps = own.doc.paragraphs[-5:]
        assert all(p.style.name == "List Paragraph" and p.paragraph_format.left_indent == Pt(36) for p in ps)
        numbering = own.doc.part.numbering_part.element
        nums = {n.get(qn("w:numId")): n for n in numbering.findall(qn("w:num"))}
        first, second = (nums[str(p._p.pPr.numPr.numId.val)] for p in (ps[0], ps[2]))
        assert first is not second and second.find(qn("w:lvlOverride")) is not None
        fmt = {a.get(qn("w:abstractNumId")): a.find(qn("w:lvl") + "/" + qn("w:numFmt")).get(qn("w:val"))
               for a in numbering.findall(qn("w:abstractNum"))}
        abstract = lambda p: nums[str(p._p.pPr.numPr.numId.val)].find(qn("w:abstractNumId")).get(qn("w:val"))
        assert fmt[abstract(ps[0])] == "decimal" and fmt[abstract(ps[-1])] == "bullet"
        own.save(str(tmp_path / "own.docx"))
        with zipfile.ZipFile(str(tmp_path / "own.docx")) as z:
            assert 'w:lang w:val="en-GB"' in z.read("word/styles.xml").decode("utf-8")
        assert _order_violations(tmp_path / "own.docx") == []


def test_footnotes_of_a_template_with_notes(campaign, tmp_path):
    """A template made from a report keeps the separators of its footnotes, not its notes; the new report numbers
    its notes from 1 in the one footnotes part, and takes the language of the report (English, the template
    German)."""
    template = document.make_template(str(campaign["REPORT_FILE"]), str(tmp_path / "template.docx"))
    assert _notes(template) == []
    with zipfile.ZipFile(template) as z:
        assert len(_note_parts(z)) == 1
    path = tmp_path / "report_en.docx"
    info = dict(network="20 kV network of a test substation", authors=["A. Author"],
                Z_E_basis="design value of the earthing systems")
    document.campaign_report(campaign["RESULTS"], str(path), language="en", template=template, raw=False,
                             bootstrap=0, **_kwargs(campaign, info=info))
    notes, refs = _notes(path), _note_refs(path)
    assert [i for i, _ in notes] == refs == list(range(1, len(refs) + 1)) and len(refs) >= 5
    assert "Design value of the earthing systems." in [t for _, t in notes]
    assert any("Format of transient recordings, IEEE C37.111 and IEC 60255-24." == t for _, t in notes)
    with zipfile.ZipFile(str(path)) as z:
        assert len(_note_parts(z)) == 1 and 'w:lang w:val="en-GB"' in z.read("word/styles.xml").decode("utf-8")
    assert _order_violations(path) == []


def test_detuning_change_explained_by_the_coil_current_balance():
    """A reactive current that changes between the tests at an equal active current: if the coil current balance of
    every single evaluation gives ΔIv/ΔΣI_L ≈ 1, the coils were adjusted during the fault (the network stayed the
    same); without the balance, or with a slope outside the range of the check, the cause stays open."""
    from types import SimpleNamespace
    check = i18n.tr("evaluation.check_stationarity", lang="de")

    def result(slope, warn):
        pl = pd.DataFrame({"check": [check], "status": ["WARNING" if warn else "OK"]})
        return SimpleNamespace(tables={"plausibility": pl},
                               values={} if slope is None else {"coil_balance_slope": slope})

    names = ["Test 1", "Test 2", "Test 3"]
    tests = {"Test 1": {"Iv": 13.0, "Iw": 31.6}, "Test 2": {"Iv": 10.3, "Iw": 31.0}, "Test 3": {"Iv": 10.1, "Iw": 31.2}}
    w = document._Writer.__new__(document._Writer)
    with i18n.language("de"):
        w.c = SimpleNamespace(subs=names, tests=tests, results=dict(zip(names, (
            result(0.98, True), result(0.96, True), result(0.99, False)))))
        text = w.detuning_text()
        assert text.startswith("Der Blindstrom ändert sich zwischen den Versuchen (13,0 / 10,3 / 10,1 A)")
        assert "in Test 1 und Test 2 auch während des Fehlers" in text
        assert "0,98 / 0,96 / 0,99" in text and "Die Spulen wurden verstellt" in text
        w.c.results["Test 3"] = result(None, False)                     # one test without a balance: named
        assert "Test 1 0,98; Test 2 0,96" in w.detuning_text()
        w.c.results["Test 2"] = result(None, True)                      # a drifting test without a balance
        assert "die Ursache ist nicht geklärt" in w.detuning_text()
        w.c.results["Test 2"] = result(3.0, True)                       # wrong ratio of a coil channel
        assert "die Ursache ist nicht geklärt" in w.detuning_text()
    with i18n.language("en"):
        w.c.results = dict(zip(names, (result(0.98, True), result(0.96, True), result(0.99, False))))
        assert "the coils were adjusted" in w.detuning_text()
    w.c.tests = {n: {"Iv": 10.0, "Iw": 31.0} for n in names}             # no change: no sentence
    assert w.detuning_text() == ""


def test_pq_measurement_in_several_periods(campaign, tmp_path):
    """A PQ measurement of two periods weeks apart: the report names both periods, the quantile of the dominant order
    per period with the assessment voltage against the highest one, the periods in the key results, and in the table
    of uncertainties I_RES(Z_E) scaled with each period alone."""
    import datetime
    ns = campaign
    first = ns["CAMPAIGN_FOLDER"] / "1_Rohdaten" / "PQ" / "UW_PQ_Export.csv"
    second = demo.make_pq_export(str(tmp_path / "pq_spring.csv"), days=14.0, seed=7,
                                 start=datetime.datetime(2026, 3, 2))
    a = pq.assess(pq.read([str(first), second["path"]]), factor=1.0)
    assert len(a.coverage["periods"]) == 2
    path = tmp_path / "periods.docx"
    document.campaign_report(ns["RESULTS"], str(path), raw=False, bootstrap=0, **_kwargs(ns, pq=a))
    text = _text(path)
    assert "in 2 Abschnitten aufgezeichnet" in text and "15.12.2025 bis" in text and "(14,0 Tage)" in text
    per = a.periods()
    dom = a.dominant_order
    assert ("Je Abschnitt beträgt das Quantil der %d. Ordnung %s V" % (dom, " / ".join(_n(v) for v in per[dom]))) in text
    assert "die Bewertungsspannung (Quantil × ks = " in text
    key = {r[0]: r for r in _table_with(path, "Größe")}
    assert key["PQ-Messung"][1].startswith("49,0 Tage in 2 Abschnitten (35,0 Tage lückenlos)")
    unc = {r[0]: r for r in _table_with(path, "Beitrag")}
    assert "Abschnitte der PQ-Messung" in unc and unc["Abschnitte der PQ-Messung"][1].startswith(
        "IRES(ZE) skaliert mit einem Abschnitt allein")
